"""PEF 등록부 엑셀 데이터를 KIIS DB에 임포트하는 스크립트.

사용법:
    cd kiis
    python -m scripts.import_pef_registry --file "../기관전용 사모집합투자기구 현황(2024.12월말 기준)_게시용.xlsx"
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert as pg_insert

# 프로젝트 루트를 sys.path에 추가
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.core.database import async_session_factory, engine
from app.models.fund import Fund, FundGP

logger = logging.getLogger(__name__)

# 펀드 유형 분류 키워드
PROJECT_KEYWORDS = ["프로젝트", "특정", "목적", "PF"]

REFERENCE_DATE = "2024.12월말"


def classify_fund_type(fund_name: str) -> str:
    """펀드명에서 blind/project를 분류한다."""
    for kw in PROJECT_KEYWORDS:
        if kw in fund_name:
            return "project"
    return "blind"


def parse_excel(file_path: str) -> list[dict]:
    """엑셀 파일을 파싱하여 딕셔너리 리스트로 반환한다."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    records: list[dict] = []
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
        # 컬럼: A(None), B(순번), C(설립근거법률), D(PEF명칭), E(등록일), F(GP1), G(GP2), H(GP3), I(총약정액)
        seq = row[1]
        if seq is None or not isinstance(seq, (int, float)):
            continue  # 빈 행 또는 푸터 스킵

        seq = int(seq)
        legal_basis = str(row[2]).strip() if row[2] else ""
        fund_name = str(row[3]).strip() if row[3] else ""
        established_raw = row[4]
        gp1 = str(row[5]).strip() if row[5] else ""
        gp2 = str(row[6]).strip() if row[6] else ""
        gp3 = str(row[7]).strip() if row[7] else ""
        total_commitment = row[8]  # 억원 단위

        if not fund_name or not gp1:
            logger.warning("Row %d: fund_name 또는 GP1 누락, 스킵", seq)
            continue

        # 등록일 파싱
        established_date: date | None = None
        if isinstance(established_raw, datetime):
            established_date = established_raw.date()
        elif isinstance(established_raw, date):
            established_date = established_raw

        # 총약정액 변환 (억원 → 원)
        total_amount: Decimal | None = None
        if total_commitment is not None:
            try:
                total_amount = Decimal(str(total_commitment)) * Decimal("100000000")
            except Exception:
                logger.warning("Row %d: 총약정액 변환 실패: %s", seq, total_commitment)

        # Co-GP 판별
        gp_names = [gp1]
        if gp2:
            gp_names.append(gp2)
        if gp3:
            gp_names.append(gp3)
        is_co_gp = len(gp_names) >= 2

        records.append(
            {
                "seq": seq,
                "fund_code": f"PEF-{seq:04d}",
                "fund_name": fund_name,
                "legal_basis": legal_basis,
                "established_date": established_date,
                "vintage_year": established_date.year if established_date else None,
                "total_amount": total_amount,
                "company_name": gp1,  # 기본 GP (기존 검색 호환)
                "fund_type": classify_fund_type(fund_name),
                "is_co_gp": is_co_gp,
                "gp_names": gp_names,
                "gp_roles": ["gp1"] + (["gp2"] if gp2 else []) + (["gp3"] if gp3 else []),
            }
        )

    wb.close()
    return records


async def import_records(records: list[dict]) -> tuple[int, int]:
    """레코드를 DB에 삽입/업데이트한다. (upsert)"""
    inserted = 0
    updated = 0

    async with async_session_factory() as session:
        for rec in records:
            # Fund upsert
            stmt = (
                pg_insert(Fund)
                .values(
                    fund_code=rec["fund_code"],
                    fund_name=rec["fund_name"],
                    company_name=rec["company_name"],
                    fund_type=rec["fund_type"],
                    legal_type="professional_private",
                    asset_class="pef",
                    fund_category="PEF",
                    total_amount=rec["total_amount"],
                    established_date=rec["established_date"],
                    vintage_year=rec["vintage_year"],
                    is_active=True,
                    is_maturity_alert=False,
                    data_source="pef_registry",
                    legal_basis=rec["legal_basis"],
                    is_co_gp=rec["is_co_gp"],
                    reference_date=REFERENCE_DATE,
                )
                .on_conflict_do_update(
                    index_elements=["fund_code"],
                    set_={
                        "fund_name": rec["fund_name"],
                        "company_name": rec["company_name"],
                        "fund_type": rec["fund_type"],
                        "total_amount": rec["total_amount"],
                        "established_date": rec["established_date"],
                        "vintage_year": rec["vintage_year"],
                        "legal_basis": rec["legal_basis"],
                        "is_co_gp": rec["is_co_gp"],
                        "reference_date": REFERENCE_DATE,
                    },
                )
                .returning(Fund.id)
            )

            result = await session.execute(stmt)
            fund_id = result.scalar_one()

            # 기존 fund_gps 삭제 후 재삽입
            existing_gps = await session.execute(select(FundGP.id).where(FundGP.fund_id == fund_id))
            if existing_gps.scalars().first() is not None:
                from sqlalchemy import delete

                await session.execute(delete(FundGP).where(FundGP.fund_id == fund_id))
                updated += 1
            else:
                inserted += 1

            # GP 관계 삽입
            for gp_name, gp_role in zip(rec["gp_names"], rec["gp_roles"], strict=False):
                gp_stmt = (
                    pg_insert(FundGP)
                    .values(
                        fund_id=fund_id,
                        gp_name=gp_name,
                        gp_role=gp_role,
                    )
                    .on_conflict_do_nothing(constraint="uq_fund_gp")
                )
                await session.execute(gp_stmt)

        await session.commit()

    return inserted, updated


async def main(file_path: str) -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

    logger.info("엑셀 파일 파싱 시작: %s", file_path)
    records = parse_excel(file_path)
    logger.info("파싱 완료: %d건", len(records))

    if not records:
        logger.error("파싱된 레코드가 없습니다. 파일 구조를 확인하세요.")
        return

    logger.info("DB 임포트 시작...")
    inserted, updated = await import_records(records)
    logger.info("임포트 완료: 신규 %d건, 갱신 %d건", inserted, updated)

    # GP 통계
    gp_set: set[str] = set()
    co_gp_count = 0
    for rec in records:
        gp_set.update(rec["gp_names"])
        if rec["is_co_gp"]:
            co_gp_count += 1
    logger.info("총 고유 GP: %d명, Co-GP 펀드: %d건", len(gp_set), co_gp_count)

    await engine.dispose()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="PEF 등록부 엑셀 → KIIS DB 임포트")
    parser.add_argument("--file", required=True, help="엑셀 파일 경로")
    args = parser.parse_args()

    asyncio.run(main(args.file))
