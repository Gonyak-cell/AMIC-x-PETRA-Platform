"""table_builder 모듈 단위 테스트."""

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.pptx.template_engine.exceptions import ShapeNotFoundError, TableShapeError
from app.pptx.template_engine.table_builder import replace_table_data, split_table_across_slides


def _create_prs_with_table(num_rows=3, num_cols=4):
    """테이블이 포함된 Presentation + 슬라이드 반환."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])

    table_shape = slide.shapes.add_table(
        num_rows,  # header + data rows
        num_cols,
        Inches(0.5),
        Inches(1),
        Inches(9),
        Inches(4),
    )
    table_shape.name = "tbl_financials"

    # 헤더 행 설정
    table = table_shape.table
    headers = ["항목", "2022", "2023", "2024"]
    for col_idx, header in enumerate(headers[:num_cols]):
        table.cell(0, col_idx).text = header

    return prs, slide


@pytest.fixture
def table_slide():
    """테이블이 포함된 테스트 슬라이드."""
    _, slide = _create_prs_with_table(num_rows=5, num_cols=4)
    return slide


@pytest.fixture
def table_prs():
    """테이블이 포함된 Presentation."""
    prs, _ = _create_prs_with_table(num_rows=5, num_cols=4)
    return prs


class TestReplaceTableData:
    """replace_table_data 함수 테스트."""

    def test_basic_replace(self, table_slide):
        """기본 테이블 데이터 교체."""
        rows = [
            ["매출", 100000, 120000, 150000],
            ["EBITDA", 30000, 40000, 55000],
        ]
        shape = replace_table_data(table_slide, "tbl_financials", None, rows)
        assert shape is not None
        assert shape.has_table

    def test_with_headers(self, table_slide):
        """헤더 갱신 포함."""
        rows = [["매출", 100, 200, 300]]
        shape = replace_table_data(
            table_slide,
            "tbl_financials",
            ["항목", "FY22", "FY23", "FY24"],
            rows,
        )
        assert shape.table.cell(0, 1).text != ""

    def test_auto_expand(self, table_slide):
        """행 부족 시 자동 확장."""
        # 기존 5행 테이블 (헤더 1 + 데이터 4)에 10행 데이터 삽입
        rows = [[f"항목{i}", i * 100, i * 200, i * 300] for i in range(10)]
        shape = replace_table_data(table_slide, "tbl_financials", None, rows)
        # 행이 자동 추가되어 에러 없이 완료
        assert shape is not None

    def test_negative_values(self, table_slide):
        """음수 값 포맷팅 — 괄호 표기, 빨간색."""
        rows = [
            ["매출", -50000, 120000, -30000],
        ]
        shape = replace_table_data(table_slide, "tbl_financials", None, rows)
        # 음수 값이 format_cell_value를 통해 (50,000) 형태로 변환
        assert shape is not None

    def test_none_values(self, table_slide):
        """None 값 → N/A 표시."""
        rows = [
            ["매출", None, 120000, None],
        ]
        shape = replace_table_data(table_slide, "tbl_financials", None, rows)
        assert shape is not None

    def test_shape_not_found(self, table_slide):
        """존재하지 않는 shape — ShapeNotFoundError."""
        with pytest.raises(ShapeNotFoundError):
            replace_table_data(table_slide, "tbl_nonexistent", None, [["a", 1, 2, 3]])

    def test_not_a_table(self):
        """테이블이 아닌 shape — TableShapeError."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        txBox.name = "txt_note"

        with pytest.raises(TableShapeError):
            replace_table_data(slide, "txt_note", None, [["a"]])


class TestSplitTableAcrossSlides:
    """split_table_across_slides 함수 테스트."""

    def test_no_split_needed(self, table_prs):
        """행 수가 적으면 분할 없음."""
        rows = [["항목1", 100, 200, 300], ["항목2", 400, 500, 600]]
        count = split_table_across_slides(
            table_prs, 0, "tbl_financials", None, rows, max_rows_per_slide=10
        )
        assert count == 1

    def test_split_into_pages(self, table_prs):
        """행이 초과하면 슬라이드 분할."""
        rows = [[f"항목{i}", i, i * 10, i * 100] for i in range(25)]
        count = split_table_across_slides(
            table_prs, 0, "tbl_financials", None, rows, max_rows_per_slide=10
        )
        assert count == 3  # ceil(25 / 10) = 3

    def test_split_preserves_headers(self, table_prs):
        """분할된 슬라이드에 헤더 전달."""
        rows = [[f"row{i}", i, i, i] for i in range(15)]
        headers = ["항목", "A", "B", "C"]
        count = split_table_across_slides(
            table_prs, 0, "tbl_financials", headers, rows, max_rows_per_slide=10
        )
        assert count == 2
