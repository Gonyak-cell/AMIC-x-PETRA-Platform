"""골든 샘플 회귀 테스트 — TM/DM 정규 입력 → PPTX 구조 안정성 검증.

> 마지막 수정: 2026-03-13 22:30:00

정규 fixture(input.json)로 PPTX를 생성하고,
manifest.json의 기대 구조와 비교하여 회귀를 감지한다.

실행 방법:
    pytest tests/test_golden_samples.py -v
    pytest tests/test_golden_samples.py --update-golden -v  # 골든 파일 갱신
"""

from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Any

import pytest

from app.pptx.memo_generator import TEMPLATE_PATH, generate_memo

# ── 경로 상수 ──────────────────────────────────────────────────────

_GOLDEN_DIR = Path(__file__).parent / "fixtures" / "golden_samples"


def _check_template() -> bool:
    """마스터 템플릿 존재 확인."""
    return TEMPLATE_PATH.exists()


TEMPLATE_READY = _check_template()
skip_no_template = pytest.mark.skipif(
    not TEMPLATE_READY,
    reason=f"마스터 템플릿 없음: {TEMPLATE_PATH}",
)


# ── 스냅샷 추출 헬퍼 ──────────────────────────────────────────────


def _shape_type_label(shape) -> str:
    """python-pptx shape를 사람이 읽을 수 있는 문자열로 변환."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    type_map = {
        MSO_SHAPE_TYPE.TABLE: "table",
        MSO_SHAPE_TYPE.PICTURE: "picture",
        MSO_SHAPE_TYPE.TEXT_BOX: "text_box",
        MSO_SHAPE_TYPE.AUTO_SHAPE: "auto_shape",
        MSO_SHAPE_TYPE.PLACEHOLDER: "placeholder",
        MSO_SHAPE_TYPE.GROUP: "group",
        MSO_SHAPE_TYPE.CHART: "chart",
    }
    return type_map.get(shape.shape_type, str(shape.shape_type))


def extract_pptx_snapshot(pptx_path: str) -> dict[str, Any]:
    """PPTX에서 결정론적 구조 스냅샷을 추출한다."""
    from pptx import Presentation

    prs = Presentation(pptx_path)
    slides_data: list[dict[str, Any]] = []

    for idx, slide in enumerate(prs.slides):
        texts: list[str] = []
        shape_types: list[str] = []
        has_table = False
        has_picture = False

        for shape in slide.shapes:
            stype = _shape_type_label(shape)
            shape_types.append(stype)

            if stype == "table":
                has_table = True
            elif stype == "picture":
                has_picture = True

            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        texts.append(text)

        shape_types.sort()
        slides_data.append(
            {
                "index": idx,
                "shape_count": len(slide.shapes),
                "text_content": texts,
                "shape_types": shape_types,
                "has_table": has_table,
                "has_picture": has_picture,
                "text_length": sum(len(t) for t in texts),
            }
        )

    return {
        "version": 1,
        "total_slides": len(prs.slides),
        "slides": slides_data,
        "total_shape_count": sum(s["shape_count"] for s in slides_data),
        "total_text_length": sum(s["text_length"] for s in slides_data),
    }


def load_golden(variant: str) -> dict[str, Any] | None:
    """골든 스냅샷 파일 로드."""
    golden_path = _GOLDEN_DIR / variant / "snapshot.json"
    if not golden_path.exists():
        return None
    return json.loads(golden_path.read_text(encoding="utf-8"))


def save_golden(variant: str, snapshot: dict[str, Any]) -> None:
    """골든 스냅샷 저장."""
    golden_path = _GOLDEN_DIR / variant / "snapshot.json"
    golden_path.parent.mkdir(parents=True, exist_ok=True)
    golden_path.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def compare_snapshots(
    actual: dict[str, Any],
    golden: dict[str, Any],
) -> list[str]:
    """실제 vs 골든 스냅샷 비교 → 차이 목록 반환."""
    diffs: list[str] = []

    if actual["total_slides"] != golden["total_slides"]:
        diffs.append(f"슬라이드 수: {actual['total_slides']} vs 골든 {golden['total_slides']}")

    min_slides = min(len(actual["slides"]), len(golden["slides"]))
    for i in range(min_slides):
        a_slide = actual["slides"][i]
        g_slide = golden["slides"][i]

        if a_slide["shape_count"] != g_slide["shape_count"]:
            diffs.append(f"슬라이드 {i}: shape 수 {a_slide['shape_count']} vs {g_slide['shape_count']}")

        if a_slide["shape_types"] != g_slide["shape_types"]:
            diffs.append(f"슬라이드 {i}: shape 타입 불일치")

        if a_slide["has_table"] != g_slide["has_table"]:
            diffs.append(f"슬라이드 {i}: 테이블 존재 {a_slide['has_table']} vs {g_slide['has_table']}")

        if a_slide["has_picture"] != g_slide["has_picture"]:
            diffs.append(f"슬라이드 {i}: 이미지 존재 {a_slide['has_picture']} vs {g_slide['has_picture']}")

    return diffs


# ── 매니페스트 검증 ──────────────────────────────────────────────


def load_manifest(variant: str) -> dict[str, Any]:
    """manifest.json 로드."""
    manifest_path = _GOLDEN_DIR / variant / "manifest.json"
    return json.loads(manifest_path.read_text(encoding="utf-8"))


def load_input(variant: str) -> dict[str, Any]:
    """input.json 로드."""
    input_path = _GOLDEN_DIR / variant / "input.json"
    return json.loads(input_path.read_text(encoding="utf-8"))


def validate_against_manifest(
    snapshot: dict[str, Any],
    manifest: dict[str, Any],
) -> list[str]:
    """스냅샷을 매니페스트 기대값과 비교 → 차이 목록."""
    issues: list[str] = []

    slide_count = snapshot["total_slides"]
    min_slides = manifest.get("expected_min_slides", 0)
    max_slides = manifest.get("expected_max_slides", 100)

    if slide_count < min_slides:
        issues.append(f"슬라이드 수 부족: {slide_count} < 최소 {min_slides}")
    if slide_count > max_slides:
        issues.append(f"슬라이드 수 초과: {slide_count} > 최대 {max_slides}")

    required = manifest.get("required_elements", {})
    if required.get("has_tables"):
        has_any_table = any(s["has_table"] for s in snapshot["slides"])
        if not has_any_table:
            issues.append("매니페스트 요구: 테이블이 1개 이상 필요하지만 없음")

    return issues


# ── 테스트 클래스 ──────────────────────────────────────────────────


@pytest.fixture
def update_golden(request) -> bool:
    """--update-golden CLI 옵션 값 (conftest.py에서 등록)."""
    return request.config.getoption("--update-golden", default=False)


class TestFixtureFiles:
    """골든 샘플 fixture 파일 존재 확인."""

    @pytest.mark.parametrize("variant", ["tm", "dm"])
    def test_input_json_exists(self, variant: str) -> None:
        """input.json이 존재하고 유효한 JSON인지."""
        data = load_input(variant)
        assert "memo_type" in data
        assert "project_code" in data
        assert "content" in data

    @pytest.mark.parametrize("variant", ["tm", "dm"])
    def test_manifest_json_exists(self, variant: str) -> None:
        """manifest.json이 존재하고 필수 필드를 포함하는지."""
        manifest = load_manifest(variant)
        assert "template_id" in manifest
        assert "doc_type" in manifest
        assert "expected_min_slides" in manifest
        assert "slides" in manifest


@skip_no_template
class TestGoldenSnapshots:
    """골든 스냅샷 회귀 테스트 — PPTX 생성 후 구조 비교."""

    @pytest.mark.parametrize("variant", ["tm", "dm"])
    def test_golden_snapshot_regression(
        self,
        variant: str,
        tmp_path: Path,
        update_golden: bool,
    ) -> None:
        """정규 입력 → PPTX 생성 → 골든 스냅샷 비교."""
        input_data = load_input(variant)
        output_path = str(tmp_path / f"golden_{variant}.pptx")

        result = generate_memo(
            memo_type=input_data["memo_type"],
            project_code=input_data["project_code"],
            output_path=output_path,
            content=input_data["content"],
        )

        assert os.path.exists(result.output_path)
        assert result.file_size_bytes > 0

        snapshot = extract_pptx_snapshot(result.output_path)

        if update_golden:
            save_golden(variant, snapshot)
            pytest.skip(f"골든 파일 갱신 완료: {variant}")

        golden = load_golden(variant)
        if golden is None:
            save_golden(variant, snapshot)
            pytest.skip(f"초기 골든 파일 생성: {variant} (다음 실행부터 비교)")

        diffs = compare_snapshots(snapshot, golden)
        assert not diffs, f"골든 스냅샷 회귀 감지 ({variant}):\n" + "\n".join(diffs)

    @pytest.mark.parametrize("variant", ["tm", "dm"])
    def test_manifest_compliance(
        self,
        variant: str,
        tmp_path: Path,
    ) -> None:
        """정규 입력 → PPTX 생성 → 매니페스트 기대값 충족."""
        input_data = load_input(variant)
        output_path = str(tmp_path / f"manifest_{variant}.pptx")

        result = generate_memo(
            memo_type=input_data["memo_type"],
            project_code=input_data["project_code"],
            output_path=output_path,
            content=input_data["content"],
        )

        snapshot = extract_pptx_snapshot(result.output_path)
        manifest = load_manifest(variant)

        issues = validate_against_manifest(snapshot, manifest)
        assert not issues, f"매니페스트 위반 ({variant}):\n" + "\n".join(issues)

    @pytest.mark.parametrize("variant", ["tm", "dm"])
    def test_slide_count_stable(
        self,
        variant: str,
        tmp_path: Path,
    ) -> None:
        """동일 입력 2회 생성 → 슬라이드 수 일치 (결정론 검증)."""
        input_data = load_input(variant)

        results = []
        for i in range(2):
            output = str(tmp_path / f"determinism_{variant}_{i}.pptx")
            r = generate_memo(
                memo_type=input_data["memo_type"],
                project_code=input_data["project_code"],
                output_path=output,
                content=input_data["content"],
            )
            results.append(r)

        assert results[0].slide_count == results[1].slide_count, (
            f"결정론 위반: 동일 입력에서 슬라이드 수 불일치 ({results[0].slide_count} vs {results[1].slide_count})"
        )

    @pytest.mark.parametrize("variant", ["tm", "dm"])
    def test_performance_metrics_present(
        self,
        variant: str,
        tmp_path: Path,
    ) -> None:
        """생성 결과에 성능 메트릭이 포함되는지 확인."""
        input_data = load_input(variant)
        output = str(tmp_path / f"metrics_{variant}.pptx")

        result = generate_memo(
            memo_type=input_data["memo_type"],
            project_code=input_data["project_code"],
            output_path=output,
            content=input_data["content"],
        )

        assert result.template_load_ms >= 0
        assert result.render_ms >= 0
        assert result.persist_ms >= 0
