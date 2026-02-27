"""통합 엔진 (engine.py) + 스키마 + 서비스 테스트."""

import base64
from io import BytesIO

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app.pptx.template_engine.engine import TemplateVisualizationResult, process_template


def _make_test_image_base64(width=100, height=80) -> str:
    """테스트용 PNG 이미지 Base64 생성."""
    img = Image.new("RGB", (width, height), "blue")
    buf = BytesIO()
    img.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode()


@pytest.fixture
def template_pptx(tmp_path):
    """통합 테스트용 템플릿 PPTX 파일 생성."""
    prs = Presentation()

    # ── 슬라이드 0: 차트 + 텍스트 ──
    slide0 = prs.slides.add_slide(prs.slide_layouts[5])

    chart_data = CategoryChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("Default", (0, 0))
    cs = slide0.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.5), Inches(0.5), Inches(4), Inches(3),
        chart_data,
    )
    cs.name = "chart_revenue"

    tx = slide0.shapes.add_textbox(Inches(5), Inches(0.5), Inches(4), Inches(1))
    tx.name = "txt_title"
    tf = tx.text_frame
    tf.paragraphs[0].text = "{{company}} — {{year}}"

    # ── 슬라이드 1: 테이블 + 이미지 ──
    slide1 = prs.slides.add_slide(prs.slide_layouts[5])

    tbl = slide1.shapes.add_table(4, 3, Inches(0.5), Inches(0.5), Inches(5), Inches(3))
    tbl.name = "tbl_financials"

    img_ph = slide1.shapes.add_shape(
        1,  # Rectangle
        Inches(6), Inches(0.5), Inches(3), Inches(3),
    )
    img_ph.name = "img_product"

    # 저장
    template_path = tmp_path / "test_template.pptx"
    prs.save(str(template_path))
    return template_path


class TestProcessTemplate:
    """process_template 통합 테스트."""

    def test_chart_replacement(self, template_pptx, tmp_path):
        """차트 데이터 교체."""
        output_path = tmp_path / "output_chart.pptx"
        instructions = {
            "slides": [
                {
                    "slide_index": 0,
                    "charts": [
                        {
                            "shape_name": "chart_revenue",
                            "categories": ["2022", "2023", "2024"],
                            "series": [
                                {"name": "매출", "values": [100, 200, 300]},
                            ],
                        }
                    ],
                }
            ]
        }
        result = process_template(str(template_pptx), str(output_path), instructions)
        assert result.charts_updated == 1
        assert result.output_path == str(output_path)
        assert output_path.exists()

    def test_text_placeholder(self, template_pptx, tmp_path):
        """텍스트 플레이스홀더 교체."""
        output_path = tmp_path / "output_text.pptx"
        instructions = {
            "slides": [
                {
                    "slide_index": 0,
                    "texts": [
                        {
                            "shape_name": "txt_title",
                            "mode": "placeholder",
                            "replacements": {
                                "company": "알파테크",
                                "year": "2026",
                            },
                        }
                    ],
                }
            ]
        }
        result = process_template(str(template_pptx), str(output_path), instructions)
        assert result.texts_replaced == 1

    def test_table_replacement(self, template_pptx, tmp_path):
        """테이블 데이터 교체."""
        output_path = tmp_path / "output_table.pptx"
        instructions = {
            "slides": [
                {
                    "slide_index": 1,
                    "tables": [
                        {
                            "shape_name": "tbl_financials",
                            "rows": [
                                ["매출", 100000, 150000],
                                ["영업이익", 20000, 35000],
                            ],
                        }
                    ],
                }
            ]
        }
        result = process_template(str(template_pptx), str(output_path), instructions)
        assert result.tables_updated == 1

    def test_image_placement(self, template_pptx, tmp_path):
        """이미지 삽입."""
        output_path = tmp_path / "output_image.pptx"
        b64 = _make_test_image_base64()
        instructions = {
            "slides": [
                {
                    "slide_index": 1,
                    "images": [
                        {
                            "shape_name": "img_product",
                            "source_base64": b64,
                        }
                    ],
                }
            ]
        }
        result = process_template(str(template_pptx), str(output_path), instructions)
        assert result.images_placed == 1

    def test_combined_operations(self, template_pptx, tmp_path):
        """차트 + 텍스트 + 테이블 + 이미지 복합 처리."""
        output_path = tmp_path / "output_combined.pptx"
        b64 = _make_test_image_base64()
        instructions = {
            "slides": [
                {
                    "slide_index": 0,
                    "charts": [
                        {
                            "shape_name": "chart_revenue",
                            "categories": ["Q1", "Q2"],
                            "series": [{"name": "매출", "values": [50, 80]}],
                        }
                    ],
                    "texts": [
                        {
                            "shape_name": "txt_title",
                            "mode": "placeholder",
                            "replacements": {"company": "베타시스템", "year": "FY25"},
                        }
                    ],
                },
                {
                    "slide_index": 1,
                    "tables": [
                        {
                            "shape_name": "tbl_financials",
                            "rows": [["매출", 500, 700]],
                        }
                    ],
                    "images": [
                        {
                            "shape_name": "img_product",
                            "source_base64": b64,
                        }
                    ],
                },
            ]
        }
        result = process_template(str(template_pptx), str(output_path), instructions)
        assert result.charts_updated == 1
        assert result.texts_replaced == 1
        assert result.tables_updated == 1
        assert result.images_placed == 1
        assert len(result.errors) == 0
        assert output_path.exists()

    def test_invalid_slide_index(self, template_pptx, tmp_path):
        """범위 초과 슬라이드 인덱스 — 에러 기록."""
        output_path = tmp_path / "output_invalid.pptx"
        instructions = {
            "slides": [
                {
                    "slide_index": 99,  # 존재하지 않는 슬라이드
                    "charts": [],
                }
            ]
        }
        result = process_template(str(template_pptx), str(output_path), instructions)
        assert len(result.errors) >= 1
        assert "99" in result.errors[0]

    def test_nonexistent_shape_error(self, template_pptx, tmp_path):
        """존재하지 않는 shape — 에러 기록 (전체 실패 아님)."""
        output_path = tmp_path / "output_shape_err.pptx"
        instructions = {
            "slides": [
                {
                    "slide_index": 0,
                    "charts": [
                        {
                            "shape_name": "chart_nonexistent",
                            "categories": ["A"],
                            "series": [{"name": "S", "values": [1]}],
                        }
                    ],
                }
            ]
        }
        result = process_template(str(template_pptx), str(output_path), instructions)
        assert len(result.errors) >= 1
        assert output_path.exists()  # 파일은 생성됨 (에러 있어도 나머지 처리)

    def test_empty_slides(self, template_pptx, tmp_path):
        """빈 slides 리스트 — 에러 없이 저장."""
        output_path = tmp_path / "output_empty.pptx"
        result = process_template(str(template_pptx), str(output_path), {"slides": []})
        assert result.charts_updated == 0
        assert result.tables_updated == 0
        assert len(result.errors) == 0
        assert output_path.exists()

    def test_template_not_found(self):
        """존재하지 않는 템플릿 — FileNotFoundError."""
        with pytest.raises(FileNotFoundError):
            process_template("/nonexistent/template.pptx", "/tmp/out.pptx", {"slides": []})


class TestTemplateVisualizationResult:
    """TemplateVisualizationResult 데이터클래스 테스트."""

    def test_defaults(self):
        """기본값 검증."""
        result = TemplateVisualizationResult()
        assert result.output_path == ""
        assert result.charts_updated == 0
        assert result.tables_updated == 0
        assert result.images_placed == 0
        assert result.texts_replaced == 0
        assert result.slides_added == 0
        assert result.errors == []
