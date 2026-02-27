"""text_replacer 모듈 단위 테스트."""

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app.pptx.template_engine.exceptions import ShapeNotFoundError, TextShapeError
from app.pptx.template_engine.text_replacer import (
    replace_placeholders_in_text,
    replace_text_preserving_format,
    replace_timeline_texts,
)


@pytest.fixture
def text_slide():
    """텍스트 shape들이 포함된 테스트 슬라이드."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])

    # 일반 텍스트 shape (2 paragraphs)
    tx = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    tx.name = "txt_body"
    tf = tx.text_frame
    tf.paragraphs[0].text = "첫 번째 문단"
    p2 = tf.add_paragraph()
    p2.text = "두 번째 문단"

    # 볼드 서식 텍스트 (서식 보존 테스트용)
    tx_bold = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1))
    tx_bold.name = "txt_bold"
    tf_bold = tx_bold.text_frame
    run = tf_bold.paragraphs[0].add_run()
    run.text = "볼드 텍스트"
    run.font.bold = True
    run.font.size = Pt(14)

    # 플레이스홀더 텍스트
    tx_ph = slide.shapes.add_textbox(Inches(6), Inches(1), Inches(3), Inches(2))
    tx_ph.name = "txt_placeholder"
    tf_ph = tx_ph.text_frame
    tf_ph.paragraphs[0].text = "회사명: {{company_name}}, 일자: {{date}}"

    # 타임라인용 연속 텍스트
    for i in range(1, 4):
        tx_tl = slide.shapes.add_textbox(
            Inches(1 + (i - 1) * 3),
            Inches(6),
            Inches(2.5),
            Inches(1),
        )
        tx_tl.name = f"txt_timeline_{i}"
        tf_tl = tx_tl.text_frame
        tf_tl.paragraphs[0].text = f"제목 {i}"
        p = tf_tl.add_paragraph()
        p.text = f"내용 {i}"

    # 테이블 shape (텍스트프레임 없음 테스트용)
    tbl = slide.shapes.add_table(2, 2, Inches(6), Inches(4), Inches(3), Inches(2))
    tbl.name = "tbl_test"

    return slide


class TestReplaceTextPreservingFormat:
    """replace_text_preserving_format 함수 테스트."""

    def test_basic_replace(self, text_slide):
        """기본 텍스트 교체."""
        shape = replace_text_preserving_format(
            text_slide, "txt_body", ["새 첫 번째 문단", "새 두 번째 문단"]
        )
        assert shape.text_frame.paragraphs[0].text == "새 첫 번째 문단"
        assert shape.text_frame.paragraphs[1].text == "새 두 번째 문단"

    def test_fewer_texts(self, text_slide):
        """texts가 기존 paragraph보다 적으면 일부만 교체."""
        shape = replace_text_preserving_format(
            text_slide, "txt_body", ["교체됨"]
        )
        assert shape.text_frame.paragraphs[0].text == "교체됨"

    def test_more_texts_adds_paragraphs(self, text_slide):
        """texts가 기존 paragraph보다 많으면 새 paragraph 추가."""
        shape = replace_text_preserving_format(
            text_slide, "txt_body", ["A", "B", "C", "D"]
        )
        paras = list(shape.text_frame.paragraphs)
        assert len(paras) >= 4

    def test_bold_preserved(self, text_slide):
        """볼드 서식 보존 확인."""
        shape = replace_text_preserving_format(
            text_slide, "txt_bold", ["새 볼드 텍스트"]
        )
        runs = list(shape.text_frame.paragraphs[0].runs)
        assert runs[0].font.bold is True

    def test_font_size_preserved(self, text_slide):
        """폰트 크기 보존 확인."""
        shape = replace_text_preserving_format(
            text_slide, "txt_bold", ["크기 보존"]
        )
        runs = list(shape.text_frame.paragraphs[0].runs)
        assert runs[0].font.size == Pt(14)

    def test_shape_not_found(self, text_slide):
        """존재하지 않는 shape — ShapeNotFoundError."""
        with pytest.raises(ShapeNotFoundError):
            replace_text_preserving_format(text_slide, "txt_nonexistent", ["a"])

    def test_no_text_frame(self, text_slide):
        """text_frame이 없는 shape — TextShapeError."""
        with pytest.raises(TextShapeError):
            replace_text_preserving_format(text_slide, "tbl_test", ["a"])


class TestReplacePlaceholdersInText:
    """replace_placeholders_in_text 함수 테스트."""

    def test_basic_placeholder(self, text_slide):
        """{{key}} 패턴 교체."""
        shape = replace_placeholders_in_text(
            text_slide,
            "txt_placeholder",
            {"company_name": "알파테크", "date": "2026.03"},
        )
        text = shape.text_frame.paragraphs[0].text
        assert "알파테크" in text
        assert "2026.03" in text
        assert "{{" not in text

    def test_no_match(self, text_slide):
        """일치하는 플레이스홀더 없음 — 에러 없이 통과."""
        shape = replace_placeholders_in_text(
            text_slide,
            "txt_placeholder",
            {"nonexistent_key": "value"},
        )
        assert shape is not None

    def test_multiple_occurrences(self):
        """같은 키가 여러 번 등장하는 경우."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tx = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        tx.name = "txt_multi"
        tx.text_frame.paragraphs[0].text = "{{name}} 대표이사 {{name}}"

        shape = replace_placeholders_in_text(
            slide, "txt_multi", {"name": "김대표"}
        )
        assert shape.text_frame.paragraphs[0].text == "김대표 대표이사 김대표"


class TestReplaceTimelineTexts:
    """replace_timeline_texts 함수 테스트."""

    def test_basic_timeline(self, text_slide):
        """타임라인 텍스트 순차 교체."""
        items = [
            {"title": "2022", "content": "시리즈A 투자"},
            {"title": "2023", "content": "매출 100억 달성"},
            {"title": "2024", "content": "해외 진출"},
        ]
        shapes = replace_timeline_texts(text_slide, "txt_timeline_", items)
        assert len(shapes) == 3

    def test_fewer_items(self, text_slide):
        """items가 shape 수보다 적으면 일부만 교체."""
        items = [{"title": "2022", "content": "테스트"}]
        shapes = replace_timeline_texts(text_slide, "txt_timeline_", items)
        assert len(shapes) == 1

    def test_empty_items(self, text_slide):
        """빈 items — 빈 리스트 반환."""
        shapes = replace_timeline_texts(text_slide, "txt_timeline_", [])
        assert len(shapes) == 0

    def test_no_matching_prefix(self, text_slide):
        """매칭되는 prefix 없음 — 빈 리스트."""
        shapes = replace_timeline_texts(text_slide, "txt_nonexistent_", [{"title": "t"}])
        assert len(shapes) == 0
