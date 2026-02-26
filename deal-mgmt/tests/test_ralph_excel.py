"""Ralph Loop Excel 통합 테스트 — Gate + Generator + LLM Judge."""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.models.enums import FinancialModelType
from app.ralph.gates.base import GateVerdict
from app.ralph.gates.excel_gate import ExcelProgrammaticGate
from app.ralph.gates.llm_judge_gate import (
    EXCEL_WEIGHTS,
    LLMJudgeGate,
)
from app.ralph.generators.excel_generator import SHEET_TITLES, RalphExcelGenerator
from app.ralph.prd_manager import load_prd

# ── Fixtures ──────────────────────────────────────────────────────────────────


@pytest.fixture
def tmp_dir():
    with tempfile.TemporaryDirectory(prefix="test_ralph_excel_") as d:
        yield Path(d)


def _create_basic_workbook(tmp_dir: Path, sheets: list[str] | None = None) -> str:
    """기본 테스트용 워크북을 생성한다."""
    wb = Workbook()
    wb.remove(wb.active)

    if sheets is None:
        sheets = ["Legend", "Input", "IS", "BS", "CF", "WACC", "DCF", "Sensitivity", "Summary"]

    for name in sheets:
        wb.create_sheet(title=name)

    path = str(tmp_dir / "test_basic.xlsx")
    wb.save(path)
    wb.close()
    return path


def _create_formula_workbook(tmp_dir: Path, formula_ratio: float = 0.4) -> str:
    """수식 비율 테스트용 워크북."""
    wb = Workbook()
    wb.remove(wb.active)

    for name in ["Legend", "IS", "Summary"]:
        ws = wb.create_sheet(title=name)
        if name == "IS":
            # IS에 데이터 + 수식 삽입
            total_cells = 50
            formula_count = int(total_cells * formula_ratio)
            for i in range(1, total_cells + 1):
                if i <= formula_count:
                    ws.cell(row=i, column=2, value=f"=SUM(C{i}:G{i})")
                else:
                    ws.cell(row=i, column=2, value=1000 * i)

    path = str(tmp_dir / "test_formula.xlsx")
    wb.save(path)
    wb.close()
    return path


def _create_bs_balanced_workbook(tmp_dir: Path, balanced: bool = True) -> str:
    """BS 균형 검증 테스트용 워크북."""
    wb = Workbook()
    wb.remove(wb.active)

    for name in ["Legend", "BS", "Summary"]:
        ws = wb.create_sheet(title=name)
        if name == "BS":
            # 행 레이블 (열 B)
            ws.cell(row=1, column=2, value="Balance Sheet")
            ws.cell(row=2, column=2, value="Total Assets")
            ws.cell(row=3, column=2, value="Total Liabilities & Equity")
            ws.cell(row=4, column=2, value="Balance Check")

            for col in range(3, 8):  # 5개 연도
                ws.cell(row=2, column=col, value=100000)
                ws.cell(row=3, column=col, value=100000 if balanced else 90000)
                # balance check = 차이 (0이면 균형)
                diff = 0 if balanced else 10000
                ws.cell(row=4, column=col, value=diff)

    path = str(tmp_dir / "test_bs.xlsx")
    wb.save(path)
    wb.close()
    return path


def _create_cross_ref_workbook(tmp_dir: Path, has_refs: bool = True) -> str:
    """교차 참조 테스트용 워크북."""
    wb = Workbook()
    wb.remove(wb.active)

    for name in ["Legend", "IS", "BS", "CF", "Summary"]:
        ws = wb.create_sheet(title=name)
        if name == "CF" and has_refs:
            ws.cell(row=1, column=2, value="Cash Flow Statement")
            ws.cell(row=2, column=2, value="=IS!C30")  # IS 참조
            ws.cell(row=3, column=2, value="=BS!C5")  # BS 참조
        elif name == "CF":
            ws.cell(row=1, column=2, value="Cash Flow Statement")
            ws.cell(row=2, column=2, value=50000)  # 참조 없음

    path = str(tmp_dir / "test_cross_ref.xlsx")
    wb.save(path)
    wb.close()
    return path


def _create_placeholder_workbook(tmp_dir: Path) -> str:
    """플레이스홀더 탐지 테스트용 워크북."""
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet(title="Input")
    ws.cell(row=1, column=1, value="Revenue Growth")
    ws.cell(row=2, column=1, value="[TBD]")
    ws.cell(row=3, column=1, value="[INSERT HERE]")
    ws.cell(row=4, column=1, value="Normal value 10%")
    ws.cell(row=5, column=1, value="[TODO]")

    path = str(tmp_dir / "test_placeholder.xlsx")
    wb.save(path)
    wb.close()
    return path


# ── ExcelProgrammaticGate 테스트 ─────────────────────────────────────────────


class TestExcelProgrammaticGate:
    """ExcelProgrammaticGate 단위 테스트."""

    def setup_method(self):
        self.gate = ExcelProgrammaticGate()

    def test_gate_name(self):
        assert self.gate.name == "excel_programmatic"

    @pytest.mark.asyncio
    async def test_basic_structure_dcf(self, tmp_dir):
        """DCF 필수 시트가 있으면 구조 점수 5점."""
        path = _create_basic_workbook(tmp_dir)
        result = await self.gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "DCF"},
        )
        assert result.gate_name == "excel_programmatic"
        assert result.weighted_score > 0
        # 구조 차원 확인
        struct_dim = next((d for d in result.dimensions if d.name == "structure"), None)
        assert struct_dim is not None
        assert struct_dim.score == 5.0

    @pytest.mark.asyncio
    async def test_missing_sheets(self, tmp_dir):
        """필수 시트 누락 시 구조 점수 하락."""
        path = _create_basic_workbook(tmp_dir, sheets=["Legend", "IS"])
        result = await self.gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "DCF"},
        )
        struct_dim = next((d for d in result.dimensions if d.name == "structure"), None)
        assert struct_dim is not None
        assert struct_dim.score < 5.0
        # 이슈에 누락 시트 언급
        missing_issues = [i for i in result.issues if "필수 워크시트 누락" in i]
        assert len(missing_issues) > 0

    @pytest.mark.asyncio
    async def test_formula_ratio_high(self, tmp_dir):
        """수식 비율 40% → 5점."""
        path = _create_formula_workbook(tmp_dir, formula_ratio=0.4)
        result = await self.gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "DCF"},
        )
        formula_dim = next((d for d in result.dimensions if d.name == "formula_ratio"), None)
        assert formula_dim is not None
        assert formula_dim.score == 5.0

    @pytest.mark.asyncio
    async def test_formula_ratio_low(self, tmp_dir):
        """수식 비율 5% → 2점."""
        path = _create_formula_workbook(tmp_dir, formula_ratio=0.05)
        result = await self.gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "DCF"},
        )
        formula_dim = next((d for d in result.dimensions if d.name == "formula_ratio"), None)
        assert formula_dim is not None
        assert formula_dim.score == 2.0

    @pytest.mark.asyncio
    async def test_bs_balanced(self, tmp_dir):
        """BS 균형 시 balance_check 5점."""
        path = _create_bs_balanced_workbook(tmp_dir, balanced=True)
        result = await self.gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "DCF"},
        )
        balance_dim = next((d for d in result.dimensions if d.name == "balance_check"), None)
        assert balance_dim is not None
        assert balance_dim.score == 5.0
        assert len(result.critical_flags) == 0

    @pytest.mark.asyncio
    async def test_bs_unbalanced_critical(self, tmp_dir):
        """BS 불균형 시 CRITICAL flag + 1점."""
        path = _create_bs_balanced_workbook(tmp_dir, balanced=False)
        result = await self.gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "DCF"},
        )
        balance_dim = next((d for d in result.dimensions if d.name == "balance_check"), None)
        assert balance_dim is not None
        assert balance_dim.score == 1.0
        assert len(result.critical_flags) > 0
        assert result.verdict == GateVerdict.FAIL

    @pytest.mark.asyncio
    async def test_cross_refs_present(self, tmp_dir):
        """교차 참조가 있으면 높은 점수."""
        path = _create_cross_ref_workbook(tmp_dir, has_refs=True)
        result = await self.gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "DCF"},
        )
        cross_dim = next((d for d in result.dimensions if d.name == "cross_validation"), None)
        assert cross_dim is not None
        assert cross_dim.score == 5.0

    @pytest.mark.asyncio
    async def test_cross_refs_missing(self, tmp_dir):
        """교차 참조가 없으면 낮은 점수."""
        path = _create_cross_ref_workbook(tmp_dir, has_refs=False)
        result = await self.gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "DCF"},
        )
        cross_dim = next((d for d in result.dimensions if d.name == "cross_validation"), None)
        assert cross_dim is not None
        assert cross_dim.score < 5.0

    @pytest.mark.asyncio
    async def test_placeholder_detection(self, tmp_dir):
        """플레이스홀더 탐지 시 completeness 점수 하락."""
        path = _create_placeholder_workbook(tmp_dir)
        result = await self.gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "DCF"},
        )
        complete_dim = next((d for d in result.dimensions if d.name == "completeness"), None)
        assert complete_dim is not None
        assert complete_dim.score < 5.0
        placeholder_issues = [i for i in result.issues if "플레이스홀더" in i]
        assert len(placeholder_issues) > 0

    @pytest.mark.asyncio
    async def test_invalid_file(self, tmp_dir):
        """잘못된 파일 → 로드 실패 이슈."""
        bad_path = str(tmp_dir / "not_excel.txt")
        Path(bad_path).write_text("not an excel file")
        result = await self.gate.evaluate(
            artifact_path=bad_path,
            prd_section={"model_type": "DCF"},
        )
        assert "Excel 로드 실패" in result.issues[0]
        assert len(result.critical_flags) > 0

    @pytest.mark.asyncio
    async def test_comps_no_bs_full_score(self, tmp_dir):
        """COMPS 모델 (BS 없음) → balance_check 5점 (만점)."""
        path = _create_basic_workbook(tmp_dir, sheets=["Legend", "GPCM", "Summary"])
        result = await self.gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "COMPS"},
        )
        balance_dim = next((d for d in result.dimensions if d.name == "balance_check"), None)
        assert balance_dim is not None
        assert balance_dim.score == 5.0

    @pytest.mark.asyncio
    async def test_six_dimensions(self, tmp_dir):
        """항상 6개 차원을 반환한다."""
        path = _create_basic_workbook(tmp_dir)
        result = await self.gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "DCF"},
        )
        assert len(result.dimensions) == 6
        expected_names = {
            "structure",
            "formula_ratio",
            "balance_check",
            "cross_validation",
            "completeness",
            "formatting",
        }
        actual_names = {d.name for d in result.dimensions}
        assert actual_names == expected_names


# ── RalphExcelGenerator 테스트 ────────────────────────────────────────────────


class TestRalphExcelGenerator:
    """RalphExcelGenerator 단위 테스트."""

    @pytest.mark.asyncio
    async def test_generate_outline_dcf(self):
        """DCF 모델 아웃라인 생성."""
        gen = RalphExcelGenerator(
            model_type=FinancialModelType.DCF,
            title="Test DCF Model",
        )
        outline = await gen.generate_outline()
        assert len(outline) > 0
        # DCF 필수 시트 확인
        ids = {s["id"] for s in outline}
        assert "IS" in ids
        assert "BS" in ids
        assert "CF" in ids
        assert "DCF" in ids
        assert "Legend" in ids
        assert "Summary" in ids

    @pytest.mark.asyncio
    async def test_generate_outline_comps(self):
        """COMPS 모델 아웃라인."""
        gen = RalphExcelGenerator(
            model_type=FinancialModelType.COMPS,
            title="Test COMPS Model",
        )
        outline = await gen.generate_outline()
        ids = {s["id"] for s in outline}
        assert "GPCM" in ids
        assert "Legend" in ids

    @pytest.mark.asyncio
    async def test_generate_outline_has_titles(self):
        """아웃라인 항목에 title이 포함된다."""
        gen = RalphExcelGenerator(
            model_type=FinancialModelType.DCF,
            title="Test",
        )
        outline = await gen.generate_outline()
        for section in outline:
            assert "id" in section
            assert "title" in section
            assert section["title"] == SHEET_TITLES.get(section["id"], section["id"])

    @pytest.mark.asyncio
    async def test_generate_section_returns_path(self):
        """generate_section은 .xlsx 파일 경로를 반환한다."""
        gen = RalphExcelGenerator(
            model_type=FinancialModelType.DCF,
            title="Test DCF",
        )
        path = await gen.generate_section(section_id="IS")
        assert path.endswith(".xlsx")
        assert Path(path).exists()

    @pytest.mark.asyncio
    async def test_generate_section_with_checklist_values(self):
        """체크리스트 값이 빌더에 전달된다."""
        gen = RalphExcelGenerator(
            model_type=FinancialModelType.DCF,
            title="Test",
            checklist_values={"매출 성장률 가정 (향후 5년)": "10"},
        )
        path = await gen.generate_section(section_id="Input")
        assert Path(path).exists()

    @pytest.mark.asyncio
    async def test_assemble_document(self):
        """assemble_document는 최종 .xlsx를 생성한다."""
        gen = RalphExcelGenerator(
            model_type=FinancialModelType.DCF,
            title="Test Assemble",
        )
        result = await gen.assemble_document(
            section_artifacts={"IS": "/tmp/fake.xlsx"},
        )
        assert result.endswith(".xlsx")
        assert Path(result).exists()

    @pytest.mark.asyncio
    async def test_assemble_document_custom_output(self, tmp_dir):
        """커스텀 output_path에 저장."""
        gen = RalphExcelGenerator(
            model_type=FinancialModelType.DCF,
            title="Test Custom",
        )
        output = str(tmp_dir / "custom_output.xlsx")
        result = await gen.assemble_document(
            section_artifacts={},
            output_path=output,
        )
        assert result == output
        assert Path(output).exists()


# ── LLMJudgeGate Excel 모드 테스트 ───────────────────────────────────────────


class TestLLMJudgeGateExcel:
    """LLM Judge Gate Excel 모드 테스트."""

    def test_excel_weights_exist(self):
        """EXCEL_WEIGHTS가 정의되어 있다."""
        assert "assumptions_quality" in EXCEL_WEIGHTS
        assert "methodology" in EXCEL_WEIGHTS
        assert "consistency" in EXCEL_WEIGHTS
        assert "professionalism" in EXCEL_WEIGHTS
        assert "scenario_depth" in EXCEL_WEIGHTS
        assert "narrative" in EXCEL_WEIGHTS

    def test_excel_weights_sum_to_one(self):
        """EXCEL_WEIGHTS 가중치 합이 1.0."""
        total = sum(w for w, _ in EXCEL_WEIGHTS.values())
        assert abs(total - 1.0) < 0.01

    def test_init_excel_mode(self):
        """doc_type='excel'로 초기화하면 EXCEL_WEIGHTS 사용."""
        gate = LLMJudgeGate(doc_type="excel")
        assert gate._weights == EXCEL_WEIGHTS

    def test_init_ldd_mode(self):
        """doc_type='ldd'는 LDD_WEIGHTS 사용."""
        from app.ralph.gates.llm_judge_gate import LDD_WEIGHTS

        gate = LLMJudgeGate(doc_type="ldd")
        assert gate._weights == LDD_WEIGHTS

    def test_init_pptx_mode(self):
        """doc_type='pptx'는 PPTX_WEIGHTS 사용."""
        from app.ralph.gates.llm_judge_gate import PPTX_WEIGHTS

        gate = LLMJudgeGate(doc_type="pptx")
        assert gate._weights == PPTX_WEIGHTS

    @pytest.mark.asyncio
    async def test_fallback_evaluate_excel(self, tmp_dir):
        """LLM 없이 Excel fallback 평가."""
        # 간단한 xlsx 생성
        wb = Workbook()
        ws = wb.active
        ws.title = "IS"
        for i in range(1, 20):
            ws.cell(row=i, column=1, value=f"Line item {i}")
            ws.cell(row=i, column=2, value=1000 * i)
        path = str(tmp_dir / "test_judge.xlsx")
        wb.save(path)
        wb.close()

        gate = LLMJudgeGate(doc_type="excel")
        result = await gate.evaluate(
            artifact_path=path,
            prd_section={},
        )
        assert result.gate_name == "llm_judge"
        assert result.weighted_score > 0
        # fallback 사용 메시지
        assert any("fallback" in i.lower() or "규칙 기반" in i for i in result.issues)

    @pytest.mark.asyncio
    async def test_extract_xlsx_text(self, tmp_dir):
        """_extract_xlsx_text로 Excel 텍스트 추출."""
        wb = Workbook()
        ws = wb.active
        ws.title = "TestSheet"
        ws.cell(row=1, column=1, value="Header A")
        ws.cell(row=1, column=2, value="Header B")
        ws.cell(row=2, column=1, value=100)
        ws.cell(row=2, column=2, value="Data")
        path = str(tmp_dir / "test_extract.xlsx")
        wb.save(path)
        wb.close()

        gate = LLMJudgeGate(doc_type="excel")
        text = gate._extract_xlsx_text(path)
        assert "TestSheet" in text
        assert "Header A" in text
        assert "100" in text
        assert "Data" in text

    @pytest.mark.asyncio
    async def test_extract_text_dispatches_xlsx(self, tmp_dir):
        """_extract_text가 .xlsx 파일에 대해 _extract_xlsx_text 호출."""
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="Test")
        path = str(tmp_dir / "dispatch.xlsx")
        wb.save(path)
        wb.close()

        gate = LLMJudgeGate(doc_type="excel")
        text = gate._extract_text(path)
        assert "Test" in text

    @pytest.mark.asyncio
    async def test_non_xlsx_returns_empty(self):
        """지원하지 않는 확장자 → 빈 문자열."""
        gate = LLMJudgeGate(doc_type="excel")
        text = gate._extract_text("/fake/path/file.csv")
        assert text == ""


# ── PRD 로드 테스트 ──────────────────────────────────────────────────────────


class TestFinancialModelPRD:
    """financial_model.json PRD 테스트."""

    def test_load_prd(self):
        """PRD 파일이 로드된다."""
        prd = load_prd("financial_model")
        assert prd["name"] == "Financial Model PRD"
        assert prd["document_type"] == "financial_model"
        assert prd["pass_threshold"] == 4.0

    def test_prd_sections(self):
        """PRD에 필수 섹션이 정의되어 있다."""
        prd = load_prd("financial_model")
        sections = prd["sections"]
        assert "Legend" in sections
        assert "Input" in sections
        assert "IS" in sections
        assert "BS" in sections
        assert "CF" in sections
        assert "Summary" in sections

    def test_prd_bs_critical(self):
        """BS 섹션에 CRITICAL 기준이 있다."""
        prd = load_prd("financial_model")
        bs = prd["sections"]["BS"]
        assert "CRITICAL" in bs["criteria"]["balance"]

    def test_prd_global_criteria(self):
        """global_criteria가 정의되어 있다."""
        prd = load_prd("financial_model")
        gc = prd["global_criteria"]
        assert "formula_ratio" in gc
        assert "balance_sheet" in gc
        assert "cross_validation" in gc

    def test_prd_required_sections(self):
        """필수 섹션이 올바르게 표시되어 있다."""
        prd = load_prd("financial_model")
        required_sections = {k for k, v in prd["sections"].items() if v.get("required")}
        assert "Legend" in required_sections
        assert "Input" in required_sections
        assert "IS" in required_sections
        assert "BS" in required_sections
        assert "CF" in required_sections
        assert "Summary" in required_sections

    def test_prd_optional_sections(self):
        """선택 섹션이 올바르게 표시되어 있다."""
        prd = load_prd("financial_model")
        optional_sections = {k for k, v in prd["sections"].items() if not v.get("required")}
        assert "WACC" in optional_sections
        assert "DCF" in optional_sections
        assert "GPCM" in optional_sections


# ── 통합 테스트: Gate + Generator ─────────────────────────────────────────────


class TestExcelGateWithGenerator:
    """ExcelProgrammaticGate + RalphExcelGenerator 통합 테스트."""

    @pytest.mark.asyncio
    async def test_generator_output_passes_gate(self):
        """Generator가 생성한 Excel이 Gate 평가를 통과한다."""
        gen = RalphExcelGenerator(
            model_type=FinancialModelType.DCF,
            title="Integration Test",
            checklist_values={
                "매출 성장률 가정 (향후 5년)": "10",
                "매출원가율 실적/가정": "60",
                "판관비율 실적/가정": "15",
            },
        )

        # 섹션 생성
        path = await gen.generate_section(section_id="IS")
        assert Path(path).exists()

        # Gate 평가
        gate = ExcelProgrammaticGate()
        result = await gate.evaluate(
            artifact_path=path,
            prd_section={"model_type": "DCF"},
        )
        assert result.weighted_score > 0
        assert len(result.dimensions) == 6

    @pytest.mark.asyncio
    async def test_full_loop_simulation(self):
        """Generator outline → section → assemble → Gate 평가 파이프라인."""
        gen = RalphExcelGenerator(
            model_type=FinancialModelType.COMPS,
            title="Pipeline Test",
        )

        # Phase 1: outline
        outline = await gen.generate_outline()
        assert len(outline) > 0

        # Phase 2: 각 섹션 생성
        artifacts = {}
        for section in outline:
            path = await gen.generate_section(section_id=section["id"])
            artifacts[section["id"]] = path

        # Phase 3: assemble
        final = await gen.assemble_document(section_artifacts=artifacts)
        assert Path(final).exists()

        # Gate 평가
        gate = ExcelProgrammaticGate()
        result = await gate.evaluate(
            artifact_path=final,
            prd_section={"model_type": "COMPS"},
        )
        assert result.gate_name == "excel_programmatic"
        # COMPS에 BS 없으므로 balance_check는 5점
        balance = next((d for d in result.dimensions if d.name == "balance_check"), None)
        assert balance is not None
        assert balance.score == 5.0
