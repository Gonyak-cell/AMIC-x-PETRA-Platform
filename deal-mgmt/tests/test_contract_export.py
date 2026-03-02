"""계약서 DOCX 내보내기 서비스 테스트."""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.services.contract_export_service import (
    _build_docx,
    _extract_tag_contents,
    _strip_tags,
)

# ── _strip_tags ──────────────────────────────────────────────────────────────


class TestStripTags:
    def test_simple_tag(self) -> None:
        assert _strip_tags("<p>hello</p>") == "hello"

    def test_nested_tags(self) -> None:
        assert _strip_tags("<div><p>inner</p></div>") == "inner"

    def test_no_tags(self) -> None:
        assert _strip_tags("plain text") == "plain text"

    def test_multiple_spaces_collapsed(self) -> None:
        assert _strip_tags("<p>a</p>  <p>b</p>") == "a b"

    def test_empty_input(self) -> None:
        assert _strip_tags("") == ""


# ── _extract_tag_contents ────────────────────────────────────────────────────


class TestExtractTagContents:
    def test_extract_p_tags(self) -> None:
        html = "<p>first</p><p>second</p>"
        assert _extract_tag_contents(html, "p") == ["first", "second"]

    def test_extract_h2(self) -> None:
        html = '<h2 class="title">제1조 (정의)</h2>'
        assert _extract_tag_contents(html, "h2") == ["제1조 (정의)"]

    def test_no_matches(self) -> None:
        html = "<p>only p tags</p>"
        assert _extract_tag_contents(html, "h1") == []

    def test_nested_content_stripped(self) -> None:
        html = "<li><strong>bold</strong> text</li>"
        assert _extract_tag_contents(html, "li") == ["bold text"]


# ── _build_docx ──────────────────────────────────────────────────────────────


class TestBuildDocx:
    def _make_html(self, title: str = "주식매매계약서", clauses: int = 3) -> str:
        sections = []
        for i in range(1, clauses + 1):
            sections.append(
                f'<section class="contract-clause"><h2>제{i}조 (조항{i})</h2><p>조항 {i}의 내용입니다.</p></section>'
            )
        body = "\n".join(sections)
        return f'<div class="contract-document"><h1>{title}</h1>{body}</div>'

    def test_returns_valid_docx_bytes(self) -> None:
        html = self._make_html()
        result = _build_docx(html, "SPA")
        assert isinstance(result, bytes)
        assert len(result) > 0

        # 유효한 DOCX 파일인지 확인
        doc = Document(io.BytesIO(result))
        assert len(doc.paragraphs) > 0

    def test_title_extracted_from_h1(self) -> None:
        html = self._make_html(title="자산양수도계약서")
        result = _build_docx(html, "BTA")
        doc = Document(io.BytesIO(result))

        # 첫 번째 paragraph이 제목
        title_text = doc.paragraphs[0].text
        assert "자산양수도계약서" in title_text

    def test_fallback_title_when_no_h1(self) -> None:
        html = "<div><section><h2>조항1</h2><p>내용</p></section></div>"
        result = _build_docx(html, "MOU")
        doc = Document(io.BytesIO(result))

        title_text = doc.paragraphs[0].text
        assert "MOU" in title_text

    def test_sections_processed(self) -> None:
        html = self._make_html(clauses=2)
        result = _build_docx(html, "SPA")
        doc = Document(io.BytesIO(result))

        # h2 제목이 포함되어야 함
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "제1조" in all_text
        assert "제2조" in all_text

    def test_no_sections_fallback(self) -> None:
        html = "<div><h1>제목</h1><h2>소제목</h2><p>내용</p></div>"
        result = _build_docx(html, "SHA")
        doc = Document(io.BytesIO(result))
        assert len(doc.paragraphs) > 0

    def test_list_items(self) -> None:
        html = "<div><h1>테스트</h1><section><ol><li>항목1</li><li>항목2</li></ol></section></div>"
        result = _build_docx(html, "SPA")
        doc = Document(io.BytesIO(result))
        all_text = " ".join(p.text for p in doc.paragraphs)
        assert "항목1" in all_text
        assert "항목2" in all_text


# ── html_to_docx (async) ────────────────────────────────────────────────────


class TestHtmlToDocxAsync:
    @pytest.mark.asyncio
    async def test_async_wrapper(self) -> None:
        from app.services.contract_export_service import html_to_docx

        html = "<div><h1>비동기 테스트</h1><section><p>내용</p></section></div>"
        result = await html_to_docx(html, "테스트")
        assert isinstance(result, bytes)
        assert len(result) > 0
