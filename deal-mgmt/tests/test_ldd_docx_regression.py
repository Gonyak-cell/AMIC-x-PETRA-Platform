from __future__ import annotations

import re
import uuid
from pathlib import Path

import pytest
from docx import Document
from docx.table import Table, _Cell
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.enums import LDDIssueLevel, LDDItemStatus, LDDReportStatus, LDDReportType, LDDSectionType, TransactionSide
from app.models.transaction import Transaction
from app.schemas.ldd_report import LDDItem, LDDReportCreate, LDDSection
from app.services import ldd_report_service

pytestmark = pytest.mark.anyio

SCOPE_NOTICE = (
    "본 문안은 현재까지 회사가 제출한 자료, 답변 및 필요한 범위의 공개 등기·공시자료를 기준으로 작성되었으며, "
    "별도 현장실사나 제3자 조회 없이 검토한 결과입니다."
)
GOVERNANCE_ANALYSIS_PREAMBLE = (
    "본 항목에서는 설립, 기관결정, 권한 배분, 이해상충 및 지배관계의 적법성과 거래 선행조치 필요 여부를 함께 검토합니다."
)
GOVERNANCE_DEAL_IMPACT_PREAMBLE = (
    "지배구조 및 권한 체계와 관련한 확인 결과는 SPA 진술보장, 선행조건, 사후 시정조치 및 PMI 실행계획에 직접 반영될 수 있습니다."
)


async def _make_txn(db: AsyncSession, *, name: str = "LDD DOCX Regression") -> Transaction:
    txn = Transaction(
        name=name,
        code_name=f"T-{uuid.uuid4().hex[:8]}",
        side=TransactionSide.BUY,
        target_company_name="Regression Holdings",
        client_name="Regression Client",
        lead_advisor_email="test@example.com",
    )
    db.add(txn)
    await db.flush()
    return txn


def _normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _table_texts(table: Table) -> list[str]:
    texts: list[str] = []
    for row in table.rows:
        for cell in row.cells:
            texts.extend(_cell_texts(cell))
    return texts


def _cell_texts(cell: _Cell) -> list[str]:
    texts = [para.text.strip() for para in cell.paragraphs if para.text.strip()]
    for nested in cell.tables:
        texts.extend(_table_texts(nested))
    return texts


def _extract_docx_text(path: str | Path) -> str:
    doc = Document(str(path))
    texts = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
    for table in doc.tables:
        texts.extend(_table_texts(table))
    return _normalize_text("\n".join(texts))


async def _create_and_render_report(
    db: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    *,
    report_type: LDDReportType,
    sections: list[LDDSection],
    title: str,
) -> tuple[Transaction, object]:
    monkeypatch.setattr(ldd_report_service, "OUTPUT_DIR", tmp_path / "generated-ldd")
    txn = await _make_txn(db, name=title)
    report = await ldd_report_service.create_ldd_report(
        db,
        transaction_id=txn.id,
        body=LDDReportCreate(
            title=title,
            report_type=report_type,
            target_company="Regression Holdings",
            dd_period="2026-03-01 ~ 2026-03-31",
            law_firm="Regression Law LLC",
            prepared_by="Regression Counsel",
            sections=sections,
        ),
        created_by_email="test@example.com",
    )
    rendered = await ldd_report_service.generate_ldd_report(db, report)
    return txn, rendered


async def test_full_ldd_docx_regression_compacts_section_preamble(
    async_session: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
):
    sections = [
        LDDSection(
            section_type=LDDSectionType.GOVERNANCE,
            title="1. Governance",
            items=[
                LDDItem(
                    item_id="CORP-02",
                    name="Board Composition Review",
                    status=LDDItemStatus.ISSUE,
                    issue_level=LDDIssueLevel.HIGH,
                    description="Board minutes for the last two years were not provided.",
                    deal_impact="Delay in confirming post-closing governance structure.",
                    recommendation="Collect the last two years of board minutes before signing.",
                    rfi_required=True,
                    rfi_number="CORP-002",
                    evidence_refs=["governance_board_minutes.pdf"],
                ),
                LDDItem(
                    item_id="CORP-04",
                    name="Executive Appointment and Term",
                    status=LDDItemStatus.ISSUE,
                    issue_level=LDDIssueLevel.MEDIUM,
                    description="Consent rights for executive appointments remain unresolved.",
                    deal_impact="May delay the target post-closing governance clean-up plan.",
                    recommendation="Obtain investor consent and document the appointment process.",
                    rfi_required=True,
                    rfi_number="CORP-004",
                    evidence_refs=["shareholders_agreement.pdf"],
                ),
            ],
        )
    ]

    _txn, report = await _create_and_render_report(
        async_session,
        monkeypatch,
        tmp_path,
        report_type=LDDReportType.FULL,
        sections=sections,
        title="FULL DOCX regression",
    )

    assert report.status == LDDReportStatus.READY
    assert report.file_path
    assert Path(report.file_path).exists()

    text = _extract_docx_text(report.file_path)

    assert "Board Composition Review" in text
    assert "Executive Appointment and Term" in text
    assert "Collect the last two years of board minutes before signing." in text
    assert "Obtain investor consent and document the appointment process." in text
    assert text.count(SCOPE_NOTICE) == 1
    assert text.count(GOVERNANCE_ANALYSIS_PREAMBLE) == 1
    assert text.count(GOVERNANCE_DEAL_IMPACT_PREAMBLE) == 1
    assert "{{" not in text
    assert "{%p" not in text


async def test_redflag_ldd_docx_regression_excludes_non_issue_items(
    async_session: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
):
    sections = [
        LDDSection(
            section_type=LDDSectionType.CONTRACTS,
            title="3. Contracts",
            items=[
                LDDItem(
                    item_id="CONTRACT-04",
                    name="Change of Control Approval",
                    status=LDDItemStatus.ISSUE,
                    issue_level=LDDIssueLevel.CRITICAL,
                    description="Prior investor approval is required before a control transfer.",
                    deal_impact="Closing cannot proceed until the waiver package is signed.",
                    recommendation="Document the waiver as an express closing condition.",
                    rfi_required=True,
                    rfi_number="CONTRACT-004",
                    evidence_refs=["sha_approval_clause.pdf"],
                ),
                LDDItem(
                    item_id="CONTRACT-07",
                    name="Routine Supplier Agreement",
                    status=LDDItemStatus.OK,
                    description="Routine supplier agreement does not present a material issue.",
                    deal_impact="No material impact on the current deal structure.",
                    recommendation="No additional action is required.",
                ),
            ],
        )
    ]

    _txn, report = await _create_and_render_report(
        async_session,
        monkeypatch,
        tmp_path,
        report_type=LDDReportType.REDFLAG,
        sections=sections,
        title="REDFLAG DOCX regression",
    )

    assert report.status == LDDReportStatus.READY
    assert report.file_path

    text = _extract_docx_text(report.file_path)

    assert "Change of Control Approval" in text
    assert "Document the waiver as an express closing condition." in text
    assert "Routine Supplier Agreement" not in text
    assert "No additional action is required." not in text


async def test_law_firm_ldd_docx_regression_renders_custom_content(
    async_session: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
):
    from app.ralph.generators.ldd.law_firm_template import LawFirmTemplateGenerator

    def _unexpected_blank_generation(*args, **kwargs):
        raise AssertionError("direct law_firm_template.docx should be used before fallback blank generation")

    monkeypatch.setattr(LawFirmTemplateGenerator, "generate_blank_template", _unexpected_blank_generation)

    sections = [
        LDDSection(
            section_type=LDDSectionType.GOVERNANCE,
            title="1. Governance",
            items=[
                LDDItem(
                    item_id="CORP-02",
                    name="Board Composition Review",
                    status=LDDItemStatus.ISSUE,
                    issue_level=LDDIssueLevel.HIGH,
                    description="Board minutes for the last two years were not provided.",
                    deal_impact="Delay in confirming post-closing governance structure.",
                    recommendation="Collect the last two years of board minutes before signing.",
                    rfi_required=True,
                    rfi_number="CORP-002",
                ),
            ],
        ),
        LDDSection(
            section_type=LDDSectionType.CONTRACTS,
            title="3. Contracts",
            items=[
                LDDItem(
                    item_id="CONTRACT-04",
                    name="Change of Control Approval",
                    status=LDDItemStatus.ISSUE,
                    issue_level=LDDIssueLevel.CRITICAL,
                    description="Prior investor approval is required before a control transfer.",
                    deal_impact="Closing cannot proceed until the waiver package is signed.",
                    recommendation="Document the waiver as an express closing condition.",
                    rfi_required=True,
                    rfi_number="CONTRACT-004",
                ),
            ],
        ),
    ]

    _txn, report = await _create_and_render_report(
        async_session,
        monkeypatch,
        tmp_path,
        report_type=LDDReportType.LAW_FIRM,
        sections=sections,
        title="LAW_FIRM DOCX regression",
    )

    assert report.status == LDDReportStatus.READY
    assert report.file_path
    assert Path(report.file_path).exists()

    text = _extract_docx_text(report.file_path)

    assert "Regression Holdings" in text
    assert "Regression Law LLC" in text
    assert "목 차" in text
    assert "서 문" in text
    assert "용 례" in text
    assert "Board Composition Review" in text
    assert "Change of Control Approval" in text
    assert "Board minutes for the last two years were not provided." in text
    assert "Prior investor approval is required before a control transfer." in text
    assert "Collect the last two years of board minutes before signing." in text
    assert "Document the waiver as an express closing condition." in text
    assert "[이곳에 텍스트 입력]" not in text


async def test_law_firm_ldd_docx_regression_applies_project_green_style(
    async_session: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
):
    sections = [
        LDDSection(
            section_type=LDDSectionType.PERMITS,
            title="2. Permits",
            items=[
                LDDItem(
                    item_id="PERMIT-01",
                    name="인허가 강한 근거",
                    status=LDDItemStatus.ISSUE,
                    issue_level=LDDIssueLevel.HIGH,
                    description="관할관청의 사전승인이 필요한 것으로 보임.",
                    deal_impact="거래종결 전 승계절차를 완료할 필요가 있습니다.",
                    recommendation="거래종결의 선행조건으로 반영해야 합니다.",
                    rfi_required=False,
                    evidence_refs=["shareholders_agreement.pdf", "permit_license.pdf"],
                    confidence=0.91,
                ),
                LDDItem(
                    item_id="PERMIT-02",
                    name="인허가 중간 근거",
                    status=LDDItemStatus.ISSUE,
                    issue_level=LDDIssueLevel.MEDIUM,
                    description="계약상 사전 동의 절차가 필요한 것으로 판단됨.",
                    deal_impact="거래종결 전 관련 동의서 확보 여부를 점검할 필요가 있습니다.",
                    recommendation="동의서 확보 계획을 거래 일정표에 반영해야 합니다.",
                    rfi_required=False,
                    evidence_refs=["management_presentation.pdf"],
                    confidence=0.62,
                ),
                LDDItem(
                    item_id="PERMIT-03",
                    name="인허가 약한 근거",
                    status=LDDItemStatus.ISSUE,
                    issue_level=LDDIssueLevel.HIGH,
                    description="신고 수리 여부를 추가로 확인할 필요가 있는 것으로 판단됨.",
                    deal_impact="자료 공백이 해소되기 전까지는 거래 구조에 미치는 영향을 보수적으로 검토할 필요가 있습니다.",
                    recommendation="추가 자료 확보 전까지는 종결 전제사실을 보수적으로 유지해야 합니다.",
                    rfi_required=True,
                    rfi_number="PERMIT-003",
                    confidence=0.2,
                ),
            ],
        )
    ]

    _txn, report = await _create_and_render_report(
        async_session,
        monkeypatch,
        tmp_path,
        report_type=LDDReportType.LAW_FIRM,
        sections=sections,
        title="LAW_FIRM style regression",
    )

    assert report.status == LDDReportStatus.READY
    assert report.file_path

    text = _extract_docx_text(report.file_path)

    assert "본 보고서는 Regression Holdings 및 관련 거래에 관하여 수행한 법률실사 결과를 정리한 것으로서" in text
    assert "관할관청의 사전승인이 필요한 것으로 판단됩니다." in text
    assert "계약상 사전 동의 절차가 필요한 것으로 보입니다." in text
    assert "신고 수리 여부를 추가로 확인할 필요가 있는 것으로 사료됩니다." in text
