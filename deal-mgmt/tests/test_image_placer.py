"""image_placer 모듈 단위 테스트."""

import base64
from io import BytesIO

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.pptx.template_engine.exceptions import ImagePlacementError, ShapeNotFoundError
from app.pptx.template_engine.image_placer import (
    _fit_and_center,
    _prepare_image_stream,
    replace_image_in_placeholder,
    replace_images_in_grid,
)


def _make_test_image_base64(width=100, height=80, color="red") -> str:
    """테스트용 PNG 이미지를 Base64로 생성."""
    img = Image.new("RGB", (width, height), color=color)
    buf = BytesIO()
    img.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode()


def _make_test_image_data_uri(width=100, height=80) -> str:
    """data URI 형식의 테스트 이미지."""
    b64 = _make_test_image_base64(width, height)
    return f"data:image/png;base64,{b64}"


@pytest.fixture
def image_slide():
    """이미지 플레이스홀더가 포함된 테스트 슬라이드."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])

    # 바운딩 박스용 도형 (직사각형)
    shape = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        Inches(1),
        Inches(1),
        Inches(4),
        Inches(3),
    )
    shape.name = "img_product"

    # 그리드용 도형들
    for i in range(1, 4):
        s = slide.shapes.add_shape(
            1,
            Inches(1 + (i - 1) * 2),
            Inches(5),
            Inches(1.5),
            Inches(1.5),
        )
        s.name = f"grid_logo_{i}"

    return slide


class TestPrepareImageStream:
    """_prepare_image_stream 내부 헬퍼 테스트."""

    def test_bytes_input(self):
        """bytes 입력 → BytesIO 변환."""
        data = b"\x89PNG\r\n\x1a\n" + b"\x00" * 100
        stream = _prepare_image_stream(data)
        assert isinstance(stream, BytesIO)

    def test_base64_input(self):
        """순수 Base64 문자열 입력."""
        b64 = _make_test_image_base64()
        stream = _prepare_image_stream(b64)
        assert isinstance(stream, BytesIO)

    def test_data_uri_input(self):
        """data:image/png;base64,... 형식."""
        uri = _make_test_image_data_uri()
        stream = _prepare_image_stream(uri)
        assert isinstance(stream, BytesIO)

    def test_invalid_string(self):
        """해석 불가 문자열 — ImagePlacementError."""
        with pytest.raises(ImagePlacementError):
            _prepare_image_stream("not-a-valid-image-or-path")


class TestFitAndCenter:
    """_fit_and_center Aspect Ratio 계산 테스트."""

    def test_wide_image_in_square_box(self):
        """가로가 긴 이미지를 정사각형 박스에 피팅."""
        img = Image.new("RGB", (200, 100), "white")
        buf = BytesIO()
        img.save(buf, format="PNG")

        _left, _top, width, height = _fit_and_center(0, 0, 1000, 1000, buf)
        # 가로가 기준 → 높이가 줄어듦
        assert width == 1000
        assert height == 500

    def test_tall_image_in_wide_box(self):
        """세로가 긴 이미지를 가로 박스에 피팅."""
        img = Image.new("RGB", (100, 200), "white")
        buf = BytesIO()
        img.save(buf, format="PNG")

        _left, _top, width, height = _fit_and_center(0, 0, 2000, 1000, buf)
        # 세로가 기준 → 너비가 줄어듦
        assert height == 1000
        assert width == 500


class TestReplaceImageInPlaceholder:
    """replace_image_in_placeholder 함수 테스트."""

    def test_basic_replace(self, image_slide):
        """기본 이미지 삽입."""
        b64 = _make_test_image_base64()
        pic = replace_image_in_placeholder(image_slide, "img_product", b64)
        assert pic is not None

    def test_data_uri(self, image_slide):
        """data URI 형식 이미지 삽입."""
        uri = _make_test_image_data_uri()
        pic = replace_image_in_placeholder(image_slide, "img_product", uri)
        assert pic is not None

    def test_shape_not_found(self, image_slide):
        """존재하지 않는 shape — ShapeNotFoundError."""
        b64 = _make_test_image_base64()
        with pytest.raises(ShapeNotFoundError):
            replace_image_in_placeholder(image_slide, "img_nonexistent", b64)


class TestReplaceImagesInGrid:
    """replace_images_in_grid 함수 테스트."""

    def test_full_grid(self, image_slide):
        """모든 그리드 슬롯에 이미지 삽입."""
        sources = [_make_test_image_base64(color=c) for c in ["red", "green", "blue"]]
        pics = replace_images_in_grid(image_slide, "grid_logo_", sources)
        assert len(pics) == 3

    def test_partial_grid(self, image_slide):
        """이미지가 그리드 슬롯보다 적은 경우."""
        sources = [_make_test_image_base64()]
        pics = replace_images_in_grid(image_slide, "grid_logo_", sources)
        assert len(pics) == 1

    def test_empty_sources(self, image_slide):
        """빈 소스 리스트 → 플레이스홀더만 삭제."""
        pics = replace_images_in_grid(image_slide, "grid_logo_", [])
        assert len(pics) == 0
