"""투입산출 계수표 시딩 — MA_ValueChain_v7.xlsx Sheet 3 → vc_industry_coefficients 테이블.

1,574×1,574 행렬을 희소(sparse) 형태로 저장. 계수 ≥ 0.001만 저장.

사용법:
    cd deal-mgmt
    python -m scripts.seed_vc_coefficients [--force] [--min-coeff 0.001]
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

from sqlalchemy import delete, func, insert, select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.orm import sessionmaker

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)

EXCEL_PATH = Path(__file__).resolve().parent.parent / "app" / "marketing" / "MA_ValueChain_v7.xlsx"
SHEET_NAME = "투입산출 계수표"
CHUNK_SIZE = 10000
DEFAULT_MIN_COEFF = Decimal("0.001")


async def _get_engine_and_session() -> tuple:
    from app.core.config import settings

    engine = create_async_engine(settings.DATABASE_URL)
    session_factory = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    return engine, session_factory


async def seed_vc_coefficients(
    session: AsyncSession,
    force: bool = False,
    min_coeff: Decimal = DEFAULT_MIN_COEFF,
) -> int:
    """MA_ValueChain_v7.xlsx Sheet 3 계수표 → vc_industry_coefficients 희소 저장."""
    from app.models.vc_industry_coefficient import VcIndustryCoefficient

    if force:
        await session.execute(delete(VcIndustryCoefficient))
        await session.commit()
        logger.info("vc_industry_coefficients 기존 데이터 삭제 완료")

    count = (await session.execute(select(func.count()).select_from(VcIndustryCoefficient))).scalar()
    if count and count > 0 and not force:
        logger.info("vc_industry_coefficients 이미 %d건 존재, 스킵", count)
        return count

    import openpyxl

    if not EXCEL_PATH.exists():
        logger.error("파일 없음: %s", EXCEL_PATH)
        return 0

    logger.info("Excel 로딩 중 (대용량 행렬): %s", EXCEL_PATH)
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True, read_only=True, keep_links=False)
    try:
        ws = wb[SHEET_NAME]
    except KeyError:
        logger.error("시트 '%s' 없음. 유효 시트: %s", SHEET_NAME, wb.sheetnames)
        wb.close()
        return 0

    # Row 3 = 열 레이블 (1,574개 업종명)
    # Column A = 행 레이블
    # 데이터: Row 4+, Column B+

    # 열 레이블 추출 (Row 3, B열부터)
    column_labels: list[str] = []
    for row_data in ws.iter_rows(min_row=3, max_row=3, values_only=True):
        for j, val in enumerate(row_data):
            if j == 0:
                continue  # A열 = 행 레이블 헤더, 스킵
            if val is not None:
                column_labels.append(str(val).strip())
            else:
                column_labels.append("")
        break

    logger.info("열 레이블 %d개 추출", len(column_labels))
    if not column_labels:
        logger.error("열 레이블을 추출할 수 없습니다")
        wb.close()
        return 0

    batch: list[dict] = []
    total = 0

    for row_data in ws.iter_rows(min_row=4, values_only=True):
        source_name = row_data[0]
        if source_name is None:
            continue
        source_name = str(source_name).strip()
        if not source_name:
            continue

        for j, val in enumerate(row_data[1:]):
            if val is None:
                continue
            try:
                coeff = Decimal(str(val))
            except (InvalidOperation, ValueError):
                continue

            if coeff < min_coeff:
                continue

            if j >= len(column_labels) or not column_labels[j]:
                continue

            target_name = column_labels[j]

            batch.append(
                {
                    "source_industry": source_name,
                    "target_industry": target_name,
                    "coefficient": coeff,
                }
            )

            if len(batch) >= CHUNK_SIZE:
                await session.execute(insert(VcIndustryCoefficient), batch)
                await session.commit()
                total += len(batch)
                batch.clear()
                if total % 50000 == 0:
                    logger.info("vc_industry_coefficients: %d건 삽입...", total)

    # 잔여 배치
    if batch:
        await session.execute(insert(VcIndustryCoefficient), batch)
        total += len(batch)

    await session.commit()
    wb.close()
    logger.info("vc_industry_coefficients 총 %d건 삽입 완료 (임계값 ≥ %s)", total, min_coeff)
    return total


async def main(force: bool = False, min_coeff: Decimal = DEFAULT_MIN_COEFF) -> None:
    if force and sys.stdin.isatty():
        confirm = input("--force: 기존 vc_industry_coefficients 데이터를 전부 삭제합니다. 계속? [y/N]: ")
        if confirm.strip().lower() != "y":
            logger.info("취소됨")
            return
    engine, session_factory = await _get_engine_and_session()
    async with session_factory() as session:
        await seed_vc_coefficients(session, force, min_coeff)
    await engine.dispose()
    logger.info("VC 계수표 시딩 완료!")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="VC 투입산출 계수표 시딩")
    parser.add_argument("--force", action="store_true", help="기존 데이터 삭제 후 재삽입")
    parser.add_argument(
        "--min-coeff",
        type=Decimal,
        default=DEFAULT_MIN_COEFF,
        help="저장할 최소 계수 임계값 (기본: 0.001)",
    )
    args = parser.parse_args()
    asyncio.run(main(force=args.force, min_coeff=args.min_coeff))
