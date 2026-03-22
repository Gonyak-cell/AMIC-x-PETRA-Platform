from __future__ import annotations

import argparse
import hashlib
import json
import re
import statistics
import subprocess
import tempfile
import unicodedata
from collections import Counter, defaultdict
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import UTC, datetime, timezone
from pathlib import Path
from typing import Any

import fitz
import yaml
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

_EASYOCR_READER = None


SECTION_ORDER = [
    "BASE",
    "GOVERNANCE",
    "CAPITAL",
    "CONTRACTS",
    "LITIGATION",
    "LABOR",
    "IP",
    "REAL_ESTATE",
    "PERMITS",
    "TAX",
    "DATA_IT",
]

SECTION_LABELS = {
    "BASE": "공통/면책/범위",
    "GOVERNANCE": "기업 일반 및 지배구조",
    "CAPITAL": "자본구조 및 주주",
    "CONTRACTS": "주요 계약",
    "LITIGATION": "소송 및 분쟁",
    "LABOR": "인사 및 노무",
    "IP": "지식재산",
    "REAL_ESTATE": "부동산 및 환경",
    "PERMITS": "인허가 및 규제",
    "TAX": "조세",
    "DATA_IT": "개인정보 및 IT",
}

SECTION_HEADING_KEYWORDS: dict[str, tuple[str, ...]] = {
    "GOVERNANCE": (
        "기업 일반",
        "회사 일반",
        "일반사항",
        "회사 개요",
        "지배구조",
        "거버넌스",
        "정관",
        "등기",
        "이사회",
        "주주총회",
        "대표이사",
        "계열회사",
        "corporate",
        "governance",
        "organization",
    ),
    "CAPITAL": (
        "자본",
        "자본구조",
        "주식",
        "주주",
        "지분",
        "전환사채",
        "신주인수권부사채",
        "우선주",
        "주주간계약",
        "스톡옵션",
        "esop",
        "equity",
        "capital",
    ),
    "CONTRACTS": (
        "계약",
        "주요 계약",
        "영업계약",
        "금융계약",
        "공급계약",
        "판매계약",
        "고객계약",
        "거래계약",
        "change of control",
        "coc",
        "commercial",
        "contract",
    ),
    "LITIGATION": (
        "소송",
        "분쟁",
        "중재",
        "가압류",
        "가처분",
        "행정처분",
        "수사",
        "litigation",
        "dispute",
        "arbitration",
        "claim",
    ),
    "LABOR": (
        "노무",
        "노동",
        "인사",
        "근로",
        "취업규칙",
        "노동조합",
        "단체협약",
        "퇴직금",
        "임직원",
        "labor",
        "employment",
        "hr",
    ),
    "IP": (
        "지식재산",
        "특허",
        "상표",
        "저작권",
        "직무발명",
        "소프트웨어",
        "기술",
        "영업비밀",
        "ip",
        "intellectual property",
        "patent",
        "trademark",
        "copyright",
    ),
    "REAL_ESTATE": (
        "부동산",
        "임대차",
        "임차",
        "전대차",
        "토지",
        "건물",
        "등기부",
        "환경",
        "real estate",
        "lease",
        "property",
    ),
    "PERMITS": (
        "인허가",
        "허가",
        "인가",
        "면허",
        "등록",
        "신고",
        "승인",
        "인증",
        "permit",
        "license",
        "regulatory",
        "규제",
    ),
    "TAX": (
        "세무",
        "조세",
        "법인세",
        "부가가치세",
        "원천세",
        "지방세",
        "tax",
        "vat",
        "withholding",
        "과세",
    ),
    "DATA_IT": (
        "개인정보",
        "정보보호",
        "보안",
        "전산",
        "시스템",
        "데이터",
        "it",
        "isms",
        "iso",
        "privacy",
        "cyber",
        "security",
    ),
    "BASE": (
        "본 보고서",
        "본 실사보고서",
        "법률의견",
        "제3자",
        "배포",
        "사용될 수 없",
        "검토 범위",
        "검토대상",
        "전제",
        "한계",
        "제공받은 자료",
        "인터뷰",
        "scope",
        "limitation",
        "confidential",
        "private & confidential",
    ),
}

SECTION_BODY_KEYWORDS: dict[str, tuple[str, ...]] = {
    "GOVERNANCE": (
        "정관",
        "법인등기",
        "법인등기부",
        "이사회",
        "주주총회",
        "대표이사",
        "사내이사",
        "감사",
        "계열회사",
        "지배구조",
        "이해충돌",
    ),
    "CAPITAL": (
        "발행주식",
        "주식수",
        "지분율",
        "신주",
        "증자",
        "감자",
        "전환사채",
        "신주인수권",
        "우선주",
        "상환전환우선주",
        "주주간계약",
        "tag along",
        "drag along",
    ),
    "CONTRACTS": (
        "거래처",
        "공급계약",
        "판매계약",
        "고객계약",
        "차입약정",
        "대출약정",
        "change of control",
        "termination",
        "해지",
        "위약",
        "라이선스계약",
    ),
    "LITIGATION": (
        "원고",
        "피고",
        "청구",
        "법원",
        "판결",
        "화해",
        "조정",
        "중재",
        "가압류",
        "가처분",
        "고소",
        "수사",
    ),
    "LABOR": (
        "근로계약",
        "취업규칙",
        "노동조합",
        "단체협약",
        "퇴직금",
        "임금",
        "체불",
        "연장근로",
        "직원",
        "임직원",
    ),
    "IP": (
        "특허",
        "상표",
        "저작권",
        "프로그램",
        "소프트웨어",
        "직무발명",
        "영업비밀",
        "라이선스",
        "기술자료",
    ),
    "REAL_ESTATE": (
        "임대차",
        "임차",
        "전대차",
        "토지",
        "건물",
        "등기부",
        "소유권",
        "지상권",
        "근저당",
        "환경오염",
        "폐기물",
    ),
    "PERMITS": (
        "인허가",
        "허가증",
        "인가",
        "면허",
        "등록증",
        "신고증",
        "승인",
        "행정처분",
        "영업허가",
        "변경허가",
        "갱신",
    ),
    "TAX": (
        "법인세",
        "부가가치세",
        "원천세",
        "지방세",
        "과세",
        "세무조사",
        "추징",
        "체납",
        "이전가격",
    ),
    "DATA_IT": (
        "개인정보",
        "개인정보처리방침",
        "정보보호",
        "정보통신망",
        "보안사고",
        "유출",
        "침해",
        "isms",
        "iso 27001",
        "전산시스템",
        "클라우드",
    ),
    "BASE": (
        "본 보고서",
        "법률의견이 아니",
        "제3자",
        "배포되거나",
        "검토 범위",
        "검토를 위하여",
        "제공받은 자료",
        "인터뷰",
        "전제로",
        "한계",
    ),
}

GENERIC_HEADING_TERMS = (
    "사실관계",
    "이슈",
    "검토 결과",
    "검토 의견",
    "분석",
    "시사점",
    "권고사항",
    "해결방안",
    "리스크",
    "현황",
    "개요",
    "요약",
    "자료",
    "첨부",
    "별첨",
    "질의응답",
    "reference",
    "analysis",
    "issues",
    "recommendation",
)
NESTED_HEADING_TERMS = (
    "사실관계",
    "이슈",
    "검토 결과",
    "검토 의견",
    "분석",
    "시사점",
    "권고사항",
    "해결방안",
    "리스크",
    "현황",
    "개요",
    "요약",
    "첨부",
    "별첨",
)

SENTENCE_SPLIT_RE = re.compile(r"(?<=[.!?])\s+")
CLAUSE_CONNECTOR_RE = re.compile(
    r"(?<=,)\s+(?=(?:다만|또한|한편|특히|따라서|아울러|나아가|그리고|그러나|즉|참고로|별도로|이 경우|이와 관련하여))"
)
LIST_MARKER_SPLIT_RE = re.compile(r"\s+(?=(?:\(?[0-9IVXivx가-힣A-Za-z]{1,4}[).]))")
DATE_PATTERNS = [
    re.compile(r"\b\d{4}[./-]\s*\d{1,2}[./-]\s*\d{1,2}\b"),
    re.compile(r"\b\d{4}\.\s*\d{1,2}\.\s*\d{1,2}\.?\b"),
    re.compile(r"\b\d{4}\s*년\s*\d{1,2}\s*월\s*\d{1,2}\s*일\b"),
]
MONEY_PATTERN = re.compile(
    r"(?:(?:KRW|USD|EUR|JPY|US\$|₩)\s*)?\d{1,3}(?:,\d{3})+(?:\.\d+)?\s*(?:원|달러|천원|백만원|만원|억원|조원|million|billion)?",
    re.IGNORECASE,
)
PERCENT_PATTERN = re.compile(r"\d+(?:\.\d+)?\s*%")
CASE_NO_PATTERN = re.compile(r"\b\d{4}\s*[가-힣]{1,6}\s*\d{1,10}\b")
REF_NO_PATTERN = re.compile(r"제\s*\d+\s*호")
SHARE_PATTERN = re.compile(r"\d{1,3}(?:,\d{3})*\s*주")
URL_PATTERN = re.compile(r"https?://\S+")
EMAIL_PATTERN = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
PLACEHOLDER_RE = re.compile(r"\{\{[^}]+\}\}")
COMPANY_PATTERNS = [
    re.compile(r"(?:주식회사|유한회사|합자회사|합명회사)\s*[A-Za-z0-9가-힣&().·,\-_/ ]{1,40}"),
    re.compile(r"㈜\s*[A-Za-z0-9가-힣&().·,\-_/ ]{1,40}"),
    re.compile(r"\b[A-Z][A-Za-z0-9&().,\- ]{2,40}\s+(?:Co\.,?\s*Ltd\.?|Inc\.?|LLC|Limited|Holdings?)\b"),
]
PAGE_NOISE_RE = re.compile(r"^(?:page\s*\d+|\d+\s*/\s*\d+|-+\s*\d+\s*-+|\d+)$", re.IGNORECASE)
ROMAN_RE = re.compile(r"^(?:[IVXLCM]+)[.)]?\s*$")
TOP_HEADING_RE = re.compile(r"^(?:제\s*\d+\s*(?:장|절|항)|[IVXLCM]+[.)]|[0-9]+[.)])")
SUB_HEADING_RE = re.compile(r"^(?:[0-9]+-[0-9]+|[0-9]+\.[0-9]+|[가-힣A-Z]\.|[가-힣A-Z]\)|\([0-9가-힣A-Z]{1,4}\))")
BRACKET_HEADING_RE = re.compile(r"^\[[^\]]{1,40}\]$")
DATE_ONLY_RE = re.compile(r"^(?:\d{4}[./-]\d{1,2}[./-]\d{1,2}|\d{4}\.\s*\d{1,2}\.\s*\d{1,2}\.?)$")
TOKEN_ONLY_RE = re.compile(r"^\(?[A-Za-z0-9가-힣ivxlcm]+\)?[.)]?$", re.IGNORECASE)


@dataclass
class TextUnit:
    source_file: Path
    relative_path: str
    format: str
    order: int
    text: str
    kind: str
    style: str | None = None
    page: int | None = None
    font_size: float | None = None
    bold: bool = False
    heading_level: int | None = None
    heading_path: tuple[str, ...] = ()
    section_hint: str | None = None
    section_id: str | None = None
    section_score: int = 0
    section_hits: tuple[str, ...] = ()


@dataclass
class FileStats:
    relative_path: str
    format: str
    unit_count: int
    paragraph_count: int
    heading_count: int
    table_rows: int
    text_hash: str


@dataclass
class CandidateAggregate:
    normalized: str
    hit_count: int = 0
    file_paths: set[str] = field(default_factory=set)
    headings: Counter[str] = field(default_factory=Counter)
    examples: list[dict[str, str]] = field(default_factory=list)

    @property
    def file_count(self) -> int:
        return len(self.file_paths)

    def add(self, unit: TextUnit, original_text: str) -> None:
        self.hit_count += 1
        self.file_paths.add(unit.relative_path)
        if unit.heading_path:
            self.headings[" > ".join(unit.heading_path)] += 1
        if len(self.examples) >= 3:
            return
        if any(example["text"] == original_text for example in self.examples):
            return
        self.examples.append(
            {
                "file": unit.relative_path,
                "heading": " > ".join(unit.heading_path),
                "text": original_text,
            }
        )


def normalize_whitespace(text: str) -> str:
    text = unicodedata.normalize("NFKC", text or "")
    text = text.replace("\u00a0", " ").replace("\u200b", " ")
    text = text.replace("\r", "\n")
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def clean_text(text: str) -> str:
    text = normalize_whitespace(text)
    return text.strip(" \t|")


def lower_text(text: str) -> str:
    return clean_text(text).lower()


def is_meaningful(text: str) -> bool:
    text = clean_text(text)
    if not text:
        return False
    if PAGE_NOISE_RE.match(text):
        return False
    alnum_count = sum(ch.isalnum() for ch in text)
    return not alnum_count < 2


def is_candidate_text(text: str) -> bool:
    text = clean_text(text)
    if len(text) < 25:
        return False
    if PAGE_NOISE_RE.match(text):
        return False
    if text.count("|") >= 4:
        return False
    return not sum(ch.isdigit() for ch in text) > len(text) * 0.45


def escape_ps_literal(path: Path) -> str:
    return str(path).replace("'", "''")


def iter_block_items(parent: DocxDocument | _Cell) -> Iterable[Paragraph | Table]:
    if isinstance(parent, DocxDocument):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def extract_docx_units(path: Path, sample_root: Path, format_hint: str | None = None) -> list[TextUnit]:
    doc = Document(path)
    relative_path = str(path.relative_to(sample_root))
    units: list[TextUnit] = []
    order = 0

    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = clean_text(block.text)
            if not is_meaningful(text):
                continue
            style = block.style.name if block.style is not None else None
            units.append(
                TextUnit(
                    source_file=path,
                    relative_path=relative_path,
                    format=format_hint or path.suffix.lower().lstrip("."),
                    order=order,
                    text=text,
                    kind="paragraph",
                    style=style,
                )
            )
            order += 1
            continue

        for row in block.rows:
            row_cells: list[str] = []
            for cell in row.cells:
                cell_text = clean_text(" ".join(clean_text(p.text) for p in cell.paragraphs if clean_text(p.text)))
                if cell_text and cell_text not in row_cells:
                    row_cells.append(cell_text)
            text = " | ".join(row_cells)
            if not is_meaningful(text):
                continue
            units.append(
                TextUnit(
                    source_file=path,
                    relative_path=relative_path,
                    format=format_hint or path.suffix.lower().lstrip("."),
                    order=order,
                    text=text,
                    kind="table_row",
                    style="table",
                )
            )
            order += 1

    return units


def convert_doc_to_docx(path: Path, destination: Path) -> None:
    script = f"""
$ErrorActionPreference = 'Stop'
$src = [string](Resolve-Path '{escape_ps_literal(path)}')
$out = [string]'{escape_ps_literal(destination)}'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {{
  $doc = $word.Documents.Open($src)
  $doc.SaveAs2($out, 16)
  $doc.Close()
  Write-Output 'OK'
}}
finally {{
  $word.Quit()
  [System.Runtime.Interopservices.Marshal]::ReleaseComObject($word) | Out-Null
}}
""".strip()
    result = subprocess.run(
        ["powershell", "-NoProfile", "-Command", script],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0 or not destination.exists():
        raise RuntimeError(
            f"failed to convert DOC file via Word COM: {path}\nstdout={result.stdout}\nstderr={result.stderr}"
        )


def extract_doc_units(path: Path, sample_root: Path) -> list[TextUnit]:
    with tempfile.TemporaryDirectory(prefix="ldd-doc-convert-") as temp_dir:
        temp_dir_path = Path(temp_dir)
        temp_source = temp_dir_path / "source.doc"
        temp_source.write_bytes(path.read_bytes())
        temp_docx = temp_dir_path / "source.docx"
        convert_doc_to_docx(temp_source, temp_docx)
        units = extract_docx_units(temp_docx, temp_docx.parent, format_hint="doc")
        relative_path = str(path.relative_to(sample_root))
        for unit in units:
            unit.source_file = path
            unit.relative_path = relative_path
        return units


def extract_pdf_units(path: Path, sample_root: Path) -> list[TextUnit]:
    relative_path = str(path.relative_to(sample_root))
    pdf = fitz.open(path)
    units: list[TextUnit] = []
    order = 0

    for page_index in range(pdf.page_count):
        page = pdf[page_index]
        page_dict = page.get_text("dict")
        sizes: list[float] = []
        lines_by_block: list[list[tuple[str, float, bool]]] = []

        for block in page_dict.get("blocks", []):
            if block.get("type") != 0:
                continue
            line_items: list[tuple[str, float, bool]] = []
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                line_text = clean_text("".join(span.get("text", "") for span in spans))
                if not is_meaningful(line_text):
                    continue
                font_size = max((float(span.get("size", 0.0)) for span in spans), default=0.0)
                bold = any("bold" in str(span.get("font", "")).lower() for span in spans)
                sizes.append(font_size)
                line_items.append((line_text, font_size, bold))
            if line_items:
                lines_by_block.append(line_items)

        median_font = statistics.median(sizes) if sizes else 10.0

        for line_items in lines_by_block:
            buffer: list[str] = []
            buffer_font: float | None = None
            buffer_bold = False

            def flush_buffer(*, _page_index: int = page_index) -> None:
                nonlocal order, buffer, buffer_font, buffer_bold
                if not buffer:
                    return
                text = clean_text(" ".join(buffer))
                if is_meaningful(text):
                    units.append(
                        TextUnit(
                            source_file=path,
                            relative_path=relative_path,
                            format="pdf",
                            order=order,
                            text=text,
                            kind="paragraph",
                            style="pdf",
                            page=_page_index + 1,
                            font_size=buffer_font,
                            bold=buffer_bold,
                        )
                    )
                    order += 1
                buffer = []
                buffer_font = None
                buffer_bold = False

            for line_text, font_size, bold in line_items:
                candidate = TextUnit(
                    source_file=path,
                    relative_path=relative_path,
                    format="pdf",
                    order=order,
                    text=line_text,
                    kind="line",
                    style="pdf-line",
                    page=page_index + 1,
                    font_size=font_size,
                    bold=bold,
                )
                level = detect_heading_level(candidate, median_font=median_font)
                if level is not None:
                    flush_buffer()
                    candidate.kind = "heading"
                    candidate.heading_level = level
                    units.append(candidate)
                    order += 1
                    continue

                if buffer:
                    projected = " ".join([*buffer, line_text])
                    if len(projected) > 900 or buffer[-1].endswith((".", "?", "!", "다.", "함.", "음.")):
                        flush_buffer()
                buffer.append(line_text)
                buffer_font = max(buffer_font or 0.0, font_size)
                buffer_bold = buffer_bold or bold
            flush_buffer()

    if units:
        return units
    return extract_pdf_units_with_ocr(path, sample_root)


def extract_pdf_units_with_ocr(path: Path, sample_root: Path) -> list[TextUnit]:
    reader = get_easyocr_reader()
    if reader is None:
        return []

    try:
        import numpy as np
    except Exception:
        return []

    relative_path = str(path.relative_to(sample_root))
    pdf = fitz.open(path)
    units: list[TextUnit] = []
    order = 0

    for page_index in range(pdf.page_count):
        page = pdf[page_index]
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)

        try:
            results = reader.readtext(image, detail=0, paragraph=True)
        except Exception:
            continue

        for text in results:
            cleaned = clean_text(text)
            if not is_meaningful(cleaned):
                continue
            units.append(
                TextUnit(
                    source_file=path,
                    relative_path=relative_path,
                    format="pdf",
                    order=order,
                    text=cleaned,
                    kind="paragraph",
                    style="ocr",
                    page=page_index + 1,
                )
            )
            order += 1

    return units


def get_easyocr_reader():
    global _EASYOCR_READER

    if _EASYOCR_READER is False:
        return None
    if _EASYOCR_READER is not None:
        return _EASYOCR_READER

    try:
        import easyocr

        _EASYOCR_READER = easyocr.Reader(["ko", "en"], gpu=False, verbose=False)
    except Exception:
        _EASYOCR_READER = False
        return None
    return _EASYOCR_READER


def extract_json_units(path: Path, sample_root: Path) -> list[TextUnit]:
    relative_path = str(path.relative_to(sample_root))
    payload = json.loads(path.read_text(encoding="utf-8"))
    units: list[TextUnit] = []
    order = 0

    def walk(node: Any, heading_path: list[str]) -> None:
        nonlocal order
        if isinstance(node, dict):
            for key, value in node.items():
                key_text = clean_text(str(key))
                if key_text:
                    units.append(
                        TextUnit(
                            source_file=path,
                            relative_path=relative_path,
                            format="json",
                            order=order,
                            text=key_text,
                            kind="heading",
                            style="json-key",
                            heading_level=min(len(heading_path) + 1, 4),
                        )
                    )
                    order += 1
                walk(value, heading_path + ([key_text] if key_text else []))
            return

        if isinstance(node, list):
            for item in node:
                walk(item, heading_path)
            return

        if isinstance(node, str):
            text = clean_text(node)
            if not is_meaningful(text):
                return
            units.append(
                TextUnit(
                    source_file=path,
                    relative_path=relative_path,
                    format="json",
                    order=order,
                    text=text,
                    kind="paragraph",
                    style="json-value",
                )
            )
            order += 1

    walk(payload, [])
    return units


def detect_heading_level(unit: TextUnit, median_font: float | None = None) -> int | None:
    if unit.heading_level is not None:
        return unit.heading_level

    text = clean_text(unit.text)
    lowered = text.lower()
    style = (unit.style or "").lower()
    font_size = unit.font_size or 0.0

    if not text:
        return None
    if len(text) > 150 and not BRACKET_HEADING_RE.match(text):
        return None
    if unit.format == "pdf" and DATE_ONLY_RE.match(text):
        return None

    if "heading" in style or "제목" in style or style.startswith("title"):
        match = re.search(r"(\d+)", style)
        if match:
            return max(1, min(int(match.group(1)), 4))
        return 1

    if style == "json-key":
        return unit.heading_level or 2

    if BRACKET_HEADING_RE.match(text):
        return 4
    if TOP_HEADING_RE.match(text) or ROMAN_RE.match(text):
        return 1 if "장" in text or ROMAN_RE.match(text) else 2
    if SUB_HEADING_RE.match(text):
        return 3

    short_line = len(text) <= 60 and not text.endswith((".", "?", "!", "다", "함"))
    has_section_keyword = any(
        keyword in lowered for keywords in SECTION_HEADING_KEYWORDS.values() for keyword in keywords
    )
    has_generic_keyword = any(term in lowered for term in GENERIC_HEADING_TERMS)

    if unit.format == "pdf":
        short_pdf_line = (
            len(text) <= 70 and len(text.split()) <= 12 and not text.endswith((".", "?", "!", "다.", "함.", "음."))
        )
        looks_structural = bool(
            BRACKET_HEADING_RE.match(text) or TOP_HEADING_RE.match(text) or SUB_HEADING_RE.match(text)
        )
        if (
            median_font
            and font_size >= median_font * 1.28
            and short_pdf_line
            and (has_section_keyword or has_generic_keyword or looks_structural)
        ):
            return 1 if font_size >= median_font * 1.55 else 2
        if unit.bold and short_pdf_line and (has_section_keyword or has_generic_keyword or looks_structural):
            return 2
        return None

    if unit.bold and short_line:
        return 2
    if has_section_keyword and short_line:
        return 2
    if has_generic_keyword and short_line:
        return 3
    return None


def classify_heading_only(text: str) -> str | None:
    text_lower = lower_text(text)
    best_section: str | None = None
    best_score = 0
    for section_id, keywords in SECTION_HEADING_KEYWORDS.items():
        if section_id == "BASE":
            continue
        score = sum(3 for keyword in keywords if keyword and keyword in text_lower)
        if score > best_score:
            best_section = section_id
            best_score = score
    return best_section if best_score >= 3 else None


def score_sections(heading_text: str, body_text: str) -> dict[str, tuple[int, list[str]]]:
    heading_lower = lower_text(heading_text)
    body_lower = lower_text(body_text)
    scores: dict[str, tuple[int, list[str]]] = {}

    for section_id in SECTION_ORDER:
        score = 0
        hits: list[str] = []
        for keyword in SECTION_HEADING_KEYWORDS.get(section_id, ()):
            if keyword in heading_lower:
                score += 4
                hits.append(f"H:{keyword}")
        for keyword in SECTION_BODY_KEYWORDS.get(section_id, ()):
            if keyword in body_lower:
                score += 1
                hits.append(f"B:{keyword}")
        scores[section_id] = (score, hits)

    return scores


def assign_heading_paths(units: list[TextUnit]) -> list[TextUnit]:
    stack: list[dict[str, Any]] = []
    assigned: list[TextUnit] = []

    for unit in units:
        level = detect_heading_level(unit)
        if level is not None:
            own_section = classify_heading_only(unit.text)
            inherit_allowed = own_section is not None or should_inherit_section(unit.text)
            inherited_section = own_section or (
                next((entry["section_hint"] for entry in reversed(stack) if entry["section_hint"]), None)
                if inherit_allowed
                else None
            )
            while stack and stack[-1]["level"] >= level:
                stack.pop()
            stack.append({"level": level, "text": unit.text, "section_hint": inherited_section})
            unit.heading_level = level
            unit.heading_path = tuple(entry["text"] for entry in stack)
            unit.section_hint = inherited_section
            assigned.append(unit)
            continue

        unit.heading_path = tuple(entry["text"] for entry in stack)
        unit.section_hint = next((entry["section_hint"] for entry in reversed(stack) if entry["section_hint"]), None)
        assigned.append(unit)

    return assigned


def should_inherit_section(text: str) -> bool:
    lowered = lower_text(text)
    return (
        bool(BRACKET_HEADING_RE.match(text))
        or bool(TOP_HEADING_RE.match(text))
        or bool(SUB_HEADING_RE.match(text))
        or any(term in lowered for term in NESTED_HEADING_TERMS)
        or (TOKEN_ONLY_RE.match(text) is not None and not DATE_ONLY_RE.match(text))
    )


def classify_unit(unit: TextUnit) -> TextUnit:
    heading_text = " > ".join(unit.heading_path)
    scores = score_sections(heading_text, unit.text)
    if unit.section_hint:
        score, hits = scores.get(unit.section_hint, (0, []))
        scores[unit.section_hint] = (score + 6, [*hits, f"STACK:{unit.section_hint}"])

    best_section = None
    best_score = 0
    best_hits: list[str] = []
    runner_up = 0

    for section_id, (score, hits) in scores.items():
        if score > best_score:
            runner_up = best_score
            best_section = section_id
            best_score = score
            best_hits = hits
        elif score > runner_up:
            runner_up = score

    if best_section == "BASE" and best_score < 4:
        best_section = None
    elif best_score < 3 or best_score - runner_up < 1:
        best_section = unit.section_hint if unit.section_hint else None
        best_score = best_score if best_section else 0

    unit.section_id = best_section
    unit.section_score = best_score
    unit.section_hits = tuple(best_hits[:8])
    return unit


def split_candidate_chunks(text: str, max_chars: int = 700) -> list[str]:
    text = clean_text(text)
    if len(text) <= max_chars:
        return [text]
    sentences = [sentence.strip() for sentence in SENTENCE_SPLIT_RE.split(text) if len(sentence.strip()) >= 20]
    if len(sentences) <= 1:
        return [text]

    chunks: list[str] = []
    current: list[str] = []
    current_len = 0
    for sentence in sentences:
        projected = current_len + len(sentence) + (1 if current else 0)
        if current and projected > max_chars:
            chunks.append(" ".join(current))
            current = [sentence]
            current_len = len(sentence)
            continue
        current.append(sentence)
        current_len = projected
    if current:
        chunks.append(" ".join(current))
    return chunks or [text]


def split_reusable_sentences(text: str) -> list[str]:
    text = clean_text(text)
    if not text:
        return []

    pieces = re.split(r"(?:\n+|(?<=[.!?])\s+|(?<=;)\s+)", text)
    sentences = [clean_text(piece) for piece in pieces if len(clean_text(piece)) >= 20]
    return sentences or ([text] if len(text) >= 20 else [])


def split_reusable_clauses(text: str) -> list[str]:
    text = clean_text(text)
    if len(text) < 60:
        return []

    queue = [text]
    splitters = [CLAUSE_CONNECTOR_RE, LIST_MARKER_SPLIT_RE]
    clauses: list[str] = []

    for splitter in splitters:
        next_queue: list[str] = []
        for item in queue:
            parts = [clean_text(part) for part in splitter.split(item) if clean_text(part)]
            if len(parts) <= 1:
                next_queue.append(item)
                continue
            next_queue.extend(parts)
        queue = next_queue

    for item in queue:
        if item == text:
            continue
        if len(item) < 25 or len(item) > 260:
            continue
        clauses.append(item)

    return clauses


def normalize_pattern(text: str) -> str:
    text = clean_text(text)
    text = URL_PATTERN.sub("{{URL}}", text)
    text = EMAIL_PATTERN.sub("{{EMAIL}}", text)
    for pattern in DATE_PATTERNS:
        text = pattern.sub("{{DATE}}", text)
    text = MONEY_PATTERN.sub("{{AMOUNT}}", text)
    text = PERCENT_PATTERN.sub("{{PERCENT}}", text)
    text = CASE_NO_PATTERN.sub("{{CASE_NO}}", text)
    text = REF_NO_PATTERN.sub("{{REF_NO}}", text)
    text = SHARE_PATTERN.sub("{{SHARE_COUNT}}", text)
    for pattern in COMPANY_PATTERNS:
        text = pattern.sub("{{COMPANY}}", text)
    text = re.sub(r"\b\d{2,}\b", "{{NUMBER}}", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def is_reusable_fragment(text: str) -> bool:
    if not is_candidate_text(text):
        return False

    normalized = normalize_pattern(text)
    if len(normalized) < 30 or len(normalized) > 320:
        return False

    placeholders = PLACEHOLDER_RE.findall(normalized)
    if len(placeholders) >= 8:
        return False
    if placeholders and (len(placeholders) * 14 > len(normalized)):
        return False

    alpha_tokens = re.findall(r"[A-Za-z가-힣]{2,}", normalized)
    if len(alpha_tokens) < 4:
        return False

    if normalized.count("{{NUMBER}}") >= 4 and len(alpha_tokens) < 6:
        return False

    stripped = normalized.strip(" -:;,.")
    if len(stripped) < 30:
        return False
    if PAGE_NOISE_RE.match(stripped):
        return False

    return True


def is_shortlist_phrase(text: str) -> bool:
    text = clean_text(text)
    if not text:
        return False
    if len(text) < 35:
        return False

    stripped = text.strip()
    if stripped[-1] in ",(":
        return False
    if stripped.endswith(("및", "또는", "관련", "대하여", "관하여", "기준으로", "수준으로", "범위에서")):
        return False

    if stripped.endswith(
        (
            ".",
            "다",
            "다.",
            "니다",
            "니다.",
            "습니다",
            "습니다.",
            "음",
            "임",
            "함",
            "됨",
            "필요",
            "보임",
            "없음",
        )
    ):
        return True

    if re.search(r"[A-Za-z0-9]\.$", stripped):
        return True

    return False


def extract_reusable_fragments(text: str) -> list[str]:
    fragments: list[str] = []
    seen: set[str] = set()

    for sentence in split_reusable_sentences(text):
        if is_reusable_fragment(sentence):
            normalized = normalize_pattern(sentence)
            if normalized not in seen:
                fragments.append(sentence)
                seen.add(normalized)

        for clause in split_reusable_clauses(sentence):
            if not is_reusable_fragment(clause):
                continue
            normalized = normalize_pattern(clause)
            if normalized in seen:
                continue
            fragments.append(clause)
            seen.add(normalized)

    return fragments


def build_text_hash(units: list[TextUnit]) -> str:
    payload = "\n".join(unit.text for unit in units if unit.kind != "heading")
    return hashlib.sha1(payload.encode("utf-8")).hexdigest()


def write_extracted_text(output_dir: Path, relative_path: str, units: list[TextUnit]) -> None:
    target = output_dir / "extracted_text" / Path(relative_path)
    target = target.with_suffix(target.suffix + ".txt")
    target.parent.mkdir(parents=True, exist_ok=True)

    lines = [
        f"# FILE: {relative_path}",
        f"# GENERATED_AT: {datetime.now(UTC).isoformat()}",
        "",
    ]
    for unit in units:
        if unit.heading_level is not None:
            lines.append(f"[H{unit.heading_level}] {unit.text}")
            continue
        section = unit.section_id or "-"
        heading = " > ".join(unit.heading_path) or "-"
        lines.append(f"[{section}] {heading}")
        lines.append(unit.text)
        lines.append("")

    target.write_text("\n".join(lines), encoding="utf-8")


def collect_candidates(units: list[TextUnit], aggregations: dict[str, dict[str, CandidateAggregate]]) -> None:
    for unit in units:
        if unit.heading_level is not None:
            continue
        if unit.kind == "table_row":
            continue
        if not is_candidate_text(unit.text):
            continue

        section_id = unit.section_id or "BASE"
        if section_id not in aggregations:
            section_id = "BASE"

        for chunk in split_candidate_chunks(unit.text):
            if not is_candidate_text(chunk):
                continue
            normalized = normalize_pattern(chunk)
            if len(normalized) < 25:
                continue
            bucket = aggregations[section_id].get(normalized)
            if bucket is None:
                bucket = CandidateAggregate(normalized=normalized)
                aggregations[section_id][normalized] = bucket
            bucket.add(unit, chunk)


def collect_phrase_candidates(units: list[TextUnit], aggregations: dict[str, dict[str, CandidateAggregate]]) -> None:
    for unit in units:
        if unit.heading_level is not None:
            continue
        if not is_candidate_text(unit.text):
            continue

        section_id = unit.section_id or "BASE"
        if section_id not in aggregations:
            section_id = "BASE"

        for fragment in extract_reusable_fragments(unit.text):
            normalized = normalize_pattern(fragment)
            for target_section in {section_id, "BASE"}:
                bucket = aggregations[target_section].get(normalized)
                if bucket is None:
                    bucket = CandidateAggregate(normalized=normalized)
                    aggregations[target_section][normalized] = bucket
                bucket.add(unit, fragment)


def write_candidates_markdown(output_dir: Path, section_id: str, candidates: dict[str, CandidateAggregate]) -> None:
    target_name = "_base_candidates.md" if section_id == "BASE" else f"{section_id.lower()}_candidates.md"
    target = output_dir / target_name
    ordered = sorted(
        candidates.values(),
        key=lambda candidate: (candidate.file_count, candidate.hit_count, len(candidate.normalized)),
        reverse=True,
    )

    lines = [f"# {SECTION_LABELS[section_id]} 후보", ""]
    if not ordered:
        lines.append("후보가 없습니다.")
        target.write_text("\n".join(lines), encoding="utf-8")
        return

    for index, candidate in enumerate(ordered[:40], start=1):
        lines.append(f"## {index}. files={candidate.file_count}, hits={candidate.hit_count}")
        lines.append("")
        lines.append("정규화 패턴")
        lines.append("```text")
        lines.append(candidate.normalized)
        lines.append("```")
        lines.append("")
        if candidate.headings:
            lines.append("주요 heading")
            for heading, count in candidate.headings.most_common(3):
                lines.append(f"- {heading} ({count})")
            lines.append("")
        if candidate.examples:
            lines.append("예시")
            for example in candidate.examples:
                lines.append(f"- {example['file']} | {example['heading'] or '-'}")
                lines.append("```text")
                lines.append(example["text"])
                lines.append("```")
            lines.append("")

    target.write_text("\n".join(lines), encoding="utf-8")


def write_candidate_yaml(output_dir: Path, section_id: str, candidates: dict[str, CandidateAggregate]) -> None:
    ordered = sorted(
        candidates.values(),
        key=lambda candidate: (candidate.file_count, candidate.hit_count, len(candidate.normalized)),
        reverse=True,
    )
    payload = {
        "section_id": section_id,
        "section_label": SECTION_LABELS[section_id],
        "generated_at": datetime.now(UTC).isoformat(),
        "candidate_count": len(ordered),
        "candidates": [
            {
                "id": f"{section_id.lower()}_{index:02d}",
                "normalized_text": candidate.normalized,
                "file_count": candidate.file_count,
                "hit_count": candidate.hit_count,
                "headings": [heading for heading, _ in candidate.headings.most_common(3)],
                "examples": candidate.examples,
            }
            for index, candidate in enumerate(ordered[:25], start=1)
        ],
    }
    target_name = "_base_draft.yaml" if section_id == "BASE" else f"{section_id.lower()}_draft.yaml"
    (output_dir / target_name).write_text(
        yaml.safe_dump(payload, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )


def write_phrase_candidates_markdown(output_dir: Path, section_id: str, candidates: dict[str, CandidateAggregate]) -> None:
    target_name = "_base_phrase_candidates.md" if section_id == "BASE" else f"{section_id.lower()}_phrase_candidates.md"
    target = output_dir / target_name
    ordered = sorted(
        candidates.values(),
        key=lambda candidate: (candidate.file_count, candidate.hit_count, len(candidate.normalized)),
        reverse=True,
    )

    lines = [f"# {SECTION_LABELS[section_id]} 재사용 문구 후보", ""]
    if not ordered:
        lines.append("후보가 없습니다.")
        target.write_text("\n".join(lines), encoding="utf-8")
        return

    for index, candidate in enumerate(ordered[:80], start=1):
        lines.append(f"## {index}. files={candidate.file_count}, hits={candidate.hit_count}")
        lines.append("")
        lines.append("정규화 패턴")
        lines.append("```text")
        lines.append(candidate.normalized)
        lines.append("```")
        lines.append("")
        if candidate.headings:
            lines.append("주요 heading")
            for heading, count in candidate.headings.most_common(3):
                lines.append(f"- {heading} ({count})")
            lines.append("")
        if candidate.examples:
            lines.append("예시")
            for example in candidate.examples:
                lines.append(f"- {example['file']} | {example['heading'] or '-'}")
                lines.append("```text")
                lines.append(example["text"])
                lines.append("```")
            lines.append("")

    target.write_text("\n".join(lines), encoding="utf-8")


def write_phrase_candidate_yaml(output_dir: Path, section_id: str, candidates: dict[str, CandidateAggregate]) -> None:
    ordered = sorted(
        candidates.values(),
        key=lambda candidate: (candidate.file_count, candidate.hit_count, len(candidate.normalized)),
        reverse=True,
    )
    payload = {
        "section_id": section_id,
        "section_label": SECTION_LABELS[section_id],
        "generated_at": datetime.now(UTC).isoformat(),
        "candidate_count": len(ordered),
        "selection_rule": "Sentence/clause-level reusable phrase mining from heading-aware corpus",
        "candidates": [
            {
                "id": f"{section_id.lower()}_phrase_{index:02d}",
                "normalized_text": candidate.normalized,
                "file_count": candidate.file_count,
                "hit_count": candidate.hit_count,
                "headings": [heading for heading, _ in candidate.headings.most_common(3)],
                "examples": candidate.examples,
            }
            for index, candidate in enumerate(ordered[:60], start=1)
        ],
    }
    target_name = "_base_phrase_draft.yaml" if section_id == "BASE" else f"{section_id.lower()}_phrase_draft.yaml"
    (output_dir / target_name).write_text(
        yaml.safe_dump(payload, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )


def write_promotion_shortlist(output_dir: Path, aggregations: dict[str, dict[str, CandidateAggregate]]) -> None:
    lines = [
        "# Promotion Shortlist",
        "",
        "file_count >= 2 이고 문장/절 단위로 반복된 재사용 후보를 섹션별로 정리한 목록입니다.",
        "",
    ]

    for section_id in SECTION_ORDER:
        ordered = sorted(
            (
                candidate
                for candidate in aggregations[section_id].values()
                if candidate.file_count >= 2
                and is_shortlist_phrase(candidate.normalized)
            ),
            key=lambda candidate: (candidate.file_count, candidate.hit_count, len(candidate.normalized)),
            reverse=True,
        )
        lines.append(f"## {section_id} / {SECTION_LABELS[section_id]}")
        lines.append("")
        if not ordered:
            lines.append("- 없음")
            lines.append("")
            continue
        for candidate in ordered[:20]:
            lines.append(f"- files={candidate.file_count}, hits={candidate.hit_count} :: {candidate.normalized}")
        lines.append("")

    (output_dir / "promotion_shortlist.md").write_text("\n".join(lines), encoding="utf-8")


def write_summary(
    output_dir: Path,
    sample_dir: Path,
    file_stats: list[FileStats],
    section_counts: Counter[str],
    failures: list[str],
    duplicates: dict[str, list[str]],
    zero_unit_files: list[str],
    phrase_counts: Counter[str],
) -> None:
    lines = [
        "# LDD Sample Analysis Summary",
        "",
        f"- Sample dir: `{sample_dir}`",
        f"- Generated at: `{datetime.now(UTC).isoformat()}`",
        f"- Files processed: `{len(file_stats)}`",
        f"- Failures: `{len(failures)}`",
        "",
        "## File Breakdown",
        "",
    ]

    by_format = Counter(stat.format for stat in file_stats)
    for ext, count in sorted(by_format.items()):
        lines.append(f"- `{ext}`: {count}")
    lines.append("")

    lines.append("## Section Assignment")
    lines.append("")
    for section_id in SECTION_ORDER:
        lines.append(f"- `{section_id}` ({SECTION_LABELS[section_id]}): {section_counts.get(section_id, 0)}")
    lines.append("")

    lines.append("## Reusable Phrase Candidates")
    lines.append("")
    for section_id in SECTION_ORDER:
        lines.append(f"- `{section_id}` ({SECTION_LABELS[section_id]}): {phrase_counts.get(section_id, 0)}")
    lines.append("")

    lines.append("## Files")
    lines.append("")
    for stat in sorted(file_stats, key=lambda item: item.relative_path.lower()):
        lines.append(
            f"- `{stat.relative_path}` | format={stat.format} | units={stat.unit_count} | paragraphs={stat.paragraph_count} | headings={stat.heading_count} | table_rows={stat.table_rows}"
        )
    lines.append("")

    if duplicates:
        lines.append("## Duplicate Text Groups")
        lines.append("")
        for paths in duplicates.values():
            lines.append(f"- {', '.join(f'`{path}`' for path in paths)}")
        lines.append("")

    if zero_unit_files:
        lines.append("## Zero-Text Files")
        lines.append("")
        lines.append("아래 파일은 텍스트 레이어가 없거나 추출 가능한 문자열이 없어 본문 분석 결과가 비어 있습니다.")
        for relative_path in zero_unit_files:
            lines.append(f"- `{relative_path}`")
        lines.append("")

    if failures:
        lines.append("## Failures")
        lines.append("")
        for failure in failures:
            lines.append(f"- {failure}")
        lines.append("")

    (output_dir / "summary.md").write_text("\n".join(lines), encoding="utf-8")


def write_manifest(output_dir: Path, file_stats: list[FileStats], failures: list[str]) -> None:
    payload = {
        "generated_at": datetime.now(UTC).isoformat(),
        "files": [stat.__dict__ for stat in file_stats],
        "failures": failures,
    }
    (output_dir / "manifest.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def resolve_units(path: Path, sample_root: Path) -> list[TextUnit]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx_units(path, sample_root)
    if suffix == ".doc":
        return extract_doc_units(path, sample_root)
    if suffix == ".pdf":
        return extract_pdf_units(path, sample_root)
    if suffix == ".json":
        return extract_json_units(path, sample_root)
    raise ValueError(f"unsupported file type: {path}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract section-aware LDD slot-fill candidates from sample reports.")
    parser.add_argument("--sample-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    sample_dir = args.sample_dir.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    supported_files = sorted(
        [
            path
            for path in sample_dir.rglob("*")
            if path.is_file() and path.suffix.lower() in {".docx", ".doc", ".pdf", ".json"}
        ]
    )

    aggregations: dict[str, dict[str, CandidateAggregate]] = {section_id: {} for section_id in SECTION_ORDER}
    phrase_aggregations: dict[str, dict[str, CandidateAggregate]] = {section_id: {} for section_id in SECTION_ORDER}
    file_stats: list[FileStats] = []
    section_counts: Counter[str] = Counter()
    phrase_counts: Counter[str] = Counter()
    failures: list[str] = []
    text_hash_index: defaultdict[str, list[str]] = defaultdict(list)
    zero_unit_files: list[str] = []

    for path in supported_files:
        try:
            units = resolve_units(path, sample_dir)
            units = assign_heading_paths(units)
            units = [classify_unit(unit) for unit in units]

            for unit in units:
                if unit.heading_level is None and unit.section_id:
                    section_counts[unit.section_id] += 1

            collect_candidates(units, aggregations)
            collect_phrase_candidates(units, phrase_aggregations)
            write_extracted_text(output_dir, str(path.relative_to(sample_dir)), units)

            stats = FileStats(
                relative_path=str(path.relative_to(sample_dir)),
                format=path.suffix.lower().lstrip("."),
                unit_count=len(units),
                paragraph_count=sum(1 for unit in units if unit.heading_level is None and unit.kind == "paragraph"),
                heading_count=sum(1 for unit in units if unit.heading_level is not None),
                table_rows=sum(1 for unit in units if unit.kind == "table_row"),
                text_hash=build_text_hash(units),
            )
            file_stats.append(stats)
            text_hash_index[stats.text_hash].append(stats.relative_path)
            if stats.unit_count == 0:
                zero_unit_files.append(stats.relative_path)
        except Exception as exc:
            failures.append(f"{path.relative_to(sample_dir)} :: {exc}")

    duplicates = {digest: paths for digest, paths in text_hash_index.items() if len(paths) > 1}

    for section_id in SECTION_ORDER:
        write_candidates_markdown(output_dir, section_id, aggregations[section_id])
        write_candidate_yaml(output_dir, section_id, aggregations[section_id])
        write_phrase_candidates_markdown(output_dir, section_id, phrase_aggregations[section_id])
        write_phrase_candidate_yaml(output_dir, section_id, phrase_aggregations[section_id])
        phrase_counts[section_id] = sum(1 for candidate in phrase_aggregations[section_id].values() if candidate.file_count >= 2)

    write_promotion_shortlist(output_dir, phrase_aggregations)
    write_summary(output_dir, sample_dir, file_stats, section_counts, failures, duplicates, zero_unit_files, phrase_counts)
    write_manifest(output_dir, file_stats, failures)

    print(f"[OK] processed files: {len(file_stats)}")
    print(f"[OK] failures: {len(failures)}")
    print(f"[OK] output: {output_dir}")
    return 0 if not failures else 1


if __name__ == "__main__":
    raise SystemExit(main())
