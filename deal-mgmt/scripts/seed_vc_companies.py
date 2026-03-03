"""Value Chain 기업개황 시딩 — MA_ValueChain_v7.xlsx Sheet 2 → vc_companies 테이블.

사용법:
    cd deal-mgmt
    python -m scripts.seed_vc_companies [--force]
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import sys
from pathlib import Path

from sqlalchemy import delete, func, insert, select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.orm import sessionmaker

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)

EXCEL_PATH = Path(__file__).resolve().parent.parent / "app" / "marketing" / "MA_ValueChain_v7.xlsx"
SHEET_NAME = "기업개황"
DATA_START_ROW = 4  # Row 4 (1-indexed) = 데이터 시작
CHUNK_SIZE = 5000


def _safe_str(val: object, max_len: int = 300) -> str | None:
    """셀 값을 문자열로 변환. None/빈값 → None, 길이 제한."""
    if val is None:
        return None
    s = str(val).strip()
    return s[:max_len] if s else None


def _safe_int(val: object) -> int | None:
    """셀 값을 int로 변환."""
    if val is None:
        return None
    try:
        return int(float(str(val).strip()))
    except (ValueError, TypeError):
        return None


async def _get_engine_and_session() -> tuple:
    from app.core.config import settings

    engine = create_async_engine(settings.DATABASE_URL)
    session_factory = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    return engine, session_factory


async def seed_vc_companies(session: AsyncSession, force: bool = False) -> int:
    """MA_ValueChain_v7.xlsx Sheet 2 기업개황 → vc_companies 벌크 삽입."""
    from app.models.vc_company import VcCompany

    if force:
        await session.execute(delete(VcCompany))
        await session.commit()
        logger.info("vc_companies 기존 데이터 삭제 완료")

    count = (await session.execute(select(func.count()).select_from(VcCompany))).scalar()
    if count and count > 0 and not force:
        logger.info("vc_companies 이미 %d건 존재, 스킵 (--force 로 재삽입)", count)
        return count

    import openpyxl

    if not EXCEL_PATH.exists():
        logger.error("파일 없음: %s", EXCEL_PATH)
        return 0

    logger.info("Excel 로딩 중 (read_only): %s", EXCEL_PATH)
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True, read_only=True, keep_links=False)
    try:
        ws = wb[SHEET_NAME]
    except KeyError:
        logger.error("시트 '%s' 없음. 유효 시트: %s", SHEET_NAME, wb.sheetnames)
        wb.close()
        return 0

    # 컬럼 매핑 (0-indexed):
    # A=회사이름, B=영문명, C=공시회사명, D=종목코드, E=대표자명
    # F=법인구분, G=법인등록번호, H=사업자등록번호, I=주소, J=홈페이지
    # K=IR홈페이지, L=업종명, M=IO부문코드, N=IO부문명, O=설립일, P=결산월

    batch: list[dict] = []
    total = 0
    skipped = 0

    for _i, row in enumerate(ws.iter_rows(min_row=DATA_START_ROW, values_only=True), start=1):
        company_name = _safe_str(row[0])
        industry_name = _safe_str(row[11])

        if not company_name or not industry_name:
            skipped += 1
            continue

        batch.append(
            {
                "company_name": company_name,
                "english_name": _safe_str(row[1]),
                "disclosure_name": _safe_str(row[2]),
                "listing_code": _safe_str(row[3], 20),
                "ceo_name": _safe_str(row[4], 200),
                "corp_type": _safe_str(row[5], 50),
                "corp_reg_no": _safe_str(row[6], 20),
                "biz_reg_no": _safe_str(row[7], 20),
                "address": _safe_str(row[8], 1000),
                "homepage": _safe_str(row[9]),
                "industry_name": industry_name,
                "io_sector_code": _safe_int(row[12]),
                "io_sector_name": _safe_str(row[13], 200),
                "founded_date": _safe_str(row[14], 10),
                "fiscal_month": _safe_str(row[15], 5),
            }
        )

        if len(batch) >= CHUNK_SIZE:
            await session.execute(insert(VcCompany), batch)
            await session.commit()
            total += len(batch)
            batch.clear()
            if total % 10000 == 0:
                logger.info("vc_companies: %d건 삽입 완료...", total)

    # 잔여 배치
    if batch:
        await session.execute(insert(VcCompany), batch)
        total += len(batch)

    await session.commit()
    wb.close()
    logger.info("vc_companies 총 %d건 삽입 완료 (스킵: %d건)", total, skipped)
    return total


async def main(force: bool = False) -> None:
    if force and sys.stdin.isatty():
        confirm = input("--force: 기존 vc_companies 데이터를 전부 삭제합니다. 계속? [y/N]: ")
        if confirm.strip().lower() != "y":
            logger.info("취소됨")
            return
    engine, session_factory = await _get_engine_and_session()
    async with session_factory() as session:
        await seed_vc_companies(session, force)
    await engine.dispose()
    logger.info("VC 기업개황 시딩 완료!")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="VC 기업개황 시딩 (MA_ValueChain_v7.xlsx)")
    parser.add_argument("--force", action="store_true", help="기존 데이터 삭제 후 재삽입")
    args = parser.parse_args()
    asyncio.run(main(force=args.force))
