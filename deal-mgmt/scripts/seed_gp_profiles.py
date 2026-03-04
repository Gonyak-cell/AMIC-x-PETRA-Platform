"""GP 프로필 시딩 — MA_GP_v3.xlsx Sheet 1 + Sheet 3 → gp_profiles 테이블.

사용법:
    cd deal-mgmt
    python -m scripts.seed_gp_profiles [--force]
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import re
import sys
import uuid
from decimal import Decimal, InvalidOperation
from pathlib import Path

from sqlalchemy import delete, func, insert, select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.orm import sessionmaker

from app.services.fi_mapping_service import normalize_gp_name

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)

EXCEL_PATH = Path(__file__).resolve().parent.parent / "data" / "MA_GP_v3.xlsx"
SHEET1_NAME = "GP별 관심 FI List"
SHEET3_NAME = "GP 연도별 활동 매트릭스"

# Sheet 1: 시작 행 (1-indexed)
DATA_START_ROW = 6  # Row 6 = 데이터 시작
# Sheet 3: 시작 행 (1-indexed)
SHEET3_HEADER_ROW = 3  # Row 3 = 헤더 (No., GP명, 최소약정, 최대약정, 2010년, ...)
SHEET3_DATA_START_ROW = 4  # Row 4 = 데이터 시작


def _safe_decimal(val: object) -> Decimal | None:
    """셀 값을 Decimal로 변환. None/빈값/비숫자 → None, 0 → None (유효값 아님)."""
    if val is None:
        return None
    try:
        d = Decimal(str(val))
        return d if d != 0 else None
    except (InvalidOperation, ValueError):
        return None


def _safe_int(val: object) -> int | None:
    """셀 값을 int로 변환. '-', None, 비숫자 → None."""
    if val is None:
        return None
    s = str(val).strip()
    if s in ("", "-", "–", "—"):
        return None
    try:
        return int(float(s))
    except (ValueError, TypeError):
        return None


def parse_portfolio(raw: str | None) -> tuple[list[str], list[str]]:
    """N열 원문을 투자분야 키워드 + 투자기업명으로 파싱.

    형식: "소비재/유통, 헬스케어 | 투자기업: 코웨이, 홈플러스"
    """
    if not raw or not raw.strip():
        return [], []

    sectors: list[str] = []
    companies: list[str] = []

    # '|' 또는 '│' 분리
    parts = re.split(r"[|│]", raw, maxsplit=1)
    sector_part = parts[0].strip()
    company_part = parts[1].strip() if len(parts) > 1 else ""

    # 투자분야: ',' 와 '/' 로 분리
    if sector_part:
        for seg in re.split(r"[,，]", sector_part):
            for sub in seg.split("/"):
                kw = sub.strip()
                if kw:
                    sectors.append(kw)

    # 투자기업: "투자기업:" 이후
    if company_part:
        company_text = re.sub(r"^투자기업\s*[:：]\s*", "", company_part).strip()
        if company_text:
            # "외 N개" 부분 제거
            company_text = re.sub(r"\s*외\s*\d+개?\s*$", "", company_text).strip()
            for c in re.split(r"[,，]", company_text):
                name = c.strip()
                if name:
                    companies.append(name)

    return sectors, companies


def _parse_sheet3_yearly(wb: object) -> dict[str, dict[str, int]]:
    """Sheet 3 '연도별 활동 매트릭스' → {GP명: {연도: 건수}} 딕셔너리.

    헤더에서 연도 열을 동적으로 파싱하여 2012년 누락 등에도 대응한다.
    """
    try:
        ws = wb[SHEET3_NAME]  # type: ignore[index]
    except KeyError:
        logger.warning("시트 '%s' 없음 — yearly_pef_counts 스킵", SHEET3_NAME)
        return {}

    # 헤더 파싱: "2010년", "2011년" 등에서 연도 추출
    year_columns: list[tuple[int, str]] = []  # (열 인덱스, 연도 문자열)
    for row in ws.iter_rows(min_row=SHEET3_HEADER_ROW, max_row=SHEET3_HEADER_ROW, values_only=True):
        for col_idx, val in enumerate(row):
            if val and isinstance(val, str) and val.endswith("년"):
                year_str = val.replace("년", "").strip()
                if year_str.isdigit():
                    year_columns.append((col_idx, year_str))
        break

    if not year_columns:
        logger.warning("시트3 헤더에서 연도 열을 찾지 못함")
        return {}

    logger.info("시트3 연도 열 %d개 감지: %s", len(year_columns), [y for _, y in year_columns])

    result: dict[str, dict[str, int]] = {}
    for row in ws.iter_rows(min_row=SHEET3_DATA_START_ROW, values_only=True):
        gp_name = row[1]  # B열
        if not gp_name or not str(gp_name).strip():
            continue
        gp_name = str(gp_name).strip()
        yearly: dict[str, int] = {}
        for col_idx, year in year_columns:
            if col_idx < len(row):
                count = _safe_int(row[col_idx])
                if count is not None and count > 0:
                    yearly[year] = count
        if yearly:
            result[gp_name] = yearly

    logger.info("시트3 GP %d개 연도별 데이터 파싱 완료", len(result))
    return result


async def _get_engine_and_session() -> tuple:
    from app.core.config import settings

    engine = create_async_engine(settings.DATABASE_URL)
    session_factory = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    return engine, session_factory


async def seed_gp_profiles(session: AsyncSession, force: bool = False) -> int:
    """MA_GP_v3.xlsx Sheet 1 파싱 → gp_profiles 테이블 벌크 삽입."""
    from app.models.gp_profile import GpProfile

    if force:
        await session.execute(delete(GpProfile))
        await session.commit()
        logger.info("gp_profiles 기존 데이터 삭제 완료")

    count = (await session.execute(select(func.count()).select_from(GpProfile))).scalar()
    if count and count > 0 and not force:
        logger.info("gp_profiles 이미 %d건 존재, 스킵 (--force 로 재삽입)", count)
        return count

    import openpyxl

    if not EXCEL_PATH.exists():
        logger.error("파일 없음: %s", EXCEL_PATH)
        return 0

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True, read_only=True, keep_links=False)
    try:
        ws = wb[SHEET1_NAME]
    except KeyError:
        logger.error("시트 '%s' 없음. 유효 시트: %s", SHEET1_NAME, wb.sheetnames)
        wb.close()
        return 0

    rows_data: list[dict] = []
    for _i, row in enumerate(ws.iter_rows(min_row=DATA_START_ROW, values_only=True), start=1):
        if len(row) < 14:
            continue
        # A=No, B=GP명, C~M=숫자, N=포트폴리오
        gp_name = row[1]  # B열
        if not gp_name or not str(gp_name).strip():
            continue

        gp_name = str(gp_name).strip()
        portfolio_raw = str(row[13]).strip() if row[13] else None
        sectors, companies = parse_portfolio(portfolio_raw)

        rows_data.append(
            {
                "id": uuid.uuid4(),
                "raw_name": gp_name,
                "normalized_name": normalize_gp_name(gp_name),
                "min_committed_capital": _safe_decimal(row[2]),  # C열
                "min_threshold": _safe_decimal(row[3]),  # D열
                "max_committed_capital": _safe_decimal(row[4]),  # E열
                "total_pef_count": _safe_int(row[5]),  # F열
                "recent_pef_count": _safe_int(row[6]),  # G열
                "recent_committed_sum": _safe_decimal(row[7]),  # H열
                "total_committed_sum": _safe_decimal(row[8]),  # I열
                "pef_count_2021": _safe_int(row[9]),  # J열
                "pef_count_2022": _safe_int(row[10]),  # K열
                "pef_count_2023": _safe_int(row[11]),  # L열
                "pef_count_2024": _safe_int(row[12]),  # M열
                "portfolio_sectors": sectors if sectors else None,
                "portfolio_companies": companies if companies else None,
                "portfolio_raw": portfolio_raw,
            }
        )

    # Sheet 3: 연도별 활동 매트릭스 병합
    yearly_map = _parse_sheet3_yearly(wb)
    for row_dict in rows_data:
        gp_name = row_dict["raw_name"]
        row_dict["yearly_pef_counts"] = yearly_map.get(gp_name)

    wb.close()

    if not rows_data:
        logger.warning("파싱된 GP 데이터 없음")
        return 0

    await session.execute(insert(GpProfile), rows_data)
    await session.commit()
    logger.info("gp_profiles %d건 삽입 완료", len(rows_data))

    # PefFundRegistry GP명 교차 매칭 통계
    await _log_gp_match_stats(session, rows_data)

    return len(rows_data)


async def _log_gp_match_stats(session: AsyncSession, gp_rows: list[dict]) -> None:
    """GpProfile ↔ PefFundRegistry GP명 매칭 통계 출력."""
    from app.models.pef_fund_registry import PefFundRegistry

    pef_count = (await session.execute(select(func.count()).select_from(PefFundRegistry))).scalar()
    if not pef_count:
        logger.info("PEF 레지스트리 비어있음 — GP 매칭 통계 스킵")
        return

    # PefFundRegistry에서 고유 GP명 추출
    result = await session.execute(select(PefFundRegistry.gp1, PefFundRegistry.gp2, PefFundRegistry.gp3))
    pef_gp_names: set[str] = set()
    for r in result:
        for gp in [r[0], r[1], r[2]]:
            if gp and gp.strip():
                pef_gp_names.add(normalize_gp_name(gp.strip()))

    gp_normalized = {row["normalized_name"] for row in gp_rows}
    matched = gp_normalized & pef_gp_names
    unmatched_gp = gp_normalized - pef_gp_names
    unmatched_pef = pef_gp_names - gp_normalized

    logger.info(
        "GP 매칭 통계: GP %d개 중 %d개 매칭 (%.1f%%), PEF GP %d개 중 %d개 미매칭",
        len(gp_normalized),
        len(matched),
        len(matched) / len(gp_normalized) * 100 if gp_normalized else 0,
        len(pef_gp_names),
        len(unmatched_pef),
    )
    if unmatched_gp:
        logger.info("GP 프로필에만 존재 (PEF 미매칭, 상위 10개): %s", sorted(unmatched_gp)[:10])
    if unmatched_pef:
        logger.info("PEF에만 존재 (GP 프로필 미매칭, 상위 10개): %s", sorted(unmatched_pef)[:10])


async def main(force: bool = False) -> None:
    if force and sys.stdin.isatty():
        confirm = input("--force: 기존 gp_profiles 데이터를 전부 삭제합니다. 계속? [y/N]: ")
        if confirm.strip().lower() != "y":
            logger.info("취소됨")
            return
    engine, session_factory = await _get_engine_and_session()
    async with session_factory() as session:
        await seed_gp_profiles(session, force)
    await engine.dispose()
    logger.info("GP 프로필 시딩 완료!")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="GP 프로필 시딩 (MA_GP_v3.xlsx)")
    parser.add_argument("--force", action="store_true", help="기존 데이터 삭제 후 재삽입")
    args = parser.parse_args()
    asyncio.run(main(force=args.force))
