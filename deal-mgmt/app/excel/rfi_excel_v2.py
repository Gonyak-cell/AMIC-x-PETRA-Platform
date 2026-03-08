"""RFI V2 동적 Excel Export — Flatten History 방식.

열 구조:
  A (숨김): item_id (UUID)
  B (숨김): version (낙관적 락)
  C (숨김): report_section_tag
  D (고정): Category
  E (고정): Priority
  F (고정): RFI Item (질문 원문)
  G (고정): Target Doc
  H (고정): History — 1~(n-1)차 이력 줄바꿈 누적
  I (편집): [n차 답변 입력란]
  J (편집): [증빙 파일명/인덱스 기재란]
  K (조건부): Internal Memo — role_type=TARGET이면 제외
"""

from __future__ import annotations

import io
import logging
import uuid

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.enums import RFIAuthorRole
from app.models.rfi_item_v2 import RFIItemV2
from app.models.rfi_thread import RFIThread

logger = logging.getLogger(__name__)


def _build_history_text(threads: list[RFIThread]) -> str:
    """스레드 이력을 줄바꿈으로 누적한 텍스트 생성."""
    lines: list[str] = []
    for t in threads:
        if not t.is_published:
            continue
        role_label = "TARGET" if t.author_role == RFIAuthorRole.TARGET else "ADVISOR"
        lines.append(f"[{t.round_num}차 {role_label}] {t.content_text}")
    return "\n".join(lines)


async def export_rfi_excel(
    db: AsyncSession,
    txn_id: uuid.UUID,
    *,
    is_advisor: bool = True,
) -> io.BytesIO:
    """RFI V2 동적 Excel Export — Flatten History 방식."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl이 설치되지 않았습니다: pip install openpyxl") from exc

    # 데이터 조회
    q = (
        select(RFIItemV2)
        .options(selectinload(RFIItemV2.threads))
        .where(
            RFIItemV2.transaction_id == txn_id,
            RFIItemV2.is_deleted.is_(False),
        )
        .order_by(RFIItemV2.created_at)
    )
    result = await db.execute(q)
    items = list(result.scalars().all())

    wb = Workbook()
    ws = wb.active
    ws.title = "RFI"

    # ── 스타일 정의 ──
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, size=10, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    locked = Protection(locked=True)
    unlocked = Protection(locked=False)

    priority_fills = {
        "HIGH": PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid"),
        "MEDIUM": PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid"),
        "LOW": PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid"),
    }

    status_fills = {
        "OPEN": PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid"),
        "ANSWERED": PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid"),
        "CLARIFICATION_NEEDED": PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid"),
        "CLOSED": PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid"),
    }

    # ── 헤더 ──
    headers = [
        "item_id",  # A — 숨김
        "version",  # B — 숨김
        "report_section",  # C — 숨김
        "Category",  # D
        "Priority",  # E
        "RFI Item",  # F
        "Target Doc",  # G
        "History",  # H — 과거 이력
        "답변 입력",  # I — 편집 가능
        "증빙 파일명",  # J — 편집 가능
    ]
    if is_advisor:
        headers.append("Internal Memo")  # K — 자문사 전용

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # ── 데이터 행 ──
    for row_idx, item in enumerate(items, 2):
        sorted_threads = sorted(
            [t for t in item.threads if t.is_published],
            key=lambda t: t.round_num,
        )
        history_text = _build_history_text(sorted_threads)

        row_data = [
            str(item.id),  # A: item_id
            item.version,  # B: version
            item.report_section_tag or "",  # C: report_section_tag
            item.category.value,  # D: Category
            item.priority.value,  # E: Priority
            item.question_text,  # F: RFI Item
            item.target_doc or "",  # G: Target Doc
            history_text,  # H: History
            "",  # I: 답변 입력 (빈 칸)
            "",  # J: 증빙 파일명 (빈 칸)
        ]
        if is_advisor:
            row_data.append(item.internal_memo or "")  # K

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = wrap_alignment

            # 고정/편집 보호
            if col_idx <= 8:  # A~H: 고정
                cell.protection = locked
            else:  # I, J, (K): 편집 가능
                cell.protection = unlocked

        # 우선순위 색상
        priority_cell = ws.cell(row=row_idx, column=5)
        fill = priority_fills.get(item.priority.value)
        if fill:
            priority_cell.fill = fill

        # 상태 표시 (Category 셀 배경으로 상태 힌트)
        status_fill = status_fills.get(item.current_status.value)
        if status_fill:
            ws.cell(row=row_idx, column=4).fill = status_fill

    # ── 열 너비 ──
    col_widths = {
        1: 5,  # A: item_id (숨김)
        2: 5,  # B: version (숨김)
        3: 5,  # C: report_section (숨김)
        4: 15,  # D: Category
        5: 10,  # E: Priority
        6: 50,  # F: RFI Item
        7: 12,  # G: Target Doc
        8: 50,  # H: History
        9: 40,  # I: 답변 입력
        10: 25,  # J: 증빙 파일명
    }
    if is_advisor:
        col_widths[11] = 30  # K: Internal Memo

    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── 숨김 열 ──
    ws.column_dimensions["A"].hidden = True
    ws.column_dimensions["B"].hidden = True
    ws.column_dimensions["C"].hidden = True

    # ── 시트 보호 (편집 가능 열만 해제) ──
    ws.protection.sheet = True
    ws.protection.password = ""  # 비밀번호 없이 보호 (사용자 편의)
    ws.protection.enable()

    # ── 필터 자동 적용 ──
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col}{len(items) + 1}"

    # ── BytesIO 반환 ──
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
