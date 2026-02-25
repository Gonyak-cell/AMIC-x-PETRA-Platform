"""IB 표준 Excel 스타일 — 색상, 폰트, 숫자 형식, 셀 스타일."""

from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side

# ── 색상 팔레트 (AMIC forest 계열) ────────────────────────────────────────

FOREST_DARK = "1B4332"
FOREST_MID = "2D6A4F"
FOREST_LIGHT = "40916C"
FOREST_PALE = "95D5B2"
FOREST_BG = "D8F3DC"

WHITE = "FFFFFF"
BLACK = "000000"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "D9D9D9"
INPUT_YELLOW = "FFFFCC"
INPUT_YELLOW_BORDER = "FFD966"
ERROR_RED = "FF4444"
LINK_BLUE = "1155CC"

# ── 폰트 ─────────────────────────────────────────────────────────────────

FONT_HEADER = Font(name="Pretendard", size=11, bold=True, color=WHITE)
FONT_SUBHEADER = Font(name="Pretendard", size=10, bold=True, color=FOREST_DARK)
FONT_NORMAL = Font(name="Pretendard", size=10, color=BLACK)
FONT_INPUT = Font(name="Pretendard", size=10, color="0000CC")  # 파란색 = 입력값
FONT_FORMULA = Font(name="Pretendard", size=10, color=BLACK)
FONT_TITLE = Font(name="SUITE", size=16, bold=True, color=FOREST_DARK)
FONT_SUBTITLE = Font(name="SUITE", size=12, bold=True, color=FOREST_MID)
FONT_SMALL = Font(name="Pretendard", size=9, color="666666")

# ── 채우기(Fill) ──────────────────────────────────────────────────────────

FILL_HEADER = PatternFill(start_color=FOREST_DARK, end_color=FOREST_DARK, fill_type="solid")
FILL_SUBHEADER = PatternFill(start_color=FOREST_PALE, end_color=FOREST_PALE, fill_type="solid")
FILL_INPUT = PatternFill(start_color=INPUT_YELLOW, end_color=INPUT_YELLOW, fill_type="solid")
FILL_FORMULA = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
FILL_LIGHT_BG = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
FILL_FOREST_BG = PatternFill(start_color=FOREST_BG, end_color=FOREST_BG, fill_type="solid")

# ── 테두리 ────────────────────────────────────────────────────────────────

THIN_BORDER = Border(
    left=Side(style="thin", color=MID_GRAY),
    right=Side(style="thin", color=MID_GRAY),
    top=Side(style="thin", color=MID_GRAY),
    bottom=Side(style="thin", color=MID_GRAY),
)
BOTTOM_BORDER = Border(bottom=Side(style="thin", color=BLACK))
THICK_BOTTOM = Border(bottom=Side(style="medium", color=BLACK))
DOUBLE_BOTTOM = Border(bottom=Side(style="double", color=BLACK))
INPUT_BORDER = Border(
    left=Side(style="thin", color=INPUT_YELLOW_BORDER),
    right=Side(style="thin", color=INPUT_YELLOW_BORDER),
    top=Side(style="thin", color=INPUT_YELLOW_BORDER),
    bottom=Side(style="thin", color=INPUT_YELLOW_BORDER),
)

# ── 정렬 ─────────────────────────────────────────────────────────────────

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)

# ── 숫자 형식 ─────────────────────────────────────────────────────────────

NUM_FMT_KRW = '#,##0'
NUM_FMT_KRW_M = '#,##0,,"백만"'
NUM_FMT_USD = '$#,##0'
NUM_FMT_PCT = '0.0%'
NUM_FMT_PCT_2 = '0.00%'
NUM_FMT_MULTIPLE = '0.0x'
NUM_FMT_DECIMAL = '#,##0.0'
NUM_FMT_INTEGER = '#,##0'
NUM_FMT_DAYS = '#,##0"일"'
NUM_FMT_YEAR = '0"Y"'

# ── 셀 스타일 적용 헬퍼 ──────────────────────────────────────────────────


def apply_header(cell, text=None):
    """헤더 셀 스타일 (진한 초록 배경 + 흰 글자)."""
    if text is not None:
        cell.value = text
    cell.font = FONT_HEADER
    cell.fill = FILL_HEADER
    cell.alignment = ALIGN_CENTER
    cell.border = THIN_BORDER


def apply_subheader(cell, text=None):
    """서브헤더 셀 스타일 (연한 초록 배경)."""
    if text is not None:
        cell.value = text
    cell.font = FONT_SUBHEADER
    cell.fill = FILL_SUBHEADER
    cell.alignment = ALIGN_LEFT
    cell.border = THIN_BORDER


def apply_input(cell, value=None, num_format=None):
    """입력 셀 스타일 (노란 배경 + 파란 글자)."""
    if value is not None:
        cell.value = value
    cell.font = FONT_INPUT
    cell.fill = FILL_INPUT
    cell.border = INPUT_BORDER
    cell.alignment = ALIGN_RIGHT
    if num_format:
        cell.number_format = num_format


def apply_formula(cell, formula=None, num_format=None):
    """수식 셀 스타일 (흰 배경 + 검정 글자)."""
    if formula is not None:
        cell.value = formula
    cell.font = FONT_FORMULA
    cell.fill = FILL_FORMULA
    cell.border = THIN_BORDER
    cell.alignment = ALIGN_RIGHT
    if num_format:
        cell.number_format = num_format


def apply_label(cell, text=None, indent=0):
    """라벨(행 이름) 셀 스타일."""
    if text is not None:
        cell.value = text
    cell.font = FONT_NORMAL
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=indent)
    cell.border = THIN_BORDER


def apply_total(cell, formula=None, num_format=None):
    """합계 행 스타일 (볼드 + 이중선)."""
    if formula is not None:
        cell.value = formula
    cell.font = Font(name="Pretendard", size=10, bold=True, color=BLACK)
    cell.border = DOUBLE_BOTTOM
    cell.alignment = ALIGN_RIGHT
    if num_format:
        cell.number_format = num_format


def apply_section_title(cell, text):
    """섹션 제목 스타일."""
    cell.value = text
    cell.font = FONT_SUBTITLE
    cell.border = THICK_BOTTOM
    cell.alignment = ALIGN_LEFT


def apply_title(cell, text):
    """워크시트 제목 스타일."""
    cell.value = text
    cell.font = FONT_TITLE
    cell.alignment = ALIGN_LEFT


def set_column_widths(ws, widths: dict[str, float]):
    """열 너비를 설정한다. widths = {'A': 30, 'B': 15, ...}"""
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width


def create_named_styles(wb):
    """워크북에 Named Style을 등록한다 (셀 스타일 재사용)."""
    styles = {
        "fm_header": NamedStyle(
            name="fm_header", font=FONT_HEADER, fill=FILL_HEADER,
            alignment=ALIGN_CENTER, border=THIN_BORDER,
        ),
        "fm_input": NamedStyle(
            name="fm_input", font=FONT_INPUT, fill=FILL_INPUT,
            border=INPUT_BORDER, alignment=ALIGN_RIGHT,
        ),
        "fm_formula": NamedStyle(
            name="fm_formula", font=FONT_FORMULA, fill=FILL_FORMULA,
            border=THIN_BORDER, alignment=ALIGN_RIGHT,
        ),
    }
    for name, style in styles.items():
        if name not in wb.named_styles:
            wb.add_named_style(style)
