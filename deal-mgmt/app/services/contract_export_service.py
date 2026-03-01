"""계약서 DOCX 내보내기 서비스 — HTML → python-docx 변환."""

from __future__ import annotations

import asyncio
import html as html_mod
import io
import logging
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

logger = logging.getLogger(__name__)

# ── 법률 문서 스타일 ────────────────────────────────────────────────────────

_FONT_NAME = "바탕"
_FONT_SIZE_BODY = Pt(11)
_FONT_SIZE_TITLE = Pt(16)
_FONT_SIZE_HEADING = Pt(13)
_LINE_SPACING = 1.5
_MARGINS = {"top": Cm(2.5), "bottom": Cm(2.5), "left": Cm(3.0), "right": Cm(2.5)}


def _setup_document(doc: Document) -> None:
    """문서 여백 및 기본 스타일을 설정한다."""
    for section in doc.sections:
        section.top_margin = _MARGINS["top"]
        section.bottom_margin = _MARGINS["bottom"]
        section.left_margin = _MARGINS["left"]
        section.right_margin = _MARGINS["right"]

    # Normal 스타일 설정
    style = doc.styles["Normal"]
    font = style.font
    font.name = _FONT_NAME
    font.size = _FONT_SIZE_BODY
    pf = style.paragraph_format
    pf.line_spacing = _LINE_SPACING
    pf.space_after = Pt(6)


def _add_title(doc: Document, text: str) -> None:
    """문서 제목을 추가한다."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = _FONT_SIZE_TITLE
    run.font.name = _FONT_NAME
    p.paragraph_format.space_after = Pt(24)


def _add_heading(doc: Document, text: str) -> None:
    """조항 제목을 추가한다."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = _FONT_SIZE_HEADING
    run.font.name = _FONT_NAME
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)


def _add_paragraph(doc: Document, text: str) -> None:
    """본문 단락을 추가한다."""
    if not text.strip():
        return
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1.0)
    run = p.add_run(text)
    run.font.name = _FONT_NAME
    run.font.size = _FONT_SIZE_BODY


def _add_list_item(doc: Document, text: str, level: int = 0) -> None:
    """번호 목록 항목을 추가한다."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0 + level * 0.5)
    run = p.add_run(text)
    run.font.name = _FONT_NAME
    run.font.size = _FONT_SIZE_BODY


# ── HTML 태그 제거 유틸리티 ─────────────────────────────────────────────────

_TAG_RE = re.compile(r"<[^>]+>")
_WS_RE = re.compile(r"\s+")


def _strip_tags(raw_html: str) -> str:
    """HTML 태그를 제거하고 텍스트만 반환한다."""
    text = _TAG_RE.sub(" ", raw_html)
    text = html_mod.unescape(text)
    text = _WS_RE.sub(" ", text)
    return text.strip()


_TAG_PATTERN_CACHE: dict[str, re.Pattern[str]] = {}


def _extract_tag_contents(html: str, tag: str) -> list[str]:
    """특정 태그의 내용을 추출한다."""
    pattern = _TAG_PATTERN_CACHE.get(tag)
    if pattern is None:
        pattern = re.compile(rf"<{tag}[^>]*>(.*?)</{tag}>", re.DOTALL | re.IGNORECASE)
        _TAG_PATTERN_CACHE[tag] = pattern
    return [_strip_tags(m.group(1)) for m in pattern.finditer(html)]


# ── 메인 변환 ───────────────────────────────────────────────────────────────

_SECTION_RE = re.compile(r"<section[^>]*>(.*?)</section>", re.DOTALL | re.IGNORECASE)

_SEQUENTIAL_TAG_RE = re.compile(
    r"<(h[23]|p|li)\b[^>]*>(.*?)</\1>",
    re.DOTALL | re.IGNORECASE,
)


def _build_docx(html: str, title: str) -> bytes:
    """HTML → DOCX 변환 (동기, 별도 스레드에서 실행)."""
    doc = Document()
    _setup_document(doc)

    # 제목 추출
    h1_texts = _extract_tag_contents(html, "h1")
    doc_title = h1_texts[0] if h1_texts else title
    _add_title(doc, doc_title)

    # 섹션 단위로 처리
    sections = _SECTION_RE.findall(html)

    if not sections:
        # 섹션 태그가 없으면 전체를 단일 블록으로 처리
        _process_block(doc, html)
    else:
        for section_html in sections:
            _process_block(doc, section_html)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _process_block(doc: Document, html_block: str) -> None:
    """HTML 블록을 출현 순서대로 DOCX 요소로 변환한다."""
    li_counter = 0

    for m in _SEQUENTIAL_TAG_RE.finditer(html_block):
        tag = m.group(1).lower()
        text = _strip_tags(m.group(2))

        if not text.strip():
            continue

        if tag == "h2":
            _add_heading(doc, text)
            li_counter = 0
        elif tag == "h3":
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.bold = True
            run.font.name = _FONT_NAME
            run.font.size = _FONT_SIZE_BODY
            li_counter = 0
        elif tag == "p":
            _add_paragraph(doc, text)
            li_counter = 0
        elif tag == "li":
            li_counter += 1
            _add_list_item(doc, f"{li_counter}. {text}")


_DOCX_TIMEOUT_SECONDS = 60.0
# NOTE: asyncio.Semaphore는 프로세스 단위 — Gunicorn 멀티 워커 환경에서는
# 워커 수 × 3 동시 변환이 가능. 현재 2-worker 기준 최대 6 동시 변환.
_DOCX_SEMAPHORE = asyncio.Semaphore(3)


async def html_to_docx(html: str, title: str) -> bytes:
    """HTML 계약서를 법률 문서 형식의 DOCX로 변환한다.

    Args:
        html: 생성된 계약서 HTML
        title: 문서 제목

    Returns:
        DOCX 파일 바이트

    Raises:
        TimeoutError: 변환이 60초를 초과한 경우
    """
    import time

    html_size = len(html)
    logger.info("DOCX 변환 시작: title=%s, html_size=%d bytes", title, html_size)
    async with _DOCX_SEMAPHORE:
        t0 = time.monotonic()
        result = await asyncio.wait_for(
            asyncio.to_thread(_build_docx, html, title),
            timeout=_DOCX_TIMEOUT_SECONDS,
        )
        elapsed = time.monotonic() - t0
        logger.info(
            "DOCX 변환 완료: title=%s, html_size=%d, docx_size=%d bytes, elapsed=%.1fs",
            title,
            html_size,
            len(result),
            elapsed,
        )
        return result
