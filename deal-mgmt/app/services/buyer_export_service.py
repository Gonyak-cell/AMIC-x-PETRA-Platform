"""매수자 Long-List Excel 내보내기 서비스."""

from __future__ import annotations

import uuid
from decimal import ROUND_HALF_UP, Decimal
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from app.models.buyer_candidate import BuyerCandidate

_TIER_LABELS: dict[str, str] = {
    "TIER_1": "Tier 1",
    "TIER_2": "Tier 2",
    "TIER_3": "Tier 3",
    "NOT_TARGET": "대상 아님",
}

_STATUS_LABELS: dict[str, str] = {
    "IDENTIFIED": "식별",
    "CONTACTED": "연락 완료",
    "NDA_SENT": "NDA 발송",
    "NDA_SIGNED": "NDA 체결",
    "CIM_SENT": "CIM 발송",
    "INTEREST_CONFIRMED": "관심 확인",
    "IOI_RECEIVED": "IOI 접수",
    "IOI_ACCEPTED": "IOI 승인",
    "DD_GRANTED": "DD 권한 부여",
    "DD_IN_PROGRESS": "DD 진행 중",
    "LOI_RECEIVED": "LOI 접수",
    "LOI_ACCEPTED": "LOI 승인",
    "SELECTED": "최종 선정",
    "REJECTED": "거절",
}

_TYPE_LABELS: dict[str, str] = {
    "STRATEGIC": "전략적 투자자",
    "FINANCIAL_SPONSOR": "재무적 투자자",
    "FAMILY_OFFICE": "패밀리오피스",
    "INDIVIDUAL": "개인",
    "OTHER": "기타",
}

_ROLE_LABELS: dict[str, str] = {
    "SOLE_BUYER": "단독 매수자",
    "CONSORTIUM_LEAD": "컨소시엄 리드",
    "CO_INVESTOR": "공동투자자",
    "FINANCING_PROVIDER": "파이낸싱 제공자",
}

_STAGE_LABELS: dict[str, str] = {
    "IDENTIFIED": "발굴",
    "EMAIL_SENT": "이메일 발송",
    "PHONE_CALL": "전화 접촉",
    "ADVISOR_MEETING": "어드바이저 미팅",
    "NDA_SIGNED": "NDA 체결",
    "TARGET_MEETING": "대상회사 미팅",
}

_HEADERS = [
    "No.",
    "회사명",
    "Tier",
    "유형",
    "역할",
    "파이프라인 상태",
    "담당자",
    "이메일",
    "전화번호",
    "IOI 금액",
    "IOI 일자",
    "LOI 금액",
    "LOI 일자",
    "최종 제안 금액",
    "DART 코드",
    "컨소시엄 관계",
    "최근 활동 단계",
    "최근 활동 일자",
    "비고",
]

_HEADER_FILL = PatternFill(start_color="1B3A5C", end_color="1B3A5C", fill_type="solid")
_HEADER_FONT = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=10)
_BODY_FONT = Font(name="맑은 고딕", size=10)
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
_ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def build_buyer_excel(
    buyers: list[BuyerCandidate],
    deal_name: str,
    *,
    marketing_latest: dict[uuid.UUID, tuple[str | None, str | None]] | None = None,
    consortium_map: dict[uuid.UUID, list[str]] | None = None,
) -> Workbook:
    """매수자 Long-List를 Excel Workbook으로 생성한다.

    Args:
        buyers: 매수자 목록
        deal_name: 딜 코드명/이름
        marketing_latest: buyer_id → (최신 stage, 최신 log_date) 매핑
        consortium_map: buyer_id → 관련 회사명 리스트
    """
    wb = Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet()
    ws.title = "Long-List"

    mkt = marketing_latest or {}
    cons = consortium_map or {}

    # 헤더 행
    for col_idx, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _ALIGN_CENTER

    # 데이터 행
    for row_idx, buyer in enumerate(buyers, start=2):
        tier_val = buyer.tier.value if buyer.tier else ""
        status_val = buyer.status.value if buyer.status else ""
        type_val = buyer.buyer_type.value if buyer.buyer_type else ""
        role_val = buyer.deal_role.value if buyer.deal_role else ""

        # 마케팅/컨소시엄 정보
        latest_stage, latest_date = mkt.get(buyer.id, (None, None))
        consortium_names = cons.get(buyer.id, [])

        row_data = [
            row_idx - 1,
            buyer.company_name,
            _TIER_LABELS.get(tier_val, tier_val),
            _TYPE_LABELS.get(type_val, type_val),
            _ROLE_LABELS.get(role_val, role_val),
            _STATUS_LABELS.get(status_val, status_val),
            buyer.contact_name or "",
            buyer.contact_email or "",
            buyer.contact_phone or "",
            int(buyer.ioi_value.quantize(Decimal("1"), rounding=ROUND_HALF_UP)) if buyer.ioi_value else "",
            buyer.ioi_date or "",
            int(buyer.loi_value.quantize(Decimal("1"), rounding=ROUND_HALF_UP)) if buyer.loi_value else "",
            buyer.loi_date or "",
            int(buyer.final_offer_value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
            if buyer.final_offer_value
            else "",
            buyer.corp_code or "",
            ", ".join(consortium_names) if consortium_names else "",
            _STAGE_LABELS.get(latest_stage or "", latest_stage or ""),
            latest_date or "",
            buyer.notes or "",
        ]
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = _BODY_FONT
            if col_idx in (1, 3, 4, 5, 6):
                cell.alignment = _ALIGN_CENTER
            elif col_idx in (10, 12, 14):
                cell.alignment = _ALIGN_RIGHT
                if isinstance(value, (int, float)):
                    cell.number_format = "#,##0"
            else:
                cell.alignment = _ALIGN_LEFT

    # 컬럼 너비 조정 (19열)
    col_widths = [5, 25, 10, 15, 15, 15, 12, 25, 15, 15, 12, 15, 12, 15, 10, 25, 15, 12, 30]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # 자동 필터
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_HEADERS))}1"

    # 시트 제목 메타
    ws.sheet_properties.tabColor = "1B3A5C"

    return wb
