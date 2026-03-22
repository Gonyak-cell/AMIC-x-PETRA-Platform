"""법무법인 표준 양식 빈 템플릿 동적 생성 (Step 1).

원본 법률실사보고서 .docx를 복제하여 플레이스홀더로 교체한 빈 템플릿을 생성한다.
python-docx + lxml 기반 XML 직접 조작으로 서식을 100% 보존.

사용법:
    gen = LawFirmTemplateGenerator(source_template_path)
    output = gen.generate_blank_template(output_path, project_code="Alpha", ...)
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from docx import Document

logger = logging.getLogger(__name__)

# ─── 스타일 식별 상수 ────────────────────────────────────────────────────────

WML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_SPACE_NS = "http://www.w3.org/XML/1998/namespace"

# 섹션 헤더 바 배경색
# 첨부 blank template는 앞 챕터와 뒤 챕터의 바 색상을 다르게 사용한다.
SECTION_BAR_FILLS = {"26382A", "385623", "1D2B20"}

# Recommendation 박스 배경색 (연한 녹색)
RECOMMENDATION_FILL = "E2EFD9"

# 데이터 표 헤더 배경색
DATA_TABLE_HEADER_FILL = "E2EFD9"

# TOC 스타일 ID — 텍스트 보존
PRESERVE_STYLES = {"11", "22", "31", "40"}

# 목차 레벨별 플레이스홀더
SECTION_PLACEHOLDERS = {
    0: "[부문명 기재]",  # 대목차 (I, II, III...)
    1: "[소주제명 기재]",  # 중목차 (1, 2, 3...)
    2: "[세부항목명 기재]",  # 소목차 (가, 나, 다...)
}

BODY_PLACEHOLDER = "[이곳에 텍스트 입력]"
COVER_PROJECT_PLACEHOLDER = "Project [프로젝트 코드명]"
COVER_DATE_PLACEHOLDER = "[YYYY]년 [MM]월 [DD]일"


# ─── XML 유틸리티 ────────────────────────────────────────────────────────────


def _get_style_id_from_xml(p_element) -> str:
    """XML p 요소에서 스타일 ID 추출."""
    pPr = p_element.find(f"{{{WML_NS}}}pPr")
    if pPr is None:
        return ""
    pStyle = pPr.find(f"{{{WML_NS}}}pStyle")
    if pStyle is None:
        return ""
    return pStyle.get(f"{{{WML_NS}}}val", "")


def _get_num_info_from_xml(p_element) -> tuple[str, str]:
    """XML p 요소에서 numId, ilvl 추출."""
    pPr = p_element.find(f"{{{WML_NS}}}pPr")
    if pPr is None:
        return "", ""
    numPr = pPr.find(f"{{{WML_NS}}}numPr")
    if numPr is None:
        return "", ""
    numId_el = numPr.find(f"{{{WML_NS}}}numId")
    ilvl_el = numPr.find(f"{{{WML_NS}}}ilvl")
    numId = numId_el.get(f"{{{WML_NS}}}val", "") if numId_el is not None else ""
    ilvl = ilvl_el.get(f"{{{WML_NS}}}val", "") if ilvl_el is not None else ""
    return numId, ilvl


def _get_text_from_xml(element) -> str:
    """XML 요소에서 모든 w:t 텍스트를 합쳐 반환."""
    texts = []
    for t in element.iter(f"{{{WML_NS}}}t"):
        if t.text:
            texts.append(t.text)
    return "".join(texts)


def _clear_runs_in_xml(p_element, placeholder: str) -> None:
    """p 요소의 run 텍스트를 placeholder로 교체 (첫 w:t만 설정, 나머지 비움)."""
    first_t = True
    for t in p_element.iter(f"{{{WML_NS}}}t"):
        if first_t:
            t.text = placeholder
            t.set(f"{{{XML_SPACE_NS}}}space", "preserve")
            first_t = False
        else:
            t.text = ""


def _clear_cell_runs_xml(tc_element, placeholder: str = "") -> None:
    """tc 요소의 모든 텍스트를 비우고 선택적으로 placeholder 삽입."""
    first_t = True
    for t in tc_element.iter(f"{{{WML_NS}}}t"):
        if first_t and placeholder:
            t.text = placeholder
            t.set(f"{{{XML_SPACE_NS}}}space", "preserve")
            first_t = False
        else:
            t.text = ""


def _get_cell_fill_xml(tc_element) -> str | None:
    """tc 요소에서 배경색(fill) 추출."""
    tcPr = tc_element.find(f"{{{WML_NS}}}tcPr")
    if tcPr is None:
        return None
    shd = tcPr.find(f"{{{WML_NS}}}shd")
    if shd is None:
        return None
    return shd.get(f"{{{WML_NS}}}fill")


def _clear_paragraph_text(paragraph, placeholder: str) -> None:
    """python-docx Paragraph 객체의 run 텍스트를 placeholder로 교체."""
    runs = paragraph.runs
    if not runs:
        return
    runs[0].text = placeholder
    for run in runs[1:]:
        run.text = ""


# ─── 메인 클래스 ─────────────────────────────────────────────────────────────


class LawFirmTemplateGenerator:
    """원본 법률실사보고서 .docx → 플레이스홀더 교체 → 빈 템플릿 생성.

    Step 1 전용: 매 요청마다 원본에서 복제하여 빈 템플릿을 생성한다.
    python-docx 로드/저장 + lxml XML 직접 조작으로 서식을 100% 보존.
    """

    def __init__(self, source_template_path: Path | str) -> None:
        self._source_path = Path(source_template_path)
        if not self._source_path.exists():
            raise FileNotFoundError(f"원본 템플릿 파일 없음: {self._source_path}")

    def generate_blank_template(
        self,
        output_path: Path | str,
        project_code: str = "[프로젝트 코드명]",
        law_firm_name: str = "[법무법인 명칭]",
        report_date: str | None = None,
    ) -> Path:
        """원본 docx를 복제하여 빈 템플릿을 생성한다.

        Args:
            output_path: 출력 파일 경로
            project_code: 표지 프로젝트 코드명
            law_firm_name: 법무법인 명칭 (헤더, 섹션 바 등)
            report_date: 보고서 날짜 (None이면 플레이스홀더 유지)

        Returns:
            생성된 파일 경로
        """
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)

        doc = Document(str(self._source_path))
        cover_project = f"Project {project_code}" if project_code != "[프로젝트 코드명]" else COVER_PROJECT_PLACEHOLDER
        cover_date = report_date if report_date else COVER_DATE_PLACEHOLDER

        self._process_body(doc, cover_project, cover_date, law_firm_name)
        self._process_headers(doc, law_firm_name)
        self._process_footnotes(doc)

        doc.save(str(output))
        logger.info("법무법인 빈 템플릿 생성: %s", output)
        return output

    def _process_body(
        self,
        doc: Document,
        cover_project: str,
        cover_date: str,
        law_firm_name: str,
    ) -> None:
        """본문 요소를 순회하며 텍스트를 플레이스홀더로 교체."""
        section_counter = 0

        for element in list(doc.element.body):
            tag = element.tag.split("}")[-1]

            if tag == "p":
                self._process_paragraph(element, cover_project, cover_date)
            elif tag == "tbl":
                section_counter = self._process_table(
                    element,
                    section_counter,
                    law_firm_name,
                )

    def _process_paragraph(
        self,
        p_element,
        cover_project: str,
        cover_date: str,
    ) -> None:
        """단일 p 요소를 처리."""
        style_id = _get_style_id_from_xml(p_element)
        numId, ilvl = _get_num_info_from_xml(p_element)
        text = _get_text_from_xml(p_element).strip()

        # TOC 스타일 보존
        if style_id in PRESERVE_STYLES:
            return

        # 빈 단락 유지
        if not text:
            return

        # Title 스타일 (표지)
        if style_id == "ae":
            if "Project" in text:
                _clear_runs_in_xml(p_element, cover_project)
            return

        # 날짜 패턴
        if re.match(r"\d{4}.*년.*\d{1,2}.*월.*\d{1,2}.*일", text):
            _clear_runs_in_xml(p_element, cover_date)
            return

        # "귀     중" 보존
        if "귀" in text and "중" in text and len(text) < 10:
            return

        # EXECUTIVE SUMMARY 보존
        if text == "EXECUTIVE SUMMARY":
            return

        # 스타일 12 = 섹션 헤더
        if style_id == "12":
            if numId:
                ilvl_int = int(ilvl) if ilvl else 0
                placeholder = SECTION_PLACEHOLDERS.get(ilvl_int, "[항목명 기재]")
                _clear_runs_in_xml(p_element, placeholder)
            else:
                _clear_runs_in_xml(p_element, BODY_PLACEHOLDER)
            return

        # List Paragraph / Normal 본문
        if style_id in ("a5", "a", "", "a0"):
            if numId:
                ilvl_int = int(ilvl) if ilvl else 0
                if ilvl_int <= 2:
                    placeholder = SECTION_PLACEHOLDERS.get(ilvl_int, "[항목명 기재]")
                else:
                    placeholder = "[세부사항 기재]"
                _clear_runs_in_xml(p_element, placeholder)
            else:
                _clear_runs_in_xml(p_element, BODY_PLACEHOLDER)
            return

        # EndnoteText, ad 등
        if style_id in ("EndnoteText1", "ad"):
            if text:
                _clear_runs_in_xml(p_element, "")
            return

        # 기타
        if text:
            _clear_runs_in_xml(p_element, BODY_PLACEHOLDER)

    def _process_table(
        self,
        tbl_element,
        section_counter: int,
        law_firm_name: str,
    ) -> int:
        """단일 tbl 요소를 처리. 업데이트된 section_counter를 반환."""
        rows = tbl_element.findall(f"{{{WML_NS}}}tr")

        # 3열 섹션 헤더 바 판별
        if len(rows) == 1:
            cells = rows[0].findall(f"{{{WML_NS}}}tc")
            if len(cells) == 3:
                fill0 = _get_cell_fill_xml(cells[0])
                if fill0 in SECTION_BAR_FILLS:
                    section_counter += 1
                    if section_counter == 1:
                        _clear_cell_runs_xml(cells[2], law_firm_name)
                    else:
                        _clear_cell_runs_xml(cells[2], "[부문명 기재]")
                    return section_counter

            # Recommendation 박스 (1열)
            if len(cells) == 1:
                fill = _get_cell_fill_xml(cells[0])
                text = _get_text_from_xml(cells[0])
                if fill == RECOMMENDATION_FILL or "Recommendation" in text:
                    _clear_cell_runs_xml(cells[0], "Recommendation: [권고사항 기재]")
                    return section_counter

        # 데이터 표 (2행 이상)
        if len(rows) >= 2:
            first_cells = rows[0].findall(f"{{{WML_NS}}}tc")
            if first_cells:
                for row_idx, row in enumerate(rows):
                    if row_idx == 0:
                        continue  # 헤더 행 유지
                    for tc in row.findall(f"{{{WML_NS}}}tc"):
                        _clear_cell_runs_xml(tc, "")

        return section_counter

    def _process_headers(self, doc: Document, law_firm_name: str) -> None:
        """머리글의 텍스트를 법무법인 명칭으로 교체."""
        for section in doc.sections:
            if section.header and section.header.paragraphs:
                for para in section.header.paragraphs:
                    if para.text.strip():
                        _clear_paragraph_text(para, law_firm_name)

    def _process_footnotes(self, doc: Document) -> None:
        """각주 텍스트를 플레이스홀더로 교체."""
        footnotes_part = None
        for rel in doc.part.rels.values():
            if "footnotes" in str(rel.reltype):
                footnotes_part = rel.target_part
                break

        if not footnotes_part:
            return

        from lxml import etree

        fn_root = etree.fromstring(footnotes_part.blob)
        ns = {"w": WML_NS}
        for footnote in fn_root.findall("w:footnote", ns):
            fn_type = footnote.get(f"{{{WML_NS}}}type")
            if fn_type in ("separator", "continuationSeparator"):
                continue
            first_found = False
            for t in footnote.iter(f"{{{WML_NS}}}t"):
                if not first_found:
                    if t.text and t.text.strip():
                        t.text = "[각주 내용 기재]"
                    first_found = True
                else:
                    t.text = ""

        footnotes_part._blob = etree.tostring(
            fn_root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )
