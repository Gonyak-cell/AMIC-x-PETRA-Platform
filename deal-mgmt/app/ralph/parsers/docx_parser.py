"""DOCX 파서 — python-docx 기반 텍스트/표 추출."""

from __future__ import annotations

import logging

from app.ralph.parsers.base import ParsedFile, ParsedTable

logger = logging.getLogger(__name__)


def parse_docx(file_path: str) -> ParsedFile:
    """DOCX 파일을 파싱한다."""
    try:
        from docx import Document
    except ImportError:
        return ParsedFile(
            source_path=file_path,
            file_type="docx",
            parse_error="python-docx가 설치되지 않았습니다",
        )

    try:
        doc = Document(file_path)
    except Exception as exc:
        logger.warning("DOCX 파싱 실패 %s: %s", file_path, exc)
        return ParsedFile(
            source_path=file_path,
            file_type="docx",
            parse_error=str(exc),
        )

    # 본문 텍스트
    all_text = [p.text for p in doc.paragraphs if p.text.strip()]

    # 표 추출
    tables: list[ParsedTable] = []
    for table in doc.tables:
        rows_data: list[list[str]] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows_data.append(cells)
        if rows_data:
            tables.append(
                ParsedTable(
                    headers=rows_data[0],
                    rows=rows_data[1:] if len(rows_data) > 1 else [],
                )
            )

    return ParsedFile(
        source_path=file_path,
        file_type="docx",
        text="\n".join(all_text),
        tables=tables,
        metadata={"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)},
    )
