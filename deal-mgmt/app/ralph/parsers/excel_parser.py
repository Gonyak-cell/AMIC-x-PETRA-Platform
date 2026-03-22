"""Excel 파서 — openpyxl 기반 .xlsx/.xls 파싱."""

from __future__ import annotations

import logging
from pathlib import Path

from app.ralph.parsers.base import ParsedFile, ParsedTable

logger = logging.getLogger(__name__)


def parse_excel(file_path: str) -> ParsedFile:
    """Excel 파일을 파싱하여 구조화 데이터를 반환한다.

    .xlsx → openpyxl, .xls → xlrd fallback.
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    try:
        if ext == ".xlsx":
            return _parse_xlsx(file_path)
        elif ext == ".xls":
            return _parse_xls(file_path)
        else:
            return ParsedFile(
                source_path=file_path,
                file_type="excel",
                parse_error=f"지원하지 않는 확장자: {ext}",
            )
    except Exception as exc:
        logger.warning("Excel 파싱 실패 %s: %s", file_path, exc)
        return ParsedFile(
            source_path=file_path,
            file_type="excel",
            parse_error=str(exc),
        )


def _parse_xlsx(file_path: str) -> ParsedFile:
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        all_text: list[str] = []
        tables: list[ParsedTable] = []
        chunks: list[dict] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            all_text.append(f"[시트: {sheet_name}]")

            rows_data: list[list[str]] = []
            for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):  # 빈 행 건너뜀
                    rows_data.append(cells)
                    row_text = " | ".join(cells)
                    all_text.append(row_text)
                    chunks.append(
                        {
                            "chunk_id": f"{sheet_name}-row-{row_index}",
                            "locator_type": "sheet_row",
                            "sheet": sheet_name,
                            "row": row_index,
                            "ordinal": len(chunks) + 1,
                            "text": row_text,
                        }
                    )

            if rows_data:
                table = ParsedTable(
                    headers=rows_data[0] if rows_data else [],
                    rows=rows_data[1:] if len(rows_data) > 1 else [],
                )
                tables.append(table)

        return ParsedFile(
            source_path=file_path,
            file_type="excel",
            text="\n".join(all_text),
            tables=tables,
            metadata={"sheet_count": len(wb.sheetnames), "chunks": chunks},
        )
    finally:
        wb.close()


def _parse_xls(file_path: str) -> ParsedFile:
    """구형 .xls 파일 파싱 (xlrd)."""
    try:
        import xlrd
    except ImportError:
        return ParsedFile(
            source_path=file_path,
            file_type="excel",
            parse_error="xlrd 패키지가 설치되지 않았습니다 (pip install xlrd)",
        )

    wb = xlrd.open_workbook(file_path)
    try:
        all_text: list[str] = []
        tables: list[ParsedTable] = []
        chunks: list[dict] = []

        for sheet_idx in range(wb.nsheets):
            ws = wb.sheet_by_index(sheet_idx)
            all_text.append(f"[시트: {ws.name}]")

            rows_data: list[list[str]] = []
            for r in range(ws.nrows):
                cells = [str(ws.cell_value(r, c)) for c in range(ws.ncols)]
                if any(cells):
                    rows_data.append(cells)
                    row_text = " | ".join(cells)
                    all_text.append(row_text)
                    chunks.append(
                        {
                            "chunk_id": f"{ws.name}-row-{r + 1}",
                            "locator_type": "sheet_row",
                            "sheet": ws.name,
                            "row": r + 1,
                            "ordinal": len(chunks) + 1,
                            "text": row_text,
                        }
                    )

            if rows_data:
                tables.append(
                    ParsedTable(
                        headers=rows_data[0] if rows_data else [],
                        rows=rows_data[1:] if len(rows_data) > 1 else [],
                    )
                )

        return ParsedFile(
            source_path=file_path,
            file_type="excel",
            text="\n".join(all_text),
            tables=tables,
            metadata={"sheet_count": wb.nsheets, "chunks": chunks},
        )
    finally:
        wb.release_resources()
