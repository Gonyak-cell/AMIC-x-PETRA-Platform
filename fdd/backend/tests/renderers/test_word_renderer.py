"""Word Renderer 테스트 — FDD-1101, FDD-1102.

테스트 ID 규칙: T-WORD-{번호}
"""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from docx import Document

from app.renderers.report_builder import (
    AlignType,
    AppendixBlock,
    AppendixItem,
    ChartBlock,
    ChartData,
    ChartType,
    ClaimBlock,
    CoverBlock,
    EvidenceRef,
    IssueBlock,
    IssueItem,
    KPIBlock,
    MethodologyBlock,
    MethodologyItem,
    ReportIR,
    ReportMetadata,
    RiskLevel,
    ScopeBlock,
    ScopeItem,
    TableBlock,
    TableColumn,
    TextBlock,
)
from app.renderers.word_renderer import (
    WORD_RENDERER_VERSION,
    WordTableStyle,
    _format_currency,
    _format_percentage,
    _get_word_styles,
    _hex_to_rgb,
    build_docx_context,
    render_word_report,
)


class TestWordRendererVersion:
    """T-WORD-01: 버전 정보 테스트."""

    def test_version_exists(self):
        """버전 문자열 존재."""
        assert WORD_RENDERER_VERSION is not None
        assert isinstance(WORD_RENDERER_VERSION, str)

    def test_version_format(self):
        """버전 포맷 (SemVer)."""
        parts = WORD_RENDERER_VERSION.split(".")
        assert len(parts) == 3
        assert all(p.isdigit() for p in parts)


class TestWordTableStyle:
    """T-WORD-02: 테이블 스타일 Enum 테스트."""

    def test_style_values(self):
        """스타일 값 확인."""
        assert WordTableStyle.HEADER == "header"
        assert WordTableStyle.BODY == "body"
        assert WordTableStyle.TOTAL == "total"
        assert WordTableStyle.SUBTOTAL == "subtotal"
        assert WordTableStyle.SEPARATOR == "separator"


class TestFormatHelpers:
    """T-WORD-03: 포맷 헬퍼 함수 테스트."""

    def test_format_currency_decimal(self):
        """Decimal 금액 포맷."""
        assert _format_currency(Decimal("1000000")) == "1,000,000"
        assert _format_currency(Decimal("-500000")) == "-500,000"

    def test_format_currency_string(self):
        """문자열 금액 포맷."""
        assert _format_currency("1000000") == "1,000,000"

    def test_format_currency_none(self):
        """None 금액 포맷."""
        assert _format_currency(None) == "-"

    def test_format_percentage_decimal(self):
        """Decimal 백분율 포맷."""
        assert _format_percentage(Decimal("0.15")) == "15.0%"

    def test_format_percentage_none(self):
        """None 백분율 포맷."""
        assert _format_percentage(None) == "-"

    def test_hex_to_rgb(self):
        """HEX → RGB 변환."""
        rgb = _hex_to_rgb("#003366")
        # RGBColor is tuple-like: (r, g, b) indexing
        assert rgb[0] == 0  # red
        assert rgb[1] == 51  # green
        assert rgb[2] == 102  # blue


class TestWordStyles:
    """T-WORD-04: Word 스타일 설정 테스트."""

    def test_get_word_styles(self):
        """스타일 설정 로드."""
        styles = _get_word_styles()
        assert "primary_color" in styles
        assert "heading_font" in styles
        assert "body_font" in styles
        assert styles["title1_size"] is not None


class TestRenderCoverBlock:
    """T-WORD-05: CoverBlock 렌더링 테스트."""

    def test_cover_block_render(self):
        """Cover 블록 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(deal_id="123", deal_name="Test Deal"),
            sections=[
                CoverBlock(
                    deal_name="Target Corp",
                    deal_type="Acquisition",
                    target_name="Target Co., Ltd.",
                    date=date(2025, 12, 31),
                    prepared_by="FDD Team",
                    confidentiality="CONFIDENTIAL",
                )
            ],
        )
        buffer = render_word_report(ir)
        assert isinstance(buffer, BytesIO)

        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Target Corp" in full_text


class TestRenderKPIBlock:
    """T-WORD-06: KPIBlock 렌더링 테스트."""

    def test_kpi_block_render(self):
        """KPI 블록 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                KPIBlock(
                    title="Key Metrics",
                    kpis=[
                        {"label": "Revenue", "value": "10,000", "unit": "백만원"},
                        {"label": "EBITDA", "value": "2,500", "unit": "백만원"},
                    ],
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        assert len(doc.tables) >= 1


class TestRenderTableBlock:
    """T-WORD-07: TableBlock 렌더링 테스트."""

    def test_table_block_basic(self):
        """기본 테이블 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                TableBlock(
                    title="QoE Bridge",
                    columns=[
                        TableColumn(
                            key="category", header="Category", align=AlignType.LEFT
                        ),
                        TableColumn(
                            key="amount",
                            header="Amount",
                            align=AlignType.RIGHT,
                            format="currency",
                        ),
                    ],
                    rows=[
                        {"category": "Revenue", "amount": "10000"},
                        {"category": "COGS", "amount": "-7000"},
                    ],
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        assert len(doc.tables) >= 1
        assert "QoE Bridge" in "\n".join(p.text for p in doc.paragraphs)

    def test_table_with_footer(self):
        """푸터 행 테이블 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                TableBlock(
                    title="Summary",
                    columns=[
                        TableColumn(key="item", header="Item"),
                        TableColumn(key="value", header="Value", format="currency"),
                    ],
                    rows=[{"item": "A", "value": "100"}],
                    footer_rows=[{"item": "Total", "value": "100"}],
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        assert len(doc.tables) >= 1


class TestRenderTextBlock:
    """T-WORD-08: TextBlock 렌더링 테스트."""

    def test_text_block_basic(self):
        """기본 텍스트 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                TextBlock(
                    title="Analysis",
                    content="This is the analysis content.",
                    bullet_points=["Point 1", "Point 2"],
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Analysis" in full_text

    def test_text_block_with_risk(self):
        """리스크 레벨 텍스트 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                TextBlock(
                    content="High risk finding",
                    risk_level=RiskLevel.HIGH,
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "HIGH" in full_text


class TestRenderClaimBlock:
    """T-WORD-09: ClaimBlock 렌더링 테스트."""

    def test_claim_block_verified(self):
        """검증된 Claim 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                ClaimBlock(
                    claim_text="Revenue increased by 10%.",
                    evidence_refs=[
                        EvidenceRef(
                            evidence_id="ev-001",
                            source_type="TB",
                            source_id="tb-001",
                            description="Trial Balance FY2025",
                        )
                    ],
                    verified=True,
                    category="QoE",
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Revenue increased" in full_text
        assert "Verified" in full_text

    def test_claim_block_unverified(self):
        """미검증 Claim 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                ClaimBlock(
                    claim_text="Unverified claim",
                    verified=False,
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Unverified" in full_text


class TestRenderIssueBlock:
    """T-WORD-10: IssueBlock 렌더링 테스트."""

    def test_issue_block_render(self):
        """이슈 블록 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                IssueBlock(
                    title="Issue Log",
                    issues=[
                        IssueItem(
                            issue_id="ISS-001",
                            category="QoE",
                            severity="high",
                            title="Non-recurring item not adjusted",
                            status="open",
                        )
                    ],
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        assert len(doc.tables) >= 1


class TestRenderMethodologyBlock:
    """T-WORD-11: MethodologyBlock 렌더링 테스트."""

    def test_methodology_block_render(self):
        """방법론 블록 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                MethodologyBlock(
                    title="Methodology",
                    introduction="This report follows FDD standards.",
                    steps=[
                        MethodologyItem(
                            step=1, title="Data Collection", description="Collect TB/GL"
                        ),
                        MethodologyItem(
                            step=2, title="Analysis", description="Perform QoE analysis"
                        ),
                    ],
                    limitations=["Limited access to management"],
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Methodology" in full_text
        assert "Data Collection" in full_text


class TestRenderScopeBlock:
    """T-WORD-12: ScopeBlock 렌더링 테스트."""

    def test_scope_block_render(self):
        """범위 블록 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                ScopeBlock(
                    title="Scope & Definitions",
                    scope_items=[
                        ScopeItem(
                            category="period", label="Analysis Period", value="FY2025"
                        ),
                        ScopeItem(
                            category="entity",
                            label="Target Entity",
                            value="Target Corp",
                        ),
                    ],
                    definitions={
                        "EBITDA": "Earnings before interest, taxes, depreciation and amortization"
                    },
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Scope" in full_text


class TestRenderAppendixBlock:
    """T-WORD-13: AppendixBlock 렌더링 테스트."""

    def test_appendix_block_render(self):
        """부록 블록 렌더링."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                AppendixBlock(
                    title="Appendix",
                    items=[
                        AppendixItem(
                            title="Data Sources",
                            content="List of data sources used in this analysis.",
                        )
                    ],
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Appendix" in full_text


class TestRenderChartBlock:
    """T-WORD-14: ChartBlock 렌더링 테스트."""

    def test_chart_block_no_image(self):
        """이미지 없는 차트 블록 (플레이스홀더)."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                ChartBlock(
                    chart_type=ChartType.WATERFALL,
                    title="EBITDA Bridge",
                    data=ChartData(
                        categories=["Start", "Adj", "End"],
                        values=[Decimal("100"), Decimal("10"), None],
                    ),
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "EBITDA Bridge" in full_text


class TestBuildDocxContext:
    """T-WORD-15: docxtpl 컨텍스트 빌더 테스트."""

    def test_context_metadata(self):
        """메타데이터 컨텍스트."""
        ir = ReportIR(
            metadata=ReportMetadata(deal_id="123", deal_name="Test Deal"),
            sections=[],
        )
        context = build_docx_context(ir)
        assert context["metadata"]["deal_id"] == "123"
        assert context["metadata"]["deal_name"] == "Test Deal"

    def test_context_cover(self):
        """Cover 컨텍스트."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                CoverBlock(deal_name="My Deal", target_name="Target Co."),
            ],
        )
        context = build_docx_context(ir)
        assert context["cover"]["deal_name"] == "My Deal"
        assert context["cover"]["target_name"] == "Target Co."

    def test_context_tables(self):
        """Table 컨텍스트."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                TableBlock(
                    title="Test Table",
                    columns=[TableColumn(key="a", header="A")],
                    rows=[{"a": "1"}],
                )
            ],
        )
        context = build_docx_context(ir)
        assert len(context["tables"]) == 1
        assert context["tables"][0]["title"] == "Test Table"


class TestFullReportRender:
    """T-WORD-16: 전체 보고서 렌더링 테스트."""

    def test_full_report(self):
        """모든 블록 포함 보고서."""
        ir = ReportIR(
            metadata=ReportMetadata(deal_id="deal-001", deal_name="Full Report Test"),
            sections=[
                CoverBlock(deal_name="Test Deal", target_name="Target"),
                KPIBlock(
                    title="Summary",
                    kpis=[{"label": "EBITDA", "value": "100", "unit": "M"}],
                ),
                TableBlock(
                    title="Bridge",
                    columns=[TableColumn(key="item", header="Item")],
                    rows=[{"item": "Revenue"}],
                ),
                TextBlock(content="Analysis notes."),
                ClaimBlock(claim_text="Finding 1", verified=True),
            ],
        )
        buffer = render_word_report(ir)

        # 파일 크기 확인
        assert buffer.getbuffer().nbytes > 0

        # 문서 열기 가능 확인
        doc = Document(buffer)
        assert len(doc.paragraphs) > 0


class TestEmptyReport:
    """T-WORD-17: 빈 보고서 렌더링 테스트."""

    def test_empty_sections(self):
        """섹션 없는 보고서."""
        ir = ReportIR(
            metadata=ReportMetadata(deal_id="empty"),
            sections=[],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        # 빈 문서도 정상 생성
        assert doc is not None


class TestEdgeCases:
    """T-WORD-18: 엣지 케이스 테스트."""

    def test_special_characters_in_text(self):
        """특수 문자 처리."""
        ir = ReportIR(
            metadata=ReportMetadata(deal_name='Test & <Special> "Chars"'),
            sections=[
                TextBlock(content="Special chars: & < > \" '"),
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        assert doc is not None

    def test_korean_text(self):
        """한글 텍스트 처리."""
        ir = ReportIR(
            metadata=ReportMetadata(deal_name="한글 테스트"),
            sections=[
                TextBlock(
                    title="분석 결과",
                    content="매출이 전년 대비 10% 증가하였습니다.",
                ),
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "분석 결과" in full_text


class TestTableWithManyRows:
    """T-WORD-19: 대용량 테이블 테스트."""

    def test_large_table(self):
        """100행 테이블."""
        rows = [{"item": f"Row {i}", "value": str(i * 100)} for i in range(100)]
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                TableBlock(
                    title="Large Table",
                    columns=[
                        TableColumn(key="item", header="Item"),
                        TableColumn(key="value", header="Value"),
                    ],
                    rows=rows,
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        # 테이블이 생성되었는지 확인
        assert len(doc.tables) >= 1
        # 행 수 확인 (헤더 + 100 데이터행)
        assert len(doc.tables[0].rows) == 101


class TestNullValues:
    """T-WORD-20: Null 값 처리 테스트."""

    def test_null_in_table(self):
        """테이블 내 None 값."""
        ir = ReportIR(
            metadata=ReportMetadata(),
            sections=[
                TableBlock(
                    title="Null Test",
                    columns=[
                        TableColumn(key="a", header="A"),
                        TableColumn(key="b", header="B", format="currency"),
                    ],
                    rows=[
                        {"a": "Item 1", "b": None},
                        {"a": None, "b": "100"},
                    ],
                )
            ],
        )
        buffer = render_word_report(ir)
        doc = Document(buffer)
        assert doc is not None
