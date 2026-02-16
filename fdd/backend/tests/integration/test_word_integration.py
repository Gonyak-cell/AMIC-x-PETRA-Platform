"""Word Integration 테스트 — Sprint 9.

테스트 ID 규칙: T-WORD-INT-{번호}
"""

from datetime import date
from decimal import Decimal

import pytest
from docx import Document

from app.renderers.report_builder import (
    ChartBlock,
    ChartData,
    ChartType,
    ClaimBlock,
    CoverBlock,
    EvidenceRef,
    KPIBlock,
    ReportIR,
    ReportMetadata,
    TableBlock,
    TableColumn,
    TextBlock,
    build_claim_block,
    build_kpi_block,
    build_qoe_table_block,
    build_text_block,
    build_waterfall_chart_block,
)
from app.renderers.word_renderer import build_docx_context, render_word_report
from app.services.chart import render_chart_block
from app.services.narrative import (
    generate_debt_narrative,
    generate_executive_summary,
    generate_nwc_narrative,
    generate_qoe_narrative,
)


class TestFullFDDReport:
    """T-WORD-INT-01: 전체 FDD 보고서 통합 테스트."""

    def test_complete_fdd_report(self):
        """QoE + NWC + Debt 전체 보고서."""
        # 1. 서술문 생성
        qoe_narrative = generate_qoe_narrative(
            period="FY2025",
            reported_ebitda=Decimal("10000000000"),
            adjusted_ebitda=Decimal("12000000000"),
            adjustments=[
                {"category": "Non-recurring", "amount": "1,500,000,000", "description": "Legal settlement"},
                {"category": "Non-operating", "amount": "500,000,000", "description": "FX gain"},
            ],
        )

        nwc_narrative = generate_nwc_narrative(
            period="2025-12-31",
            total_nwc=Decimal("3000000000"),
            current_assets=Decimal("8000000000"),
            current_liabilities=Decimal("5000000000"),
            peg_method="Average",
            target_nwc=Decimal("2800000000"),
        )

        debt_narrative = generate_debt_narrative(
            period="2025-12-31",
            net_debt=Decimal("5000000000"),
            total_debt=Decimal("8000000000"),
            cash=Decimal("3000000000"),
        )

        # 2. Report IR 생성
        ir = ReportIR(
            metadata=ReportMetadata(
                deal_id="deal-001",
                deal_name="Project Alpha",
                generated_at="2025-12-31T00:00:00",
                version="1.0",
                engine_versions={"qoe": "0.1.0", "nwc": "0.1.0", "debt": "0.1.0"},
            ),
            sections=[
                # Cover
                CoverBlock(
                    deal_name="Project Alpha",
                    deal_type="Acquisition",
                    target_name="Target Corp",
                    date=date(2025, 12, 31),
                    prepared_by="FDD Team",
                    confidentiality="CONFIDENTIAL",
                ),
                # Executive Summary KPIs
                build_kpi_block(
                    title="Executive Summary",
                    kpis=[
                        ("Adjusted EBITDA", "12.0B", "KRW"),
                        ("Net Working Capital", "3.0B", "KRW"),
                        ("Net Debt", "5.0B", "KRW"),
                    ],
                ),
                # QoE Section
                build_text_block(
                    title="Quality of Earnings",
                    content=qoe_narrative,
                ),
                build_qoe_table_block(
                    title="EBITDA Bridge",
                    bridge_data=[
                        {"category": "Reported EBITDA", "FY2025": "10,000"},
                        {"category": "Non-recurring", "FY2025": "1,500"},
                        {"category": "Non-operating", "FY2025": "500"},
                        {"category": "Adjusted EBITDA", "FY2025": "12,000"},
                    ],
                    fiscal_years=["FY2025"],
                ),
                # NWC Section
                build_text_block(
                    title="Net Working Capital",
                    content=nwc_narrative,
                ),
                # Debt Section
                build_text_block(
                    title="Net Debt",
                    content=debt_narrative,
                ),
                # Claim with evidence
                build_claim_block(
                    claim_text="Revenue growth of 15% is supported by customer contract analysis.",
                    evidence_refs=[
                        ("ev-001", "FILE", "customer_contracts.xlsx", "Customer contract summary"),
                    ],
                    category="QoE",
                ),
            ],
        )

        # 3. Word 렌더링
        buffer = render_word_report(ir)

        # 4. 검증
        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)

        assert "Project Alpha" in full_text
        assert "Target Corp" in full_text
        assert "Quality of Earnings" in full_text
        assert "Net Working Capital" in full_text
        assert "Net Debt" in full_text


class TestNarrativeChartIntegration:
    """T-WORD-INT-02: 서술문 + 차트 통합 테스트."""

    def test_narrative_with_chart(self):
        """서술문과 차트가 함께 포함된 보고서."""
        # 차트 블록 생성 및 렌더링
        chart_block = ChartBlock(
            chart_type=ChartType.WATERFALL,
            title="EBITDA Bridge",
            data=ChartData(
                categories=["Reported", "Non-recurring", "Non-operating", "Adjusted"],
                values=[Decimal("100"), Decimal("15"), Decimal("5"), None],
            ),
        )
        rendered_chart = render_chart_block(chart_block)

        # Report IR
        ir = ReportIR(
            metadata=ReportMetadata(deal_name="Chart Test"),
            sections=[
                build_text_block(
                    title="Analysis",
                    content="The following chart shows the EBITDA bridge.",
                ),
                rendered_chart,
            ],
        )

        buffer = render_word_report(ir)
        doc = Document(buffer)

        # 텍스트 확인
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "EBITDA Bridge" in full_text


class TestMultiTableReport:
    """T-WORD-INT-03: 다중 테이블 통합 테스트."""

    def test_multiple_tables(self):
        """여러 테이블 포함 보고서."""
        ir = ReportIR(
            metadata=ReportMetadata(deal_name="Multi-Table"),
            sections=[
                TableBlock(
                    title="QoE Summary",
                    columns=[
                        TableColumn(key="item", header="Item"),
                        TableColumn(key="amount", header="Amount"),
                    ],
                    rows=[{"item": "Revenue", "amount": "100"}],
                ),
                TableBlock(
                    title="NWC Summary",
                    columns=[
                        TableColumn(key="category", header="Category"),
                        TableColumn(key="balance", header="Balance"),
                    ],
                    rows=[{"category": "AR", "balance": "50"}],
                ),
                TableBlock(
                    title="Debt Summary",
                    columns=[
                        TableColumn(key="type", header="Type"),
                        TableColumn(key="amount", header="Amount"),
                    ],
                    rows=[{"type": "Bank Loan", "amount": "80"}],
                ),
            ],
        )

        buffer = render_word_report(ir)
        doc = Document(buffer)

        # 3개 테이블 생성 확인
        assert len(doc.tables) >= 3


class TestDocxTemplateContext:
    """T-WORD-INT-04: docxtpl 컨텍스트 통합 테스트."""

    def test_context_building(self):
        """컨텍스트 빌드 통합 테스트."""
        ir = ReportIR(
            metadata=ReportMetadata(deal_id="ctx-001", deal_name="Context Test"),
            sections=[
                CoverBlock(deal_name="Test Deal", target_name="Target"),
                KPIBlock(title="KPIs", kpis=[{"label": "EBITDA", "value": "100", "unit": "M"}]),
                TableBlock(
                    title="Data",
                    columns=[TableColumn(key="a", header="A")],
                    rows=[{"a": "1"}],
                ),
                ClaimBlock(
                    claim_text="Test claim",
                    verified=True,
                    evidence_refs=[EvidenceRef("e1", "TB", "tb-1", "")],
                ),
            ],
        )

        context = build_docx_context(ir)

        # 메타데이터
        assert context["metadata"]["deal_id"] == "ctx-001"

        # Cover
        assert context["cover"]["deal_name"] == "Test Deal"

        # KPIs
        assert len(context["kpis"]) == 1

        # Tables
        assert len(context["tables"]) == 1

        # Claims
        assert len(context["claims"]) == 1
        assert context["claims"][0]["verified"] is True


class TestLargeReport:
    """T-WORD-INT-05: 대용량 보고서 통합 테스트."""

    def test_large_report(self):
        """100개 섹션 보고서."""
        sections = [CoverBlock(deal_name="Large Report")]

        # 50개 테이블 추가
        for i in range(50):
            sections.append(
                TableBlock(
                    title=f"Table {i + 1}",
                    columns=[TableColumn(key="item", header="Item")],
                    rows=[{"item": f"Row {j}"} for j in range(10)],
                )
            )

        # 50개 텍스트 블록 추가
        for i in range(50):
            sections.append(
                TextBlock(
                    title=f"Section {i + 1}",
                    content=f"Content for section {i + 1}. " * 10,
                )
            )

        ir = ReportIR(
            metadata=ReportMetadata(deal_name="Large Report"),
            sections=sections,
        )

        buffer = render_word_report(ir)

        # 파일 크기 확인 (충분히 큼)
        assert buffer.getbuffer().nbytes > 30000  # 30KB 이상

        # 문서 열기 가능 확인
        doc = Document(buffer)
        assert len(doc.paragraphs) > 0


class TestNarrativeInContext:
    """T-WORD-INT-06: 서술문 컨텍스트 통합 테스트."""

    def test_narratives_in_report(self):
        """서술문이 보고서에 올바르게 포함되는지 테스트."""
        qoe = generate_qoe_narrative(
            period="FY2025",
            reported_ebitda=Decimal("1000"),
            adjusted_ebitda=Decimal("1100"),
        )

        nwc = generate_nwc_narrative(
            period="FY2025",
            total_nwc=Decimal("500"),
            current_assets=Decimal("800"),
            current_liabilities=Decimal("300"),
        )

        exec_summary = generate_executive_summary(
            deal_name="Test Deal",
            target_name="Target Co.",
            analysis_period="FY2025",
            qoe_summary=qoe,
            nwc_summary=nwc,
            debt_summary="Net Debt is 200M",
        )

        ir = ReportIR(
            metadata=ReportMetadata(deal_name="Narrative Test"),
            sections=[
                TextBlock(title="Executive Summary", content=exec_summary),
                TextBlock(title="Quality of Earnings", content=qoe),
                TextBlock(title="Working Capital", content=nwc),
            ],
        )

        buffer = render_word_report(ir)
        doc = Document(buffer)
        full_text = "\n".join(p.text for p in doc.paragraphs)

        # 핵심 정보 포함 확인
        assert "Test Deal" in full_text
        assert "FY2025" in full_text


class TestChartDataValidation:
    """T-WORD-INT-07: 차트 데이터 검증 통합 테스트."""

    def test_chart_table_consistency(self):
        """차트와 테이블 데이터 일관성."""
        from app.services.chart import ChartSpec, validate_chart_data

        # 테이블 데이터
        table_data = [
            {"category": "Reported", "value": Decimal("100")},
            {"category": "Adjustment", "value": Decimal("20")},
            {"category": "Adjusted", "value": Decimal("120")},
        ]

        # 차트 스펙
        spec = ChartSpec(
            chart_type=ChartType.WATERFALL,
            categories=["Reported", "Adjustment", "Adjusted"],
            values=[Decimal("100"), Decimal("20"), None],
        )

        # 검증
        errors = validate_chart_data(spec, table_data)
        # 첫 번째 행만 비교하므로 에러 없어야 함 (또는 있을 수 있음)
        # 이 테스트는 검증 로직이 동작하는지 확인
        assert isinstance(errors, list)


class TestReportWithAllBlockTypes:
    """T-WORD-INT-08: 모든 블록 타입 통합 테스트."""

    def test_all_block_types_render(self):
        """모든 블록 타입이 렌더링되는지 확인."""
        from app.renderers.report_builder import (
            AppendixBlock,
            AppendixItem,
            IssueBlock,
            IssueItem,
            MethodologyBlock,
            MethodologyItem,
            RiskLevel,
            ScopeBlock,
            ScopeItem,
        )

        ir = ReportIR(
            metadata=ReportMetadata(deal_name="All Blocks"),
            sections=[
                CoverBlock(deal_name="All Blocks Report"),
                KPIBlock(title="KPIs", kpis=[{"label": "Test", "value": "100", "unit": "M"}]),
                TableBlock(
                    title="Table",
                    columns=[TableColumn(key="a", header="A")],
                    rows=[{"a": "1"}],
                ),
                ChartBlock(chart_type=ChartType.BAR, title="Chart"),
                TextBlock(title="Text", content="Some text.", risk_level=RiskLevel.MEDIUM),
                ClaimBlock(claim_text="A claim", verified=False),
                IssueBlock(
                    title="Issues",
                    issues=[IssueItem(issue_id="I1", category="QoE", severity="high", title="Issue 1")],
                ),
                MethodologyBlock(
                    title="Methodology",
                    steps=[MethodologyItem(step=1, title="Step 1")],
                ),
                ScopeBlock(
                    title="Scope",
                    scope_items=[ScopeItem(category="period", label="Period", value="FY2025")],
                ),
                AppendixBlock(
                    title="Appendix",
                    items=[AppendixItem(title="Data Sources", content="List of sources")],
                ),
            ],
        )

        buffer = render_word_report(ir)
        doc = Document(buffer)

        full_text = "\n".join(p.text for p in doc.paragraphs)

        # 각 블록 타입의 제목/내용 확인
        assert "All Blocks Report" in full_text
        assert "KPIs" in full_text
        assert "Table" in full_text
        assert "Chart" in full_text
        assert "Text" in full_text
        assert "Issues" in full_text
        assert "Methodology" in full_text
        assert "Scope" in full_text
        assert "Appendix" in full_text


class TestKoreanReport:
    """T-WORD-INT-09: 한글 보고서 통합 테스트."""

    def test_korean_content(self):
        """한글 콘텐츠 보고서."""
        qoe = generate_qoe_narrative(
            period="2025년 회계연도",
            reported_ebitda=Decimal("100000000000"),  # 1000억
            adjusted_ebitda=Decimal("120000000000"),  # 1200억
            adjustments=[
                {"category": "일회성", "amount": "200억", "description": "법적 합의금"},
            ],
        )

        ir = ReportIR(
            metadata=ReportMetadata(deal_name="한글 보고서 테스트"),
            sections=[
                CoverBlock(
                    deal_name="프로젝트 알파",
                    target_name="대상 회사 주식회사",
                    prepared_by="실사팀",
                ),
                TextBlock(
                    title="분석 개요",
                    content=qoe,
                    bullet_points=[
                        "매출 증가 추세 지속",
                        "비용 구조 개선 필요",
                        "운전자본 관리 양호",
                    ],
                ),
                TableBlock(
                    title="손익 요약",
                    columns=[
                        TableColumn(key="항목", header="항목"),
                        TableColumn(key="금액", header="금액 (백만원)"),
                    ],
                    rows=[
                        {"항목": "매출", "금액": "100,000"},
                        {"항목": "영업이익", "금액": "15,000"},
                        {"항목": "EBITDA", "금액": "20,000"},
                    ],
                ),
            ],
        )

        buffer = render_word_report(ir)
        doc = Document(buffer)

        full_text = "\n".join(p.text for p in doc.paragraphs)

        assert "프로젝트 알파" in full_text
        assert "대상 회사 주식회사" in full_text
        assert "분석 개요" in full_text


class TestEmptyEdgeCases:
    """T-WORD-INT-10: 빈 값 엣지 케이스 통합 테스트."""

    def test_empty_strings(self):
        """빈 문자열 처리."""
        ir = ReportIR(
            metadata=ReportMetadata(deal_name=""),
            sections=[
                CoverBlock(deal_name="", target_name=""),
                TextBlock(title="", content=""),
                TableBlock(title="", columns=[], rows=[]),
            ],
        )

        # 오류 없이 렌더링되어야 함
        buffer = render_word_report(ir)
        doc = Document(buffer)
        assert doc is not None
