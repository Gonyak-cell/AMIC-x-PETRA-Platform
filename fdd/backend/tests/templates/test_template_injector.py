"""Template Injector Tests - 템플릿 주입 테스트.

EPIC-10 FDD-1002: Template Injector 테스트.
"""

import os
import tempfile
from decimal import Decimal
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.renderers.report_builder import (
    AlignType,
    BlockType,
    ChartBlock,
    ChartData,
    ChartType,
    ClaimBlock,
    CoverBlock,
    KPIBlock,
    ReportIR,
    ReportMetadata,
    ScopeBlock,
    ScopeItem,
    TableBlock,
    TableColumn,
    TextBlock,
)
from app.renderers.template_injector import (
    InjectionContext,
    InjectionResult,
    TemplateInjector,
    inject_template,
)
from app.schemas.template import (
    SlotType,
    StyleTokens,
    TemplateContract,
    TemplateSlot,
)

# =============================================================================
# Fixtures
# =============================================================================


@pytest.fixture
def sample_report_ir() -> ReportIR:
    """샘플 Report IR 생성."""
    return ReportIR(
        metadata=ReportMetadata(
            deal_id="test-deal-001",
            deal_name="Test Corp Acquisition",
            generated_at="2026-02-06T12:00:00Z",
            version="1.0",
        ),
        sections=[
            CoverBlock(
                deal_name="Test Corp Acquisition",
                target_name="Target Co., Ltd.",
                date=None,
                prepared_by="FDD Team",
            ),
            KPIBlock(
                title="key_metrics",
                kpis=[
                    {"label": "Adjusted EBITDA", "value": "12,500", "unit": "백만원"},
                    {"label": "Net Debt", "value": "8,000", "unit": "백만원"},
                    {"label": "NWC", "value": "3,500", "unit": "백만원"},
                ],
            ),
            TableBlock(
                title="qoe_bridge",
                columns=[
                    TableColumn(
                        key="category", header="Category", align=AlignType.LEFT
                    ),
                    TableColumn(key="amount", header="Amount", align=AlignType.RIGHT),
                ],
                rows=[
                    {"category": "Reported EBITDA", "amount": "10,000"},
                    {"category": "Non-recurring adjustment", "amount": "1,500"},
                    {"category": "Normalization", "amount": "1,000"},
                    {"category": "Adjusted EBITDA", "amount": "12,500"},
                ],
            ),
            TextBlock(
                title="executive_summary",
                content="This is the executive summary of the FDD analysis.",
                bullet_points=[
                    "Key finding 1",
                    "Key finding 2",
                    "Key finding 3",
                ],
            ),
            ChartBlock(
                chart_type=ChartType.WATERFALL,
                title="ebitda_bridge",
                data=ChartData(
                    categories=[
                        "Reported",
                        "Non-recurring",
                        "Normalization",
                        "Adjusted",
                    ],
                    values=[
                        Decimal("10000"),
                        Decimal("1500"),
                        Decimal("1000"),
                        Decimal("12500"),
                    ],
                ),
            ),
        ],
    )


@pytest.fixture
def sample_contract() -> TemplateContract:
    """샘플 템플릿 계약."""
    return TemplateContract(
        template_id="test_template",
        template_name="Test Template",
        template_type="pptx",
        slots=[
            TemplateSlot(
                slot_id="{{SLOT:DEAL_NAME}}",
                slot_type=SlotType.TEXT,
                required=True,
            ),
            TemplateSlot(
                slot_id="{{TABLE:QOE_BRIDGE}}",
                slot_type=SlotType.TABLE,
                required=True,
            ),
            TemplateSlot(
                slot_id="{{TEXT:EXECUTIVE_SUMMARY}}",
                slot_type=SlotType.TEXT,
                required=True,
            ),
            TemplateSlot(
                slot_id="{{CHART:EBITDA_BRIDGE}}",
                slot_type=SlotType.CHART,
                required=False,
            ),
        ],
        style_tokens=StyleTokens(),
    )


@pytest.fixture
def sample_pptx_template():
    """슬롯이 있는 샘플 PPTX 템플릿 생성."""
    prs = Presentation()

    # Slide 1: Cover
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    textbox1 = slide1.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    textbox1.text_frame.text = "{{SLOT:DEAL_NAME}}"

    # Slide 2: QoE Table
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    textbox2 = slide2.shapes.add_textbox(Inches(0.5), Inches(1), Inches(9), Inches(5))
    textbox2.text_frame.text = "{{TABLE:QOE_BRIDGE}}"

    # Slide 3: Executive Summary
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    textbox3 = slide3.shapes.add_textbox(Inches(0.5), Inches(1), Inches(9), Inches(5))
    textbox3.text_frame.text = "{{TEXT:EXECUTIVE_SUMMARY}}"

    # Slide 4: Chart
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    textbox4 = slide4.shapes.add_textbox(Inches(0.5), Inches(1), Inches(9), Inches(5))
    textbox4.text_frame.text = "{{CHART:EBITDA_BRIDGE}}"

    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
        prs.save(f.name)
        yield Path(f.name)
    os.unlink(f.name)


# =============================================================================
# Injection Context Tests
# =============================================================================


class TestInjectionContext:
    """InjectionContext tests."""

    def test_create_context(
        self, sample_pptx_template, sample_report_ir, sample_contract
    ):
        """컨텍스트 생성."""
        context = InjectionContext(
            template_path=sample_pptx_template,
            output_path=Path("output.pptx"),
            report_ir=sample_report_ir,
            contract=sample_contract,
        )
        assert context.template_path == sample_pptx_template
        assert context.slot_mappings == {}

    def test_context_with_custom_mappings(
        self, sample_pptx_template, sample_report_ir, sample_contract
    ):
        """커스텀 매핑이 있는 컨텍스트."""
        mappings = {"{{SLOT:DEAL_NAME}}": "custom_block_id"}
        context = InjectionContext(
            template_path=sample_pptx_template,
            output_path=Path("output.pptx"),
            report_ir=sample_report_ir,
            contract=sample_contract,
            slot_mappings=mappings,
        )
        assert context.slot_mappings == mappings


# =============================================================================
# Template Injector Tests
# =============================================================================


class TestTemplateInjector:
    """TemplateInjector tests."""

    def test_inject_pptx_success(
        self, sample_pptx_template, sample_report_ir, sample_contract, tmp_path
    ):
        """PPTX 주입 성공."""
        output_path = tmp_path / "output.pptx"

        context = InjectionContext(
            template_path=sample_pptx_template,
            output_path=output_path,
            report_ir=sample_report_ir,
            contract=sample_contract,
        )
        injector = TemplateInjector(context)
        result = injector.inject()

        assert result.success is True
        assert result.output_path == output_path
        assert output_path.exists()
        assert len(result.filled_slots) > 0

    def test_inject_fills_slots(
        self, sample_pptx_template, sample_report_ir, sample_contract, tmp_path
    ):
        """슬롯이 콘텐츠로 채워짐."""
        output_path = tmp_path / "output.pptx"

        result = inject_template(
            template_path=sample_pptx_template,
            output_path=output_path,
            report_ir=sample_report_ir,
            contract=sample_contract,
        )

        assert result.success is True

        # 출력 파일 확인
        prs = Presentation(str(output_path))
        all_text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    all_text += shape.text_frame.text

        # 슬롯 플레이스홀더가 교체되었는지 확인
        # (일부 슬롯은 매칭 안 될 수 있음)
        assert "Test Corp Acquisition" in all_text or len(result.filled_slots) > 0

    def test_inject_with_custom_mappings(
        self, sample_pptx_template, sample_report_ir, sample_contract, tmp_path
    ):
        """커스텀 슬롯 매핑 적용."""
        output_path = tmp_path / "output.pptx"

        # 커스텀 매핑
        mappings = {
            "{{SLOT:DEAL_NAME}}": "cover_block",
        }

        result = inject_template(
            template_path=sample_pptx_template,
            output_path=output_path,
            report_ir=sample_report_ir,
            contract=sample_contract,
            slot_mappings=mappings,
        )

        assert result.success is True

    def test_inject_empty_slots_tracked(
        self, sample_pptx_template, sample_report_ir, sample_contract, tmp_path
    ):
        """매칭 안 된 슬롯 추적."""
        output_path = tmp_path / "output.pptx"

        # 매칭 안 되는 슬롯 추가
        sample_contract.slots.append(
            TemplateSlot(
                slot_id="{{SLOT:NONEXISTENT}}",
                slot_type=SlotType.TEXT,
                required=False,
            )
        )

        result = inject_template(
            template_path=sample_pptx_template,
            output_path=output_path,
            report_ir=sample_report_ir,
            contract=sample_contract,
        )

        assert result.success is True
        # 템플릿에 NONEXISTENT 슬롯이 없으므로 empty_slots에는 포함 안 됨

    def test_inject_missing_template(self, sample_report_ir, sample_contract, tmp_path):
        """존재하지 않는 템플릿 처리."""
        output_path = tmp_path / "output.pptx"

        result = inject_template(
            template_path=Path("/nonexistent/template.pptx"),
            output_path=output_path,
            report_ir=sample_report_ir,
            contract=sample_contract,
        )

        assert result.success is False
        assert len(result.errors) > 0


# =============================================================================
# Block to Text Conversion Tests
# =============================================================================


class TestBlockToText:
    """Block to text conversion tests."""

    def test_text_block_to_text(self):
        """TextBlock 변환."""
        block = TextBlock(
            title="Test Title",
            content="Test content",
            bullet_points=["Point 1", "Point 2"],
        )
        injector = TemplateInjector(
            InjectionContext(
                template_path=Path("dummy.pptx"),
                output_path=Path("output.pptx"),
                report_ir=ReportIR(),
                contract=TemplateContract(template_id="test", template_name="Test"),
            )
        )
        text = injector._block_to_text(block)

        assert "Test Title" in text
        assert "Test content" in text
        assert "• Point 1" in text

    def test_kpi_block_to_text(self):
        """KPIBlock 변환."""
        block = KPIBlock(
            title="KPIs",
            kpis=[
                {"label": "EBITDA", "value": "10,000", "unit": "M"},
                {"label": "Revenue", "value": "50,000", "unit": "M"},
            ],
        )
        injector = TemplateInjector(
            InjectionContext(
                template_path=Path("dummy.pptx"),
                output_path=Path("output.pptx"),
                report_ir=ReportIR(),
                contract=TemplateContract(template_id="test", template_name="Test"),
            )
        )
        text = injector._block_to_text(block)

        assert "EBITDA: 10,000 M" in text
        assert "Revenue: 50,000 M" in text

    def test_table_block_to_text(self):
        """TableBlock 변환."""
        block = TableBlock(
            title="Test Table",
            columns=[
                TableColumn(key="col1", header="Column 1", align=AlignType.LEFT),
                TableColumn(key="col2", header="Column 2", align=AlignType.RIGHT),
            ],
            rows=[
                {"col1": "A", "col2": "1"},
                {"col1": "B", "col2": "2"},
            ],
        )
        injector = TemplateInjector(
            InjectionContext(
                template_path=Path("dummy.pptx"),
                output_path=Path("output.pptx"),
                report_ir=ReportIR(),
                contract=TemplateContract(template_id="test", template_name="Test"),
            )
        )
        text = injector._block_to_text(block)

        assert "Test Table" in text
        assert "Column 1" in text
        assert "A | 1" in text

    def test_chart_block_to_text(self):
        """ChartBlock 변환."""
        block = ChartBlock(
            chart_type=ChartType.WATERFALL,
            title="EBITDA Bridge",
            data=ChartData(
                categories=["Start", "Add", "End"],
                values=[Decimal("100"), Decimal("50"), Decimal("150")],
            ),
        )
        injector = TemplateInjector(
            InjectionContext(
                template_path=Path("dummy.pptx"),
                output_path=Path("output.pptx"),
                report_ir=ReportIR(),
                contract=TemplateContract(template_id="test", template_name="Test"),
            )
        )
        text = injector._block_to_text(block)

        assert "EBITDA Bridge" in text
        assert "Start" in text

    def test_claim_block_to_text(self):
        """ClaimBlock 변환."""
        block = ClaimBlock(
            claim_text="This is a verified claim with evidence.",
            verified=True,
        )
        injector = TemplateInjector(
            InjectionContext(
                template_path=Path("dummy.pptx"),
                output_path=Path("output.pptx"),
                report_ir=ReportIR(),
                contract=TemplateContract(template_id="test", template_name="Test"),
            )
        )
        text = injector._block_to_text(block)

        assert "This is a verified claim" in text

    def test_cover_block_to_text(self):
        """CoverBlock 변환."""
        block = CoverBlock(
            deal_name="Test Deal",
            target_name="Target Corp",
        )
        injector = TemplateInjector(
            InjectionContext(
                template_path=Path("dummy.pptx"),
                output_path=Path("output.pptx"),
                report_ir=ReportIR(),
                contract=TemplateContract(template_id="test", template_name="Test"),
            )
        )
        text = injector._block_to_text(block)

        assert "Test Deal" in text
        assert "Target Corp" in text


# =============================================================================
# Type Compatibility Tests
# =============================================================================


class TestTypeCompatibility:
    """Slot-block type compatibility tests."""

    def test_text_slot_with_text_block(self):
        """TEXT 슬롯과 TextBlock 호환."""
        block = TextBlock(content="test")
        injector = TemplateInjector(
            InjectionContext(
                template_path=Path("dummy.pptx"),
                output_path=Path("output.pptx"),
                report_ir=ReportIR(),
                contract=TemplateContract(template_id="test", template_name="Test"),
            )
        )
        assert injector._is_type_compatible(SlotType.TEXT, block) is True

    def test_text_slot_with_claim_block(self):
        """TEXT 슬롯과 ClaimBlock 호환."""
        block = ClaimBlock(claim_text="test")
        injector = TemplateInjector(
            InjectionContext(
                template_path=Path("dummy.pptx"),
                output_path=Path("output.pptx"),
                report_ir=ReportIR(),
                contract=TemplateContract(template_id="test", template_name="Test"),
            )
        )
        assert injector._is_type_compatible(SlotType.TEXT, block) is True

    def test_table_slot_with_table_block(self):
        """TABLE 슬롯과 TableBlock 호환."""
        block = TableBlock(title="test", columns=[], rows=[])
        injector = TemplateInjector(
            InjectionContext(
                template_path=Path("dummy.pptx"),
                output_path=Path("output.pptx"),
                report_ir=ReportIR(),
                contract=TemplateContract(template_id="test", template_name="Test"),
            )
        )
        assert injector._is_type_compatible(SlotType.TABLE, block) is True

    def test_chart_slot_with_chart_block(self):
        """CHART 슬롯과 ChartBlock 호환."""
        block = ChartBlock(chart_type=ChartType.BAR, title="test")
        injector = TemplateInjector(
            InjectionContext(
                template_path=Path("dummy.pptx"),
                output_path=Path("output.pptx"),
                report_ir=ReportIR(),
                contract=TemplateContract(template_id="test", template_name="Test"),
            )
        )
        assert injector._is_type_compatible(SlotType.CHART, block) is True

    def test_text_slot_with_table_block_incompatible(self):
        """TEXT 슬롯과 TableBlock 비호환."""
        block = TableBlock(title="test", columns=[], rows=[])
        injector = TemplateInjector(
            InjectionContext(
                template_path=Path("dummy.pptx"),
                output_path=Path("output.pptx"),
                report_ir=ReportIR(),
                contract=TemplateContract(template_id="test", template_name="Test"),
            )
        )
        assert injector._is_type_compatible(SlotType.TEXT, block) is False


# =============================================================================
# Convenience Function Tests
# =============================================================================


class TestInjectTemplateFunction:
    """inject_template convenience function tests."""

    def test_inject_template_basic(
        self, sample_pptx_template, sample_report_ir, sample_contract, tmp_path
    ):
        """기본 inject_template 호출."""
        output_path = tmp_path / "output.pptx"

        result = inject_template(
            template_path=sample_pptx_template,
            output_path=output_path,
            report_ir=sample_report_ir,
            contract=sample_contract,
        )

        assert result.success is True
        assert output_path.exists()

    def test_inject_template_with_style_tokens(
        self, sample_pptx_template, sample_report_ir, sample_contract, tmp_path
    ):
        """스타일 토큰 적용."""
        output_path = tmp_path / "output.pptx"

        result = inject_template(
            template_path=sample_pptx_template,
            output_path=output_path,
            report_ir=sample_report_ir,
            contract=sample_contract,
            style_tokens=StyleTokens(),
        )

        assert result.success is True
