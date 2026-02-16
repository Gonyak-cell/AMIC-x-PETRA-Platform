"""Template Contract Tests - 템플릿 계약 스키마 테스트.

EPIC-10 FDD-1001: Template Contract v1 테스트.
"""

import io
import os
import tempfile

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.schemas.template import (
    ColorMapping,
    DetectedSlot,
    FontMapping,
    SizeMapping,
    SlotConstraints,
    SlotStatus,
    SlotType,
    StyleTokens,
    TemplateContract,
    TemplateSlot,
    TemplateValidationResult,
    ValidationIssue,
)
from app.services.template import template_service


# =============================================================================
# Schema Tests
# =============================================================================


class TestSlotType:
    """SlotType enum tests."""

    def test_slot_type_values(self):
        """슬롯 타입 enum 값 확인."""
        assert SlotType.TEXT == "text"
        assert SlotType.TABLE == "table"
        assert SlotType.CHART == "chart"
        assert SlotType.IMAGE == "image"


class TestTemplateSlot:
    """TemplateSlot schema tests."""

    def test_valid_slot_text(self):
        """유효한 TEXT 슬롯 생성."""
        slot = TemplateSlot(
            slot_id="{{SLOT:QOE_SUMMARY}}",
            slot_type=SlotType.TEXT,
            required=True,
            description="QoE 요약 섹션",
        )
        assert slot.slot_id == "{{SLOT:QOE_SUMMARY}}"
        assert slot.slot_type == SlotType.TEXT
        assert slot.required is True

    def test_valid_slot_table(self):
        """유효한 TABLE 슬롯 생성."""
        slot = TemplateSlot(
            slot_id="{{TABLE:QOE_BRIDGE}}",
            slot_type=SlotType.TABLE,
            constraints=SlotConstraints(max_rows=50),
        )
        assert slot.slot_id == "{{TABLE:QOE_BRIDGE}}"
        assert slot.slot_type == SlotType.TABLE
        assert slot.constraints.max_rows == 50

    def test_valid_slot_chart(self):
        """유효한 CHART 슬롯 생성."""
        slot = TemplateSlot(
            slot_id="{{CHART:EBITDA_TREND}}",
            slot_type=SlotType.CHART,
            required=False,
        )
        assert slot.slot_type == SlotType.CHART
        assert slot.required is False

    def test_valid_slot_image(self):
        """유효한 IMAGE 슬롯 생성."""
        slot = TemplateSlot(
            slot_id="{{IMAGE:LOGO}}",
            slot_type=SlotType.IMAGE,
            constraints=SlotConstraints(max_width=2.0, max_height=1.0),
        )
        assert slot.slot_type == SlotType.IMAGE
        assert slot.constraints.max_width == 2.0

    def test_invalid_slot_id_format(self):
        """잘못된 슬롯 ID 형식 거부."""
        with pytest.raises(Exception):  # Pydantic validation error
            TemplateSlot(
                slot_id="INVALID_SLOT_ID",  # Missing {{ }}
                slot_type=SlotType.TEXT,
            )

    def test_slot_constraints(self):
        """슬롯 제약 조건 테스트."""
        constraints = SlotConstraints(
            max_rows=100,
            max_chars=5000,
            max_width=8.0,
            max_height=6.0,
            allowed_formats=["png", "jpg"],
        )
        assert constraints.max_rows == 100
        assert constraints.allowed_formats == ["png", "jpg"]


class TestStyleTokens:
    """StyleTokens schema tests."""

    def test_default_style_tokens(self):
        """기본 스타일 토큰 생성."""
        tokens = StyleTokens()
        assert tokens.colors.primary == "#003366"
        assert tokens.fonts.heading == "Arial Bold"
        assert tokens.sizes.title == 24

    def test_custom_color_mapping(self):
        """커스텀 색상 매핑."""
        colors = ColorMapping(
            primary="#FF0000",
            secondary="#00FF00",
            accent="#0000FF",
        )
        tokens = StyleTokens(colors=colors)
        assert tokens.colors.primary == "#FF0000"

    def test_custom_font_mapping(self):
        """커스텀 폰트 매핑."""
        fonts = FontMapping(
            heading="맑은 고딕 Bold",
            body="맑은 고딕",
        )
        tokens = StyleTokens(fonts=fonts)
        assert tokens.fonts.heading == "맑은 고딕 Bold"

    def test_custom_size_mapping(self):
        """커스텀 크기 매핑."""
        sizes = SizeMapping(title=28, heading1=20, body=11)
        tokens = StyleTokens(sizes=sizes)
        assert tokens.sizes.title == 28


class TestTemplateContract:
    """TemplateContract schema tests."""

    def test_minimal_contract(self):
        """최소 필수 필드만으로 계약 생성."""
        contract = TemplateContract(
            template_id="test_template_001",
            template_name="Test Template",
        )
        assert contract.contract_version == "1.0"
        assert contract.template_id == "test_template_001"
        assert contract.slots == []

    def test_full_contract(self):
        """모든 필드가 있는 계약 생성."""
        slots = [
            TemplateSlot(slot_id="{{SLOT:COVER}}", slot_type=SlotType.TEXT),
            TemplateSlot(slot_id="{{TABLE:QOE_BRIDGE}}", slot_type=SlotType.TABLE),
            TemplateSlot(slot_id="{{CHART:TREND}}", slot_type=SlotType.CHART),
        ]
        contract = TemplateContract(
            template_id="full_template_001",
            template_name="Full Template",
            template_type="pptx",
            slots=slots,
            style_tokens=StyleTokens(),
            validation_rules=["all_slots_required", "style_consistency"],
            metadata={"author": "FDD Team", "version": "1.0"},
        )
        assert len(contract.slots) == 3
        assert contract.template_type == "pptx"
        assert "all_slots_required" in contract.validation_rules

    def test_contract_with_docx_type(self):
        """DOCX 타입 계약 생성."""
        contract = TemplateContract(
            template_id="word_template",
            template_name="Word Template",
            template_type="docx",
        )
        assert contract.template_type == "docx"


# =============================================================================
# Validation Tests
# =============================================================================


class TestContractValidation:
    """Template contract validation tests."""

    def test_validate_valid_contract(self):
        """유효한 계약 검증."""
        contract = TemplateContract(
            template_id="valid_template",
            template_name="Valid Template",
            slots=[
                TemplateSlot(slot_id="{{SLOT:TEST}}", slot_type=SlotType.TEXT, required=True),
            ],
        )
        issues = template_service.validate_template_contract(contract)
        errors = [i for i in issues if i.severity == "error"]
        assert len(errors) == 0

    def test_validate_contract_duplicate_slots(self):
        """중복 슬롯 감지."""
        contract = TemplateContract(
            template_id="dup_template",
            template_name="Duplicate Slots Template",
            slots=[
                TemplateSlot(slot_id="{{SLOT:SAME}}", slot_type=SlotType.TEXT),
                TemplateSlot(slot_id="{{SLOT:SAME}}", slot_type=SlotType.TEXT),  # Duplicate
            ],
        )
        issues = template_service.validate_template_contract(contract)
        assert any(i.code == "DUPLICATE_SLOT_IN_CONTRACT" for i in issues)

    def test_validate_contract_no_required_slots_warning(self):
        """필수 슬롯 없음 경고."""
        contract = TemplateContract(
            template_id="no_required",
            template_name="No Required Slots",
            slots=[
                TemplateSlot(slot_id="{{SLOT:OPTIONAL}}", slot_type=SlotType.TEXT, required=False),
            ],
        )
        issues = template_service.validate_template_contract(contract)
        warnings = [i for i in issues if i.severity == "warning"]
        assert any(i.code == "NO_REQUIRED_SLOTS" for i in warnings)


# =============================================================================
# Slot Detection Tests (PPTX)
# =============================================================================


@pytest.fixture
def sample_pptx_with_slots():
    """슬롯이 있는 샘플 PPTX 파일 생성."""
    prs = Presentation()

    # Slide 1: Cover with text slot
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    textbox1 = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    textbox1.text_frame.text = "{{SLOT:DEAL_NAME}}"

    # Slide 2: QoE table slot
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    textbox2 = slide2.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4))
    textbox2.text_frame.text = "{{TABLE:QOE_BRIDGE}}"

    # Slide 3: Chart slot
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    textbox3 = slide3.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4))
    textbox3.text_frame.text = "{{CHART:EBITDA_TREND}}"

    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
        prs.save(f.name)
        yield f.name
    os.unlink(f.name)


@pytest.fixture
def sample_pptx_no_slots():
    """슬롯이 없는 샘플 PPTX 파일 생성."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    textbox.text_frame.text = "This is a regular slide without slots"

    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
        prs.save(f.name)
        yield f.name
    os.unlink(f.name)


class TestSlotDetection:
    """PPTX slot detection tests."""

    def test_detect_slots_from_pptx(self, sample_pptx_with_slots):
        """PPTX에서 슬롯 탐지."""
        from pathlib import Path

        slots = template_service.detect_slots_from_pptx(Path(sample_pptx_with_slots))

        assert len(slots) == 3
        slot_ids = [s.slot_id for s in slots]
        assert "{{SLOT:DEAL_NAME}}" in slot_ids
        assert "{{TABLE:QOE_BRIDGE}}" in slot_ids
        assert "{{CHART:EBITDA_TREND}}" in slot_ids

    def test_detect_slots_types(self, sample_pptx_with_slots):
        """슬롯 타입 올바르게 감지."""
        from pathlib import Path

        slots = template_service.detect_slots_from_pptx(Path(sample_pptx_with_slots))

        slot_dict = {s.slot_id: s for s in slots}
        assert slot_dict["{{SLOT:DEAL_NAME}}"].slot_type == SlotType.TEXT
        assert slot_dict["{{TABLE:QOE_BRIDGE}}"].slot_type == SlotType.TABLE
        assert slot_dict["{{CHART:EBITDA_TREND}}"].slot_type == SlotType.CHART

    def test_detect_slots_location(self, sample_pptx_with_slots):
        """슬롯 위치 정보 포함."""
        from pathlib import Path

        slots = template_service.detect_slots_from_pptx(Path(sample_pptx_with_slots))

        for slot in slots:
            assert slot.location.startswith("Slide")

    def test_detect_no_slots(self, sample_pptx_no_slots):
        """슬롯 없는 파일에서 빈 리스트 반환."""
        from pathlib import Path

        slots = template_service.detect_slots_from_pptx(Path(sample_pptx_no_slots))
        assert slots == []


# =============================================================================
# File Validation Tests
# =============================================================================


class TestFileValidation:
    """Template file validation tests."""

    def test_validate_valid_pptx(self, sample_pptx_with_slots):
        """유효한 PPTX 파일 검증."""
        from pathlib import Path

        result = template_service.validate_template_file(
            Path(sample_pptx_with_slots), template_type="pptx"
        )

        assert result.is_valid is True
        assert len(result.detected_slots) == 3
        assert result.contract is not None

    def test_validate_pptx_no_slots_warning(self, sample_pptx_no_slots):
        """슬롯 없는 PPTX 파일 경고."""
        from pathlib import Path

        result = template_service.validate_template_file(
            Path(sample_pptx_no_slots), template_type="pptx"
        )

        # 에러는 아니지만 경고는 있음
        assert result.is_valid is True
        warnings = [i for i in result.issues if i.severity == "warning"]
        assert any(i.code == "NO_SLOTS_FOUND" for i in warnings)

    def test_validate_missing_file(self):
        """존재하지 않는 파일 검증 실패."""
        from pathlib import Path

        result = template_service.validate_template_file(
            Path("/nonexistent/file.pptx"), template_type="pptx"
        )

        assert result.is_valid is False
        assert any(i.code == "FILE_NOT_FOUND" for i in result.issues)

    def test_validate_wrong_extension(self, sample_pptx_with_slots):
        """잘못된 확장자 검증 실패."""
        from pathlib import Path

        result = template_service.validate_template_file(
            Path(sample_pptx_with_slots), template_type="docx"  # Wrong type
        )

        assert result.is_valid is False
        assert any(i.code == "INVALID_EXTENSION" for i in result.issues)

    def test_validate_with_expected_slots(self, sample_pptx_with_slots):
        """기대 슬롯 검증."""
        from pathlib import Path

        # 있는 슬롯 기대 - 통과
        result = template_service.validate_template_file(
            Path(sample_pptx_with_slots),
            template_type="pptx",
            expected_slots=["{{SLOT:DEAL_NAME}}"],
        )
        assert result.is_valid is True

        # 없는 슬롯 기대 - 실패
        result = template_service.validate_template_file(
            Path(sample_pptx_with_slots),
            template_type="pptx",
            expected_slots=["{{SLOT:NONEXISTENT}}"],
        )
        assert result.is_valid is False
        assert any(i.code == "MISSING_EXPECTED_SLOT" for i in result.issues)
