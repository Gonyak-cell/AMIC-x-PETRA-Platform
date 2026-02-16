"""Template API Tests - 템플릿 API 엔드포인트 테스트.

EPIC-10: Template API 테스트.
"""

import io
import os
import tempfile
import uuid

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app.schemas.template import (
    SlotType,
    StyleTokens,
    TemplateContract,
    TemplateSlot,
)


# =============================================================================
# Fixtures
# =============================================================================


@pytest.fixture
def sample_pptx_bytes():
    """슬롯이 있는 샘플 PPTX 파일 바이트 생성."""
    prs = Presentation()

    # Slide 1: Cover with text slot
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    textbox1 = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    textbox1.text_frame.text = "{{SLOT:DEAL_NAME}}"

    # Slide 2: QoE table slot
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    textbox2 = slide2.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4))
    textbox2.text_frame.text = "{{TABLE:QOE_BRIDGE}}"

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer.read()


@pytest.fixture
def sample_pptx_no_slots_bytes():
    """슬롯이 없는 샘플 PPTX 파일 바이트."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    textbox.text_frame.text = "Regular slide content"

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer.read()


@pytest.fixture
def sample_contract():
    """샘플 템플릿 계약."""
    return TemplateContract(
        template_id="test_template_001",
        template_name="Test Template",
        template_type="pptx",
        slots=[
            TemplateSlot(
                slot_id="{{SLOT:DEAL_NAME}}",
                slot_type=SlotType.TEXT,
                required=True,
            ),
            TemplateSlot(
                slot_id="{{TABLE:QOE_BRIDGE}}",
                slot_type=SlotType.TABLE,
                required=True,
            ),
        ],
        style_tokens=StyleTokens(),
    )


# =============================================================================
# Validation API Tests
# =============================================================================


class TestValidateTemplateAPI:
    """Template validation API tests."""

    def test_validate_pptx_success(self, client: TestClient, sample_pptx_bytes):
        """PPTX 검증 성공."""
        response = client.post(
            "/api/v1/templates/validate",
            files={"file": ("test.pptx", io.BytesIO(sample_pptx_bytes), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            params={"template_type": "pptx"},
        )
        assert response.status_code == 200
        data = response.json()
        assert data["is_valid"] is True
        assert len(data["detected_slots"]) == 2
        assert data["contract"] is not None

    def test_validate_pptx_no_slots(self, client: TestClient, sample_pptx_no_slots_bytes):
        """슬롯 없는 PPTX 검증."""
        response = client.post(
            "/api/v1/templates/validate",
            files={"file": ("test.pptx", io.BytesIO(sample_pptx_no_slots_bytes), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            params={"template_type": "pptx"},
        )
        assert response.status_code == 200
        data = response.json()
        # 경고는 있지만 여전히 유효
        assert data["is_valid"] is True
        assert any(i["code"] == "NO_SLOTS_FOUND" for i in data["issues"])

    def test_validate_with_expected_slots(self, client: TestClient, sample_pptx_bytes):
        """기대 슬롯 검증."""
        # 있는 슬롯 기대
        response = client.post(
            "/api/v1/templates/validate",
            files={"file": ("test.pptx", io.BytesIO(sample_pptx_bytes), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            params={
                "template_type": "pptx",
                "expected_slots": ["{{SLOT:DEAL_NAME}}"],
            },
        )
        assert response.status_code == 200
        assert response.json()["is_valid"] is True

        # 없는 슬롯 기대
        response = client.post(
            "/api/v1/templates/validate",
            files={"file": ("test.pptx", io.BytesIO(sample_pptx_bytes), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            params={
                "template_type": "pptx",
                "expected_slots": ["{{SLOT:NONEXISTENT}}"],
            },
        )
        assert response.status_code == 200
        assert response.json()["is_valid"] is False


class TestValidateContractAPI:
    """Contract validation API tests."""

    def test_validate_contract_success(self, client: TestClient, sample_contract):
        """계약 검증 성공."""
        response = client.post(
            "/api/v1/templates/validate-contract",
            json=sample_contract.model_dump(),
        )
        assert response.status_code == 200
        issues = response.json()
        errors = [i for i in issues if i["severity"] == "error"]
        assert len(errors) == 0

    def test_validate_contract_duplicate_slots(self, client: TestClient):
        """중복 슬롯 검증 실패."""
        contract = TemplateContract(
            template_id="dup_test",
            template_name="Duplicate Test",
            slots=[
                TemplateSlot(slot_id="{{SLOT:SAME}}", slot_type=SlotType.TEXT),
                TemplateSlot(slot_id="{{SLOT:SAME}}", slot_type=SlotType.TEXT),
            ],
        )
        response = client.post(
            "/api/v1/templates/validate-contract",
            json=contract.model_dump(),
        )
        assert response.status_code == 200
        issues = response.json()
        assert any(i["code"] == "DUPLICATE_SLOT_IN_CONTRACT" for i in issues)


class TestDetectSlotsAPI:
    """Slot detection API tests."""

    def test_detect_slots_pptx(self, client: TestClient, sample_pptx_bytes):
        """PPTX 슬롯 탐지."""
        response = client.post(
            "/api/v1/templates/detect-slots",
            files={"file": ("test.pptx", io.BytesIO(sample_pptx_bytes), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            params={"template_type": "pptx"},
        )
        assert response.status_code == 200
        slots = response.json()
        assert len(slots) == 2
        slot_ids = [s["slot_id"] for s in slots]
        assert "{{SLOT:DEAL_NAME}}" in slot_ids
        assert "{{TABLE:QOE_BRIDGE}}" in slot_ids


# =============================================================================
# CRUD API Tests
# =============================================================================


class TestCreateTemplateAPI:
    """Template create API tests."""

    def test_create_template(self, client: TestClient, sample_contract):
        """템플릿 생성."""
        response = client.post(
            "/api/v1/templates",
            json={
                "template_name": "New Template",
                "template_type": "pptx",
                "contract": sample_contract.model_dump(),
                "description": "Test template description",
            },
        )
        assert response.status_code == 201
        data = response.json()
        assert data["template_name"] == "New Template"
        assert data["status"] == "draft"
        assert len(data["contract"]["slots"]) == 2

    def test_create_template_invalid_contract(self, client: TestClient):
        """잘못된 계약으로 생성 실패."""
        contract = TemplateContract(
            template_id="invalid",
            template_name="Invalid",
            slots=[
                TemplateSlot(slot_id="{{SLOT:SAME}}", slot_type=SlotType.TEXT),
                TemplateSlot(slot_id="{{SLOT:SAME}}", slot_type=SlotType.TEXT),
            ],
        )
        response = client.post(
            "/api/v1/templates",
            json={
                "template_name": "Invalid Template",
                "template_type": "pptx",
                "contract": contract.model_dump(),
            },
        )
        assert response.status_code == 400  # ValidationError


class TestListTemplatesAPI:
    """Template list API tests."""

    def test_list_templates_empty(self, client: TestClient):
        """빈 템플릿 목록."""
        response = client.get("/api/v1/templates")
        assert response.status_code == 200
        assert response.json() == []

    def test_list_templates_with_data(self, client: TestClient, sample_contract):
        """템플릿 목록 조회."""
        # 템플릿 생성
        client.post(
            "/api/v1/templates",
            json={
                "template_name": "Template 1",
                "template_type": "pptx",
                "contract": sample_contract.model_dump(),
            },
        )
        # template_id 중복 방지를 위해 새 contract 생성
        contract2 = sample_contract.model_copy()
        contract2.template_id = "test_template_002"
        client.post(
            "/api/v1/templates",
            json={
                "template_name": "Template 2",
                "template_type": "pptx",
                "contract": contract2.model_dump(),
            },
        )

        response = client.get("/api/v1/templates")
        assert response.status_code == 200
        data = response.json()
        assert len(data) == 2

    def test_list_templates_filter_status(self, client: TestClient, sample_contract):
        """상태별 필터링."""
        # 템플릿 생성
        create_resp = client.post(
            "/api/v1/templates",
            json={
                "template_name": "Draft Template",
                "template_type": "pptx",
                "contract": sample_contract.model_dump(),
            },
        )
        template_id = create_resp.json()["id"]

        # draft 상태 필터
        response = client.get("/api/v1/templates", params={"status": "draft"})
        assert response.status_code == 200
        assert len(response.json()) == 1

        # active 상태 필터 (없음)
        response = client.get("/api/v1/templates", params={"status": "active"})
        assert response.status_code == 200
        assert len(response.json()) == 0


class TestGetTemplateAPI:
    """Template get API tests."""

    def test_get_template(self, client: TestClient, sample_contract):
        """템플릿 조회."""
        # 생성
        create_resp = client.post(
            "/api/v1/templates",
            json={
                "template_name": "Get Test",
                "template_type": "pptx",
                "contract": sample_contract.model_dump(),
            },
        )
        template_uuid = create_resp.json()["id"]

        # 조회
        response = client.get(f"/api/v1/templates/{template_uuid}")
        assert response.status_code == 200
        data = response.json()
        assert data["template_name"] == "Get Test"

    def test_get_template_not_found(self, client: TestClient):
        """존재하지 않는 템플릿 조회."""
        fake_uuid = str(uuid.uuid4())
        response = client.get(f"/api/v1/templates/{fake_uuid}")
        assert response.status_code == 404

    def test_get_template_by_id(self, client: TestClient, sample_contract):
        """template_id로 조회."""
        client.post(
            "/api/v1/templates",
            json={
                "template_name": "By ID Test",
                "template_type": "pptx",
                "contract": sample_contract.model_dump(),
            },
        )

        response = client.get(f"/api/v1/templates/by-id/{sample_contract.template_id}")
        assert response.status_code == 200
        assert response.json()["template_name"] == "By ID Test"


class TestUpdateTemplateAPI:
    """Template update API tests."""

    def test_update_template(self, client: TestClient, sample_contract):
        """템플릿 업데이트."""
        create_resp = client.post(
            "/api/v1/templates",
            json={
                "template_name": "Original Name",
                "template_type": "pptx",
                "contract": sample_contract.model_dump(),
            },
        )
        template_uuid = create_resp.json()["id"]

        # 업데이트
        response = client.patch(
            f"/api/v1/templates/{template_uuid}",
            json={"template_name": "Updated Name", "description": "New description"},
        )
        assert response.status_code == 200
        data = response.json()
        assert data["template_name"] == "Updated Name"
        assert data["description"] == "New description"


class TestDeleteTemplateAPI:
    """Template delete API tests."""

    def test_delete_template(self, client: TestClient, sample_contract):
        """템플릿 삭제."""
        create_resp = client.post(
            "/api/v1/templates",
            json={
                "template_name": "To Delete",
                "template_type": "pptx",
                "contract": sample_contract.model_dump(),
            },
        )
        template_uuid = create_resp.json()["id"]

        # 삭제
        response = client.delete(f"/api/v1/templates/{template_uuid}")
        assert response.status_code == 204

        # 조회 시 404
        response = client.get(f"/api/v1/templates/{template_uuid}")
        assert response.status_code == 404


class TestTemplateSlots:
    """Template slots API tests."""

    def test_get_template_slots(self, client: TestClient, sample_contract):
        """템플릿 슬롯 목록 조회."""
        create_resp = client.post(
            "/api/v1/templates",
            json={
                "template_name": "Slots Test",
                "template_type": "pptx",
                "contract": sample_contract.model_dump(),
            },
        )
        template_uuid = create_resp.json()["id"]

        response = client.get(f"/api/v1/templates/{template_uuid}/slots")
        assert response.status_code == 200
        slots = response.json()
        assert len(slots) == 2


class TestTemplateUploadAPI:
    """Template upload API tests."""

    def test_upload_template(self, client: TestClient, sample_pptx_bytes):
        """템플릿 파일 업로드."""
        response = client.post(
            "/api/v1/templates/upload",
            files={"file": ("upload_test.pptx", io.BytesIO(sample_pptx_bytes), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            params={
                "template_name": "Uploaded Template",
                "description": "Uploaded via API",
                "template_type": "pptx",
            },
        )
        assert response.status_code == 201
        data = response.json()
        assert data["template_name"] == "Uploaded Template"
        assert data["file_path"] is not None
        assert len(data["contract"]["slots"]) == 2

    def test_upload_template_wrong_extension(self, client: TestClient, sample_pptx_bytes):
        """잘못된 확장자로 업로드 실패."""
        response = client.post(
            "/api/v1/templates/upload",
            files={"file": ("test.docx", io.BytesIO(sample_pptx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            params={
                "template_name": "Wrong Extension",
                "template_type": "pptx",
            },
        )
        assert response.status_code == 400


class TestTemplateStatusAPI:
    """Template status change API tests."""

    def test_activate_template(self, client: TestClient, sample_contract):
        """템플릿 활성화."""
        create_resp = client.post(
            "/api/v1/templates",
            json={
                "template_name": "To Activate",
                "template_type": "pptx",
                "contract": sample_contract.model_dump(),
            },
        )
        template_uuid = create_resp.json()["id"]

        response = client.post(f"/api/v1/templates/{template_uuid}/activate")
        assert response.status_code == 200
        assert response.json()["status"] == "active"

    def test_archive_template(self, client: TestClient, sample_contract):
        """템플릿 아카이브."""
        create_resp = client.post(
            "/api/v1/templates",
            json={
                "template_name": "To Archive",
                "template_type": "pptx",
                "contract": sample_contract.model_dump(),
            },
        )
        template_uuid = create_resp.json()["id"]

        response = client.post(f"/api/v1/templates/{template_uuid}/archive")
        assert response.status_code == 200
        assert response.json()["status"] == "archived"
