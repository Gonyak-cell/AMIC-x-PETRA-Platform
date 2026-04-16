"""Schema validation per upload type.

Validates:
1. Required fields present in headers
2. Money columns contain numeric values (not strings)
3. Date columns are parseable
4. TB-specific: debit+credit or balance column must exist
"""

import uuid
from datetime import datetime

from openpyxl import load_workbook

from app.models.upload import UploadType, UploadValidationError, ValidationSeverity
from app.services.ingestion.header_map import REQUIRED_FIELDS
from app.services.ingestion.type_detector import find_header_row

MONEY_FIELDS = {"debit", "credit", "balance", "amount"}
DATE_FIELDS = {"entry_date", "due_date", "maturity_date", "lease_start"}


def validate_upload(
    file_path: str,
    upload_type: UploadType,
    upload_file_id: uuid.UUID | None = None,
    sample_rows: int = 100,
) -> list[UploadValidationError]:
    """Validate an Excel file against the schema for the given type.

    Returns a list of UploadValidationError model instances (not yet persisted).
    """
    errors: list[UploadValidationError] = []
    wb = load_workbook(file_path, read_only=True, data_only=True)
    sheet = wb.active

    # 1. Read and normalize headers
    header_row_idx, raw_headers, normalized = find_header_row(sheet)

    if not raw_headers:
        errors.append(
            UploadValidationError(
                upload_file_id=upload_file_id,
                severity=ValidationSeverity.ERROR,
                error_code="VAL-000",
                message="File has no headers (first row is empty)",
                suggestion="Ensure the first row contains column headers",
            )
        )
        wb.close()
        return errors

    header_set = {h for h in normalized if h is not None}

    # 2. Check required fields
    required = REQUIRED_FIELDS.get(upload_type.value, [])
    for field in required:
        if field not in header_set:
            errors.append(
                UploadValidationError(
                    upload_file_id=upload_file_id,
                    severity=ValidationSeverity.ERROR,
                    error_code="VAL-001",
                    field_name=field,
                    message=f"Required field '{field}' not found in headers",
                    suggestion=f"Add a column with header matching '{field}' (Korean or English)",
                )
            )

    # 3. TB-specific: must have debit+credit or balance
    if upload_type == UploadType.TB:
        has_debit_credit = "debit" in header_set and "credit" in header_set
        has_balance = "balance" in header_set
        if not has_debit_credit and not has_balance:
            errors.append(
                UploadValidationError(
                    upload_file_id=upload_file_id,
                    severity=ValidationSeverity.ERROR,
                    error_code="VAL-002",
                    message="TB must have either debit+credit columns or a balance column",
                    suggestion="Add '차변'/'대변' columns or a '잔액' column",
                )
            )

    # 4. GL-specific: must have debit+credit or amount
    if upload_type == UploadType.GL:
        has_debit_credit = "debit" in header_set and "credit" in header_set
        has_amount = "amount" in header_set
        if not has_debit_credit and not has_amount:
            errors.append(
                UploadValidationError(
                    upload_file_id=upload_file_id,
                    severity=ValidationSeverity.ERROR,
                    error_code="VAL-003",
                    message="GL must have debit+credit columns or an amount column",
                    suggestion="Add '차변'/'대변' columns or a '금액' column",
                )
            )

    # 5. Sample-row validation (check data types in first N rows)
    money_col_indices = [i for i, h in enumerate(normalized) if h in MONEY_FIELDS]
    date_col_indices = [i for i, h in enumerate(normalized) if h in DATE_FIELDS]

    # Track which columns have already warned (one warning per column is enough)
    warned_money_cols: set[int] = set()
    warned_date_cols: set[int] = set()

    row_num = 1
    for row in sheet.iter_rows(
        min_row=header_row_idx + 1,
        max_row=header_row_idx + sample_rows,
        values_only=True,
    ):
        row_num += 1

        # Money field type check
        for col_idx in money_col_indices:
            if col_idx in warned_money_cols:
                continue
            if col_idx < len(row) and row[col_idx] is not None:
                val = row[col_idx]
                if isinstance(val, str):
                    cleaned = val.replace(",", "").replace(" ", "").strip()
                    if cleaned and cleaned != "-":
                        try:
                            float(cleaned)
                        except ValueError:
                            errors.append(
                                UploadValidationError(
                                    upload_file_id=upload_file_id,
                                    severity=ValidationSeverity.WARNING,
                                    error_code="VAL-010",
                                    field_name=normalized[col_idx],
                                    row_number=row_num,
                                    message=f"Money field contains non-numeric string: '{val}'",
                                    suggestion="Ensure all money cells are numeric (remove currency symbols, text)",
                                )
                            )
                            warned_money_cols.add(col_idx)

        # Date field type check
        for col_idx in date_col_indices:
            if col_idx in warned_date_cols:
                continue
            if col_idx < len(row) and row[col_idx] is not None:
                val = row[col_idx]
                if isinstance(val, str):
                    parseable = False
                    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
                        try:
                            datetime.strptime(val.strip(), fmt)
                            parseable = True
                            break
                        except ValueError:
                            continue
                    if not parseable:
                        errors.append(
                            UploadValidationError(
                                upload_file_id=upload_file_id,
                                severity=ValidationSeverity.WARNING,
                                error_code="VAL-011",
                                field_name=normalized[col_idx],
                                row_number=row_num,
                                message=f"Date field has unparseable format: '{val}'",
                                suggestion="Use YYYY-MM-DD, YYYY/MM/DD, or YYYY.MM.DD format",
                            )
                        )
                        warned_date_cols.add(col_idx)

    wb.close()
    return errors
