"""Auto-detect upload file type from Excel headers and sheet names.

Algorithm:
1. Scan the first rows of each sheet for the most likely header row
2. Normalize headers via HEADER_SYNONYMS
3. Score each UploadType by field signature overlap (0.7) + keyword match (0.3)
4. Return highest-scoring type with confidence
"""

from dataclasses import dataclass
from decimal import Decimal

from openpyxl import load_workbook

from app.models.upload import UploadType
from app.services.ingestion.header_map import (
    HEADER_SYNONYMS,
    TYPE_KEYWORDS,
    TYPE_SIGNATURES,
)


@dataclass
class DetectionResult:
    detected_type: UploadType | None
    confidence: Decimal | None
    all_scores: dict[str, float]
    normalized_headers: list[str]


def score_normalized_headers(
    raw_headers: list[str],
    normalized_set: set[str],
    sheet_title: str,
) -> dict[str, float]:
    """Score normalized headers against each supported upload type."""
    scores: dict[str, float] = {}
    for type_name, signature_fields in TYPE_SIGNATURES.items():
        overlap = normalized_set & signature_fields
        field_score = len(overlap) / len(signature_fields) if signature_fields else 0.0

        keyword_score = 0.0
        sheet_name_lower = sheet_title.lower()
        for kw in TYPE_KEYWORDS.get(type_name, []):
            if kw in sheet_name_lower:
                keyword_score = 1.0
                break
            if any(kw in h.lower() for h in raw_headers if h):
                keyword_score = 0.5
                break

        scores[type_name] = field_score * 0.7 + keyword_score * 0.3
    return scores


def normalize_header(raw: str) -> str | None:
    """Normalize a raw header cell to a canonical field name."""
    if not raw:
        return None
    cleaned = str(raw).strip().lower().replace(" ", "_")
    # Try cleaned version first, then original stripped
    return HEADER_SYNONYMS.get(cleaned) or HEADER_SYNONYMS.get(str(raw).strip())


def find_header_row(
    sheet,
    *,
    max_scan_rows: int = 20,
) -> tuple[int, list[str], list[str]]:
    """Return the best header row index, raw headers, and normalized headers."""
    best_row_idx = 1
    best_raw_headers: list[str] = []
    best_normalized: list[str] = []
    best_score = -1.0

    for row_idx, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True),
        start=1,
    ):
        raw_headers = [str(cell) if cell is not None else "" for cell in row]
        if not any(h.strip() for h in raw_headers):
            continue

        normalized = [normalize_header(h) for h in raw_headers]
        normalized_set = {h for h in normalized if h is not None}
        if not normalized_set:
            continue

        scores = score_normalized_headers(raw_headers, normalized_set, sheet.title)
        score = max(scores.values(), default=0.0)
        if score > best_score:
            best_row_idx = row_idx
            best_raw_headers = raw_headers
            best_normalized = normalized
            best_score = score

    return best_row_idx, best_raw_headers, best_normalized


def detect_upload_type(file_path: str) -> DetectionResult:
    """Detect the upload type from an Excel file."""
    wb = load_workbook(file_path, read_only=True, data_only=True)

    best_type: str | None = None
    best_score = 0.0
    all_scores: dict[str, float] = {}
    best_headers: list[str] = []

    for sheet in wb.worksheets:
        header_row_idx, raw_headers, normalized = find_header_row(sheet)
        if not raw_headers:
            continue

        normalized_set = {h for h in normalized if h is not None}
        scores = score_normalized_headers(raw_headers, normalized_set, sheet.title)

        for type_name, score in scores.items():
            all_scores[f"{sheet.title}:row{header_row_idx}:{type_name}"] = round(score, 4)

            if score > best_score:
                best_score = score
                best_type = type_name
                best_headers = [h for h in normalized if h is not None]

    wb.close()

    if best_type and best_score >= 0.3:
        return DetectionResult(
            detected_type=UploadType(best_type),
            confidence=Decimal(str(round(min(best_score, 1.0), 4))),
            all_scores=all_scores,
            normalized_headers=best_headers,
        )

    return DetectionResult(
        detected_type=None,
        confidence=None,
        all_scores=all_scores,
        normalized_headers=best_headers,
    )
