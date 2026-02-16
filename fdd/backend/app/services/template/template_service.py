"""Template Service - 템플릿 관리 서비스.

EPIC-10: PPT/Word 템플릿 검증, 슬롯 탐지, CRUD 기능 제공.
"""

import hashlib
import logging
import re
import uuid
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import ErrorCode
from app.core.exceptions import NotFoundError, ValidationError
from app.models.template import Template, TemplateStatus, TemplateType
from app.schemas.template import (
    DetectedSlot,
    SlotType,
    StyleTokens,
    TemplateContract,
    TemplateCreate,
    TemplateRead,
    TemplateSlot,
    TemplateUpdate,
    TemplateValidationResult,
    ValidationIssue,
)

logger = logging.getLogger(__name__)

# 슬롯 패턴 정규식
SLOT_PATTERN = re.compile(r"\{\{(SLOT|TABLE|CHART|TEXT|IMAGE):([\w_]+)\}\}")

# 슬롯 타입 매핑
SLOT_TYPE_MAP = {
    "SLOT": SlotType.TEXT,  # 기본 슬롯은 TEXT로 처리
    "TABLE": SlotType.TABLE,
    "CHART": SlotType.CHART,
    "TEXT": SlotType.TEXT,
    "IMAGE": SlotType.IMAGE,
}


# =============================================================================
# Slot Detection
# =============================================================================


def detect_slots_from_pptx(file_path: Path) -> list[DetectedSlot]:
    """PPTX 파일에서 슬롯 탐지.

    Args:
        file_path: PPTX 파일 경로

    Returns:
        탐지된 슬롯 목록
    """
    try:
        from pptx import Presentation
    except ImportError:
        logger.warning("python-pptx not installed, returning empty slots")
        return []

    detected: list[DetectedSlot] = []
    prs = Presentation(str(file_path))

    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            text = shape.text_frame.text
            matches = SLOT_PATTERN.findall(text)

            for slot_prefix, slot_name in matches:
                slot_id = f"{{{{{slot_prefix}:{slot_name}}}}}"
                slot_type = SLOT_TYPE_MAP.get(slot_prefix, SlotType.TEXT)

                detected.append(
                    DetectedSlot(
                        slot_id=slot_id,
                        slot_type=slot_type,
                        location=f"Slide {slide_idx}",
                        raw_text=text,
                        shape_id=str(shape.shape_id) if hasattr(shape, "shape_id") else None,
                    )
                )

    return detected


def detect_slots_from_docx(file_path: Path) -> list[DetectedSlot]:
    """DOCX 파일에서 슬롯 탐지.

    Args:
        file_path: DOCX 파일 경로

    Returns:
        탐지된 슬롯 목록
    """
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx not installed, returning empty slots")
        return []

    detected: list[DetectedSlot] = []
    doc = Document(str(file_path))

    for para_idx, para in enumerate(doc.paragraphs, start=1):
        text = para.text
        matches = SLOT_PATTERN.findall(text)

        for slot_prefix, slot_name in matches:
            slot_id = f"{{{{{slot_prefix}:{slot_name}}}}}"
            slot_type = SLOT_TYPE_MAP.get(slot_prefix, SlotType.TEXT)

            detected.append(
                DetectedSlot(
                    slot_id=slot_id,
                    slot_type=slot_type,
                    location=f"Paragraph {para_idx}",
                    raw_text=text,
                    shape_id=None,
                )
            )

    # 테이블 내 슬롯 탐지
    for table_idx, table in enumerate(doc.tables, start=1):
        for row_idx, row in enumerate(table.rows, start=1):
            for cell_idx, cell in enumerate(row.cells, start=1):
                text = cell.text
                matches = SLOT_PATTERN.findall(text)

                for slot_prefix, slot_name in matches:
                    slot_id = f"{{{{{slot_prefix}:{slot_name}}}}}"
                    slot_type = SLOT_TYPE_MAP.get(slot_prefix, SlotType.TEXT)

                    detected.append(
                        DetectedSlot(
                            slot_id=slot_id,
                            slot_type=slot_type,
                            location=f"Table {table_idx}, Row {row_idx}, Cell {cell_idx}",
                            raw_text=text,
                            shape_id=None,
                        )
                    )

    return detected


# =============================================================================
# Validation
# =============================================================================


def validate_template_file(
    file_path: Path,
    template_type: str = "pptx",
    expected_slots: list[str] | None = None,
) -> TemplateValidationResult:
    """템플릿 파일 검증.

    Args:
        file_path: 템플릿 파일 경로
        template_type: 파일 타입 (pptx, docx)
        expected_slots: 기대하는 슬롯 ID 목록 (선택)

    Returns:
        검증 결과
    """
    issues: list[ValidationIssue] = []

    # 파일 존재 확인
    if not file_path.exists():
        return TemplateValidationResult(
            is_valid=False,
            detected_slots=[],
            issues=[
                ValidationIssue(
                    severity="error",
                    code="FILE_NOT_FOUND",
                    message=f"Template file not found: {file_path}",
                )
            ],
        )

    # 파일 확장자 확인
    ext = file_path.suffix.lower()
    if template_type == "pptx" and ext != ".pptx":
        issues.append(
            ValidationIssue(
                severity="error",
                code="INVALID_EXTENSION",
                message=f"Expected .pptx file, got {ext}",
            )
        )
    elif template_type == "docx" and ext != ".docx":
        issues.append(
            ValidationIssue(
                severity="error",
                code="INVALID_EXTENSION",
                message=f"Expected .docx file, got {ext}",
            )
        )

    if issues:
        return TemplateValidationResult(is_valid=False, detected_slots=[], issues=issues)

    # 슬롯 탐지
    if template_type == "pptx":
        detected_slots = detect_slots_from_pptx(file_path)
    else:
        detected_slots = detect_slots_from_docx(file_path)

    # 슬롯 없음 경고
    if not detected_slots:
        issues.append(
            ValidationIssue(
                severity="warning",
                code="NO_SLOTS_FOUND",
                message="No template slots found in the file",
            )
        )

    # 중복 슬롯 검사
    slot_ids = [s.slot_id for s in detected_slots]
    duplicates = [s for s in set(slot_ids) if slot_ids.count(s) > 1]
    for dup in duplicates:
        issues.append(
            ValidationIssue(
                severity="warning",
                code="DUPLICATE_SLOT",
                message=f"Duplicate slot found: {dup}",
                slot_id=dup,
            )
        )

    # 기대 슬롯 검증
    if expected_slots:
        found_ids = set(slot_ids)
        for expected in expected_slots:
            if expected not in found_ids:
                issues.append(
                    ValidationIssue(
                        severity="error",
                        code="MISSING_EXPECTED_SLOT",
                        message=f"Expected slot not found: {expected}",
                        slot_id=expected,
                    )
                )

    # 에러가 있으면 유효하지 않음
    has_errors = any(i.severity == "error" for i in issues)

    # 템플릿 계약 생성 (검증 성공 시)
    contract = None
    if not has_errors:
        template_id = f"template_{uuid.uuid4().hex[:8]}"
        slots = [
            TemplateSlot(
                slot_id=s.slot_id,
                slot_type=s.slot_type,
                required=True,
                description=f"Auto-detected from {s.location}",
            )
            for s in detected_slots
        ]
        # 중복 제거
        unique_slots = {s.slot_id: s for s in slots}
        contract = TemplateContract(
            template_id=template_id,
            template_name=file_path.stem,
            template_type=template_type,
            slots=list(unique_slots.values()),
            style_tokens=StyleTokens(),
            validation_rules=[],
        )

    return TemplateValidationResult(
        is_valid=not has_errors,
        detected_slots=detected_slots,
        issues=issues,
        contract=contract,
    )


def validate_template_contract(contract: TemplateContract) -> list[ValidationIssue]:
    """템플릿 계약 검증.

    Args:
        contract: 검증할 템플릿 계약

    Returns:
        검증 이슈 목록
    """
    issues: list[ValidationIssue] = []

    # 슬롯 ID 형식 검증
    for slot in contract.slots:
        if not SLOT_PATTERN.match(slot.slot_id):
            issues.append(
                ValidationIssue(
                    severity="error",
                    code="INVALID_SLOT_ID_FORMAT",
                    message=f"Invalid slot ID format: {slot.slot_id}",
                    slot_id=slot.slot_id,
                )
            )

    # 중복 슬롯 검사
    slot_ids = [s.slot_id for s in contract.slots]
    if len(slot_ids) != len(set(slot_ids)):
        duplicates = [s for s in set(slot_ids) if slot_ids.count(s) > 1]
        for dup in duplicates:
            issues.append(
                ValidationIssue(
                    severity="error",
                    code="DUPLICATE_SLOT_IN_CONTRACT",
                    message=f"Duplicate slot in contract: {dup}",
                    slot_id=dup,
                )
            )

    # 필수 슬롯 확인 (최소 1개 필요)
    required_slots = [s for s in contract.slots if s.required]
    if contract.slots and not required_slots:
        issues.append(
            ValidationIssue(
                severity="warning",
                code="NO_REQUIRED_SLOTS",
                message="No required slots defined in contract",
            )
        )

    return issues


# =============================================================================
# File Hashing
# =============================================================================


def compute_file_hash(file_path: Path) -> str:
    """파일 해시 계산 (SHA-256).

    Args:
        file_path: 파일 경로

    Returns:
        SHA-256 해시 문자열
    """
    sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            sha256.update(chunk)
    return sha256.hexdigest()


# =============================================================================
# CRUD Operations
# =============================================================================


def create_template(
    db: Session,
    template_data: TemplateCreate,
    file_path: Path | None = None,
    created_by: str = "system",
) -> Template:
    """템플릿 생성.

    Args:
        db: DB 세션
        template_data: 생성 데이터
        file_path: 템플릿 파일 경로 (선택)
        created_by: 생성자

    Returns:
        생성된 Template 모델
    """
    # 계약 검증
    issues = validate_template_contract(template_data.contract)
    errors = [i for i in issues if i.severity == "error"]
    if errors:
        raise ValidationError(
            ErrorCode.TEMPLATE_CONTRACT_INVALID,
            f"Template contract validation failed: {errors[0].message}",
        )

    # 파일 해시 계산
    file_hash = None
    if file_path and file_path.exists():
        file_hash = compute_file_hash(file_path)

    template = Template(
        template_id=template_data.contract.template_id,
        template_name=template_data.template_name,
        template_type=TemplateType(template_data.template_type),
        file_path=str(file_path) if file_path else None,
        file_hash=file_hash,
        contract=template_data.contract.model_dump(),
        status=TemplateStatus.DRAFT,
        description=template_data.description,
        created_by=created_by,
    )

    db.add(template)
    db.flush()
    logger.info(f"Created template: {template.template_id} (id={template.id})")

    return template


def get_template(db: Session, template_uuid: uuid.UUID) -> Template:
    """템플릿 조회 (UUID).

    Args:
        db: DB 세션
        template_uuid: 템플릿 UUID

    Returns:
        Template 모델

    Raises:
        NotFoundError: 템플릿이 없을 경우
    """
    template = db.get(Template, template_uuid)
    if not template:
        raise NotFoundError(resource="Template", resource_id=str(template_uuid))
    return template


def get_template_by_template_id(db: Session, template_id: str) -> Template:
    """템플릿 조회 (template_id).

    Args:
        db: DB 세션
        template_id: 템플릿 ID

    Returns:
        Template 모델

    Raises:
        NotFoundError: 템플릿이 없을 경우
    """
    stmt = select(Template).where(Template.template_id == template_id)
    template = db.scalars(stmt).first()
    if not template:
        raise NotFoundError(resource="Template", resource_id=template_id)
    return template


def list_templates(
    db: Session,
    status: TemplateStatus | None = None,
    template_type: str | None = None,
    limit: int = 100,
    offset: int = 0,
) -> list[Template]:
    """템플릿 목록 조회.

    Args:
        db: DB 세션
        status: 상태 필터 (선택)
        template_type: 타입 필터 (선택)
        limit: 최대 결과 수
        offset: 시작 오프셋

    Returns:
        Template 목록
    """
    stmt = select(Template).order_by(Template.created_at.desc())

    if status:
        stmt = stmt.where(Template.status == status)
    if template_type:
        stmt = stmt.where(Template.template_type == TemplateType(template_type))

    stmt = stmt.limit(limit).offset(offset)
    return list(db.scalars(stmt).all())


def update_template(
    db: Session,
    template_uuid: uuid.UUID,
    update_data: TemplateUpdate,
) -> Template:
    """템플릿 업데이트.

    Args:
        db: DB 세션
        template_uuid: 템플릿 UUID
        update_data: 업데이트 데이터

    Returns:
        업데이트된 Template 모델
    """
    template = get_template(db, template_uuid)

    if update_data.template_name is not None:
        template.template_name = update_data.template_name
    if update_data.description is not None:
        template.description = update_data.description
    if update_data.status is not None:
        template.status = TemplateStatus(update_data.status.value)
    if update_data.style_tokens is not None:
        # 기존 contract에 스타일 토큰만 업데이트
        contract_dict = template.contract.copy()
        contract_dict["style_tokens"] = update_data.style_tokens.model_dump()
        template.contract = contract_dict

    db.flush()
    logger.info(f"Updated template: {template.template_id}")

    return template


def delete_template(db: Session, template_uuid: uuid.UUID) -> None:
    """템플릿 삭제.

    Args:
        db: DB 세션
        template_uuid: 템플릿 UUID
    """
    template = get_template(db, template_uuid)
    db.delete(template)
    db.flush()
    logger.info(f"Deleted template: {template.template_id}")


def template_to_read(template: Template) -> TemplateRead:
    """Template 모델을 TemplateRead 스키마로 변환.

    Args:
        template: Template 모델

    Returns:
        TemplateRead 스키마
    """
    return TemplateRead(
        id=template.id,
        template_id=template.template_id,
        template_name=template.template_name,
        template_type=template.template_type.value,
        file_path=template.file_path,
        contract=TemplateContract(**template.contract),
        status=template.status,
        description=template.description,
        created_by=template.created_by,
        created_at=template.created_at,
        updated_at=template.updated_at,
    )
