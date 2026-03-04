"""Word 보고서 렌더러.

Report IR을 Word(DOCX) 문서로 변환합니다.
docxtpl 템플릿 엔진을 사용하여 Jinja2 문법 기반 템플릿을 렌더링합니다.
"""

from decimal import Decimal
from enum import Enum
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docxtpl import DocxTemplate

from app.renderers.design_system import get_colors, get_fonts
from app.renderers.report_builder import (
    AlignType,
    AppendixBlock,
    ChartBlock,
    ClaimBlock,
    CoverBlock,
    IssueBlock,
    KPIBlock,
    MethodologyBlock,
    ReportBlock,
    ReportIR,
    RiskLevel,
    ScopeBlock,
    TableBlock,
    TextBlock,
)

WORD_RENDERER_VERSION = "0.1.0"


# =============================================================================
# Word Styles
# =============================================================================


class WordTableStyle(str, Enum):
    """Word 표 스타일."""

    HEADER = "header"  # 진한 배경, 흰색 글자
    BODY = "body"  # 흰색 배경, 검은 글자
    TOTAL = "total"  # 밑줄 강조
    SUBTOTAL = "subtotal"  # 회색 배경
    SEPARATOR = "separator"  # 빈 행


def _get_word_styles() -> dict[str, Any]:
    """Word 스타일 설정을 반환합니다."""
    colors = get_colors()
    fonts = get_fonts()

    return {
        "primary_color": colors.get("primary", "#003366"),
        "positive_color": colors.get("positive", "#2E7D32"),
        "negative_color": colors.get("negative", "#E0301E"),
        "neutral_color": colors.get("neutral", "#666666"),
        "background_light": colors.get("background_light", "#F5F5F5"),
        "heading_font": fonts.get("heading", "맑은 고딕"),
        "body_font": fonts.get("body", "맑은 고딕"),
        "title1_size": Pt(18),
        "title2_size": Pt(14),
        "body_size": Pt(10),
        "table_header_size": Pt(9),
        "table_body_size": Pt(9),
    }


def _hex_to_rgb(hex_color: str) -> RGBColor:
    """HEX 색상을 RGBColor로 변환합니다."""
    hex_color = hex_color.lstrip("#")
    r, g, b = int(hex_color[:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
    return RGBColor(r, g, b)


def _format_currency(value: str | Decimal | None) -> str:
    """금액을 포맷팅합니다."""
    if value is None:
        return "-"
    if isinstance(value, str):
        try:
            value = Decimal(value)
        except Exception:
            return value
    return f"{value:,.0f}"


def _format_percentage(value: str | Decimal | None) -> str:
    """백분율을 포맷팅합니다."""
    if value is None:
        return "-"
    if isinstance(value, str):
        try:
            value = Decimal(value)
        except Exception:
            return value
    return f"{value:.1%}"


# =============================================================================
# Table Rendering
# =============================================================================


def _set_cell_shading(cell: Any, hex_color: str) -> None:
    """셀 배경색을 설정합니다."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color.lstrip("#")}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def _apply_table_style(
    table: Table, style: WordTableStyle = WordTableStyle.BODY
) -> None:
    """테이블에 스타일을 적용합니다."""
    styles = _get_word_styles()

    # 테이블 기본 설정
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # 헤더 행 스타일
    if len(table.rows) > 0:
        header_row = table.rows[0]
        for cell in header_row.cells:
            _set_cell_shading(cell, styles["primary_color"])
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
                    run.font.size = styles["table_header_size"]
                    run.font.name = styles["body_font"]


def _render_table_block(doc: Document, block: TableBlock) -> None:
    """TableBlock을 Word 테이블로 렌더링합니다."""
    styles = _get_word_styles()

    # 제목
    if block.title:
        heading = doc.add_heading(block.title, level=2)
        for run in heading.runs:
            run.font.color.rgb = _hex_to_rgb(styles["primary_color"])
            run.font.name = styles["heading_font"]

    # 테이블 생성
    col_count = len(block.columns)
    row_count = (
        len(block.rows) + (1 if block.show_header else 0) + len(block.footer_rows)
    )

    if col_count == 0 or row_count == 0:
        return

    table = doc.add_table(rows=row_count, cols=col_count)
    table.style = "Table Grid"

    # 헤더 행
    row_idx = 0
    if block.show_header:
        header_row = table.rows[row_idx]
        for col_idx, col_def in enumerate(block.columns):
            cell = header_row.cells[col_idx]
            cell.text = col_def.header
        row_idx += 1

    # 데이터 행
    for row_data in block.rows:
        data_row = table.rows[row_idx]
        for col_idx, col_def in enumerate(block.columns):
            cell = data_row.cells[col_idx]
            raw_value = row_data.get(col_def.key, "")

            # 포맷팅
            if col_def.format == "currency":
                cell.text = _format_currency(raw_value)
            elif col_def.format == "percentage":
                cell.text = _format_percentage(raw_value)
            else:
                cell.text = str(raw_value) if raw_value is not None else ""

            # 정렬
            for paragraph in cell.paragraphs:
                if col_def.align == AlignType.RIGHT:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif col_def.align == AlignType.CENTER:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        row_idx += 1

    # 푸터 행
    for footer_data in block.footer_rows:
        footer_row = table.rows[row_idx]
        for col_idx, col_def in enumerate(block.columns):
            cell = footer_row.cells[col_idx]
            raw_value = footer_data.get(col_def.key, "")
            if col_def.format == "currency":
                cell.text = _format_currency(raw_value)
            else:
                cell.text = str(raw_value) if raw_value is not None else ""
            # 굵게
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
        row_idx += 1

    # 스타일 적용
    _apply_table_style(table)

    # 빈 줄 추가
    doc.add_paragraph()


# =============================================================================
# Block Renderers
# =============================================================================


def _render_cover_block(doc: Document, block: CoverBlock) -> None:
    """CoverBlock을 렌더링합니다."""
    styles = _get_word_styles()

    # 표지 제목
    title = doc.add_heading(block.deal_name or "FDD Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = _hex_to_rgb(styles["primary_color"])
        run.font.size = Pt(28)

    # 부제목 (대상 회사)
    if block.target_name:
        subtitle = doc.add_paragraph(block.target_name)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in subtitle.runs:
            run.font.size = Pt(16)
            run.font.color.rgb = _hex_to_rgb(styles["neutral_color"])

    # 딜 타입
    if block.deal_type:
        deal_type = doc.add_paragraph(block.deal_type)
        deal_type.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 날짜
    if block.date:
        date_para = doc.add_paragraph(block.date.strftime("%Y년 %m월 %d일"))
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 작성자
    if block.prepared_by:
        prepared = doc.add_paragraph(f"Prepared by: {block.prepared_by}")
        prepared.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 기밀 표시
    if block.confidentiality:
        conf = doc.add_paragraph(block.confidentiality)
        conf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in conf.runs:
            run.font.bold = True
            run.font.color.rgb = _hex_to_rgb(styles["negative_color"])

    doc.add_page_break()


def _render_kpi_block(doc: Document, block: KPIBlock) -> None:
    """KPIBlock을 렌더링합니다."""
    styles = _get_word_styles()

    if block.title:
        heading = doc.add_heading(block.title, level=2)
        for run in heading.runs:
            run.font.color.rgb = _hex_to_rgb(styles["primary_color"])

    # KPI 테이블로 렌더링
    if block.kpis:
        table = doc.add_table(rows=2, cols=len(block.kpis))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # 레이블 행
        for col_idx, kpi in enumerate(block.kpis):
            cell = table.rows[0].cells[col_idx]
            cell.text = kpi.get("label", "")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = _hex_to_rgb(styles["neutral_color"])

        # 값 행
        for col_idx, kpi in enumerate(block.kpis):
            cell = table.rows[1].cells[col_idx]
            value = kpi.get("value", "")
            unit = kpi.get("unit", "")
            cell.text = f"{value} {unit}".strip()
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(18)
                    run.font.bold = True
                    run.font.color.rgb = _hex_to_rgb(styles["primary_color"])

    doc.add_paragraph()


def _render_chart_block(doc: Document, block: ChartBlock) -> None:
    """ChartBlock을 렌더링합니다 (이미지 삽입)."""
    styles = _get_word_styles()

    if block.title:
        heading = doc.add_heading(block.title, level=2)
        for run in heading.runs:
            run.font.color.rgb = _hex_to_rgb(styles["primary_color"])

    # base64 이미지가 있으면 삽입
    if block.image_base64:
        import base64

        image_data = base64.b64decode(block.image_base64)
        image_stream = BytesIO(image_data)
        doc.add_picture(image_stream, width=Inches(6.0))
    else:
        # 이미지가 없으면 플레이스홀더 텍스트
        placeholder = doc.add_paragraph("[Chart: Image not available]")
        placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()


def _render_text_block(doc: Document, block: TextBlock) -> None:
    """TextBlock을 렌더링합니다."""
    styles = _get_word_styles()

    if block.title:
        heading = doc.add_heading(block.title, level=2)
        for run in heading.runs:
            run.font.color.rgb = _hex_to_rgb(styles["primary_color"])

    # 리스크 레벨 표시
    if block.risk_level:
        risk_para = doc.add_paragraph()
        risk_run = risk_para.add_run(f"[{block.risk_level.value.upper()}]")
        if block.risk_level == RiskLevel.HIGH:
            risk_run.font.color.rgb = _hex_to_rgb(styles["negative_color"])
        elif block.risk_level == RiskLevel.MEDIUM:
            risk_run.font.color.rgb = RGBColor(255, 165, 0)  # Orange
        else:
            risk_run.font.color.rgb = _hex_to_rgb(styles["positive_color"])
        risk_run.font.bold = True

    # 본문 텍스트
    if block.content:
        para = doc.add_paragraph(block.content)
        if block.highlight:
            for run in para.runs:
                run.font.bold = True

    # 불릿 포인트
    for point in block.bullet_points:
        doc.add_paragraph(point, style="List Bullet")

    doc.add_paragraph()


def _render_claim_block(doc: Document, block: ClaimBlock) -> None:
    """ClaimBlock을 렌더링합니다."""
    styles = _get_word_styles()

    # 검증 상태 표시
    status_text = "✓ Verified" if block.verified else "⚠ Unverified"
    status_color = (
        styles["positive_color"] if block.verified else styles["negative_color"]
    )

    # 카테고리 표시
    if block.category:
        cat_para = doc.add_paragraph()
        cat_run = cat_para.add_run(f"[{block.category}] ")
        cat_run.font.bold = True
        cat_run.font.color.rgb = _hex_to_rgb(styles["primary_color"])

        status_run = cat_para.add_run(status_text)
        status_run.font.color.rgb = _hex_to_rgb(status_color)

    # Claim 텍스트
    claim_para = doc.add_paragraph(block.claim_text)

    # 근거 참조
    if block.evidence_refs:
        doc.add_paragraph("Evidence:", style="List Bullet")
        for ref in block.evidence_refs:
            ref_text = f"{ref.source_type}: {ref.source_id}"
            if ref.description:
                ref_text += f" - {ref.description}"
            doc.add_paragraph(ref_text, style="List Bullet 2")

    doc.add_paragraph()


def _render_issue_block(doc: Document, block: IssueBlock) -> None:
    """IssueBlock을 렌더링합니다."""
    styles = _get_word_styles()

    if block.title:
        heading = doc.add_heading(block.title, level=2)
        for run in heading.runs:
            run.font.color.rgb = _hex_to_rgb(styles["primary_color"])

    # 이슈 테이블
    if block.issues:
        table = doc.add_table(rows=len(block.issues) + 1, cols=5)
        table.style = "Table Grid"

        # 헤더
        headers = ["ID", "Category", "Severity", "Issue", "Status"]
        for col_idx, header in enumerate(headers):
            table.rows[0].cells[col_idx].text = header

        # 데이터
        for row_idx, issue in enumerate(block.issues, start=1):
            if not block.show_resolved and issue.status == "resolved":
                continue
            table.rows[row_idx].cells[0].text = issue.issue_id
            table.rows[row_idx].cells[1].text = issue.category
            table.rows[row_idx].cells[2].text = issue.severity
            table.rows[row_idx].cells[3].text = issue.title
            table.rows[row_idx].cells[4].text = issue.status

        _apply_table_style(table)

    doc.add_paragraph()


def _render_methodology_block(doc: Document, block: MethodologyBlock) -> None:
    """MethodologyBlock을 렌더링합니다."""
    styles = _get_word_styles()

    if block.title:
        heading = doc.add_heading(block.title, level=2)
        for run in heading.runs:
            run.font.color.rgb = _hex_to_rgb(styles["primary_color"])

    # 소개
    if block.introduction:
        doc.add_paragraph(block.introduction)

    # 단계
    for step in block.steps:
        step_para = doc.add_paragraph()
        step_run = step_para.add_run(f"Step {step.step}: {step.title}")
        step_run.font.bold = True
        if step.description:
            doc.add_paragraph(step.description)

    # 제한 사항
    if block.limitations:
        doc.add_heading("Limitations", level=3)
        for limitation in block.limitations:
            doc.add_paragraph(limitation, style="List Bullet")

    doc.add_paragraph()


def _render_scope_block(doc: Document, block: ScopeBlock) -> None:
    """ScopeBlock을 렌더링합니다."""
    styles = _get_word_styles()

    if block.title:
        heading = doc.add_heading(block.title, level=2)
        for run in heading.runs:
            run.font.color.rgb = _hex_to_rgb(styles["primary_color"])

    # 범위 항목 테이블
    if block.scope_items:
        table = doc.add_table(rows=len(block.scope_items), cols=2)
        table.style = "Table Grid"
        for row_idx, item in enumerate(block.scope_items):
            table.rows[row_idx].cells[0].text = item.label
            table.rows[row_idx].cells[1].text = item.value
            for paragraph in table.rows[row_idx].cells[0].paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

    # 정의
    if block.definitions:
        doc.add_heading("Definitions", level=3)
        for term, definition in block.definitions.items():
            para = doc.add_paragraph()
            term_run = para.add_run(f"{term}: ")
            term_run.font.bold = True
            para.add_run(definition)

    doc.add_paragraph()


def _render_appendix_block(doc: Document, block: AppendixBlock) -> None:
    """AppendixBlock을 렌더링합니다."""
    styles = _get_word_styles()

    doc.add_page_break()

    if block.title:
        heading = doc.add_heading(block.title, level=1)
        for run in heading.runs:
            run.font.color.rgb = _hex_to_rgb(styles["primary_color"])

    for item in block.items:
        # 항목 제목
        item_heading = doc.add_heading(item.title, level=2)

        # 내용
        if item.content:
            doc.add_paragraph(item.content)

        # 테이블 데이터
        if item.table_data:
            if len(item.table_data) > 0:
                keys = list(item.table_data[0].keys())
                table = doc.add_table(rows=len(item.table_data) + 1, cols=len(keys))
                table.style = "Table Grid"

                # 헤더
                for col_idx, key in enumerate(keys):
                    table.rows[0].cells[col_idx].text = key

                # 데이터
                for row_idx, row_data in enumerate(item.table_data, start=1):
                    for col_idx, key in enumerate(keys):
                        table.rows[row_idx].cells[col_idx].text = str(
                            row_data.get(key, "")
                        )

                _apply_table_style(table)


# =============================================================================
# Main Renderer Functions
# =============================================================================


def _render_block(doc: Document, block: ReportBlock) -> None:
    """블록 타입에 따라 적절한 렌더러를 호출합니다."""
    if isinstance(block, CoverBlock):
        _render_cover_block(doc, block)
    elif isinstance(block, KPIBlock):
        _render_kpi_block(doc, block)
    elif isinstance(block, TableBlock):
        _render_table_block(doc, block)
    elif isinstance(block, ChartBlock):
        _render_chart_block(doc, block)
    elif isinstance(block, TextBlock):
        _render_text_block(doc, block)
    elif isinstance(block, ClaimBlock):
        _render_claim_block(doc, block)
    elif isinstance(block, IssueBlock):
        _render_issue_block(doc, block)
    elif isinstance(block, MethodologyBlock):
        _render_methodology_block(doc, block)
    elif isinstance(block, ScopeBlock):
        _render_scope_block(doc, block)
    elif isinstance(block, AppendixBlock):
        _render_appendix_block(doc, block)


def render_word_report(report_ir: ReportIR, output_path: Path | None = None) -> BytesIO:
    """Report IR을 Word 문서로 렌더링합니다.

    Args:
        report_ir: ReportIR 인스턴스
        output_path: 저장할 파일 경로 (None이면 BytesIO 반환만)

    Returns:
        Word 문서 BytesIO 버퍼
    """
    # 새 문서 생성
    doc = Document()

    # 문서 스타일 설정
    styles = _get_word_styles()

    # 기본 폰트 설정
    style = doc.styles["Normal"]
    font = style.font
    font.name = styles["body_font"]
    font.size = styles["body_size"]

    # 각 블록 렌더링
    for block in report_ir.sections:
        _render_block(doc, block)

    # 저장
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    if output_path:
        with open(output_path, "wb") as f:
            f.write(buffer.getvalue())
        buffer.seek(0)

    return buffer


def render_word_from_template(
    template_path: Path,
    context: dict[str, Any],
    output_path: Path | None = None,
) -> BytesIO:
    """docxtpl 템플릿으로 Word 문서를 렌더링합니다.

    Args:
        template_path: 템플릿 파일 경로
        context: 템플릿 변수 딕셔너리
        output_path: 저장할 파일 경로 (None이면 BytesIO 반환만)

    Returns:
        Word 문서 BytesIO 버퍼
    """
    doc = DocxTemplate(template_path)
    doc.render(context)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    if output_path:
        with open(output_path, "wb") as f:
            f.write(buffer.getvalue())
        buffer.seek(0)

    return buffer


def build_docx_context(report_ir: ReportIR) -> dict[str, Any]:
    """Report IR을 docxtpl 템플릿 컨텍스트로 변환합니다.

    Args:
        report_ir: ReportIR 인스턴스

    Returns:
        템플릿 컨텍스트 딕셔너리
    """
    context: dict[str, Any] = {
        "metadata": {
            "deal_id": report_ir.metadata.deal_id,
            "deal_name": report_ir.metadata.deal_name,
            "generated_at": report_ir.metadata.generated_at,
            "version": report_ir.metadata.version,
            "engine_versions": report_ir.metadata.engine_versions,
        },
        "sections": [],
        "cover": None,
        "kpis": [],
        "tables": [],
        "charts": [],
        "claims": [],
        "issues": [],
    }

    for block in report_ir.sections:
        if isinstance(block, CoverBlock):
            context["cover"] = {
                "deal_name": block.deal_name,
                "deal_type": block.deal_type,
                "target_name": block.target_name,
                "date": block.date.strftime("%Y년 %m월 %d일") if block.date else "",
                "prepared_by": block.prepared_by,
                "confidentiality": block.confidentiality,
            }
        elif isinstance(block, KPIBlock):
            context["kpis"].append(
                {
                    "title": block.title,
                    "items": block.kpis,
                }
            )
        elif isinstance(block, TableBlock):
            context["tables"].append(
                {
                    "title": block.title,
                    "columns": [
                        {"key": c.key, "header": c.header} for c in block.columns
                    ],
                    "rows": block.rows,
                    "footer_rows": block.footer_rows,
                }
            )
        elif isinstance(block, ClaimBlock):
            context["claims"].append(
                {
                    "text": block.claim_text,
                    "verified": block.verified,
                    "category": block.category,
                    "evidence_refs": [
                        {
                            "source_type": ref.source_type,
                            "source_id": ref.source_id,
                            "description": ref.description,
                        }
                        for ref in block.evidence_refs
                    ],
                }
            )
        elif isinstance(block, IssueBlock):
            context["issues"].append(
                {
                    "title": block.title,
                    "items": [
                        {
                            "id": issue.issue_id,
                            "category": issue.category,
                            "severity": issue.severity,
                            "title": issue.title,
                            "status": issue.status,
                        }
                        for issue in block.issues
                    ],
                }
            )

        # 모든 섹션을 리스트에 추가 (타입 정보 포함)
        context["sections"].append(
            {
                "type": block.type.value if hasattr(block, "type") else "unknown",
                "data": _block_to_dict(block),
            }
        )

    return context


def _block_to_dict(block: ReportBlock) -> dict[str, Any]:
    """블록을 딕셔너리로 변환합니다."""
    result: dict[str, Any] = {}

    for field_name in block.__dataclass_fields__:
        value = getattr(block, field_name)
        if isinstance(value, Enum):
            result[field_name] = value.value
        elif isinstance(value, Decimal):
            result[field_name] = str(value)
        elif hasattr(value, "__dataclass_fields__"):
            result[field_name] = _block_to_dict(value)  # type: ignore
        elif isinstance(value, list):
            result[field_name] = [
                _block_to_dict(item) if hasattr(item, "__dataclass_fields__") else item  # type: ignore
                for item in value
            ]
        else:
            result[field_name] = value

    return result
