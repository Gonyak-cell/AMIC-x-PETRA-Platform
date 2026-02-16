"""Template Injector - PPT/Word 템플릿에 콘텐츠 주입.

EPIC-10 FDD-1002: 고객 PPT/Word 템플릿에 FDD 콘텐츠를 주입합니다.
슬롯 위치에 Report IR 블록을 매핑하여 콘텐츠를 채워 넣습니다.
"""

import logging
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

from app.renderers.report_builder import (
    BlockType,
    ChartBlock,
    ClaimBlock,
    CoverBlock,
    IssueBlock,
    KPIBlock,
    MethodologyBlock,
    ReportBlock,
    ReportIR,
    ScopeBlock,
    TableBlock,
    TextBlock,
)
from app.schemas.template import (
    SlotBlockMapping,
    SlotStatus,
    SlotType,
    StyleTokens,
    TemplateContract,
)

logger = logging.getLogger(__name__)

# 슬롯 패턴 정규식
SLOT_PATTERN = re.compile(r"\{\{(SLOT|TABLE|CHART|TEXT|IMAGE):([\w_]+)\}\}")


@dataclass
class InjectionContext:
    """주입 컨텍스트."""

    template_path: Path
    output_path: Path
    report_ir: ReportIR
    contract: TemplateContract
    slot_mappings: dict[str, str] = field(default_factory=dict)  # slot_id → block_id
    style_tokens: StyleTokens | None = None


@dataclass
class InjectionResult:
    """주입 결과."""

    success: bool
    output_path: Path | None = None
    filled_slots: list[str] = field(default_factory=list)
    empty_slots: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


class TemplateInjector:
    """PPT/Word 템플릿 주입기.

    고객 템플릿의 슬롯 위치에 Report IR 블록을 주입합니다.
    """

    def __init__(self, context: InjectionContext):
        """초기화.

        Args:
            context: 주입 컨텍스트
        """
        self.context = context
        self._slot_to_block: dict[str, ReportBlock | None] = {}
        self._errors: list[str] = []
        self._warnings: list[str] = []

    def inject(self) -> InjectionResult:
        """템플릿에 콘텐츠 주입.

        Returns:
            주입 결과
        """
        # 슬롯-블록 매핑 구축
        self._build_slot_block_mapping()

        # 템플릿 타입에 따라 주입
        template_type = self.context.contract.template_type
        if template_type == "pptx":
            return self._inject_pptx()
        elif template_type == "docx":
            return self._inject_docx()
        else:
            return InjectionResult(
                success=False,
                errors=[f"Unsupported template type: {template_type}"],
            )

    def _build_slot_block_mapping(self) -> None:
        """슬롯-블록 매핑 구축."""
        # Report IR에서 블록 맵 생성
        block_map: dict[str, ReportBlock] = {}
        for block in self.context.report_ir.sections:
            block_id = self._get_block_id(block)
            if block_id:
                block_map[block_id] = block

        # 각 슬롯에 대해 블록 매핑
        for slot in self.context.contract.slots:
            slot_id = slot.slot_id

            # 커스텀 매핑 우선
            if slot_id in self.context.slot_mappings:
                block_id = self.context.slot_mappings[slot_id]
                self._slot_to_block[slot_id] = block_map.get(block_id)
            # 기본 매핑 사용
            elif slot.default_block_id and slot.default_block_id in block_map:
                self._slot_to_block[slot_id] = block_map[slot.default_block_id]
            # 슬롯 이름으로 자동 매핑 시도
            else:
                block = self._auto_match_block(slot_id, slot.slot_type, block_map)
                self._slot_to_block[slot_id] = block

    def _get_block_id(self, block: ReportBlock) -> str | None:
        """블록에서 ID 추출."""
        if hasattr(block, "title") and block.title:
            # 타이틀을 ID로 사용
            return block.title.lower().replace(" ", "_")
        # 타입 기반 기본 ID
        return f"{block.type.value}_{id(block)}"

    def _auto_match_block(
        self,
        slot_id: str,
        slot_type: SlotType,
        block_map: dict[str, ReportBlock],
    ) -> ReportBlock | None:
        """슬롯과 블록 자동 매칭.

        슬롯 이름과 블록 ID/타이틀을 비교하여 자동 매칭합니다.
        """
        # 슬롯 ID에서 이름 추출 (예: {{TABLE:QOE_BRIDGE}} → qoe_bridge)
        match = SLOT_PATTERN.match(slot_id)
        if not match:
            return None

        slot_name = match.group(2).lower()

        # 블록 이름에서 매칭 시도
        for block_id, block in block_map.items():
            if slot_name in block_id.lower():
                # 타입 호환성 검사
                if self._is_type_compatible(slot_type, block):
                    return block

        return None

    def _is_type_compatible(self, slot_type: SlotType, block: ReportBlock) -> bool:
        """슬롯 타입과 블록 타입 호환성 검사."""
        type_map = {
            SlotType.TEXT: (BlockType.TEXT, BlockType.CLAIM, BlockType.SCOPE, BlockType.METHODOLOGY),
            SlotType.TABLE: (BlockType.TABLE, BlockType.ISSUE),
            SlotType.CHART: (BlockType.CHART,),
            SlotType.IMAGE: (BlockType.CHART, BlockType.COVER),  # 차트 이미지 포함
        }
        allowed_types = type_map.get(slot_type, ())
        return block.type in allowed_types

    def _inject_pptx(self) -> InjectionResult:
        """PPTX 템플릿에 주입."""
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
        except ImportError:
            return InjectionResult(
                success=False,
                errors=["python-pptx not installed"],
            )

        filled_slots: list[str] = []
        empty_slots: list[str] = []

        try:
            prs = Presentation(str(self.context.template_path))

            # 각 슬라이드의 각 도형에서 슬롯 찾기
            for slide in prs.slides:
                for shape in slide.shapes:
                    if not shape.has_text_frame:
                        continue

                    text = shape.text_frame.text
                    matches = SLOT_PATTERN.findall(text)

                    for slot_prefix, slot_name in matches:
                        slot_id = f"{{{{{slot_prefix}:{slot_name}}}}}"
                        block = self._slot_to_block.get(slot_id)

                        if block:
                            # 블록 콘텐츠로 교체
                            new_text = self._block_to_text(block)
                            self._replace_slot_text(shape, slot_id, new_text)
                            filled_slots.append(slot_id)
                        else:
                            # 슬롯 비워두기 (또는 플레이스홀더 유지)
                            empty_slots.append(slot_id)

            # 저장
            prs.save(str(self.context.output_path))

            return InjectionResult(
                success=True,
                output_path=self.context.output_path,
                filled_slots=filled_slots,
                empty_slots=empty_slots,
                errors=self._errors,
                warnings=self._warnings,
            )

        except Exception as e:
            logger.exception(f"PPTX injection failed: {e}")
            return InjectionResult(
                success=False,
                errors=[f"PPTX injection failed: {str(e)}"],
            )

    def _inject_docx(self) -> InjectionResult:
        """DOCX 템플릿에 주입."""
        try:
            from docx import Document
        except ImportError:
            return InjectionResult(
                success=False,
                errors=["python-docx not installed"],
            )

        filled_slots: list[str] = []
        empty_slots: list[str] = []

        try:
            doc = Document(str(self.context.template_path))

            # 문단에서 슬롯 찾기
            for para in doc.paragraphs:
                self._process_docx_paragraph(para, filled_slots, empty_slots)

            # 테이블 내 슬롯 찾기
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            self._process_docx_paragraph(para, filled_slots, empty_slots)

            # 저장
            doc.save(str(self.context.output_path))

            return InjectionResult(
                success=True,
                output_path=self.context.output_path,
                filled_slots=filled_slots,
                empty_slots=empty_slots,
                errors=self._errors,
                warnings=self._warnings,
            )

        except Exception as e:
            logger.exception(f"DOCX injection failed: {e}")
            return InjectionResult(
                success=False,
                errors=[f"DOCX injection failed: {str(e)}"],
            )

    def _process_docx_paragraph(
        self,
        para: Any,
        filled_slots: list[str],
        empty_slots: list[str],
    ) -> None:
        """DOCX 문단 처리."""
        text = para.text
        matches = SLOT_PATTERN.findall(text)

        for slot_prefix, slot_name in matches:
            slot_id = f"{{{{{slot_prefix}:{slot_name}}}}}"
            block = self._slot_to_block.get(slot_id)

            if block:
                new_text = self._block_to_text(block)
                para.text = para.text.replace(slot_id, new_text)
                filled_slots.append(slot_id)
            else:
                empty_slots.append(slot_id)

    def _replace_slot_text(self, shape: Any, slot_id: str, new_text: str) -> None:
        """도형 내 슬롯 텍스트 교체."""
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if slot_id in run.text:
                        run.text = run.text.replace(slot_id, new_text)

    def _block_to_text(self, block: ReportBlock) -> str:
        """블록을 텍스트로 변환.

        복잡한 블록(테이블, 차트)은 단순 텍스트 요약으로 변환합니다.
        실제 PPT 렌더링은 pptx-service를 통해 수행됩니다.
        """
        if isinstance(block, TextBlock):
            parts = []
            if block.title:
                parts.append(block.title)
            if block.content:
                parts.append(block.content)
            if block.bullet_points:
                parts.extend([f"• {bp}" for bp in block.bullet_points])
            return "\n".join(parts)

        elif isinstance(block, ClaimBlock):
            return block.claim_text

        elif isinstance(block, KPIBlock):
            kpi_texts = []
            for kpi in block.kpis:
                label = kpi.get("label", "")
                value = kpi.get("value", "")
                unit = kpi.get("unit", "")
                kpi_texts.append(f"{label}: {value} {unit}")
            return "\n".join(kpi_texts)

        elif isinstance(block, TableBlock):
            # 테이블은 간단한 텍스트 요약
            header = " | ".join([c.header for c in block.columns])
            rows = []
            for row in block.rows[:5]:  # 처음 5행만
                row_vals = [str(row.get(c.key, "")) for c in block.columns]
                rows.append(" | ".join(row_vals))
            if len(block.rows) > 5:
                rows.append(f"... ({len(block.rows) - 5} more rows)")
            return f"{block.title}\n{header}\n" + "\n".join(rows)

        elif isinstance(block, ChartBlock):
            # 차트는 타이틀과 데이터 요약
            if block.data:
                data_summary = ", ".join(
                    [f"{c}: {v}" for c, v in zip(block.data.categories[:3], block.data.values[:3], strict=False)]
                )
                return f"{block.title}\n{data_summary}"
            return block.title

        elif isinstance(block, CoverBlock):
            return f"{block.deal_name}\n{block.target_name}"

        elif isinstance(block, ScopeBlock):
            items = [f"{item.label}: {item.value}" for item in block.scope_items]
            return f"{block.title}\n" + "\n".join(items)

        elif isinstance(block, MethodologyBlock):
            steps = [f"{s.step}. {s.title}" for s in block.steps]
            return f"{block.title}\n" + "\n".join(steps)

        elif isinstance(block, IssueBlock):
            issues = [f"[{i.severity.upper()}] {i.title}" for i in block.issues[:5]]
            return f"{block.title}\n" + "\n".join(issues)

        else:
            return f"[{block.type.value} block]"


# =============================================================================
# Convenience Functions
# =============================================================================


def inject_template(
    template_path: Path,
    output_path: Path,
    report_ir: ReportIR,
    contract: TemplateContract,
    slot_mappings: dict[str, str] | None = None,
    style_tokens: StyleTokens | None = None,
) -> InjectionResult:
    """템플릿에 콘텐츠 주입.

    Args:
        template_path: 입력 템플릿 경로
        output_path: 출력 파일 경로
        report_ir: Report IR
        contract: 템플릿 계약
        slot_mappings: 커스텀 슬롯-블록 매핑 (선택)
        style_tokens: 스타일 토큰 (선택)

    Returns:
        주입 결과
    """
    context = InjectionContext(
        template_path=template_path,
        output_path=output_path,
        report_ir=report_ir,
        contract=contract,
        slot_mappings=slot_mappings or {},
        style_tokens=style_tokens,
    )
    injector = TemplateInjector(context)
    return injector.inject()
