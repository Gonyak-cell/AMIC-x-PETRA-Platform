"""Horizontal Bar 네이티브 PPTX 차트 빌더.

> 마지막 수정: 2026-02-27 10:28:00

입력 스키마 (Plotly hbar와 동일):
    {
        "categories": ["경쟁사A", "경쟁사B", "당사", "경쟁사C"],
        "values": [120, 100, 95, 80],
        "highlight": "당사",  # 강조할 항목 (optional)
    }
"""

from __future__ import annotations

from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from src.chart_engine.config import ChartConfig
from src.chart_engine.exceptions import ChartDataError
from src.chart_engine.pptx_native.base import NativeChartBuilder
from src.chart_engine.pptx_native.styling import (
    _hex_to_rgb,
    apply_axis_style,
    apply_data_labels,
)


class HBarBuilder(NativeChartBuilder):
    """BAR_CLUSTERED (horizontal) 네이티브 차트 빌더."""

    def build(
        self,
        slide: Any,
        data: dict[str, Any],
        *,
        title: str = "",
        left: float = 0.5,
        top: float = 1.5,
        width: float = 9.0,
        height: float = 4.5,
        config: ChartConfig | None = None,
    ) -> Any:
        cfg = self._get_config(config)
        c = cfg.colors

        categories = data.get("categories", [])
        values = data.get("values", [])
        highlight = data.get("highlight", "")

        if not categories or not values:
            raise ChartDataError("hbar", "categories와 values는 필수입니다.")

        chart_data = CategoryChartData()
        chart_data.categories = categories
        chart_data.add_series("", tuple(values))

        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
            chart_data,
        )
        chart = chart_shape.chart

        if title:
            chart.has_title = True
            chart.chart_title.text_frame.text = title

        # 포인트별 색상: highlight 항목은 accent, 나머지 primary
        plot = chart.plots[0]
        series = plot.series[0]
        for idx, cat in enumerate(categories):
            point = series.points[idx]
            fill = point.format.fill
            fill.solid()
            if cat == highlight:
                fill.fore_color.rgb = _hex_to_rgb(c.accent)
            else:
                fill.fore_color.rgb = _hex_to_rgb(c.primary)

        # 범례 숨김 (단일 시리즈)
        chart.has_legend = False

        apply_axis_style(chart, cfg)
        apply_data_labels(chart, number_format="#,##0", config=cfg)

        return chart_shape
