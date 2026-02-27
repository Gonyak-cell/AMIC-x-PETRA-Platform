"""Stacked Bar 네이티브 PPTX 차트 빌더.

> 마지막 수정: 2026-02-27 10:28:00

입력 스키마 (Plotly stacked_bar와 동일):
    {
        "categories": ["2021", "2022", "2023"],
        "series": [
            {"name": "IT서비스", "values": [50, 60, 70]},
            {"name": "SI", "values": [30, 35, 40]},
        ]
    }
"""

from __future__ import annotations

from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from src.chart_engine.config import ChartConfig
from src.chart_engine.exceptions import ChartDataError
from src.chart_engine.pptx_native.base import NativeChartBuilder
from src.chart_engine.pptx_native.styling import apply_amic_chart_style


class StackedBarBuilder(NativeChartBuilder):
    """COLUMN_STACKED 네이티브 차트 빌더."""

    def build(
        self,
        slide: Any,
        data: dict[str, Any],
        *,
        title: str = "",
        left: float = 0.5,
        top: float = 1.5,
        width: float = 9.0,
        height: float = 4.5,
        config: ChartConfig | None = None,
    ) -> Any:
        cfg = self._get_config(config)

        categories = data.get("categories", [])
        series_list = data.get("series", [])

        if not categories or not series_list:
            raise ChartDataError("stacked_bar", "categories와 series는 필수입니다.")

        n_cats = len(categories)
        for s in series_list:
            if len(s.get("values", [])) != n_cats:
                raise ChartDataError(
                    "stacked_bar",
                    f"series '{s.get('name', '')}' values 길이가 "
                    f"categories 길이({n_cats})와 다릅니다.",
                )

        chart_data = CategoryChartData()
        chart_data.categories = categories
        for s in series_list:
            chart_data.add_series(
                s.get("name", ""),
                tuple(s.get("values", [])),
            )

        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_STACKED,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
            chart_data,
        )
        chart = chart_shape.chart

        if title:
            chart.has_title = True
            chart.chart_title.text_frame.text = title

        apply_amic_chart_style(chart, cfg)

        # gap width 조정 (기본값 150% → 80%로 간결하게)
        chart.plots[0].gap_width = 80

        return chart_shape
