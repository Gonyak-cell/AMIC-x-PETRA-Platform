"""Line 네이티브 PPTX 차트 빌더.

> 마지막 수정: 2026-02-27 10:28:00

입력 스키마 (Plotly line와 동일):
    {
        "x": ["2020", "2021", "2022", "2023", "2024"],
        "series": [
            {"name": "매출", "values": [100, 120, 140, 160, 180]},
            {"name": "영업이익", "values": [10, 15, 20, 25, 30]},
        ]
    }
"""

from __future__ import annotations

from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from src.chart_engine.config import ChartConfig
from src.chart_engine.exceptions import ChartDataError
from src.chart_engine.pptx_native.base import NativeChartBuilder
from src.chart_engine.pptx_native.styling import apply_amic_chart_style


class LineChartBuilder(NativeChartBuilder):
    """LINE_MARKERS 네이티브 차트 빌더."""

    def build(
        self,
        slide: Any,
        data: dict[str, Any],
        *,
        title: str = "",
        left: float = 0.5,
        top: float = 1.5,
        width: float = 9.0,
        height: float = 4.5,
        config: ChartConfig | None = None,
    ) -> Any:
        cfg = self._get_config(config)

        x = data.get("x", [])
        series_list = data.get("series", [])

        if not x or not series_list:
            raise ChartDataError("line", "x와 series는 필수입니다.")

        chart_data = CategoryChartData()
        chart_data.categories = x
        for s in series_list:
            chart_data.add_series(
                s.get("name", ""),
                tuple(s.get("values", [])),
            )

        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.LINE_MARKERS,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
            chart_data,
        )
        chart = chart_shape.chart

        if title:
            chart.has_title = True
            chart.chart_title.text_frame.text = title

        apply_amic_chart_style(chart, cfg)

        return chart_shape
