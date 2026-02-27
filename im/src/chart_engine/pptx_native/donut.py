"""Donut 네이티브 PPTX 차트 빌더.

> 마지막 수정: 2026-02-27 10:28:00

입력 스키마 (Plotly donut와 동일):
    {
        "labels": ["당사", "경쟁사A", "경쟁사B", "기타"],
        "values": [30, 25, 20, 25],
    }
"""

from __future__ import annotations

from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION
from pptx.util import Inches

from src.chart_engine.config import ChartConfig
from src.chart_engine.exceptions import ChartDataError
from src.chart_engine.pptx_native.base import NativeChartBuilder
from src.chart_engine.pptx_native.styling import (
    apply_legend,
    apply_series_colors,
)


class DonutBuilder(NativeChartBuilder):
    """DOUGHNUT 네이티브 차트 빌더."""

    def build(
        self,
        slide: Any,
        data: dict[str, Any],
        *,
        title: str = "",
        left: float = 0.5,
        top: float = 1.5,
        width: float = 9.0,
        height: float = 4.5,
        config: ChartConfig | None = None,
    ) -> Any:
        cfg = self._get_config(config)

        labels = data.get("labels", [])
        values = data.get("values", [])

        if not labels or not values:
            raise ChartDataError("donut", "labels와 values는 필수입니다.")

        chart_data = CategoryChartData()
        chart_data.categories = labels
        chart_data.add_series("", tuple(values))

        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
            chart_data,
        )
        chart = chart_shape.chart

        if title:
            chart.has_title = True
            chart.chart_title.text_frame.text = title

        # 도넛 포인트별 색상 적용 (시리즈 색상이 아닌 포인트별 색상)
        apply_series_colors(chart, cfg)

        # 데이터 레이블: 카테고리명 + 백분율
        plot = chart.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_category_name = True
        dl.show_percentage = True
        dl.show_value = False
        dl.number_format = "0%"
        dl.number_format_is_linked = False
        try:
            dl.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
        except ValueError:
            pass

        apply_legend(chart, cfg)

        return chart_shape
