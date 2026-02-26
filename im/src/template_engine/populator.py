"""템플릿 팝퓰레이터 — 마스터 PPTX에 콘텐츠를 삽입하는 오케스트레이터.

TemplateContent를 받아 마스터 템플릿을 복사한 후,
각 슬라이드의 shape에 텍스트/차트/테이블 데이터를 삽입하여
최종 PPTX를 생성한다.
"""

from __future__ import annotations

import logging
import shutil
import uuid
from pathlib import Path

from pptx import Presentation

from src.template_engine.content_injector import ContentInjector
from src.template_engine.schemas import TemplateContent

logger = logging.getLogger(__name__)


class TemplatePopulator:
    """마스터 PPTX 템플릿에 콘텐츠를 삽입하는 오케스트레이터."""

    def __init__(
        self,
        template_path: Path,
        injector: ContentInjector | None = None,
        output_dir: str | Path = "",
    ) -> None:
        """
        Args:
            template_path: 마스터 PPTX 템플릿 절대 경로.
            injector: ContentInjector 인스턴스 (기본: 새로 생성).
            output_dir: 출력 디렉토리 (기본: 템플릿과 같은 디렉토리).
        """
        if not template_path.exists():
            raise FileNotFoundError(f"마스터 템플릿 없음: {template_path}")

        self.template_path = template_path
        self.injector = injector or ContentInjector()
        self.output_dir = Path(output_dir) if output_dir else template_path.parent

    def populate(self, content: TemplateContent) -> Path:
        """마스터 템플릿에 콘텐츠를 삽입하여 새 PPTX를 생성한다.

        Args:
            content: Ralph Loop이 생성한 TemplateContent.

        Returns:
            생성된 PPTX 파일의 절대 경로.
        """
        # 1. 출력 경로 결정
        output_name = (
            f"{content.project_name} - {content.memo_type} - "
            f"{uuid.uuid4().hex[:8]}.pptx"
        )
        # 파일명에 사용 불가한 문자 제거
        output_name = "".join(
            c for c in output_name if c not in r'\/:*?"<>|'
        )
        output_path = self.output_dir / output_name
        self.output_dir.mkdir(parents=True, exist_ok=True)

        # 2. 템플릿 복사
        shutil.copy2(self.template_path, output_path)
        logger.info("템플릿 복사: %s → %s", self.template_path.name, output_path.name)

        # 3. PPTX 로드
        prs = Presentation(str(output_path))

        # 4. 커버 슬라이드 처리 (항상 slide 0)
        self._populate_cover(prs, content)

        # 5. 콘텐츠 슬라이드 처리
        for slide_content in content.slides:
            self._populate_slide(prs, slide_content)

        # 6. 불필요 슬라이드 제거 (인덱스 역순으로 제거)
        if content.excluded_slides:
            self._remove_slides(prs, content.excluded_slides)

        # 7. 저장
        prs.save(str(output_path))
        logger.info(
            "PPTX 생성 완료: %s (%d 슬라이드)",
            output_path.name,
            len(prs.slides),
        )

        return output_path

    # ── 내부 메서드 ──────────────────────────────────────────────────────

    def _populate_cover(
        self, prs: Presentation, content: TemplateContent,
    ) -> None:
        """커버 슬라이드의 프로젝트명, 날짜를 교체한다."""
        if not prs.slides:
            return

        cover = prs.slides[0]
        for shape in cover.shapes:
            if not shape.has_text_frame:
                continue

            text = shape.text_frame.text.strip()

            # 날짜 패턴 감지 (영문 월 + 연도, e.g. "January 2025", "February 2026")
            if _looks_like_date(text):
                self.injector.inject_text(shape, content.date)
                logger.debug("커버 날짜 교체: %r → %r", text[:30], content.date)

            # 메모 유형 패턴 감지
            elif any(kw in text.upper() for kw in (
                "TEASER", "DISCUSSION MEMO", "INFORMATION MEMORANDUM",
                "CONFIDENTIAL MEMORANDUM",
            )):
                memo_label = {
                    "TM": "TEASER MEMORANDUM",
                    "DM": "DISCUSSION MEMO",
                }.get(content.memo_type, content.memo_type)
                self.injector.inject_text(shape, memo_label)

    def _populate_slide(
        self, prs: Presentation, slide_content: object,
    ) -> None:
        """단일 슬라이드의 shape에 콘텐츠를 삽입한다."""
        from src.template_engine.schemas import SlideContent

        sc: SlideContent = slide_content  # type: ignore[assignment]

        if sc.slide_idx >= len(prs.slides):
            logger.warning(
                "슬라이드 인덱스 %d 범위 초과 (총 %d) — 스킵",
                sc.slide_idx,
                len(prs.slides),
            )
            return

        slide = prs.slides[sc.slide_idx]

        # shape 이름 → shape 객체 매핑
        shape_map = {s.name: s for s in slide.shapes}

        # 타이틀 교체 (PLACEHOLDER idx=0 = TITLE)
        if sc.title:
            title_set = False
            for shape in slide.shapes:
                if shape.has_text_frame and hasattr(shape, "placeholder_format"):
                    pf = shape.placeholder_format
                    if pf is not None and pf.idx == 0:
                        self.injector.inject_text(shape, sc.title)
                        title_set = True
                        break
            if not title_set:
                logger.debug(
                    "슬라이드 %d: TITLE placeholder (idx=0) 없음 — 타이틀 교체 스킵",
                    sc.slide_idx,
                )

        # 텍스트 shape 교체
        for shape_name, text in sc.texts.items():
            shape = shape_map.get(shape_name)
            if shape is None:
                logger.debug("텍스트 shape %r 없음 — 스킵", shape_name)
                continue
            self.injector.inject_text(shape, text)

        # 차트 데이터 교체
        for shape_name, chart_content in sc.charts.items():
            shape = shape_map.get(shape_name)
            if shape is None:
                logger.debug("차트 shape %r 없음 — 스킵", shape_name)
                continue
            self.injector.inject_chart_data(shape, chart_content)

        # 테이블 데이터 교체
        for shape_name, table_content in sc.tables.items():
            shape = shape_map.get(shape_name)
            if shape is None:
                logger.debug("테이블 shape %r 없음 — 스킵", shape_name)
                continue
            self.injector.inject_table_data(shape, table_content)

    @staticmethod
    def _remove_slides(prs: Presentation, indices: list[int]) -> None:
        """지정 인덱스의 슬라이드를 제거한다 (역순)."""
        from pptx.oxml.ns import qn

        slide_list = prs.slides._sldIdLst  # type: ignore[attr-defined]
        slide_ids = list(slide_list)

        for idx in sorted(set(indices), reverse=True):
            if 0 <= idx < len(slide_ids):
                rId = slide_ids[idx].get(qn("r:id"))
                prs.part.drop_rel(rId)
                slide_list.remove(slide_ids[idx])
                logger.debug("슬라이드 %d 제거", idx)


# ── 헬퍼 ─────────────────────────────────────────────────────────────────────

_MONTH_NAMES = {
    "january", "february", "march", "april", "may", "june",
    "july", "august", "september", "october", "november", "december",
}


def _looks_like_date(text: str) -> bool:
    """문자열이 날짜 패턴인지 판별한다. (e.g. 'January 2025', '2026년 2월')"""
    lower = text.lower().strip()
    # 영문 월 + 연도
    parts = lower.split()
    if len(parts) == 2 and parts[0] in _MONTH_NAMES and parts[1].isdigit():
        return True
    # 한국어 패턴
    if "년" in lower and "월" in lower:
        return True
    return False
