"""IM PPTX Programmatic Gate — python-pptx 기반 IM 디자인 검증.

6개 검증 레이어:
1. 슬라이드 구조 검증 (필수 슬라이드, 순서)
2. 텍스트 완전성 (플레이스홀더 부재)
3. 재무 테이블 수치 정합성
4. 차트 데이터 유효성
5. 디자인 일관성 (AMIC 5단계 그린, SUITE/Pretendard)
6. 정보 밀도 (슬라이드당 텍스트/차트/테이블 비율)
"""

from __future__ import annotations

import re
import time
from typing import Any

from src.ralph.gates.base import DimensionScore, GateResult, QualityGate

# ── AMIC IM 디자인 상수 ──────────────────────────────────────────────────────

ALLOWED_FONTS = {"SUITE", "Pretendard", "Noto Sans KR", "NanumGothic", "Arial", "Calibri"}

AMIC_COLORS = {
    "0F3A32",  # Signature Green
    "1C8F57",  # Solid Green
    "26C260",  # Highlight Green
    "A3E96B",  # Fresh Green
    "E6FDD6",  # Light Green
    "000000", "FFFFFF", "777777", "333333",
    "BC2C1A",  # Negative
    "EF6C00",  # Caution
    "F4F6F8",  # Cool Grey BG (RGB로 변환 시)
    "757575",  # Gray medium
    "E0E0E0",  # Gray border
    "666666", "999999",
}

REQUIRED_SLIDES: dict[str, list[str]] = {
    "IM": ["Cover", "Disclaimer", "Table of Contents"],
    "TM": ["Cover", "Disclaimer"],
    "DM": ["Cover", "Disclaimer"],
}

_PLACEHOLDER_PATTERNS = [
    r"\[INSERT\s*(?:HERE)?\]",
    r"\[TBD\]",
    r"\[XXX\]",
    r"\[TODO\]",
    r"\[금액\]",
    r"\[프로젝트명\]",
    r"Lorem\s+ipsum",
]


class IMPPTXProgrammaticGate(QualityGate):
    """IM PPTX 프로그래밍 검증 게이트 — AMIC 디자인 시스템 기반."""

    @property
    def name(self) -> str:
        return "im_pptx_programmatic"

    async def evaluate(
        self,
        artifact_path: str,
        prd_section: dict[str, Any],
        source_data: dict[str, Any] | None = None,
    ) -> GateResult:
        start = time.perf_counter_ns()
        issues: list[str] = []
        suggestions: list[str] = []
        critical_flags: list[str] = []

        try:
            from pptx import Presentation
            prs = Presentation(artifact_path)
        except Exception as exc:
            return self._timed_result(
                start, [], [f"PPTX 로드 실패: {exc}"], [], [f"PPTX 파일 손상: {exc}"],
            )

        memo_type = prd_section.get("memo_type", "IM")

        # 1. 슬라이드 구조 검증
        struct_score, struct_issues = self._check_structure(prs, memo_type, prd_section)
        issues.extend(struct_issues)

        # 2. 텍스트 완전성
        complete_score, complete_issues = self._check_completeness(prs)
        issues.extend(complete_issues)
        if any("CRITICAL" in i for i in complete_issues):
            critical_flags.extend([i for i in complete_issues if "CRITICAL" in i])

        # 3. 재무 테이블 정합성
        table_score, table_issues = self._check_table_totals(prs)
        issues.extend(table_issues)

        # 4. 차트 데이터 유효성
        chart_score, chart_issues = self._check_charts(prs)
        issues.extend(chart_issues)

        # 5. 디자인 일관성 (AMIC 전용)
        design_score, design_issues = self._check_design(prs)
        issues.extend(design_issues)

        # 6. 정보 밀도
        density_score, density_issues = self._check_information_density(prs)
        issues.extend(density_issues)

        dimensions = [
            DimensionScore("structure", "슬라이드 구조", struct_score, 0.15),
            DimensionScore("completeness", "텍스트 완전성", complete_score, 0.20),
            DimensionScore("table_integrity", "테이블 정합성", table_score, 0.15),
            DimensionScore("chart_validity", "차트 유효성", chart_score, 0.15),
            DimensionScore("design", "디자인 일관성", design_score, 0.20),
            DimensionScore("info_density", "정보 밀도", density_score, 0.15),
        ]

        return self._timed_result(
            start, dimensions, issues, suggestions, critical_flags,
        )

    # ── 검증 레이어 ──────────────────────────────────────────────────────────

    def _check_structure(self, prs: Any, memo_type: str, prd: dict) -> tuple[float, list[str]]:
        """슬라이드 구조 검증."""
        issues: list[str] = []
        slide_count = len(prs.slides)

        min_slides = prd.get("min_slides", 5)
        if slide_count < min_slides:
            issues.append(f"슬라이드 수 부족: {slide_count}장 (최소 {min_slides}장)")

        max_slides = prd.get("max_slides", 100)
        if slide_count > max_slides:
            issues.append(f"슬라이드 수 초과: {slide_count}장 (최대 {max_slides}장)")

        # 필수 슬라이드 존재 확인
        slide_titles = []
        for slide in prs.slides:
            if slide.shapes.title:
                slide_titles.append(slide.shapes.title.text.strip())
            else:
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        slide_titles.append(shape.text_frame.text.strip())
                        break
                else:
                    slide_titles.append("")

        required = REQUIRED_SLIDES.get(memo_type, [])
        for req in required:
            if not any(req.lower() in t.lower() for t in slide_titles):
                issues.append(f"필수 슬라이드 누락: '{req}'")

        score = max(1.0, 5.0 - len(issues) * 1.0)
        return score, issues

    def _check_completeness(self, prs: Any) -> tuple[float, list[str]]:
        """플레이스홀더 패턴 탐지."""
        issues: list[str] = []
        placeholder_count = 0

        for slide_idx, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text
                for pattern in _PLACEHOLDER_PATTERNS:
                    matches = re.findall(pattern, text, re.IGNORECASE)
                    if matches:
                        placeholder_count += len(matches)
                        issues.append(
                            f"CRITICAL: 슬라이드 {slide_idx} — 플레이스홀더 '{matches[0]}'"
                        )

        score = 5.0 if placeholder_count == 0 else max(1.0, 5.0 - placeholder_count)
        return score, issues

    def _check_table_totals(self, prs: Any) -> tuple[float, list[str]]:
        """재무 테이블 합계행 정합성."""
        issues: list[str] = []
        table_count = 0
        error_count = 0

        for slide_idx, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                if not shape.has_table:
                    continue
                table = shape.table
                table_count += 1

                rows = len(table.rows)
                cols = len(table.columns)
                if rows < 3 or cols < 2:
                    continue

                last_row_text = [table.cell(rows - 1, c).text.strip() for c in range(cols)]
                if not any(kw in last_row_text[0].lower() for kw in ["합계", "total", "소계", "계"]):
                    continue

                for c in range(1, cols):
                    total_text = table.cell(rows - 1, c).text.strip()
                    total_val = self._parse_number(total_text)
                    if total_val is None:
                        continue

                    item_sum = 0.0
                    parseable = True
                    for r in range(1, rows - 1):
                        val = self._parse_number(table.cell(r, c).text)
                        if val is None:
                            parseable = False
                            break
                        item_sum += val

                    if parseable and abs(total_val - item_sum) > 0.01:
                        error_count += 1
                        issues.append(
                            f"슬라이드 {slide_idx} 테이블: 열 {c+1} 합계={total_val}, "
                            f"항목합={item_sum} (차이: {abs(total_val - item_sum):.2f})"
                        )

        if table_count == 0:
            return 5.0, []

        score = max(1.0, 5.0 - error_count * 2.0)
        return score, issues

    def _check_charts(self, prs: Any) -> tuple[float, list[str]]:
        """차트 데이터 유효성."""
        issues: list[str] = []
        chart_count = 0

        for slide_idx, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                if not shape.has_chart:
                    continue
                chart = shape.chart
                chart_count += 1

                for plot in chart.plots:
                    for series in plot.series:
                        values = list(series.values)
                        if not values or all(v is None for v in values):
                            issues.append(f"슬라이드 {slide_idx}: 차트 시리즈에 데이터 없음")

                if not chart.has_title:
                    issues.append(f"슬라이드 {slide_idx}: 차트 제목 없음")

        if chart_count == 0:
            return 5.0, []

        score = max(1.0, 5.0 - len(issues) * 1.0)
        return score, issues

    def _check_design(self, prs: Any) -> tuple[float, list[str]]:
        """AMIC 디자인 일관성 검증."""
        issues: list[str] = []
        font_violations = 0
        color_violations = 0

        for slide_idx, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        # 폰트 확인 (None = 테마 상속, 허용)
                        if run.font.name and run.font.name not in ALLOWED_FONTS:
                            font_violations += 1
                            if font_violations <= 3:
                                issues.append(
                                    f"슬라이드 {slide_idx}: 비허용 폰트 '{run.font.name}'"
                                )

                        # 색상 확인
                        if run.font.color and run.font.color.rgb:
                            rgb = str(run.font.color.rgb)
                            if rgb not in AMIC_COLORS:
                                color_violations += 1
                                if color_violations <= 3:
                                    issues.append(
                                        f"슬라이드 {slide_idx}: AMIC 팔레트 외 색상 #{rgb}"
                                    )

        total_violations = font_violations + color_violations
        score = max(1.0, 5.0 - total_violations * 0.3)
        return score, issues

    def _check_information_density(self, prs: Any) -> tuple[float, list[str]]:
        """슬라이드당 정보 밀도 검증."""
        issues: list[str] = []
        slide_count = len(prs.slides)
        if slide_count == 0:
            return 5.0, []

        empty_slides = 0
        dense_slides = 0

        for slide_idx, slide in enumerate(prs.slides, 1):
            char_count = 0
            has_chart = False
            has_table = False

            for shape in slide.shapes:
                if shape.has_text_frame:
                    char_count += len(shape.text_frame.text)
                if shape.has_chart:
                    has_chart = True
                if shape.has_table:
                    has_table = True

            # 커버, 면책, 구분 슬라이드 제외 (처음 3개)
            if slide_idx <= 3:
                continue

            if char_count < 30 and not has_chart and not has_table:
                empty_slides += 1
                if empty_slides <= 3:
                    issues.append(f"슬라이드 {slide_idx}: 콘텐츠 부족 (텍스트 {char_count}자)")

            if char_count > 800 and not has_table:
                dense_slides += 1
                if dense_slides <= 3:
                    issues.append(f"슬라이드 {slide_idx}: 텍스트 과밀 ({char_count}자)")

        problem_ratio = (empty_slides + dense_slides) / max(slide_count - 3, 1)
        score = max(1.0, 5.0 - problem_ratio * 5.0)
        return score, issues

    @staticmethod
    def _parse_number(text: str) -> float | None:
        """한국어 숫자 파싱 (간략 버전)."""
        if not text:
            return None
        cleaned = text.replace(",", "").replace(" ", "").replace("원", "").replace("억", "")
        cleaned = cleaned.replace("백만", "").replace("천", "")
        try:
            return float(cleaned)
        except ValueError:
            return None
