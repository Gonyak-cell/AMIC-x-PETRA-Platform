"""Gate A: 템플릿 파일 무결성 사전 검증 (IM 모듈).

> 마지막 수정: 2026-03-13 21:33:36

렌더링 시작 전에 호출하여 템플릿 파일이 유효한지 검증한다.
실패 시 즉시 abort하여 불필요한 렌더링을 방지한다.
"""

from __future__ import annotations

import logging
import time
from pathlib import Path
from typing import Any

from src.quality_gate.base import DimensionScore, GateResult, GateVerdict, QualityGate

logger = logging.getLogger(__name__)


class TemplatePreflight(QualityGate):
    """Gate A: 템플릿 파일 무결성 사전 검증.

    렌더링 시작 전에 호출하여 템플릿 파일이 유효한지 검증한다.
    실패 시 즉시 abort하여 불필요한 렌더링을 방지한다.

    5개 검증 차원:
    1. template_exists (weight=0.3): 파일 존재 여부
    2. template_parseable (weight=0.25): python-pptx 파싱 가능 여부
    3. layout_available (weight=0.2): 레이아웃 이름 매칭
    4. shapes_present (weight=0.15): 필수 shape 존재 여부
    5. font_available (weight=0.1): 브랜드 폰트 사용 가능 여부
    """

    @property
    def name(self) -> str:
        return "template_preflight"

    async def evaluate(
        self,
        artifact_path: str,
        prd_section: dict[str, Any],
        source_data: dict[str, Any] | None = None,
    ) -> GateResult:
        """템플릿 무결성을 사전 검증한다.

        Args:
            artifact_path: 템플릿 파일 경로.
            prd_section: PRD 섹션 (memo_type 포함).
            source_data: template_spec이 담긴 딕셔너리.

        Returns:
            GateResult: 검증 결과.
        """
        start_ns = time.perf_counter_ns()
        issues: list[str] = []
        suggestions: list[str] = []
        critical_flags: list[str] = []
        dimensions: list[DimensionScore] = []

        # template_spec이 없으면 하위 호환을 위해 PASS 반환
        if not source_data or "template_spec" not in source_data:
            return GateResult(
                gate_name=self.name,
                verdict=GateVerdict.PASS,
                weighted_score=5.0,
                issues=["template_spec 미제공 — 프리플라이트 생략"],
            )

        spec = source_data["template_spec"]
        template_path = Path(artifact_path)

        # ── 차원 1: template_exists (weight=0.3) ──
        exists_score, exists_feedback = self._check_exists(
            template_path,
            issues,
            critical_flags,
        )
        dimensions.append(
            DimensionScore(
                name="template_exists",
                label="파일 존재",
                score=exists_score,
                weight=0.3,
                feedback=exists_feedback,
            )
        )

        # 파일이 없으면 나머지 검증 불가 — 즉시 반환
        if exists_score == 0.0:
            for dim_name, label, weight in [
                ("template_parseable", "파싱 가능", 0.25),
                ("layout_available", "레이아웃 존재", 0.2),
                ("shapes_present", "필수 Shape", 0.15),
                ("font_available", "폰트 가용", 0.1),
            ]:
                dimensions.append(
                    DimensionScore(
                        name=dim_name,
                        label=label,
                        score=0.0,
                        weight=weight,
                        feedback="템플릿 파일 미존재로 검증 불가",
                    )
                )
            return self._timed_result(
                start_ns,
                dimensions,
                issues,
                suggestions,
                critical_flags,
            )

        # ── 차원 2: template_parseable (weight=0.25) ──
        prs, parse_score, parse_feedback = self._check_parseable(
            template_path,
            issues,
            critical_flags,
        )
        dimensions.append(
            DimensionScore(
                name="template_parseable",
                label="파싱 가능",
                score=parse_score,
                weight=0.25,
                feedback=parse_feedback,
            )
        )

        # 파싱 불가면 나머지 검증 불가
        if prs is None:
            for dim_name, label, weight in [
                ("layout_available", "레이아웃 존재", 0.2),
                ("shapes_present", "필수 Shape", 0.15),
                ("font_available", "폰트 가용", 0.1),
            ]:
                dimensions.append(
                    DimensionScore(
                        name=dim_name,
                        label=label,
                        score=0.0,
                        weight=weight,
                        feedback="템플릿 파싱 실패로 검증 불가",
                    )
                )
            return self._timed_result(
                start_ns,
                dimensions,
                issues,
                suggestions,
                critical_flags,
            )

        # ── 차원 3: layout_available (weight=0.2) ──
        layout_score, layout_feedback = self._check_layouts(
            prs,
            spec,
            issues,
            critical_flags,
        )
        dimensions.append(
            DimensionScore(
                name="layout_available",
                label="레이아웃 존재",
                score=layout_score,
                weight=0.2,
                feedback=layout_feedback,
            )
        )

        # ── 차원 4: shapes_present (weight=0.15) ──
        shapes_score, shapes_feedback = self._check_shapes(
            prs,
            spec,
            issues,
            suggestions,
        )
        dimensions.append(
            DimensionScore(
                name="shapes_present",
                label="필수 Shape",
                score=shapes_score,
                weight=0.15,
                feedback=shapes_feedback,
            )
        )

        # ── 차원 5: font_available (weight=0.1) ──
        font_score, font_feedback = self._check_fonts(
            spec,
            issues,
            suggestions,
        )
        dimensions.append(
            DimensionScore(
                name="font_available",
                label="폰트 가용",
                score=font_score,
                weight=0.1,
                feedback=font_feedback,
            )
        )

        return self._timed_result(
            start_ns,
            dimensions,
            issues,
            suggestions,
            critical_flags,
        )

    # ── 개별 검증 메서드 ──────────────────────────────────────────

    def _check_exists(
        self,
        path: Path,
        issues: list[str],
        critical_flags: list[str],
    ) -> tuple[float, str]:
        """템플릿 파일 존재 여부를 확인한다."""
        if path.exists() and path.is_file():
            return 5.0, "템플릿 파일 존재 확인"
        logger.error("템플릿 파일 미존재: %s", path)
        critical_flags.append(f"템플릿 파일 미존재: {path.name}")
        issues.append(f"템플릿 파일이 존재하지 않습니다: {path.name}")
        return 0.0, f"파일 미존재: {path.name}"

    def _check_parseable(
        self,
        path: Path,
        issues: list[str],
        critical_flags: list[str],
    ) -> tuple[Any, float, str]:
        """python-pptx로 파싱 가능한지 확인한다.

        Returns:
            (Presentation | None, score, feedback)
        """
        try:
            from pptx import Presentation

            prs = Presentation(str(path))
            return prs, 5.0, "PPTX 파싱 성공"
        except Exception as exc:
            logger.error("템플릿 파싱 실패: %s", exc, exc_info=True)
            critical_flags.append("템플릿 파싱 실패")
            issues.append("python-pptx로 파싱할 수 없습니다")
            return None, 0.0, "파싱 실패"

    def _check_layouts(
        self,
        prs: Any,
        spec: Any,
        issues: list[str],
        critical_flags: list[str],
    ) -> tuple[float, str]:
        """SlideSpec.layout_name이 템플릿 slide_layouts에 존재하는지 확인한다."""
        if not hasattr(spec, "slides") or not spec.slides:
            return 5.0, "SlideSpec 미정의 — 레이아웃 검증 생략"

        available_layouts: set[str] = set()
        for layout in prs.slide_layouts:
            available_layouts.add(layout.name)

        missing: list[str] = []
        for slide_spec in spec.slides:
            if slide_spec.layout_name not in available_layouts:
                missing.append(slide_spec.layout_name)

        if not missing:
            return 5.0, f"모든 레이아웃 존재 ({len(spec.slides)}개)"

        critical_flags.append(
            f"레이아웃 미존재: {missing} (가용: {sorted(available_layouts)})"
        )
        issues.append(f"템플릿에 없는 레이아웃: {missing}")

        # 부분 존재 시 비례 점수
        total = len(spec.slides)
        found = total - len(missing)
        score = (found / total) * 5.0 if total > 0 else 0.0
        return score, f"{found}/{total} 레이아웃 존재"

    def _check_shapes(
        self,
        prs: Any,
        spec: Any,
        issues: list[str],
        suggestions: list[str],
    ) -> tuple[float, str]:
        """required_shapes가 해당 레이아웃에 존재하는지 확인한다."""
        if not hasattr(spec, "slides") or not spec.slides:
            return 5.0, "SlideSpec 미정의 — Shape 검증 생략"

        # 레이아웃별 shape 이름 수집
        layout_shapes: dict[str, set[str]] = {}
        for layout in prs.slide_layouts:
            shapes_in_layout: set[str] = set()
            for ph in layout.placeholders:
                if ph.name:
                    shapes_in_layout.add(ph.name)
            layout_shapes[layout.name] = shapes_in_layout

        total_required = 0
        total_found = 0
        missing_details: list[str] = []

        for slide_spec in spec.slides:
            if not slide_spec.required_shapes:
                continue
            available = layout_shapes.get(slide_spec.layout_name, set())
            for shape_name in slide_spec.required_shapes:
                total_required += 1
                if shape_name in available:
                    total_found += 1
                else:
                    missing_details.append(
                        f"레이아웃 '{slide_spec.layout_name}'에 "
                        f"shape '{shape_name}' 없음"
                    )

        if total_required == 0:
            return 5.0, "필수 Shape 미정의 — 검증 생략"

        if not missing_details:
            return 5.0, f"모든 필수 Shape 존재 ({total_required}개)"

        for detail in missing_details:
            issues.append(detail)
            suggestions.append(f"누락된 shape 추가 필요: {detail}")

        score = (total_found / total_required) * 5.0
        return score, f"{total_found}/{total_required} Shape 존재"

    def _check_fonts(
        self,
        spec: Any,
        issues: list[str],
        suggestions: list[str],
    ) -> tuple[float, str]:
        """brand_tokens의 폰트가 시스템에 사용 가능한지 확인한다.

        폰트 미존재는 critical이 아닌 warning 수준이다.
        """
        if not hasattr(spec, "brand_tokens") or not spec.brand_tokens:
            return 5.0, "브랜드 토큰 미정의 — 폰트 검증 생략"

        font_keys = [k for k in spec.brand_tokens if "font" in k.lower()]
        if not font_keys:
            return 5.0, "폰트 토큰 없음 — 검증 생략"

        font_names: list[str] = [spec.brand_tokens[k] for k in font_keys]
        missing_fonts: list[str] = []

        for font_name in font_names:
            if not self._is_font_available(font_name):
                missing_fonts.append(font_name)

        if not missing_fonts:
            return 5.0, f"모든 브랜드 폰트 사용 가능 ({len(font_names)}개)"

        for font in missing_fonts:
            issues.append(f"폰트 미설치 (warning): {font}")
            suggestions.append(f"폰트 '{font}' 설치 권장")

        # 폰트 미존재는 경미한 감점
        found = len(font_names) - len(missing_fonts)
        score = max(3.0, (found / len(font_names)) * 5.0)
        return score, f"폰트 {found}/{len(font_names)} 가용 (미설치: {missing_fonts})"

    @staticmethod
    def _is_font_available(font_name: str) -> bool:
        """시스템에 폰트가 설치되어 있는지 확인한다.

        matplotlib.font_manager가 있으면 정확한 검색,
        없으면 True를 반환하여 false negative를 방지한다.
        """
        try:
            from matplotlib.font_manager import findSystemFonts

            system_fonts = findSystemFonts()
            font_lower = font_name.lower()
            return any(font_lower in f.lower() for f in system_fonts)
        except ImportError:
            # matplotlib 미설치 — 폰트 검증 생략 (낙관적)
            logger.debug(
                "matplotlib 미설치 — 폰트 '%s' 검증 생략 (가용 간주)",
                font_name,
            )
            return True
