"""PPTX 슬라이드 팩토리 — 레이아웃별 슬라이드 생성 + 플레이스홀더 값 주입.

template_manager.py의 TemplateManager를 사용하여
5종 레이아웃 기반 슬라이드를 생성한다:
  - BLANK: 면책 등 완전 빈 슬라이드
  - FOREST: 배경 이미지 + 녹색 오버레이 (커버/TOC 간지/연락처)
  - BLANK_PGNO: 페이지 번호만 (재무제표 full-width)
  - MAIN: 제목바 + 각주 + 페이지번호 (일반 콘텐츠)
  - MAIN_w/Andersen: MAIN + 공동 브랜딩
"""

from __future__ import annotations

import logging
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from src.design_renderer.assets import image_path
from src.design_renderer.design_tokens import IMDesignTokens
from src.design_renderer.pptx_engine.font_helper import set_font_with_ea
from src.design_renderer.pptx_engine.template_manager import TemplateManager

logger = logging.getLogger(__name__)


def cm_to_inches(cm: float) -> float:
    """센티미터를 인치로 변환."""
    return cm / 2.54


class SlideFactory:
    """슬라이드 생성 팩토리.

    TemplateManager를 통해 레이아웃을 참조하고,
    각 유형의 슬라이드를 표준화된 API로 생성한다.
    """

    def __init__(
        self,
        manager: TemplateManager,
        prs: Any | None = None,
        tokens: IMDesignTokens | None = None,
    ) -> None:
        self._manager = manager
        self._tokens = tokens or manager.tokens
        self._prs = prs

    @property
    def prs(self) -> Any:
        if self._prs is None:
            raise RuntimeError(
                "Presentation이 없습니다. set_presentation()을 호출하세요."
            )
        return self._prs

    def set_presentation(self, prs: Any) -> None:
        self._prs = prs

    def add_blank_slide(self, prs: Any | None = None) -> Any:
        """빈 BLANK 슬라이드를 추가한다."""
        target = prs or self.prs
        return self._manager.add_slide("blank", target)

    def add_forest_slide(self, prs: Any | None = None) -> Any:
        """FOREST 배경 슬라이드 추가.

        FOREST 레이아웃 사용 + forest_cover.jpg 배경 이미지를 삽입한다.
        커버, TOC 간지, 연락처 슬라이드에 사용.

        Returns:
            생성된 Slide 객체 (배경 이미지 포함).
        """
        target = prs or self.prs
        slide = self._manager.add_slide("forest", target)
        self._apply_forest_background(slide)
        return slide

    def add_blank_pgno_slide(self, prs: Any | None = None) -> Any:
        """BLANK_PGNO 슬라이드 추가 (페이지 번호만).

        재무제표 등 full-width 콘텐츠에 사용.
        섹션 바 없이 전체 영역 활용.

        Returns:
            생성된 Slide 객체.
        """
        target = prs or self.prs
        return self._manager.add_slide("blank_pgno", target)

    def _apply_forest_background(self, slide: Any) -> None:
        """슬라이드에 FOREST 배경 이미지 삽입.

        forest_cover.jpg를 슬라이드 전체 크기로 삽입하고
        가장 뒤(z-order 0)로 이동한다.
        """
        try:
            bg_path = image_path("forest_cover")
            if not bg_path.exists():
                logger.warning(f"FOREST 배경 이미지 미존재: {bg_path}")
                return
        except ValueError:
            logger.warning("forest_cover 이미지 미등록")
            return

        t = self._tokens
        try:
            pic = slide.shapes.add_picture(
                str(bg_path),
                Inches(0),
                Inches(0),
                Inches(t.layout.page_width),
                Inches(t.layout.page_height),
            )
            # 배경을 z-order 맨 뒤로 이동
            sp_tree = slide.shapes._spTree
            sp_tree.remove(pic._element)
            sp_tree.insert(2, pic._element)  # index 2 = spTree의 첫 shape 위치
        except Exception as e:
            logger.warning(f"FOREST 배경 삽입 실패: {e}")

    def add_cover_slide(
        self,
        *,
        project_name: str = "",
        subtitle: str = "",
        date: str = "",
    ) -> Any:
        """표지 슬라이드 생성.

        FOREST 레이아웃 + 배경 이미지. 프로젝트명, 부제, 날짜, AMIC 로고 배치.
        TM/DM 원본 실측: FOREST 배경 위 흰색 텍스트.

        Args:
            project_name: 프로젝트명 (예: "Project TITAN").
            subtitle: 부제 (예: "Information Memorandum").
            date: 날짜 문자열.

        Returns:
            생성된 Slide 객체.
        """
        t = self._tokens
        c = t.colors
        typo = t.typography
        f = t.font_sizes

        slide = self.add_forest_slide()

        # 프로젝트명 — FOREST 배경 위 흰색 텍스트 (토큰 좌표)
        dp = t.dual_panel
        title_shape = slide.shapes.add_textbox(
            Inches(cm_to_inches(dp.cover_margin_x)),
            Inches(cm_to_inches(dp.cover_title_y)),
            Inches(cm_to_inches(20.0)),
            Inches(cm_to_inches(3.0)),
        )
        tf = title_shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = project_name
        set_font_with_ea(run, typo.font_heading)
        run.font.size = Pt(f.cover_title)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(
            c.text_white.lstrip("#")
        )

        # 부제 + 날짜 (토큰 좌표)
        if subtitle or date:
            sub_shape = slide.shapes.add_textbox(
                Inches(cm_to_inches(dp.cover_margin_x)),
                Inches(cm_to_inches(dp.cover_subtitle_y)),
                Inches(cm_to_inches(20.0)),
                Inches(cm_to_inches(2.0)),
            )
            stf = sub_shape.text_frame
            stf.word_wrap = True
            if subtitle:
                p = stf.paragraphs[0]
                run = p.add_run()
                run.text = subtitle
                set_font_with_ea(run, typo.font_cover_subtitle)
                run.font.size = Pt(f.cover_subtitle)
                run.font.color.rgb = RGBColor.from_string(
                    c.text_white.lstrip("#")
                )
            if date:
                p = stf.add_paragraph() if subtitle else stf.paragraphs[0]
                run = p.add_run()
                run.text = date
                set_font_with_ea(run, typo.font_cover_subtitle)
                run.font.size = Pt(f.cover_date)
                run.font.color.rgb = RGBColor.from_string(
                    c.text_white.lstrip("#")
                )

        # AMIC 로고 (흰색, 하단)
        self._add_logo(slide, logo_type="white", bottom=True)

        return slide

    def add_disclaimer_slide(
        self,
        *,
        text: str = "",
        title: str = "Disclaimer",
    ) -> Any:
        """면책조항 슬라이드 생성.

        BLANK 레이아웃 사용. 정형 텍스트 1~2 슬라이드.

        Args:
            text: 면책조항 텍스트. 빈 문자열이면 기본 문구 사용.
            title: 슬라이드 제목.

        Returns:
            생성된 Slide 객체.
        """
        t = self._tokens
        c = t.colors
        typo = t.typography
        f = t.font_sizes
        lay = t.layout

        if not text:
            text = t.footer_note

        slide = self._manager.add_slide("blank", self.prs)

        # 제목
        title_shape = slide.shapes.add_textbox(
            Inches(lay.margin_left),
            Inches(lay.margin_top),
            Inches(lay.content_width),
            Inches(0.5),
        )
        tf = title_shape.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.name = typo.font_heading
        run.font.size = Pt(f.slide_title)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(c.primary.lstrip("#"))

        # 면책조항 본문
        body_shape = slide.shapes.add_textbox(
            Inches(lay.margin_left),
            Inches(lay.content_top),
            Inches(lay.content_width),
            Inches(lay.content_height),
        )
        btf = body_shape.text_frame
        btf.word_wrap = True
        p = btf.paragraphs[0]
        run = p.add_run()
        run.text = text
        run.font.name = typo.font_body
        run.font.size = Pt(f.footnote)
        run.font.color.rgb = RGBColor.from_string(
            c.text_secondary.lstrip("#")
        )

        return slide

    def add_content_slide(
        self,
        *,
        title: str = "",
    ) -> Any:
        """콘텐츠 슬라이드 생성.

        MAIN 레이아웃 사용. 제목(idx=11) 설정 후 빈 콘텐츠 영역 반환.
        호출자가 shape_builder 등으로 콘텐츠를 추가한다.

        Args:
            title: 슬라이드 제목.

        Returns:
            생성된 Slide 객체.
        """
        t = self._tokens

        slide = self._manager.add_slide("main", self.prs)

        # 제목 플레이스홀더 설정
        if title:
            self._manager.set_placeholder_text(
                slide, t.layout.ph_title_idx, title
            )

        return slide

    def add_contact_slide(
        self,
        *,
        contacts: list[dict[str, str]] | None = None,
        title: str = "Contact",
    ) -> Any:
        """연락처/종료 슬라이드 생성.

        FOREST 레이아웃 + 배경 이미지. 담당자 정보 + AMIC 로고 배치.

        Args:
            contacts: 연락처 리스트. 각 항목은 dict:
                {"name": "...", "title": "...", "email": "...", "phone": "..."}
            title: 슬라이드 제목.

        Returns:
            생성된 Slide 객체.
        """
        t = self._tokens
        c = t.colors
        typo = t.typography
        f = t.font_sizes
        lay = t.layout

        slide = self.add_forest_slide()

        # 제목 (FOREST 배경 위 흰색)
        title_shape = slide.shapes.add_textbox(
            Inches(lay.page_width / 2 - 2),
            Inches(1.5),
            Inches(4),
            Inches(0.6),
        )
        tf = title_shape.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        set_font_with_ea(run, typo.font_heading)
        run.font.size = Pt(f.toc_section_title)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(c.text_white.lstrip("#"))

        # 연락처 정보 (FOREST 배경 위 흰색)
        if contacts:
            y_pos = 2.5
            for contact in contacts:
                info_shape = slide.shapes.add_textbox(
                    Inches(lay.page_width / 2 - 2.5),
                    Inches(y_pos),
                    Inches(5),
                    Inches(0.9),
                )
                itf = info_shape.text_frame
                itf.word_wrap = True

                # 이름
                p_name = itf.paragraphs[0]
                p_name.alignment = PP_ALIGN.CENTER
                r = p_name.add_run()
                r.text = contact.get("name", "")
                set_font_with_ea(r, typo.font_body)
                r.font.size = Pt(f.summary_text)
                r.font.bold = True
                r.font.color.rgb = RGBColor.from_string(
                    c.text_white.lstrip("#")
                )

                # 직함
                if contact.get("title"):
                    p_title = itf.add_paragraph()
                    p_title.alignment = PP_ALIGN.CENTER
                    r = p_title.add_run()
                    r.text = contact["title"]
                    set_font_with_ea(r, typo.font_body)
                    r.font.size = Pt(f.footnote)
                    r.font.color.rgb = RGBColor.from_string(
                        c.text_white.lstrip("#")
                    )

                # 이메일/전화
                details = []
                if contact.get("email"):
                    details.append(contact["email"])
                if contact.get("phone"):
                    details.append(contact["phone"])
                if details:
                    p_detail = itf.add_paragraph()
                    p_detail.alignment = PP_ALIGN.CENTER
                    r = p_detail.add_run()
                    r.text = " | ".join(details)
                    set_font_with_ea(r, typo.font_body)
                    r.font.size = Pt(f.body)
                    r.font.color.rgb = RGBColor.from_string(
                        c.text_white.lstrip("#")
                    )

                y_pos += 1.1

        # AMIC 로고 (흰색, 하단)
        self._add_logo(slide, logo_type="white", bottom=True)

        return slide

    def _add_logo(
        self,
        slide: Any,
        *,
        logo_type: str = "dark",
        bottom: bool = True,
    ) -> None:
        """슬라이드에 AMIC 로고 삽입.

        Args:
            slide: Slide 객체.
            logo_type: "dark" | "white"
            bottom: True이면 하단 중앙 배치.
        """
        name = f"amic_logo_{logo_type}"
        try:
            logo = image_path(name)
            if not logo.exists():
                logger.warning(f"로고 파일 미존재: {logo}")
                return
        except ValueError:
            return

        lay = self._tokens.layout
        logo_width = Inches(1.8)
        logo_height = Inches(0.5)

        if bottom:
            left = Inches(lay.page_width / 2) - logo_width // 2
            top = Inches(lay.page_height - 0.8)
        else:
            left = Inches(lay.margin_left)
            top = Inches(lay.margin_top)

        try:
            slide.shapes.add_picture(
                str(logo), left, top, logo_width, logo_height
            )
        except Exception as e:
            logger.warning(f"로고 삽입 실패: {e}")
