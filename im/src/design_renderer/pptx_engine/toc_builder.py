"""TOC 구분 슬라이드 생성 — BLANK 레이아웃, 현재 섹션 ALL CAPS 하이라이트.

섹션 목록을 2열 테이블 (번호+섹션명)로 배치하고,
현재 활성 섹션을 AMIC 컬러로 강조한다.

디자인 토큰 연동: dual_panel.toc_x/y/width/height, font_sizes.toc_heading 참조.
"""

from __future__ import annotations

import logging
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from src.design_renderer.design_tokens import DEFAULT_TOKENS, IMDesignTokens
from src.design_renderer.pptx_engine.font_helper import set_font_with_ea
from src.design_renderer.pptx_engine.slide_factory import cm_to_inches

logger = logging.getLogger(__name__)


# 섹션 ID → 한글 표시명 매핑
SECTION_DISPLAY_NAMES: dict[str, str] = {
    "deal_overview": "거래 개요",
    "executive_summary": "Executive Summary",
    "investment_highlights": "투자 포인트",
    "company_overview": "회사 개요",
    "business_model": "사업 모델",
    "market_overview": "시장 분석",
    "business_overview": "사업부별 분석",
    "value_creation": "가치 창출 계획",
    "growth_strategy": "성장 전략",
    "financial_analysis": "재무 분석",
    "management_team": "경영진 소개",
    "shareholder_structure": "주주 구성",
    "transaction_structure": "거래 구조",
    "appendix": "부록",
    # TM (Teaser Memorandum) 전용
    "target_positioning": "Target Positioning",
    "market_outlook": "Market Outlook",
    "demand_driver": "Key Demand Driver",
    "supply_driver": "Key Supply Driver",
    "target_overview": "Target Overview",
    "target_highlights": "Target Highlights",
    "proforma_plan": "Pro-Forma 사업계획",
    "proforma_financials": "Pro-Forma 재무제표",
    # DM (Discussion Memorandum) 전용
    "dm_market_trends": "Market & Transaction Trends",
    "dm_deal_structure": "Deal Structure Considerations",
    "dm_investment_thesis": "Investment Thesis",
    "dm_valuation": "Valuation Analysis",
    "dm_risk_assessment": "Risk Assessment",
    "dm_summary": "Summary & Recommendations",
}


def build_toc_slide(
    slide: Any,
    sections: list[str],
    current_section: str,
    *,
    tokens: IMDesignTokens | None = None,
    section_names: dict[str, str] | None = None,
) -> Any:
    """TOC 구분 슬라이드 콘텐츠 생성.

    BLANK 레이아웃 슬라이드에 "TABLE OF CONTENTS" 제목과
    섹션 목록을 배치한다. 현재 섹션은 ALL CAPS + AMIC 컬러.

    좌표는 design_tokens.dual_panel에서 참조 (cm → inches 변환).

    Args:
        slide: Slide 객체 (BLANK 레이아웃).
        sections: 표시할 섹션 ID 리스트.
        current_section: 현재 활성 섹션 ID.
        tokens: 디자인 토큰.
        section_names: 섹션 ID→표시명 매핑 오버라이드.

    Returns:
        slide 객체.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    names = section_names or SECTION_DISPLAY_NAMES
    c = tokens.colors
    t = tokens.typography
    f = tokens.font_sizes
    dp = tokens.dual_panel

    # TOC에서 제외할 섹션
    excluded = {"cover", "disclaimer", "toc_divider", "contact"}
    toc_sections = [s for s in sections if s not in excluded]

    if not toc_sections:
        return slide

    # 제목 좌표 (디자인 토큰 기반: TOC 테이블 위)
    title_x = cm_to_inches(dp.toc_x)
    title_y = cm_to_inches(dp.toc_y) - 1.0  # 테이블 위 약 1인치
    title_w = cm_to_inches(dp.toc_width)

    title_shape = slide.shapes.add_textbox(
        Inches(title_x),
        Inches(title_y),
        Inches(title_w),
        Inches(0.6),
    )
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "TABLE OF CONTENTS"
    set_font_with_ea(run, t.font_heading)
    run.font.size = Pt(f.toc_heading)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(c.primary.lstrip("#"))

    # 섹션 목록 (디자인 토큰 기반)
    list_top = cm_to_inches(dp.toc_y)
    item_height = cm_to_inches(dp.toc_height) / max(len(toc_sections), 1)
    item_height = min(item_height, 0.45)  # 최대 높이 제한

    for idx, section_id in enumerate(toc_sections):
        display_name = names.get(section_id, section_id)
        is_current = section_id == current_section
        y = list_top + idx * item_height

        # 번호
        num_shape = slide.shapes.add_textbox(
            Inches(title_x),
            Inches(y),
            Inches(0.5),
            Inches(item_height),
        )
        ntf = num_shape.text_frame
        p_num = ntf.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        r_num = p_num.add_run()
        r_num.text = f"{idx + 1:02d}"
        set_font_with_ea(r_num, t.font_mono)
        r_num.font.size = Pt(f.summary_text)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor.from_string(
            (c.accent if is_current else c.text_secondary).lstrip("#")
        )

        # 섹션명
        name_shape = slide.shapes.add_textbox(
            Inches(title_x + 0.7),
            Inches(y),
            Inches(title_w - 0.7),
            Inches(item_height),
        )
        stf = name_shape.text_frame
        p_name = stf.paragraphs[0]
        r_name = p_name.add_run()
        r_name.text = display_name.upper() if is_current else display_name
        set_font_with_ea(r_name, t.font_body)
        r_name.font.size = Pt(f.summary_text)
        r_name.font.bold = is_current
        r_name.font.color.rgb = RGBColor.from_string(
            (c.primary if is_current else c.text_secondary).lstrip("#")
        )

        # 현재 섹션 하이라이트 바
        if is_current:
            _add_highlight_bar(slide, y, item_height, tokens)

    return slide


def build_toc_slide_html(
    sections: list[str],
    current_section: str,
    *,
    tokens: IMDesignTokens | None = None,
    section_names: dict[str, str] | None = None,
) -> str:
    """HTML TOC 구분 슬라이드 생성."""
    if tokens is None:
        tokens = DEFAULT_TOKENS

    from html import escape as html_escape

    names = section_names or SECTION_DISPLAY_NAMES
    excluded = {"cover", "disclaimer", "toc_divider", "contact"}
    toc_sections = [s for s in sections if s not in excluded]

    items_html = ""
    for idx, section_id in enumerate(toc_sections):
        display_name = names.get(section_id, section_id)
        is_current = section_id == current_section
        item_class = "toc-item toc-current" if is_current else "toc-item"
        text = display_name.upper() if is_current else display_name

        items_html += (
            f'<li class="{item_class}">'
            f'<span class="toc-number">{idx + 1:02d}</span>'
            f"<span>{html_escape(text)}</span>"
            f"</li>\n"
        )

    return f"""<div class="slide slide-toc">
    <div class="toc-heading">TABLE OF CONTENTS</div>
    <ul class="toc-list">
        {items_html}
    </ul>
</div>"""


def _add_highlight_bar(
    slide: Any,
    y: float,
    height: float,
    tokens: IMDesignTokens,
) -> None:
    """현재 섹션 하이라이트용 좌측 바."""
    from pptx.enum.shapes import MSO_SHAPE

    dp = tokens.dual_panel
    bar_x = cm_to_inches(dp.toc_x) - 0.1  # 번호 왼쪽

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(bar_x),
        Inches(y),
        Inches(0.08),
        Inches(height),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(
        tokens.colors.accent.lstrip("#")
    )
    bar.line.fill.background()


# ---------------------------------------------------------------------------
# TM (Teaser Memorandum) 전용 고정 TOC
# ---------------------------------------------------------------------------


def build_tm_toc_slide(
    slide: Any,
    current_group: str,
    *,
    page_numbers: dict[str, int] | None = None,
    tokens: IMDesignTokens | None = None,
) -> Any:
    """TM 전용 고정 4그룹 TOC 슬라이드 생성.

    TEASER_TOC_GROUPS 기반 고정 구조를 항상 4그룹 모두 표시하고,
    current_group에 해당하는 그룹을 하이라이트한다.

    좌표는 design_tokens에서 참조 (toc_heading, dual_panel 등).

    Args:
        slide: Slide 객체 (BLANK 레이아웃).
        current_group: 현재 활성 그룹 key.
        page_numbers: 그룹 key → 시작 페이지 번호 매핑.
        tokens: 디자인 토큰.

    Returns:
        slide 객체.
    """
    from src.design_renderer.im_document import TEASER_TOC_GROUPS

    if tokens is None:
        tokens = DEFAULT_TOKENS

    c = tokens.colors
    t = tokens.typography
    f = tokens.font_sizes
    dp = tokens.dual_panel
    page_numbers = page_numbers or {}

    # 제목 좌표 (디자인 토큰 기반)
    title_x = cm_to_inches(dp.toc_x)
    title_y = cm_to_inches(dp.toc_title_y)
    title_w = cm_to_inches(dp.toc_width)

    title_shape = slide.shapes.add_textbox(
        Inches(title_x), Inches(title_y), Inches(title_w), Inches(0.6),
    )
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "TABLE OF CONTENTS"
    set_font_with_ea(run, t.font_heading)
    run.font.size = Pt(f.toc_heading)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(c.primary.lstrip("#"))

    # 4그룹 렌더링 (디자인 토큰 기반 Y 좌표)
    list_start_y = cm_to_inches(dp.toc_y)  # 실측: 6.024cm
    group_height = cm_to_inches(dp.toc_height) / 4  # 4그룹 균등 분배

    for idx, group in enumerate(TEASER_TOC_GROUPS):
        group_key = group["key"]
        group_title = group["title"]
        subsections = group["subsections"]
        is_current = group_key == current_group
        page_num = page_numbers.get(group_key)

        group_y = list_start_y + idx * group_height

        # 하이라이트 바 (현재 그룹)
        if is_current:
            _add_highlight_bar(slide, group_y, group_height - 0.15, tokens)

        # 그룹 번호
        num_shape = slide.shapes.add_textbox(
            Inches(title_x), Inches(group_y), Inches(0.5), Inches(0.4),
        )
        ntf = num_shape.text_frame
        p_num = ntf.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        r_num = p_num.add_run()
        r_num.text = f"({idx + 1})"
        set_font_with_ea(r_num, t.font_mono)
        r_num.font.size = Pt(f.summary_text)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor.from_string(
            (c.accent if is_current else c.text_secondary).lstrip("#")
        )

        # 그룹 제목
        name_shape = slide.shapes.add_textbox(
            Inches(title_x + 0.7), Inches(group_y),
            Inches(title_w - 2.2), Inches(0.4),
        )
        stf = name_shape.text_frame
        p_name = stf.paragraphs[0]
        r_name = p_name.add_run()
        r_name.text = group_title.upper() if is_current else group_title
        set_font_with_ea(r_name, t.font_body)
        r_name.font.size = Pt(f.summary_text)
        r_name.font.bold = is_current
        r_name.font.color.rgb = RGBColor.from_string(
            (c.primary if is_current else c.text_secondary).lstrip("#")
        )

        # 페이지 번호
        if page_num is not None:
            pg_shape = slide.shapes.add_textbox(
                Inches(title_x + title_w - 1.5), Inches(group_y),
                Inches(1.5), Inches(0.4),
            )
            ptf = pg_shape.text_frame
            p_pg = ptf.paragraphs[0]
            p_pg.alignment = PP_ALIGN.RIGHT
            r_pg = p_pg.add_run()
            r_pg.text = f"p.{page_num}"
            set_font_with_ea(r_pg, t.font_mono)
            r_pg.font.size = Pt(11)
            r_pg.font.bold = is_current
            r_pg.font.color.rgb = RGBColor.from_string(
                (c.accent if is_current else c.text_secondary).lstrip("#")
            )

        # 하위 섹션 목록
        sub_y = group_y + 0.4
        for sub_idx, (_, sub_name) in enumerate(subsections):
            sub_shape = slide.shapes.add_textbox(
                Inches(title_x + 1.0),
                Inches(sub_y + sub_idx * 0.22),
                Inches(title_w - 1.0),
                Inches(0.22),
            )
            sub_tf = sub_shape.text_frame
            p_sub = sub_tf.paragraphs[0]
            r_sub = p_sub.add_run()
            r_sub.text = f"- {sub_name}"
            set_font_with_ea(r_sub, t.font_body)
            r_sub.font.size = Pt(9)
            r_sub.font.color.rgb = RGBColor.from_string(
                (c.text_body if is_current else c.text_secondary).lstrip("#")
            )

    return slide


def build_tm_toc_slide_html(
    current_group: str,
    *,
    page_numbers: dict[str, int] | None = None,
    tokens: IMDesignTokens | None = None,
) -> str:
    """TM 전용 고정 4그룹 TOC HTML 슬라이드 생성."""
    from html import escape as html_escape

    from src.design_renderer.im_document import TEASER_TOC_GROUPS

    if tokens is None:
        tokens = DEFAULT_TOKENS

    c = tokens.colors
    page_numbers = page_numbers or {}

    groups_html = ""
    for idx, group in enumerate(TEASER_TOC_GROUPS):
        group_key = group["key"]
        group_title = group["title"]
        subsections = group["subsections"]
        is_current = group_key == current_group
        page_num = page_numbers.get(group_key)

        border_left = f"3px solid {c.accent}" if is_current else "3px solid transparent"
        bg = c.bg_light_green if is_current else "transparent"
        title_color = c.primary if is_current else c.text_secondary
        title_weight = "bold" if is_current else "normal"
        title_text = html_escape(
            group_title.upper() if is_current else group_title
        )

        pg_html = ""
        if page_num is not None:
            pg_color = c.accent if is_current else c.text_secondary
            pg_html = (
                f'<span style="font-family:\'IBM Plex Mono\',monospace;'
                f'font-size:10pt;color:{pg_color};float:right;">'
                f"p.{page_num}</span>"
            )

        subs_html = ""
        for _, sub_name in subsections:
            sub_color = c.text_body if is_current else c.text_secondary
            subs_html += (
                f'<div style="font-size:9pt;color:{sub_color};'
                f'margin-left:1em;line-height:1.5;">'
                f"- {html_escape(sub_name)}</div>"
            )

        groups_html += (
            f'<div style="border-left:{border_left};background:{bg};'
            f'padding:0.5em 0.8em;margin-bottom:0.6em;border-radius:2px;">'
            f'<div style="font-size:12pt;font-weight:{title_weight};'
            f'color:{title_color};">'
            f'<span style="font-family:\'IBM Plex Mono\',monospace;'
            f'margin-right:0.5em;">({idx + 1})</span>'
            f"{title_text}{pg_html}</div>"
            f"{subs_html}</div>"
        )

    return (
        f'<div class="slide slide-toc">'
        f'<div class="toc-heading" style="font-size:20pt;font-weight:bold;'
        f'color:{c.primary};margin-bottom:1em;">TABLE OF CONTENTS</div>'
        f"{groups_html}</div>"
    )
