"""PPTX 마스터 템플릿 관리 — amic_im_template.pptx 로드 및 레이아웃 참조.

create_template.py가 생성한 템플릿 파일을 로드하고,
COVER/BLANK/MAIN 레이아웃 참조와 플레이스홀더 접근 API를 제공한다.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches

from src.design_renderer.design_tokens import DEFAULT_TOKENS, IMDesignTokens
from src.design_renderer.pptx_engine.create_template import (
    create_im_template,
    get_layout_by_purpose,
)

logger = logging.getLogger(__name__)

_ASSETS_DIR = Path(__file__).resolve().parents[1] / "assets"
_DEFAULT_TEMPLATE = _ASSETS_DIR / "templates" / "amic_im_template.pptx"


class TemplateManager:
    """PPTX 마스터 템플릿 관리자.

    템플릿 파일을 로드하고 레이아웃 참조를 제공한다.
    템플릿이 존재하지 않으면 자동 생성한다.
    """

    def __init__(
        self,
        template_path: Path | None = None,
        tokens: IMDesignTokens | None = None,
    ) -> None:
        self._tokens = tokens or DEFAULT_TOKENS
        self._template_path = template_path or _DEFAULT_TEMPLATE
        self._prs: Presentation | None = None

    @property
    def template_path(self) -> Path:
        return self._template_path

    @property
    def tokens(self) -> IMDesignTokens:
        return self._tokens

    def _ensure_template(self) -> Path:
        """템플릿 파일이 없으면 자동 생성."""
        if not self._template_path.exists():
            logger.info("템플릿 파일 미존재, 자동 생성 시작")
            create_im_template(
                output_path=self._template_path,
                tokens=self._tokens,
            )
        return self._template_path

    def new_presentation(self) -> Presentation:
        """템플릿 기반 새 Presentation 생성.

        매번 새 Presentation을 반환한다 (멀티 문서 생성 안전).

        Returns:
            빈 슬라이드의 새 Presentation 객체.
        """
        path = self._ensure_template()
        prs = Presentation(str(path))

        # 슬라이드 크기 보장
        prs.slide_width = Inches(self._tokens.layout.page_width)
        prs.slide_height = Inches(self._tokens.layout.page_height)

        self._prs = prs
        return prs

    def get_layout(self, purpose: str, prs: Presentation | None = None) -> Any:
        """목적별 레이아웃 참조 반환.

        Args:
            purpose: "cover" | "blank" | "main"
            prs: Presentation 객체. None이면 마지막 생성된 것 사용.

        Returns:
            SlideLayout 객체.

        Raises:
            ValueError: 알 수 없는 purpose.
            RuntimeError: Presentation이 없을 때.
        """
        presentation = prs or self._prs
        if presentation is None:
            raise RuntimeError(
                "Presentation이 없습니다. new_presentation()을 먼저 호출하세요."
            )
        return get_layout_by_purpose(presentation, purpose)

    def get_cover_layout(self, prs: Presentation | None = None) -> Any:
        """COVER 레이아웃 반환."""
        return self.get_layout("cover", prs)

    def get_blank_layout(self, prs: Presentation | None = None) -> Any:
        """BLANK 레이아웃 반환 (TOC 구분자/면책/연락처)."""
        return self.get_layout("blank", prs)

    def get_main_layout(self, prs: Presentation | None = None) -> Any:
        """MAIN 레이아웃 반환 (콘텐츠 슬라이드)."""
        return self.get_layout("main", prs)

    def add_slide(self, purpose: str, prs: Presentation | None = None) -> Any:
        """레이아웃 목적에 따라 새 슬라이드 추가.

        Args:
            purpose: "cover" | "blank" | "main"
            prs: Presentation 객체.

        Returns:
            새로 추가된 Slide 객체.
        """
        presentation = prs or self._prs
        if presentation is None:
            raise RuntimeError(
                "Presentation이 없습니다. new_presentation()을 먼저 호출하세요."
            )
        layout = self.get_layout(purpose, presentation)
        return presentation.slides.add_slide(layout)

    @staticmethod
    def get_placeholder(slide: Any, idx: int) -> Any | None:
        """슬라이드에서 특정 idx의 플레이스홀더 반환.

        Args:
            slide: Slide 객체.
            idx: 플레이스홀더 인덱스 (11=제목, 12=각주, 13=페이지번호).

        Returns:
            플레이스홀더 Shape 또는 None.
        """
        for shape in slide.placeholders:
            if shape.placeholder_format.idx == idx:
                return shape
        return None

    @staticmethod
    def set_placeholder_text(slide: Any, idx: int, text: str) -> bool:
        """플레이스홀더 텍스트 설정.

        Args:
            slide: Slide 객체.
            idx: 플레이스홀더 인덱스.
            text: 설정할 텍스트.

        Returns:
            성공 여부.
        """
        ph = TemplateManager.get_placeholder(slide, idx)
        if ph is None:
            return False
        ph.text = text
        return True
