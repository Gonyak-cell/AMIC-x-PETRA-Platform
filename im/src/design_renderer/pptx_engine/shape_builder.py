"""PPTX Shape 빌더 — TITAN/COVENANT shape 패턴 재현.

슬라이드에 표준화된 shape(텍스트박스, 서브헤더, 테이블, KPI, 차트)를
추가하는 API를 제공한다.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from src.design_renderer.design_tokens import DEFAULT_TOKENS, IMDesignTokens
from src.design_renderer.pptx_engine.font_helper import set_font_with_ea

logger = logging.getLogger(__name__)


def shape_bottom_inches(shape: Any) -> float:
    """shape 하단 위치를 inches로 반환 — (top + height) / 914400 EMU."""
    return (shape.top + shape.height) / 914400


def add_summary_textbox(
    slide: Any,
    text: str,
    *,
    left: float | None = None,
    top: float | None = None,
    width: float | None = None,
    height: float = 0.5,
    tokens: IMDesignTokens | None = None,
) -> Any:
    """요약/설명 텍스트박스 추가 (14pt Bold).

    슬라이드 상단 요약 영역에 사용된다.

    Args:
        slide: Slide 인스턴스.
        text: 요약 텍스트.
        left, top, width, height: 위치/크기 (inches).
        tokens: 디자인 토큰.

    Returns:
        생성된 TextBox shape.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    lay = tokens.layout
    c = tokens.colors
    t = tokens.typography
    f = tokens.font_sizes

    shape = slide.shapes.add_textbox(
        Inches(left if left is not None else lay.content_left),
        Inches(top if top is not None else lay.content_top),
        Inches(width if width is not None else lay.content_width),
        Inches(height),
    )

    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_font_with_ea(run, t.font_body)
    run.font.size = Pt(f.summary_text)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(c.text_body.lstrip("#"))

    return shape


def add_sub_header_bar(
    slide: Any,
    text: str,
    *,
    left: float | None = None,
    top: float | None = None,
    width: float | None = None,
    height: float = 0.35,
    tokens: IMDesignTokens | None = None,
) -> Any:
    """컬러 서브헤더 바 추가 (AMIC 다크그린 bg + 12pt Bold White).

    Args:
        slide: Slide 인스턴스.
        text: 서브헤더 텍스트.
        left, top, width, height: 위치/크기 (inches).
        tokens: 디자인 토큰.

    Returns:
        생성된 shape.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    lay = tokens.layout
    c = tokens.colors
    t = tokens.typography
    f = tokens.font_sizes

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left if left is not None else lay.content_left),
        Inches(top if top is not None else lay.content_top),
        Inches(width if width is not None else lay.content_width),
        Inches(height),
    )

    # 배경색
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(c.primary.lstrip("#"))
    shape.line.fill.background()

    # 텍스트
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    set_font_with_ea(run, t.font_body)
    run.font.size = Pt(f.sub_header_bar)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(c.text_white.lstrip("#"))

    return shape


def add_body_textbox(
    slide: Any,
    text: str,
    *,
    left: float | None = None,
    top: float | None = None,
    width: float | None = None,
    height: float = 2.0,
    font_size: int | None = None,
    tokens: IMDesignTokens | None = None,
) -> Any:
    """본문 텍스트박스 추가 (10pt).

    Args:
        slide: Slide 인스턴스.
        text: 본문 텍스트.
        left, top, width, height: 위치/크기 (inches).
        font_size: 폰트 크기 오버라이드. None이면 body (10pt).
        tokens: 디자인 토큰.

    Returns:
        생성된 TextBox shape.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    lay = tokens.layout
    c = tokens.colors
    t = tokens.typography
    f = tokens.font_sizes

    shape = slide.shapes.add_textbox(
        Inches(left if left is not None else lay.content_left),
        Inches(top if top is not None else lay.content_top + 0.5),
        Inches(width if width is not None else lay.content_width),
        Inches(height),
    )

    tf = shape.text_frame
    tf.word_wrap = True

    # 텍스트를 줄바꿈 기준으로 분리
    for i, line in enumerate(text.split("\n")):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        run = p.add_run()
        run.text = line
        set_font_with_ea(run, t.font_body)
        run.font.size = Pt(font_size if font_size is not None else f.body)
        run.font.color.rgb = RGBColor.from_string(c.text_body.lstrip("#"))

    return shape


def add_bullet_list(
    slide: Any,
    items: list[str],
    *,
    left: float | None = None,
    top: float | None = None,
    width: float | None = None,
    height: float = 3.0,
    tokens: IMDesignTokens | None = None,
) -> Any:
    """불릿 리스트 텍스트박스 추가.

    Args:
        slide: Slide 인스턴스.
        items: 불릿 아이템 리스트.
        left, top, width, height: 위치/크기 (inches).
        tokens: 디자인 토큰.

    Returns:
        생성된 TextBox shape.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    lay = tokens.layout
    c = tokens.colors
    t = tokens.typography
    f = tokens.font_sizes

    shape = slide.shapes.add_textbox(
        Inches(left if left is not None else lay.content_left),
        Inches(top if top is not None else lay.content_top + 0.5),
        Inches(width if width is not None else lay.content_width),
        Inches(height),
    )

    tf = shape.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        run.text = f"• {item}"
        set_font_with_ea(run, t.font_body)
        run.font.size = Pt(f.body)
        run.font.color.rgb = RGBColor.from_string(c.text_body.lstrip("#"))

    return shape


def add_chart_image(
    slide: Any,
    image_path_or_bytes: str | Path | bytes,
    *,
    left: float | None = None,
    top: float | None = None,
    width: float | None = None,
    height: float = 3.5,
    tokens: IMDesignTokens | None = None,
) -> Any:
    """차트 이미지 삽입.

    chart_embed.py가 생성한 PNG를 슬라이드에 삽입한다.

    Args:
        slide: Slide 인스턴스.
        image_path_or_bytes: 이미지 파일 경로 또는 바이트.
        left, top, width, height: 위치/크기 (inches).
        tokens: 디자인 토큰.

    Returns:
        생성된 Picture shape.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    lay = tokens.layout

    pos_left = Inches(left if left is not None else lay.content_left)
    pos_top = Inches(top if top is not None else lay.content_top + 0.5)
    pos_width = Inches(width if width is not None else lay.content_width)
    pos_height = Inches(height)

    if isinstance(image_path_or_bytes, bytes):
        from io import BytesIO

        stream = BytesIO(image_path_or_bytes)
        return slide.shapes.add_picture(
            stream, pos_left, pos_top, pos_width, pos_height
        )

    return slide.shapes.add_picture(
        str(image_path_or_bytes), pos_left, pos_top, pos_width, pos_height
    )


def add_chart_or_image(
    slide: Any,
    chart: Any,
    *,
    left: float | None = None,
    top: float | None = None,
    width: float | None = None,
    height: float = 3.5,
    tokens: IMDesignTokens | None = None,
) -> Any:
    """ChartData에서 네이티브 PPTX 차트 또는 이미지를 자동 선택하여 삽입.

    chart_type이 네이티브 전환 대상(stacked_bar, donut, line, hbar)이면
    편집 가능한 PPTX 차트로, 그 외에는 기존 이미지 방식으로 삽입한다.

    Args:
        slide: Slide 인스턴스.
        chart: ChartData 인스턴스 (chart_type, title, data).
        left, top, width, height: 위치/크기 (inches).
        tokens: 디자인 토큰.

    Returns:
        생성된 shape (Chart 또는 Picture).
    """
    from src.chart_engine.config import chart_config_from_design_tokens
    from src.chart_engine.pptx_native import get_native_builder

    if tokens is None:
        tokens = DEFAULT_TOKENS

    lay = tokens.layout
    pos_left = left if left is not None else lay.content_left
    pos_top = top if top is not None else lay.content_top + 0.5
    pos_width = width if width is not None else lay.content_width
    pos_height = height

    chart_type = getattr(chart, "chart_type", "")
    builder = get_native_builder(chart_type)

    if builder is not None:
        cfg = chart_config_from_design_tokens(tokens)
        chart_data_dict = getattr(chart, "data", {}) or {}
        chart_title = getattr(chart, "title", "") or ""
        return builder.build(
            slide,
            chart_data_dict,
            title=chart_title,
            left=pos_left,
            top=pos_top,
            width=pos_width,
            height=pos_height,
            config=cfg,
        )

    # 네이티브 미지원 → 기존 이미지 경로
    chart_data_dict = getattr(chart, "data", {}) or {}
    img = chart_data_dict.get("image_bytes") or chart_data_dict.get("image_path")
    if img:
        return add_chart_image(
            slide,
            img,
            left=pos_left,
            top=pos_top,
            width=pos_width,
            height=pos_height,
            tokens=tokens,
        )
    return None


def add_kpi_grid(
    slide: Any,
    kpis: list[dict[str, Any]],
    *,
    left: float | None = None,
    top: float | None = None,
    width: float | None = None,
    cols: int = 4,
    tokens: IMDesignTokens | None = None,
    number_config: Any | None = None,
) -> list[Any]:
    """KPI 카드 그리드 추가.

    kpi_card.py의 render_kpi_grid_pptx 래퍼.

    Args:
        slide: Slide 인스턴스.
        kpis: KPI 데이터 리스트.
        left, top, width: 그리드 위치/크기.
        cols: 열 수.
        tokens: 디자인 토큰.
        number_config: 숫자 표기 설정.

    Returns:
        생성된 shape 리스트.
    """
    from src.design_renderer.components.kpi_card import render_kpi_grid_pptx

    return render_kpi_grid_pptx(
        slide,
        kpis,
        left=left,
        top=top,
        width=width,
        cols=cols,
        tokens=tokens,
        number_config=number_config,
    )


def add_section_bar(
    slide: Any,
    text: str,
    *,
    left: float | None = None,
    top: float | None = None,
    width: float | None = None,
    height: float | None = None,
    tokens: IMDesignTokens | None = None,
) -> Any:
    """섹션 타이틀 바 추가 (design_tokens.colors.section_bar_bg 참조).

    TM/DM 실측: #0F3A32 배경 + 12pt Bold 흰색 텍스트.

    Args:
        slide: Slide 인스턴스.
        text: 섹션 제목.
        left, top, width, height: 위치/크기 (inches). None이면 토큰 기본값.
        tokens: 디자인 토큰.

    Returns:
        생성된 shape.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    lay = tokens.layout
    c = tokens.colors
    t = tokens.typography
    f = tokens.font_sizes
    dp = tokens.dual_panel

    _left = left if left is not None else lay.content_left
    _top = top if top is not None else dp.section_bar_y / 2.54  # cm→inches
    _width = width if width is not None else lay.content_width
    _height = height if height is not None else dp.section_bar_height / 2.54

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(_left),
        Inches(_top),
        Inches(_width),
        Inches(_height),
    )

    # 배경색 — design_tokens.colors.section_bar_bg
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(c.section_bar_bg.lstrip("#"))
    shape.line.fill.background()

    # 텍스트
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    set_font_with_ea(run, t.font_body)
    run.font.size = Pt(f.sub_header_bar)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(c.text_white.lstrip("#"))

    return shape


def add_dual_section_bars(
    slide: Any,
    left_title: str,
    right_title: str,
    *,
    tokens: IMDesignTokens | None = None,
) -> tuple[Any, Any]:
    """듀얼 패널 섹션 바 쌍 추가.

    좌표는 design_tokens.dual_panel에서 참조.

    Returns:
        (좌측 shape, 우측 shape) 튜플.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    dp = tokens.dual_panel
    bar_y = dp.section_bar_y / 2.54
    bar_h = dp.section_bar_height / 2.54
    panel_w = dp.panel_width / 2.54

    left_bar = add_section_bar(
        slide,
        left_title,
        left=dp.left_x / 2.54,
        top=bar_y,
        width=panel_w,
        height=bar_h,
        tokens=tokens,
    )
    right_bar = add_section_bar(
        slide,
        right_title,
        left=dp.right_x / 2.54,
        top=bar_y,
        width=panel_w,
        height=bar_h,
        tokens=tokens,
    )

    return left_bar, right_bar


def add_financial_table(
    slide: Any,
    *,
    headers: list[str],
    rows: list[dict[str, Any]],
    left: float | None = None,
    top: float | None = None,
    tokens: IMDesignTokens | None = None,
    number_config: Any | None = None,
    show_cagr: bool = False,
) -> Any:
    """재무 테이블 추가.

    financial_table.py의 render_financial_table_pptx 래퍼.

    Args:
        slide: Slide 인스턴스.
        headers: 열 헤더.
        rows: 행 데이터.
        left, top: 위치.
        tokens: 디자인 토큰.
        number_config: 숫자 표기 설정.
        show_cagr: CAGR 열 표시.

    Returns:
        생성된 Table shape.
    """
    from src.design_renderer.components.financial_table import (
        render_financial_table_pptx,
    )

    return render_financial_table_pptx(
        slide,
        headers=headers,
        rows=rows,
        left=left,
        top=top,
        tokens=tokens,
        number_config=number_config,
        show_cagr=show_cagr,
    )
