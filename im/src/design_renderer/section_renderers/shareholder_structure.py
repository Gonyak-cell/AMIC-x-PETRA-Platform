"""주주 구성(Shareholder Structure) 섹션 렌더러.

주주별 지분율, 주식 수, 카테고리 등 주주 구성 정보를 테이블로 표시.
"""

from __future__ import annotations

import logging
from html import escape as html_escape
from typing import Any

from src.design_renderer.design_tokens import DEFAULT_TOKENS, IMDesignTokens
from src.design_renderer.im_document import IMDocumentData
from src.design_renderer.pdf_output.html_builder import build_slide_html
from src.design_renderer.section_renderers import register_renderer
from src.design_renderer.section_renderers.base import BaseSectionRenderer

logger = logging.getLogger(__name__)


@register_renderer
class ShareholderStructureRenderer(BaseSectionRenderer):
    """주주 구성 슬라이드 렌더러."""

    section_id = "shareholder_structure"

    def render_html(
        self,
        data: IMDocumentData,
        *,
        tokens: IMDesignTokens | None = None,
    ) -> list[str]:
        tokens = tokens or DEFAULT_TOKENS
        c = tokens.colors

        shareholders = data.shareholders
        narrative = html_escape(
            data.narratives.get("shareholder_structure", "")
        )

        narrative_html = ""
        if narrative:
            narrative_html = (
                f'<p style="font-size:10pt;color:{c.text_body};'
                f'line-height:1.6;margin-bottom:1em;">{narrative}</p>'
            )

        table_html = ""
        if shareholders:
            rows = ""
            for i, sh in enumerate(shareholders):
                bg = c.table_alt_row_bg if i % 2 == 1 else c.bg_white
                name = html_escape(sh.name)
                category = html_escape(sh.category) if sh.category else ""
                pct = f"{sh.stake_pct * 100:.1f}%" if sh.stake_pct else "N/A"
                shares = (
                    f"{sh.share_count:,}" if sh.share_count is not None else ""
                )

                # 지분율 바
                bar_width = min(sh.stake_pct * 100, 100) if sh.stake_pct else 0
                bar_html = (
                    f'<div style="width:80px;height:12px;'
                    f'background:{c.bg_cool_grey};border-radius:6px;'
                    f'display:inline-block;vertical-align:middle;'
                    f'margin-left:6px;">'
                    f'<div style="width:{bar_width}%;height:100%;'
                    f'background:{c.accent};border-radius:6px;"></div></div>'
                )

                rows += (
                    f'<tr style="background:{bg};">'
                    f'<td style="padding:5px 10px;font-size:9pt;'
                    f'font-weight:bold;color:{c.primary};">{name}</td>'
                    f'<td style="padding:5px 10px;font-size:9pt;'
                    f'color:{c.text_secondary};">{category}</td>'
                    f'<td style="padding:5px 10px;font-size:9pt;'
                    f"text-align:right;font-family:'IBM Plex Mono',monospace;"
                    f'color:{c.text_body};">{pct}{bar_html}</td>'
                    f'<td style="padding:5px 10px;font-size:9pt;'
                    f"text-align:right;font-family:'IBM Plex Mono',monospace;"
                    f'color:{c.text_body};">{shares}</td></tr>'
                )

            table_html = (
                f'<table style="border-collapse:collapse;width:100%;">'
                f"<thead><tr>"
                f'<th style="padding:6px 10px;background:{c.table_header_bg};'
                f'color:{c.text_white};font-size:9pt;text-align:left;">주주명</th>'
                f'<th style="padding:6px 10px;background:{c.table_header_bg};'
                f'color:{c.text_white};font-size:9pt;text-align:left;">구분</th>'
                f'<th style="padding:6px 10px;background:{c.table_header_bg};'
                f'color:{c.text_white};font-size:9pt;text-align:right;">지분율</th>'
                f'<th style="padding:6px 10px;background:{c.table_header_bg};'
                f'color:{c.text_white};font-size:9pt;text-align:right;">주식 수</th>'
                f"</tr></thead>"
                f"<tbody>{rows}</tbody></table>"
            )

        content = f"""
        {narrative_html}
        {table_html}
        """

        slide = build_slide_html(
            content,
            title="주주 구성",
            slide_class="slide-shareholder-structure",
            tokens=tokens,
        )
        return [slide]

    def render_pptx(
        self,
        factory: Any,
        data: IMDocumentData,
        *,
        prs: Any,
        tokens: IMDesignTokens | None = None,
    ) -> list[Any]:
        tokens = tokens or DEFAULT_TOKENS
        lay = tokens.layout

        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        from src.design_renderer.pptx_engine.shape_builder import (
            add_summary_textbox,
        )

        slide = factory.add_content_slide(title="주주 구성")
        shareholders = data.shareholders
        narrative = data.narratives.get("shareholder_structure", "")

        y = lay.content_top
        t = tokens.typography
        c = tokens.colors
        f = tokens.font_sizes

        if narrative:
            add_summary_textbox(slide, narrative, top=y, tokens=tokens)
            y += 0.7

        if shareholders:
            # 테이블 직접 생성
            cols = 3
            rows_count = len(shareholders) + 1  # 헤더 + 데이터
            tbl_width = lay.content_width
            tbl = slide.shapes.add_table(
                rows_count,
                cols,
                Inches(lay.content_left),
                Inches(y),
                Inches(tbl_width),
                Inches(rows_count * 0.3),
            ).table

            # 열 너비
            tbl.columns[0].width = Inches(tbl_width * 0.4)
            tbl.columns[1].width = Inches(tbl_width * 0.3)
            tbl.columns[2].width = Inches(tbl_width * 0.3)

            # 헤더
            headers = ["주주명", "구분", "지분율"]
            for col_idx, header in enumerate(headers):
                cell = tbl.cell(0, col_idx)
                cell.text = header
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.alignment = (
                        PP_ALIGN.LEFT if col_idx < 2 else PP_ALIGN.RIGHT
                    )
                    for run in paragraph.runs:
                        run.font.name = t.font_body
                        run.font.size = Pt(f.footnote)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(
                            c.text_white.lstrip("#")
                        )
                # 헤더 배경색
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(
                    c.table_header_bg.lstrip("#")
                )

            # 데이터
            for r_idx, sh in enumerate(shareholders, 1):
                tbl.cell(r_idx, 0).text = sh.name
                tbl.cell(r_idx, 1).text = sh.category or ""
                pct = (
                    f"{sh.stake_pct * 100:.1f}%"
                    if sh.stake_pct
                    else "N/A"
                )
                tbl.cell(r_idx, 2).text = pct

                for col_idx in range(cols):
                    cell = tbl.cell(r_idx, col_idx)
                    cell.text_frame.paragraphs[0].alignment = (
                        PP_ALIGN.LEFT if col_idx < 2 else PP_ALIGN.RIGHT
                    )
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = (
                                t.font_mono if col_idx == 2 else t.font_body
                            )
                            run.font.size = Pt(f.body)
                            run.font.color.rgb = RGBColor.from_string(
                                c.text_body.lstrip("#")
                            )
                    # 줄무늬 배경
                    if r_idx % 2 == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(
                            c.table_alt_row_bg.lstrip("#")
                        )

        return [slide]
