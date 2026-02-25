"""경영진(Management Team) 섹션 렌더러.

경영진 프로필 카드를 그리드 형식으로 표시.
"""

from __future__ import annotations

import logging
from html import escape as html_escape
from typing import Any

from src.design_renderer.design_tokens import DEFAULT_TOKENS, IMDesignTokens
from src.design_renderer.im_document import IMDocumentData
from src.design_renderer.pdf_output.html_builder import build_slide_html
from src.design_renderer.pptx_engine.font_helper import set_font_with_ea
from src.design_renderer.section_renderers import register_renderer
from src.design_renderer.section_renderers.base import BaseSectionRenderer

logger = logging.getLogger(__name__)


@register_renderer
class ManagementTeamRenderer(BaseSectionRenderer):
    """경영진 슬라이드 렌더러."""

    section_id = "management_team"

    def render_html(
        self,
        data: IMDocumentData,
        *,
        tokens: IMDesignTokens | None = None,
    ) -> list[str]:
        tokens = tokens or DEFAULT_TOKENS
        c = tokens.colors

        members = data.management_team
        narrative = html_escape(data.narratives.get("management_team", ""))

        narrative_html = ""
        if narrative:
            narrative_html = (
                f'<p style="font-size:10pt;color:{c.text_body};'
                f'line-height:1.6;margin-bottom:1em;">{narrative}</p>'
            )

        cards_html = ""
        for member in members:
            name = html_escape(member.name)
            title = html_escape(member.title)
            role = html_escape(member.role) if member.role else ""

            career_html = ""
            if member.career:
                items = "".join(
                    f'<li style="font-size:8pt;color:{c.text_secondary};'
                    f'margin-bottom:2px;">{html_escape(c_item)}</li>'
                    for c_item in member.career
                )
                career_html = (
                    f'<ul style="margin:0.3em 0 0 0;'
                    f'padding-left:1em;">{items}</ul>'
                )

            role_badge = ""
            if role:
                role_badge = (
                    f'<span style="display:inline-block;padding:2px 8px;'
                    f'background:{c.accent};color:{c.text_white};'
                    f'border-radius:10px;font-size:8pt;font-weight:bold;'
                    f'margin-left:6px;">{role}</span>'
                )

            cards_html += (
                f'<div style="padding:0.8em;background:{c.bg_cool_grey};'
                f'border-radius:4px;border-top:3px solid {c.primary};">'
                f'<div style="font-size:12pt;font-weight:bold;'
                f'color:{c.primary};">{name}{role_badge}</div>'
                f'<div style="font-size:9pt;color:{c.text_secondary};'
                f'margin-top:2px;">{title}</div>'
                f"{career_html}</div>"
            )

        grid_html = ""
        if cards_html:
            cols = min(len(members), 3)
            grid_html = (
                f'<div style="display:grid;grid-template-columns:'
                f"repeat({cols}, 1fr);gap:0.8em;"
                f'margin-top:0.5em;">{cards_html}</div>'
            )

        content = f"""
        {narrative_html}
        {grid_html}
        """

        slide = build_slide_html(
            content,
            title="경영진",
            slide_class="slide-management-team",
            tokens=tokens,
        )
        return [slide]

    def render_pptx(
        self,
        factory: Any,
        data: IMDocumentData,
        *,
        prs: Any,
        tokens: IMDesignTokens | None = None,
    ) -> list[Any]:
        tokens = tokens or DEFAULT_TOKENS
        lay = tokens.layout

        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        slide = factory.add_content_slide(title="경영진")
        members = data.management_team
        t = tokens.typography
        c = tokens.colors
        f = tokens.font_sizes

        if not members:
            return [slide]

        # 프로필 카드 배치 (최대 3열)
        cols = min(len(members), 3)
        card_width = (lay.content_width - 0.3 * (cols - 1)) / cols
        y = lay.content_top

        for idx, member in enumerate(members):
            col = idx % cols
            row = idx // cols
            x = lay.content_left + col * (card_width + 0.3)
            card_y = y + row * 1.8

            # 이름 + 직위
            shape = slide.shapes.add_textbox(
                Inches(x),
                Inches(card_y),
                Inches(card_width),
                Inches(1.5),
            )
            tf = shape.text_frame
            tf.word_wrap = True

            # 이름
            p_name = tf.paragraphs[0]
            r = p_name.add_run()
            r.text = member.name
            if member.role:
                r.text += f" ({member.role})"
            set_font_with_ea(r, t.font_body)
            r.font.size = Pt(f.summary_text)
            r.font.bold = True
            r.font.color.rgb = RGBColor.from_string(c.primary.lstrip("#"))

            # 직위
            p_title = tf.add_paragraph()
            r = p_title.add_run()
            r.text = member.title
            set_font_with_ea(r, t.font_body)
            r.font.size = Pt(f.footnote)
            r.font.color.rgb = RGBColor.from_string(
                c.text_secondary.lstrip("#")
            )

            # 경력
            for career_item in member.career:
                p_career = tf.add_paragraph()
                r = p_career.add_run()
                r.text = f"• {career_item}"
                set_font_with_ea(r, t.font_body)
                r.font.size = Pt(f.small_label)
                r.font.color.rgb = RGBColor.from_string(
                    c.text_body.lstrip("#")
                )

        return [slide]
