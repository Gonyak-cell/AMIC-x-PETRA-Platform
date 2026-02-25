"""폴백 슬라이드 렌더러 — 섹션 실패 시 플레이스홀더 슬라이드 생성.

> 마지막 수정: 2026-02-11 10:30:00

개별 섹션 렌더링 실패 시, 빈 슬라이드 대신 "데이터 없음" 안내 슬라이드를
삽입하여 전체 PPTX의 구조적 일관성을 유지한다.
"""

from __future__ import annotations

import logging
from typing import Any

from src.design_renderer.design_tokens import DEFAULT_TOKENS, IMDesignTokens
from src.design_renderer.pptx_engine.font_helper import set_font_with_ea

logger = logging.getLogger(__name__)


def render_fallback_slide_pptx(
    factory: Any,
    section_id: str,
    error: str,
    *,
    prs: Any,
    tokens: IMDesignTokens | None = None,
    show_error_detail: bool = False,
) -> Any:
    """섹션 실패 시 폴백 슬라이드를 생성한다.

    Args:
        factory: SlideFactory 인스턴스.
        section_id: 실패한 섹션 ID.
        error: 에러 메시지.
        prs: Presentation 인스턴스.
        tokens: 디자인 토큰.
        show_error_detail: True이면 에러 상세 표시 (개발용). 프로덕션에서는 False.

    Returns:
        생성된 Slide.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    slide = factory.add_content_slide(title=section_id.replace("_", " ").title())
    c = tokens.colors
    t = tokens.typography

    # 안내 메시지 텍스트박스
    lay = tokens.layout
    txbox = slide.shapes.add_textbox(
        Inches(lay.content_left),
        Inches(lay.content_top + 1.0),
        Inches(lay.content_width),
        Inches(2.0),
    )
    tf = txbox.text_frame
    tf.word_wrap = True

    # 제목: "데이터 없음"
    p_title = tf.paragraphs[0]
    p_title.alignment = PP_ALIGN.CENTER
    run_title = p_title.add_run()
    run_title.text = "데이터를 불러올 수 없습니다"
    set_font_with_ea(run_title, t.font_body)
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor.from_string(
        c.text_secondary.lstrip("#")
    )

    # 상세: 프로덕션에서는 에러 메시지 숨김
    p_detail = tf.add_paragraph()
    p_detail.alignment = PP_ALIGN.CENTER
    run_detail = p_detail.add_run()
    if show_error_detail:
        error_msg = error if len(error) <= 200 else error[:197] + "..."
        run_detail.text = f"[{section_id}] {error_msg}"
    else:
        run_detail.text = "해당 섹션의 데이터를 처리 중입니다."
    set_font_with_ea(run_detail, t.font_body)
    run_detail.font.size = Pt(9)
    run_detail.font.color.rgb = RGBColor.from_string(
        c.gray_medium.lstrip("#")
    )

    logger.info(f"폴백 슬라이드 생성: {section_id}")
    return slide


def render_fallback_slide_html(
    section_id: str,
    error: str,
    *,
    tokens: IMDesignTokens | None = None,
    show_error_detail: bool = False,
) -> str:
    """섹션 실패 시 폴백 HTML 슬라이드를 생성한다.

    Args:
        section_id: 실패한 섹션 ID.
        error: 에러 메시지.
        tokens: 디자인 토큰.
        show_error_detail: True이면 에러 상세 표시 (개발용).

    Returns:
        HTML 문자열.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    from html import escape as html_escape

    c = tokens.colors
    title = section_id.replace("_", " ").title()

    if show_error_detail:
        error_safe = html_escape(error if len(error) <= 200 else error[:197] + "...")
        detail_html = f"[{html_escape(section_id)}] {error_safe}"
    else:
        detail_html = "해당 섹션의 데이터를 처리 중입니다."

    return f"""<div class="slide fallback-slide">
    <h2 class="slide-title">{html_escape(title)}</h2>
    <div style="text-align:center; padding:2em;">
        <p style="color:{c.text_secondary}; font-size:16pt; font-weight:bold;">
            데이터를 불러올 수 없습니다
        </p>
        <p style="color:{c.gray_medium}; font-size:9pt;">
            {detail_html}
        </p>
    </div>
</div>"""
