"""재무제표 테이블 컴포넌트 — TITAN/COVENANT 패턴 듀얼 렌더링.

PPTX: col1=2.6"(라벨) + 7×0.91"(기간), row=0.26"
HTML: css_generator.py의 .financial-table 클래스 사용

계정계층, 다년도, CAGR 열을 지원하며,
NumberFormatConfig로 숫자 포맷팅을 일괄 적용한다.
"""

from __future__ import annotations

import logging
from html import escape as html_escape
from typing import Any

from src.design_renderer.components.number_formatter import (
    format_currency,
    format_growth_indicator,
    format_percentage,
)
from src.design_renderer.design_tokens import DEFAULT_TOKENS, IMDesignTokens
from src.design_renderer.im_document import NumberFormatConfig

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# 데이터 구조
# ---------------------------------------------------------------------------

# 행 스타일 타입
ROW_NORMAL = "normal"
ROW_HEADER = "header"
ROW_TOTAL = "total"  # 합계 행 (Bold, 상하 border)
ROW_SUBTOTAL = "subtotal"  # 소계 행 (SemiBold)
ROW_INDENT_1 = "indent1"  # 1단계 들여쓰기
ROW_INDENT_2 = "indent2"  # 2단계 들여쓰기


# ---------------------------------------------------------------------------
# PPTX 렌더링
# ---------------------------------------------------------------------------


def render_financial_table_pptx(
    slide: Any,
    *,
    headers: list[str],
    rows: list[dict[str, Any]],
    left: float | None = None,
    top: float | None = None,
    tokens: IMDesignTokens | None = None,
    number_config: NumberFormatConfig | None = None,
    show_cagr: bool = False,
) -> Any:
    """PPTX 슬라이드에 TITAN/COVENANT 재무 테이블 추가.

    Args:
        slide: pptx.slide.Slide 인스턴스.
        headers: 열 헤더 리스트 (예: ["", "2022", "2023", "2024E", "CAGR"]).
        rows: 행 데이터 리스트. 각 행은 dict:
            {"label": "매출액", "values": [100, 120, 150], "style": "normal",
             "cagr": 0.224}
        left, top: 테이블 위치 (inches). None이면 layout 기본값.
        tokens: 디자인 토큰.
        number_config: 숫자 표기 설정.
        show_cagr: True이면 CAGR 열 표시.

    Returns:
        생성된 Table shape 객체.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    from pptx.util import Inches

    lay = tokens.layout
    c = tokens.colors
    t = tokens.typography
    f = tokens.font_sizes

    tbl_left = Inches(left if left is not None else lay.content_left)
    tbl_top = Inches(top if top is not None else lay.content_top + 0.5)

    n_cols = len(headers)
    n_rows = len(rows) + 1  # +1 for header row

    # 열 너비 계산
    col_widths = _calculate_col_widths(n_cols, lay, show_cagr)
    tbl_width = sum(col_widths)
    tbl_height = Inches(lay.table_row_height * n_rows)

    # 테이블 생성
    table_shape = slide.shapes.add_table(
        n_rows, n_cols, tbl_left, tbl_top, tbl_width, tbl_height
    )
    table = table_shape.table

    # 열 너비 설정
    for col_idx, width in enumerate(col_widths):
        table.columns[col_idx].width = width

    # 헤더 행
    for col_idx, header_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = header_text
        _style_cell_pptx(
            cell,
            bg_hex=c.table_header_bg,
            text_hex=c.text_white,
            font_name=t.font_body,
            font_size=f.footnote,
            bold=True,
            align="center" if col_idx > 0 else "left",
        )

    # 데이터 행
    for row_idx, row_data in enumerate(rows):
        actual_row = row_idx + 1
        label = row_data.get("label", "")
        values = row_data.get("values", [])
        style = row_data.get("style", ROW_NORMAL)
        cagr = row_data.get("cagr")

        # 라벨 셀
        label_cell = table.cell(actual_row, 0)
        indent_prefix = ""
        if style == ROW_INDENT_1:
            indent_prefix = "  "
        elif style == ROW_INDENT_2:
            indent_prefix = "    "

        label_cell.text = indent_prefix + label
        is_total = style in (ROW_TOTAL, ROW_SUBTOTAL)
        _style_cell_pptx(
            label_cell,
            text_hex=c.primary if is_total else c.text_dark,
            font_name=t.font_body,
            font_size=f.body,
            bold=is_total,
            align="left",
        )

        # 짝수 행 배경
        if actual_row % 2 == 0 and not is_total:
            _set_cell_bg(label_cell, c.table_alt_row_bg)

        # 데이터 셀
        for col_idx, value in enumerate(values):
            cell_col = col_idx + 1
            if cell_col >= n_cols:
                break
            cell = table.cell(actual_row, cell_col)
            cell.text = _format_cell_value(value, number_config)
            _style_cell_pptx(
                cell,
                text_hex=c.primary if is_total else c.text_body,
                font_name=t.font_mono,
                font_size=f.body,
                bold=is_total,
                align="right",
            )
            if actual_row % 2 == 0 and not is_total:
                _set_cell_bg(cell, c.table_alt_row_bg)

        # CAGR 열
        if show_cagr and cagr is not None:
            cagr_col = n_cols - 1
            cagr_cell = table.cell(actual_row, cagr_col)
            cagr_text, cagr_color = format_growth_indicator(cagr, number_config)
            cagr_cell.text = format_percentage(cagr, number_config, show_sign=True)
            _style_cell_pptx(
                cagr_cell,
                text_hex=cagr_color,
                font_name=t.font_mono,
                font_size=f.body,
                bold=False,
                align="right",
            )

    return table_shape


def _calculate_col_widths(
    n_cols: int,
    layout: Any,
    show_cagr: bool,
) -> list[int]:
    """열 너비 계산 (EMU 단위)."""
    from pptx.util import Inches

    widths: list[int] = []
    # 첫 번째 열: 라벨
    widths.append(Inches(layout.table_label_col_width))
    # 나머지 열: 데이터
    data_cols = n_cols - 1
    for _ in range(data_cols):
        widths.append(Inches(layout.table_data_col_width))
    return widths


def _format_cell_value(
    value: Any,
    config: NumberFormatConfig | None,
) -> str:
    """셀 값을 포맷팅."""
    if value is None:
        return "N/A" if config is None else config.na_display
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float)):
        return format_currency(value, config, show_unit=False)
    return str(value)


def _style_cell_pptx(
    cell: Any,
    *,
    bg_hex: str | None = None,
    text_hex: str = "#3D3D3D",
    font_name: str = "Pretendard",
    font_size: int = 10,
    bold: bool = False,
    align: str = "left",
) -> None:
    """PPTX 테이블 셀 스타일 적용."""
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    if bg_hex:
        _set_cell_bg(cell, bg_hex)

    # 여백 최소화
    cell.margin_left = Pt(4)
    cell.margin_right = Pt(4)
    cell.margin_top = Pt(2)
    cell.margin_bottom = Pt(2)

    alignment_map = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }

    for para in cell.text_frame.paragraphs:
        para.alignment = alignment_map.get(align, PP_ALIGN.LEFT)
        for run in para.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(text_hex.lstrip("#"))


def _set_cell_bg(cell: Any, hex_color: str) -> None:
    """셀 배경색 설정."""
    from pptx.oxml.ns import qn
    from lxml import etree

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    # 기존 solidFill 제거
    for old in tc_pr.findall(qn("a:solidFill")):
        tc_pr.remove(old)
    solid_fill = etree.SubElement(tc_pr, qn("a:solidFill"))
    srgb = etree.SubElement(solid_fill, qn("a:srgbClr"))
    srgb.set("val", hex_color.lstrip("#"))


# ---------------------------------------------------------------------------
# HTML 렌더링
# ---------------------------------------------------------------------------


def render_financial_table_html(
    *,
    headers: list[str],
    rows: list[dict[str, Any]],
    tokens: IMDesignTokens | None = None,
    number_config: NumberFormatConfig | None = None,
    show_cagr: bool = False,
    table_class: str = "financial-table",
) -> str:
    """HTML 재무 테이블 생성.

    css_generator.py의 .financial-table 클래스를 사용한다.

    Args:
        headers: 열 헤더 리스트.
        rows: 행 데이터 리스트 (PPTX와 동일 형식).
        tokens: 디자인 토큰.
        number_config: 숫자 표기 설정.
        show_cagr: CAGR 열 표시 여부.
        table_class: CSS 클래스명.

    Returns:
        <table> HTML 문자열.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    c = tokens.colors

    # 헤더
    thead = "<thead><tr>"
    for h in headers:
        thead += f"<th>{html_escape(h)}</th>"
    thead += "</tr></thead>"

    # 바디
    tbody = "<tbody>"
    for row_data in rows:
        label = row_data.get("label", "")
        values = row_data.get("values", [])
        style = row_data.get("style", ROW_NORMAL)
        cagr = row_data.get("cagr")

        row_class = ""
        if style == ROW_TOTAL:
            row_class = ' class="total-row"'
        elif style == ROW_SUBTOTAL:
            row_class = ' class="subtotal-row"'

        tbody += f"<tr{row_class}>"

        # 라벨 셀
        indent_style = ""
        if style == ROW_INDENT_1:
            indent_style = ' style="padding-left: 24px;"'
        elif style == ROW_INDENT_2:
            indent_style = ' style="padding-left: 40px;"'

        tbody += f"<td{indent_style}>{html_escape(label)}</td>"

        # 데이터 셀
        n_data = len(headers) - 1 - (1 if show_cagr else 0)
        for i, value in enumerate(values[:n_data]):
            formatted = html_escape(_format_cell_value(value, number_config))
            tbody += f"<td>{formatted}</td>"

        # CAGR 열
        if show_cagr:
            if cagr is not None:
                cagr_pct = format_percentage(cagr, number_config, show_sign=True)
                color = (
                    c.positive
                    if cagr > 0
                    else c.negative
                    if cagr < 0
                    else c.text_secondary
                )
                tbody += f'<td style="color: {color};">{html_escape(cagr_pct)}</td>'
            else:
                na = number_config.na_display if number_config else "N/A"
                tbody += f"<td>{html_escape(na)}</td>"

        tbody += "</tr>"

    tbody += "</tbody>"

    return f'<table class="{table_class}">{thead}{tbody}</table>'
