"""헤더/푸터/구분선 공통 요소 — PPTX idx=11/12/13 플레이스홀더 패턴.

source_citation.py와 연동하여 각주를 자동 배치한다.
PPTX와 HTML 듀얼 렌더링을 지원한다.
"""

from __future__ import annotations

import logging
from html import escape as html_escape
from typing import Any

from src.design_renderer.components.source_citation import (
    render_footnote_html,
    render_footnote_pptx,
)
from src.design_renderer.design_tokens import DEFAULT_TOKENS, IMDesignTokens
from src.design_renderer.im_document import SourceCitation

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# PPTX 렌더링
# ---------------------------------------------------------------------------


def render_slide_title_pptx(
    slide: Any,
    title: str,
    *,
    tokens: IMDesignTokens | None = None,
) -> bool:
    """PPTX 슬라이드 제목(idx=11) 설정.

    Args:
        slide: pptx.slide.Slide 인스턴스.
        title: 슬라이드 제목 텍스트.
        tokens: 디자인 토큰.

    Returns:
        성공 여부.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    ph = _find_placeholder(slide, tokens.layout.ph_title_idx)
    if ph is None:
        return False

    ph.text = ""
    tf = ph.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = tokens.typography.font_heading
    run.font.size = Pt(tokens.font_sizes.slide_title)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(
        tokens.colors.primary.lstrip("#")
    )

    return True


def render_page_number_pptx(
    slide: Any,
    page_number: int,
    total_pages: int,
    *,
    tokens: IMDesignTokens | None = None,
) -> bool:
    """PPTX 슬라이드 페이지 번호(idx=13) 설정.

    Args:
        slide: Slide 인스턴스.
        page_number: 현재 페이지 번호.
        total_pages: 총 페이지 수.
        tokens: 디자인 토큰.

    Returns:
        성공 여부.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    ph = _find_placeholder(slide, tokens.layout.ph_page_number_idx)
    if ph is None:
        return False

    ph.text = ""
    tf = ph.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"{page_number} / {total_pages}"
    run.font.name = tokens.typography.font_mono
    run.font.size = Pt(tokens.font_sizes.page_number)
    run.font.color.rgb = RGBColor.from_string(
        tokens.colors.text_secondary.lstrip("#")
    )

    return True


def render_footer_pptx(
    slide: Any,
    *,
    page_number: int | None = None,
    total_pages: int | None = None,
    citations: list[tuple[int, SourceCitation]] | None = None,
    tokens: IMDesignTokens | None = None,
) -> None:
    """PPTX 슬라이드 푸터 (각주 + 페이지번호) 일괄 렌더링.

    Args:
        slide: Slide 인스턴스.
        page_number: 현재 페이지 번호.
        total_pages: 총 페이지 수.
        citations: 각주 데이터 (source_citation 형식).
        tokens: 디자인 토큰.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    # 각주
    if citations:
        render_footnote_pptx(slide, citations, tokens=tokens)

    # 페이지 번호
    if page_number is not None and total_pages is not None:
        render_page_number_pptx(
            slide, page_number, total_pages, tokens=tokens
        )


def render_horizontal_line_pptx(
    slide: Any,
    *,
    left: float = 0.5,
    top: float = 1.0,
    width: float = 9.83,
    tokens: IMDesignTokens | None = None,
) -> Any:
    """PPTX 슬라이드에 수평 구분선 추가.

    Args:
        slide: Slide 인스턴스.
        left, top, width: 위치/크기 (inches).
        tokens: 디자인 토큰.

    Returns:
        생성된 Shape 객체.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Pt(1.5),  # 1.5pt 두께
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(
        tokens.colors.gray_border.lstrip("#")
    )
    shape.line.fill.background()

    return shape


# ---------------------------------------------------------------------------
# HTML 렌더링
# ---------------------------------------------------------------------------


def render_slide_title_html(
    title: str,
    *,
    tokens: IMDesignTokens | None = None,
) -> str:
    """HTML 슬라이드 제목 생성.

    Args:
        title: 슬라이드 제목 텍스트.
        tokens: 디자인 토큰.

    Returns:
        <div class="slide-title">...</div> HTML.
    """
    return f'<div class="slide-title">{html_escape(title)}</div>'


def render_footer_html(
    *,
    page_number: int | None = None,
    total_pages: int | None = None,
    citations: list[tuple[int, SourceCitation]] | None = None,
    tokens: IMDesignTokens | None = None,
) -> str:
    """HTML 슬라이드 푸터 (각주 + 페이지번호) 생성.

    Args:
        page_number: 현재 페이지 번호.
        total_pages: 총 페이지 수.
        citations: 각주 데이터.
        tokens: 디자인 토큰.

    Returns:
        <div class="slide-footer">...</div> HTML.
    """
    if tokens is None:
        tokens = DEFAULT_TOKENS

    footnote_html = ""
    if citations:
        footnote_html = render_footnote_html(citations, tokens=tokens)

    page_html = ""
    if page_number is not None and total_pages is not None:
        page_html = f"{page_number} / {total_pages}"

    return f"""<div class="slide-footer">
    <div class="slide-footnote">{footnote_html}</div>
    <div class="slide-page-number">{page_html}</div>
</div>"""


def render_horizontal_line_html() -> str:
    """HTML 수평 구분선."""
    return '<hr style="border: none; border-top: 1.5px solid #E0E0E0; margin: 8px 0;">'


# ---------------------------------------------------------------------------
# 헬퍼
# ---------------------------------------------------------------------------


def _find_placeholder(slide: Any, idx: int) -> Any | None:
    """슬라이드에서 특정 idx의 플레이스홀더 반환."""
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            return shape
    return None
