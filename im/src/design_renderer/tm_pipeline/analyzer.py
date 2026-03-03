"""PPTX 구조 분석기 — 기존 PPTX 파일의 레이아웃/디자인 분석.

샘플 PPTX 파일을 분석하여 레이아웃, 색상, 폰트, 슬라이드 구조를
추출한다. 디자인 토큰 검증 및 역공학에 활용.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)


@dataclass
class SlideAnalysis:
    """개별 슬라이드 분석 결과."""

    index: int
    layout_name: str = ""
    shape_count: int = 0
    text_runs: int = 0
    fonts_used: list[str] = field(default_factory=list)
    colors_used: list[str] = field(default_factory=list)
    has_chart: bool = False
    has_table: bool = False
    has_image: bool = False


@dataclass
class PptxAnalysis:
    """PPTX 전체 분석 결과."""

    file_path: str = ""
    slide_count: int = 0
    slide_width_inches: float = 0.0
    slide_height_inches: float = 0.0
    slides: list[SlideAnalysis] = field(default_factory=list)
    all_fonts: list[str] = field(default_factory=list)
    all_colors: list[str] = field(default_factory=list)
    layout_names: list[str] = field(default_factory=list)


def analyze_pptx(pptx_path: str | Path) -> PptxAnalysis:
    """PPTX 파일 구조 분석.

    Args:
        pptx_path: 분석할 PPTX 파일 경로.

    Returns:
        PptxAnalysis 분석 결과.

    Raises:
        FileNotFoundError: 파일이 존재하지 않을 때.
    """
    path = Path(pptx_path)
    if not path.exists():
        raise FileNotFoundError("PPTX 파일을 찾을 수 없습니다")

    try:
        from pptx import Presentation
    except ImportError:
        logger.warning("python-pptx 미설치 — 분석 건너뜀")
        return PptxAnalysis(file_path=str(path))

    prs = Presentation(str(path))
    result = PptxAnalysis(
        file_path=str(path),
        slide_count=len(prs.slides),
        slide_width_inches=prs.slide_width / 914400,  # EMU → inches
        slide_height_inches=prs.slide_height / 914400,
    )

    all_fonts: set[str] = set()
    all_colors: set[str] = set()
    layout_names: set[str] = set()

    for idx, slide in enumerate(prs.slides):
        slide_analysis = SlideAnalysis(index=idx)

        if slide.slide_layout:
            slide_analysis.layout_name = slide.slide_layout.name
            layout_names.add(slide.slide_layout.name)

        slide_analysis.shape_count = len(slide.shapes)
        slide_fonts: set[str] = set()

        slide_colors: set[str] = set()
        for shape in slide.shapes:
            if shape.has_chart:
                slide_analysis.has_chart = True
            if shape.has_table:
                slide_analysis.has_table = True
            if hasattr(shape, "image"):
                slide_analysis.has_image = True

            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        slide_analysis.text_runs += 1
                        if run.font and run.font.name:
                            slide_fonts.add(run.font.name)
                        if run.font and run.font.color and run.font.color.rgb:
                            slide_colors.add(str(run.font.color.rgb))

        slide_analysis.fonts_used = sorted(slide_fonts)
        slide_analysis.colors_used = sorted(slide_colors)
        all_fonts.update(slide_fonts)
        all_colors.update(slide_colors)
        result.slides.append(slide_analysis)

    result.all_fonts = sorted(all_fonts)
    result.all_colors = sorted(all_colors)
    result.layout_names = sorted(layout_names)

    return result
