"""VDR Analysis Service — VDR 문서 파싱 및 체크리스트 아이템 값 추출.

> 마지막 수정: 2026-02-25 21:00:00

VDR(Virtual Data Room) 문서(Excel, PDF)를 파싱하여
체크리스트 아이템에 매핑되는 데이터를 자동 추출한다.

한국어/영어 듀얼 키워드 매칭으로 재무 데이터, 회사 정보,
딜 구조 등을 식별한다.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# 한국어/영어 듀얼 키워드 매핑
# ---------------------------------------------------------------------------

KEYWORD_MAP: dict[str, list[str]] = {
    "revenue": ["매출", "매출액", "Revenue", "Sales", "매출 합계", "Total Revenue"],
    "operating_income": ["영업이익", "Operating Income", "Operating Profit"],
    "ebitda": ["EBITDA", "상각전영업이익"],
    "net_income": ["순이익", "당기순이익", "Net Income", "Net Profit"],
    "total_assets": ["총자산", "자산총계", "Total Assets"],
    "total_equity": ["자본총계", "Total Equity"],
    "total_debt": ["총차입금", "Total Debt", "차입금"],
    "cash": ["현금", "현금성자산", "Cash", "Cash and Equivalents"],
    "cogs": ["매출원가", "Cost of Goods Sold", "COGS"],
    "gross_profit": ["매출총이익", "Gross Profit"],
    "employee_count": ["임직원", "종업원", "직원수", "Employees"],
    "ceo_name": ["대표이사", "대표자", "CEO"],
    "company_name": ["상호", "회사명", "Company Name"],
    "founded_date": ["설립일", "설립연도", "Founded"],
}

# 연도 탐색 패턴
YEAR_PATTERNS: list[str] = [
    "2022", "2023", "2024", "2025", "2026",
    "FY2022", "FY2023", "FY2024", "FY2025", "FY2026",
    "'22", "'23", "'24", "'25", "'26",
]

# 연도 정규화 매핑 (약칭 → 4자리 연도)
_YEAR_NORMALIZE: dict[str, str] = {
    "'22": "2022", "'23": "2023", "'24": "2024", "'25": "2025", "'26": "2026",
    "FY2022": "2022", "FY2023": "2023", "FY2024": "2024",
    "FY2025": "2025", "FY2026": "2026",
}

# 재무 키워드 (연도별 매핑 대상)
_FINANCIAL_KEYS: set[str] = {
    "revenue", "operating_income", "ebitda", "net_income",
    "total_assets", "total_equity", "total_debt", "cash",
    "cogs", "gross_profit",
}

# 숫자 추출 정규식
_NUMBER_RE = re.compile(r"[-+]?[\d,]+(?:\.\d+)?")
_YEAR_4DIGIT_RE = re.compile(r"(?:FY)?(\d{4})|'(\d{2})")


# ---------------------------------------------------------------------------
# 추출 결과 데이터클래스
# ---------------------------------------------------------------------------


@dataclass
class ExtractionResult:
    """단일 추출 결과.

    Attributes:
        field_key: 체크리스트 필드 키 (예: "revenue", "company_name").
        value: 추출된 값 문자열.
        confidence: 신뢰도 (0.0~1.0).
        source_location: 소스 위치 (예: "Sheet1:B15" 또는 "Page 3, Table 2, Row 5").
        source_doc_name: 소스 문서 파일명.
        fiscal_year: 연도 (재무 데이터인 경우, 예: "2023").
    """

    field_key: str
    value: str
    confidence: float = 0.0
    source_location: str = ""
    source_doc_name: str = ""
    fiscal_year: str | None = None


# ---------------------------------------------------------------------------
# VDR Analysis Service
# ---------------------------------------------------------------------------


class VdrAnalysisService:
    """VDR 문서를 파싱하여 체크리스트 아이템 값을 추출하는 서비스.

    Excel(openpyxl)과 PDF(PyMuPDF) 파서를 사용하여 문서를 분석하고,
    한국어/영어 듀얼 키워드 매핑으로 필요한 데이터를 식별한다.

    Usage::

        service = VdrAnalysisService()
        results = service.extract_from_vdr_documents(
            vdr_doc_paths=["/vdr-shared/deal-123/financial_summary.xlsx"],
            checklist_items=["revenue", "operating_income", "company_name"],
        )
    """

    def __init__(
        self,
        keyword_map: dict[str, list[str]] | None = None,
    ) -> None:
        """초기화.

        Args:
            keyword_map: 커스텀 키워드 매핑. None이면 기본 KEYWORD_MAP 사용.
        """
        self._keyword_map = keyword_map or KEYWORD_MAP

    def extract_from_vdr_documents(
        self,
        vdr_doc_paths: list[str],
        checklist_items: list[str],
    ) -> list[ExtractionResult]:
        """VDR 문서 목록을 파싱하여 체크리스트 아이템에 추출값을 매핑한다.

        Args:
            vdr_doc_paths: VDR 문서 파일 경로 목록.
            checklist_items: 추출 대상 체크리스트 필드 키 목록.

        Returns:
            ExtractionResult 리스트 (중복 필드는 신뢰도 높은 것 우선).
        """
        all_results: list[ExtractionResult] = []
        target_keys = set(checklist_items)

        for doc_path in vdr_doc_paths:
            path = Path(doc_path)
            if not path.exists():
                logger.warning("VDR 문서를 찾을 수 없음: %s", doc_path)
                continue

            try:
                suffix = path.suffix.lower()
                if suffix in (".xlsx", ".xls"):
                    results = self._parse_excel(path, target_keys)
                elif suffix == ".pdf":
                    results = self._parse_pdf(path, target_keys)
                else:
                    logger.warning(
                        "지원하지 않는 파일 형식: %s (%s)", path.name, suffix,
                    )
                    continue

                all_results.extend(results)
            except Exception as exc:
                logger.error(
                    "VDR 문서 파싱 실패: %s — %s", path.name, exc,
                )

        # 필드별 최고 신뢰도 결과만 유지
        return self._deduplicate_results(all_results)

    # ------------------------------------------------------------------
    # Excel 파서
    # ------------------------------------------------------------------

    def _parse_excel(
        self,
        path: Path,
        target_keys: set[str],
    ) -> list[ExtractionResult]:
        """openpyxl로 Excel 시트를 파싱하여 재무 데이터를 추출한다.

        행/열을 탐색하여 키워드 매칭 → 해당 행의 값을 연도별로 추출한다.

        Args:
            path: Excel 파일 경로.
            target_keys: 추출 대상 필드 키 집합.

        Returns:
            ExtractionResult 리스트.
        """
        try:
            import openpyxl
        except ImportError:
            raise RuntimeError(
                "openpyxl 패키지 미설치. pip install openpyxl 필요."
            )

        results: list[ExtractionResult] = []
        doc_name = path.name

        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)

        for ws in wb.worksheets:
            sheet_name = ws.title

            # 1) 헤더 행에서 연도 열 탐색
            year_columns: dict[int, str] = {}  # col_idx → year_str
            rows = list(ws.iter_rows(min_row=1, values_only=False))

            if not rows:
                continue

            # 상위 5행에서 연도 헤더 탐색
            for row_idx, row in enumerate(rows[:5], start=1):
                for cell in row:
                    if cell.value is None:
                        continue
                    cell_str = str(cell.value).strip()
                    year = self._normalize_year(cell_str)
                    if year:
                        year_columns[cell.column] = year

            # 2) 키워드 행 탐색
            for row_idx, row in enumerate(rows, start=1):
                if not row:
                    continue

                first_cell = row[0]
                if first_cell.value is None:
                    continue

                cell_text = str(first_cell.value).strip()
                matched_key = self._match_keyword(cell_text, target_keys)

                if matched_key is None:
                    continue

                if matched_key in _FINANCIAL_KEYS and year_columns:
                    # 연도별 값 추출
                    for cell in row[1:]:
                        if cell.column in year_columns and cell.value is not None:
                            year = year_columns[cell.column]
                            value_str = self._clean_number(str(cell.value))
                            if value_str:
                                col_letter = openpyxl.utils.get_column_letter(
                                    cell.column,
                                )
                                results.append(ExtractionResult(
                                    field_key=matched_key,
                                    value=value_str,
                                    confidence=0.85,
                                    source_location=f"{sheet_name}:{col_letter}{row_idx}",
                                    source_doc_name=doc_name,
                                    fiscal_year=year,
                                ))
                else:
                    # 비재무 데이터: 인접 셀 값 추출
                    for cell in row[1:]:
                        if cell.value is not None:
                            value_str = str(cell.value).strip()
                            if value_str:
                                col_letter = openpyxl.utils.get_column_letter(
                                    cell.column,
                                )
                                results.append(ExtractionResult(
                                    field_key=matched_key,
                                    value=value_str,
                                    confidence=0.80,
                                    source_location=f"{sheet_name}:{col_letter}{row_idx}",
                                    source_doc_name=doc_name,
                                ))
                            break  # 첫 번째 유효한 값만 사용

        wb.close()
        return results

    # ------------------------------------------------------------------
    # PDF 파서
    # ------------------------------------------------------------------

    def _parse_pdf(
        self,
        path: Path,
        target_keys: set[str],
    ) -> list[ExtractionResult]:
        """PyMuPDF find_tables()로 테이블을 추출하고,
        텍스트 블록에서 회사/딜 정보를 추출한다.

        Args:
            path: PDF 파일 경로.
            target_keys: 추출 대상 필드 키 집합.

        Returns:
            ExtractionResult 리스트.
        """
        try:
            import fitz  # PyMuPDF
        except ImportError:
            raise RuntimeError(
                "PyMuPDF(fitz) 패키지 미설치. pip install PyMuPDF 필요."
            )

        results: list[ExtractionResult] = []
        doc_name = path.name

        doc = fitz.open(str(path))

        for page_num, page in enumerate(doc, start=1):
            # --- 테이블 추출 ---
            try:
                tables = page.find_tables()
                for table_idx, table in enumerate(tables, start=1):
                    table_data = table.extract()
                    table_results = self._extract_from_table(
                        table_data,
                        target_keys,
                        page_num=page_num,
                        table_idx=table_idx,
                        doc_name=doc_name,
                    )
                    results.extend(table_results)
            except Exception as exc:
                logger.warning(
                    "PDF 테이블 추출 실패: %s page %d — %s",
                    doc_name, page_num, exc,
                )

            # --- 텍스트 블록에서 비재무 데이터 추출 ---
            text = page.get_text("text")
            text_results = self._extract_from_text(
                text,
                target_keys,
                page_num=page_num,
                doc_name=doc_name,
            )
            results.extend(text_results)

        doc.close()
        return results

    # ------------------------------------------------------------------
    # 공통 추출 유틸리티
    # ------------------------------------------------------------------

    def _extract_from_table(
        self,
        table_data: list[list[str | None]],
        target_keys: set[str],
        *,
        page_num: int,
        table_idx: int,
        doc_name: str,
    ) -> list[ExtractionResult]:
        """2D 테이블 데이터에서 키워드 매칭으로 값을 추출한다.

        Args:
            table_data: 2D 문자열 리스트 (행/열).
            target_keys: 추출 대상 필드 키 집합.
            page_num: 페이지 번호.
            table_idx: 테이블 인덱스.
            doc_name: 문서 파일명.

        Returns:
            ExtractionResult 리스트.
        """
        results: list[ExtractionResult] = []

        if not table_data or len(table_data) < 2:
            return results

        # 헤더 행에서 연도 열 탐색
        header = table_data[0]
        year_columns: dict[int, str] = {}
        for col_idx, cell in enumerate(header):
            if cell is None:
                continue
            year = self._normalize_year(str(cell).strip())
            if year:
                year_columns[col_idx] = year

        # 데이터 행에서 키워드 매칭
        for row_idx, row in enumerate(table_data[1:], start=2):
            if not row or row[0] is None:
                continue

            first_cell = str(row[0]).strip()
            matched_key = self._match_keyword(first_cell, target_keys)

            if matched_key is None:
                continue

            if matched_key in _FINANCIAL_KEYS and year_columns:
                for col_idx, year in year_columns.items():
                    if col_idx < len(row) and row[col_idx] is not None:
                        value_str = self._clean_number(str(row[col_idx]))
                        if value_str:
                            results.append(ExtractionResult(
                                field_key=matched_key,
                                value=value_str,
                                confidence=0.75,
                                source_location=(
                                    f"Page {page_num}, Table {table_idx}, "
                                    f"Row {row_idx}"
                                ),
                                source_doc_name=doc_name,
                                fiscal_year=year,
                            ))
            else:
                for col_idx in range(1, len(row)):
                    if row[col_idx] is not None:
                        value_str = str(row[col_idx]).strip()
                        if value_str:
                            results.append(ExtractionResult(
                                field_key=matched_key,
                                value=value_str,
                                confidence=0.70,
                                source_location=(
                                    f"Page {page_num}, Table {table_idx}, "
                                    f"Row {row_idx}"
                                ),
                                source_doc_name=doc_name,
                            ))
                            break

        return results

    def _extract_from_text(
        self,
        text: str,
        target_keys: set[str],
        *,
        page_num: int,
        doc_name: str,
    ) -> list[ExtractionResult]:
        """페이지 텍스트에서 비재무 데이터를 추출한다.

        "대표이사: 홍길동", "설립일 : 2005년" 같은 패턴을 매칭한다.

        Args:
            text: 페이지 전체 텍스트.
            target_keys: 추출 대상 필드 키 집합.
            page_num: 페이지 번호.
            doc_name: 문서 파일명.

        Returns:
            ExtractionResult 리스트.
        """
        results: list[ExtractionResult] = []

        # 재무 데이터는 테이블에서 추출하므로 텍스트에서는 비재무만 대상
        non_financial_keys = target_keys - _FINANCIAL_KEYS

        for line in text.split("\n"):
            line = line.strip()
            if not line:
                continue

            for field_key in non_financial_keys:
                keywords = self._keyword_map.get(field_key, [])
                for keyword in keywords:
                    if keyword in line:
                        # "키워드: 값" 또는 "키워드 값" 패턴
                        value = self._extract_value_after_keyword(
                            line, keyword,
                        )
                        if value:
                            results.append(ExtractionResult(
                                field_key=field_key,
                                value=value,
                                confidence=0.65,
                                source_location=f"Page {page_num}, Text",
                                source_doc_name=doc_name,
                            ))
                            break  # 동일 line에서 같은 field_key 중복 방지

        return results

    def _match_keyword(
        self,
        text: str,
        target_keys: set[str],
    ) -> str | None:
        """텍스트가 키워드 맵의 어떤 필드에 매칭되는지 확인한다.

        Args:
            text: 셀 또는 행 텍스트.
            target_keys: 추출 대상 필드 키 집합.

        Returns:
            매칭된 field_key 또는 None.
        """
        text_lower = text.lower().strip()

        for field_key, keywords in self._keyword_map.items():
            if field_key not in target_keys:
                continue
            for keyword in keywords:
                kw_lower = keyword.lower()
                if kw_lower == text_lower or kw_lower in text_lower:
                    return field_key

        return None

    @staticmethod
    def _normalize_year(text: str) -> str | None:
        """텍스트에서 연도를 추출하고 4자리로 정규화한다.

        Args:
            text: 셀 텍스트 (예: "FY2023", "'23", "2023년").

        Returns:
            4자리 연도 문자열 또는 None.
        """
        text = text.strip()

        # 직접 매핑 체크
        if text in _YEAR_NORMALIZE:
            return _YEAR_NORMALIZE[text]

        # 패턴 리스트 직접 체크
        for pattern in YEAR_PATTERNS:
            if pattern in text:
                if pattern in _YEAR_NORMALIZE:
                    return _YEAR_NORMALIZE[pattern]
                # 4자리 숫자 패턴
                if re.match(r"^\d{4}$", pattern):
                    return pattern

        # 정규식으로 4자리 연도 추출
        match = re.search(r"(20\d{2})", text)
        if match:
            return match.group(1)

        # 약칭 패턴 ('22, '23 등)
        match = re.search(r"'(\d{2})", text)
        if match:
            return f"20{match.group(1)}"

        return None

    @staticmethod
    def _clean_number(value: str) -> str:
        """문자열에서 숫자 값을 정제한다.

        콤마 제거, 단위 제거 등.

        Args:
            value: 원본 값 문자열.

        Returns:
            정제된 숫자 문자열, 또는 빈 문자열.
        """
        value = value.strip()

        # 괄호 표기법 처리: (100) → -100
        if value.startswith("(") and value.endswith(")"):
            value = "-" + value[1:-1]

        # 숫자 추출
        match = _NUMBER_RE.search(value)
        if match:
            num_str = match.group().replace(",", "")
            try:
                float(num_str)
                return num_str
            except ValueError:
                pass

        return ""

    @staticmethod
    def _extract_value_after_keyword(line: str, keyword: str) -> str:
        """키워드 뒤의 값을 추출한다.

        "대표이사: 홍길동" → "홍길동"
        "설립일 2005년" → "2005년"

        Args:
            line: 텍스트 라인.
            keyword: 키워드 문자열.

        Returns:
            추출된 값 문자열, 또는 빈 문자열.
        """
        idx = line.find(keyword)
        if idx == -1:
            return ""

        after = line[idx + len(keyword):].strip()

        # 구분자 제거 (":", "：", "-", "–")
        for sep in (":", "：", "-", "–", "|"):
            if after.startswith(sep):
                after = after[len(sep):].strip()
                break

        # 다음 구분자 또는 줄 끝까지
        for sep in ("\t", "|", "  "):
            sep_idx = after.find(sep)
            if sep_idx > 0:
                after = after[:sep_idx].strip()
                break

        return after.strip()

    @staticmethod
    def _deduplicate_results(
        results: list[ExtractionResult],
    ) -> list[ExtractionResult]:
        """필드별로 중복을 제거하고 최고 신뢰도 결과만 유지한다.

        재무 데이터(연도별)는 (field_key, fiscal_year)를 키로 사용하고,
        비재무 데이터는 field_key만 키로 사용한다.

        Args:
            results: 전체 추출 결과 리스트.

        Returns:
            중복 제거된 결과 리스트.
        """
        best: dict[tuple[str, str | None], ExtractionResult] = {}

        for r in results:
            key = (r.field_key, r.fiscal_year)
            if key not in best or r.confidence > best[key].confidence:
                best[key] = r

        return list(best.values())
