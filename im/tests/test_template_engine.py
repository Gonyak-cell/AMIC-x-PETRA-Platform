"""Template Engine 통합 테스트.

registry, shape_mapper, content_injector, populator를 검증한다.
샘플 PPTX 템플릿이 im/templates/에 존재해야 한다.
"""

from __future__ import annotations

import shutil
import tempfile
from pathlib import Path

import pytest
from pptx import Presentation

from src.template_engine.content_injector import ContentInjector
from src.template_engine.populator import TemplatePopulator
from src.template_engine.registry import (
    get_available_variants,
    get_template_path,
    list_templates,
)
from src.template_engine.schemas import (
    ChartContent,
    SeriesData,
    SlideContent,
    TableContent,
    TemplateContent,
)
from src.template_engine.shape_mapper import (
    analyze_template,
    get_charts,
    get_tables,
    get_text_shapes,
)

# ── im/templates/ 디렉토리 경로 ─────────────────────────────────────────────

_IM_ROOT = Path(__file__).resolve().parent.parent
_TEMPLATES_DIR = _IM_ROOT / "templates"

# 샘플 템플릿 존재 여부 확인
_HAS_NX3_DM = (_TEMPLATES_DIR / "NX3 Games - DM - 260116.pptx").exists()
_HAS_SPICY_TM = (_TEMPLATES_DIR / "SPICY - TM - 260219 vSHARE.pptx").exists()
_HAS_YTN_DM = (_TEMPLATES_DIR / "YTN - Structure DM - 260116.pptx").exists()
_HAS_SWITCH_TM = (_TEMPLATES_DIR / "SWITCH - TM - 260119.pptx").exists()

needs_nx3 = pytest.mark.skipif(not _HAS_NX3_DM, reason="NX3 DM 템플릿 없음")
needs_spicy = pytest.mark.skipif(not _HAS_SPICY_TM, reason="SPICY TM 템플릿 없음")
needs_ytn = pytest.mark.skipif(not _HAS_YTN_DM, reason="YTN DM 템플릿 없음")
needs_switch = pytest.mark.skipif(not _HAS_SWITCH_TM, reason="SWITCH TM 템플릿 없음")


# ═══════════════════════════════════════════════════════════════════════════════
# Registry 테스트
# ═══════════════════════════════════════════════════════════════════════════════


class TestTemplateRegistry:
    """템플릿 레지스트리 테스트."""

    @needs_spicy
    def test_get_tm_default(self) -> None:
        """TM default 템플릿 경로를 반환한다."""
        path = get_template_path("TM", "default")
        assert path.exists()
        assert "SPICY" in path.name

    @needs_switch
    def test_get_tm_andersen(self) -> None:
        """TM andersen variant 경로를 반환한다."""
        path = get_template_path("TM", "andersen")
        assert path.exists()
        assert "SWITCH" in path.name

    @needs_nx3
    def test_get_dm_default(self) -> None:
        """DM default 템플릿 경로를 반환한다."""
        path = get_template_path("DM", "default")
        assert path.exists()
        assert "NX3" in path.name

    @needs_spicy
    def test_teaser_alias_resolves_to_tm(self) -> None:
        """TEASER는 TM의 별칭이다."""
        path = get_template_path("TEASER", "default")
        assert "SPICY" in path.name

    def test_unknown_type_raises_key_error(self) -> None:
        """등록되지 않은 유형은 KeyError."""
        with pytest.raises(KeyError):
            get_template_path("UNKNOWN", "default")

    @needs_nx3
    def test_unknown_variant_falls_back_to_default(self) -> None:
        """등록되지 않은 variant는 default로 폴백."""
        path = get_template_path("DM", "nonexistent_variant")
        assert path.exists()

    def test_list_templates_returns_all(self) -> None:
        """등록된 모든 템플릿을 반환한다."""
        templates = list_templates()
        assert len(templates) >= 4

    def test_get_available_variants_tm(self) -> None:
        """TM의 사용 가능한 variant 목록을 반환한다."""
        variants = get_available_variants("TM")
        assert "default" in variants
        assert "andersen" in variants

    def test_get_available_variants_dm(self) -> None:
        """DM의 사용 가능한 variant 목록을 반환한다."""
        variants = get_available_variants("DM")
        assert "default" in variants


# ═══════════════════════════════════════════════════════════════════════════════
# Shape Mapper 테스트
# ═══════════════════════════════════════════════════════════════════════════════


class TestShapeMapper:
    """Shape 분석 테스트."""

    @needs_nx3
    def test_analyze_nx3_dm(self) -> None:
        """NX3 DM 템플릿의 shape를 분석한다."""
        path = get_template_path("DM", "default")
        slots = analyze_template(path)
        assert len(slots) > 0
        # 7슬라이드에서 최소 20개 이상의 shape
        assert len(slots) >= 20

    @needs_nx3
    def test_get_charts_nx3_dm(self) -> None:
        """NX3 DM에서 차트 3개를 찾는다."""
        path = get_template_path("DM", "default")
        charts = get_charts(path)
        assert len(charts) == 3
        assert all(c.shape_type == "chart" for c in charts)

    @needs_nx3
    def test_get_tables_nx3_dm(self) -> None:
        """NX3 DM에서 테이블 3개를 찾는다."""
        path = get_template_path("DM", "default")
        tables = get_tables(path)
        assert len(tables) == 3

    @needs_spicy
    def test_analyze_spicy_tm(self) -> None:
        """SPICY TM 템플릿의 shape를 분석한다."""
        path = get_template_path("TM", "default")
        slots = analyze_template(path)
        assert len(slots) > 0

    @needs_spicy
    def test_get_charts_spicy_tm(self) -> None:
        """SPICY TM에서 차트 10개를 찾는다."""
        path = get_template_path("TM", "default")
        charts = get_charts(path)
        assert len(charts) == 10

    @needs_ytn
    def test_ytn_dm_no_charts(self) -> None:
        """YTN DM에는 차트가 없다."""
        path = get_template_path("DM", "deal_structure")
        charts = get_charts(path)
        assert len(charts) == 0

    @needs_ytn
    def test_ytn_dm_has_tables(self) -> None:
        """YTN DM에는 테이블 3개가 있다."""
        path = get_template_path("DM", "deal_structure")
        tables = get_tables(path)
        assert len(tables) == 3


# ═══════════════════════════════════════════════════════════════════════════════
# Content Injector 테스트
# ═══════════════════════════════════════════════════════════════════════════════


class TestContentInjector:
    """콘텐츠 인젝터 테스트."""

    @needs_nx3
    def test_inject_chart_data_preserves_native(self) -> None:
        """차트 데이터 교체 후에도 네이티브 차트가 유지된다."""
        path = get_template_path("DM", "default")
        prs = Presentation(str(path))

        # Slide 2 (idx 1)의 차트 찾기
        chart_shape = None
        for shape in prs.slides[1].shapes:
            if shape.has_chart:
                chart_shape = shape
                break

        assert chart_shape is not None

        # 새 데이터로 교체
        content = ChartContent(
            categories=["2020", "2021", "2022", "2023", "2024"],
            series=[
                SeriesData(name="거래금액", values=[100, 150, 200, 180, 220]),
                SeriesData(name="YoY", values=[0.1, 0.5, 0.33, -0.1, 0.22]),
            ],
        )
        injector = ContentInjector()
        injector.inject_chart_data(chart_shape, content)

        # 검증: 차트가 여전히 네이티브
        assert chart_shape.has_chart
        assert len(chart_shape.chart.series) == 2
        assert list(chart_shape.chart.series[0].values)[:3] == [100.0, 150.0, 200.0]

    @needs_ytn
    def test_inject_table_data_preserves_format(self) -> None:
        """테이블 데이터 교체 후 서식이 유지된다."""
        path = get_template_path("DM", "deal_structure")
        prs = Presentation(str(path))

        # Slide 5 (idx 4)의 테이블 찾기
        table_shape = None
        for shape in prs.slides[4].shapes:
            if shape.has_table:
                table_shape = shape
                break

        assert table_shape is not None

        # 새 데이터로 교체
        content = TableContent(
            rows=[
                ["항목 1", "설명 A", "비고 1"],
                ["항목 2", "설명 B", "비고 2"],
            ],
        )
        injector = ContentInjector()
        injector.inject_table_data(table_shape, content)

        # 검증: 데이터 교체 확인
        table = table_shape.table
        assert table.cell(1, 0).text == "항목 1"
        assert table.cell(2, 0).text == "항목 2"

    @needs_ytn
    def test_inject_text_preserves_format(self) -> None:
        """텍스트 교체 후 shape가 여전히 텍스트를 가진다."""
        path = get_template_path("DM", "deal_structure")
        prs = Presentation(str(path))

        # Slide 1 (idx 0)의 TextBox 3 찾기
        text_shape = None
        for shape in prs.slides[0].shapes:
            if shape.has_text_frame and "TextBox 3" in shape.name:
                text_shape = shape
                break

        assert text_shape is not None

        injector = ContentInjector()
        injector.inject_text(text_shape, "새로운 프로젝트명")

        assert text_shape.has_text_frame


# ═══════════════════════════════════════════════════════════════════════════════
# Template Populator E2E 테스트
# ═══════════════════════════════════════════════════════════════════════════════


class TestTemplatePopulator:
    """TemplatePopulator 통합 테스트."""

    @needs_nx3
    def test_populate_nx3_dm_e2e(self, tmp_path: Path) -> None:
        """NX3 DM 템플릿에 콘텐츠를 삽입하여 새 PPTX를 생성한다."""
        tpl_path = get_template_path("DM", "default")
        populator = TemplatePopulator(
            template_path=tpl_path,
            output_dir=tmp_path,
        )

        content = TemplateContent(
            project_name="Test Project Alpha",
            date="February 2026",
            memo_type="DM",
            slides=[
                SlideContent(
                    slide_idx=1,
                    charts={
                        "차트 12": ChartContent(
                            categories=["2020", "2021", "2022"],
                            series=[
                                SeriesData(name="Deal Value", values=[100, 200, 150]),
                            ],
                        ),
                    },
                ),
            ],
            excluded_slides=[],
        )

        output_path = populator.populate(content)

        # 파일 존재 확인
        assert output_path.exists()
        assert output_path.suffix == ".pptx"

        # PPTX 열기 검증
        prs = Presentation(str(output_path))
        assert len(prs.slides) == 7  # 원본 슬라이드 수 유지

        # 차트 데이터 검증
        chart_found = False
        for shape in prs.slides[1].shapes:
            if shape.has_chart:
                assert list(shape.chart.series[0].values) == [100.0, 200.0, 150.0]
                chart_found = True
                break
        assert chart_found

    @needs_ytn
    def test_populate_ytn_dm_e2e(self, tmp_path: Path) -> None:
        """YTN DM 템플릿에 콘텐츠를 삽입한다."""
        tpl_path = get_template_path("DM", "deal_structure")
        populator = TemplatePopulator(
            template_path=tpl_path,
            output_dir=tmp_path,
        )

        content = TemplateContent(
            project_name="Project Beta",
            date="March 2026",
            memo_type="DM",
            slides=[
                SlideContent(
                    slide_idx=4,
                    tables={
                        "표 8": TableContent(
                            rows=[
                                ["법률", "주주간 차등은 적법", "주의사항 A"],
                                ["세무", "과세 이슈 없음", "주의사항 B"],
                            ],
                        ),
                    },
                ),
            ],
        )

        output_path = populator.populate(content)
        assert output_path.exists()

        prs = Presentation(str(output_path))
        assert len(prs.slides) == 6

    @needs_nx3
    def test_populate_with_excluded_slides(self, tmp_path: Path) -> None:
        """excluded_slides로 슬라이드를 제거한다."""
        tpl_path = get_template_path("DM", "default")
        populator = TemplatePopulator(
            template_path=tpl_path,
            output_dir=tmp_path,
        )

        content = TemplateContent(
            project_name="Test Exclusion",
            date="February 2026",
            memo_type="DM",
            excluded_slides=[5, 6],  # 마지막 2 슬라이드 제거
        )

        output_path = populator.populate(content)
        prs = Presentation(str(output_path))
        assert len(prs.slides) == 5  # 7 - 2 = 5


# ═══════════════════════════════════════════════════════════════════════════════
# Schemas 테스트
# ═══════════════════════════════════════════════════════════════════════════════


class TestSchemas:
    """스키마 데이터클래스 테스트."""

    def test_template_content_defaults(self) -> None:
        """TemplateContent의 기본값이 올바르다."""
        tc = TemplateContent(
            project_name="Test",
            date="2026-02",
            memo_type="TM",
        )
        assert tc.slides == []
        assert tc.excluded_slides == []

    def test_chart_content_creation(self) -> None:
        """ChartContent를 생성할 수 있다."""
        cc = ChartContent(
            categories=["2022", "2023"],
            series=[SeriesData(name="매출", values=[100, 200])],
        )
        assert len(cc.series) == 1
        assert cc.series[0].values == [100, 200]

    def test_table_content_optional_headers(self) -> None:
        """TableContent의 headers는 선택사항이다."""
        tc = TableContent(rows=[["a", "b"]])
        assert tc.headers is None

        tc2 = TableContent(rows=[["a"]], headers=["Col1"])
        assert tc2.headers == ["Col1"]
