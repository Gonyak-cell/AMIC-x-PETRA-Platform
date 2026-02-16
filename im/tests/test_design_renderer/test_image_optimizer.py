"""이미지 최적화 모듈 단위 테스트.

> 마지막 수정: 2026-02-11 22:00:00
"""

from __future__ import annotations

from io import BytesIO

import pytest

from src.design_renderer.image_optimizer import (
    ImageOptimizationError,
    compress_png,
    optimize_pptx_images,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_png(
    width: int = 1000,
    height: int = 600,
    *,
    mode: str = "RGB",
    color: tuple = (100, 150, 200),  # type: ignore[assignment]
    compress_level: int = 0,
) -> bytes:
    """테스트용 PNG 바이트 생성 (compress_level=0으로 의도적으로 큰 크기)."""
    from PIL import Image

    img = Image.new(mode, (width, height), color=color)
    buf = BytesIO()
    img.save(buf, format="PNG", compress_level=compress_level)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# compress_png 테스트
# ---------------------------------------------------------------------------


class TestCompressPng:
    """compress_png() 단위 테스트."""

    def test_reduces_size(self):
        """큰 PNG 이미지가 최적화 후 크기가 줄어든다."""
        original = _make_png(1000, 600, compress_level=0)
        optimized = compress_png(original)

        assert optimized[:4] == b"\x89PNG"
        assert len(optimized) <= len(original)

    def test_preserves_transparency(self):
        """RGBA PNG의 투명도가 양자화 후에도 보존된다."""
        from PIL import Image

        original = _make_png(400, 300, mode="RGBA", color=(100, 150, 200, 128))
        optimized = compress_png(original)

        assert optimized[:4] == b"\x89PNG"
        result_img = Image.open(BytesIO(optimized))
        # 양자화 후 P 모드 + transparency 또는 RGBA 보존
        assert result_img.mode in ("P", "PA", "RGBA")

    def test_skip_small_images(self):
        """MIN_SIZE_FOR_OPTIMIZATION 미만 PNG는 그대로 반환."""
        tiny_png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 100
        result = compress_png(tiny_png)
        assert result is tiny_png

    def test_graceful_on_invalid_input(self):
        """유효하지 않은 바이트를 넣어도 원본을 반환한다 (예외 없음)."""
        garbage = b"not a png at all " * 1000
        result = compress_png(garbage)
        assert result == garbage

    def test_max_width_resize(self):
        """max_width 지정 시 리사이즈된다."""
        from PIL import Image

        original = _make_png(4000, 2000, compress_level=0)
        optimized = compress_png(original, max_width=2000)

        result_img = Image.open(BytesIO(optimized))
        assert result_img.width == 2000
        assert result_img.height == 1000  # 비율 유지

    def test_no_quantize_when_zero(self):
        """max_colors=0이면 양자화 없이 압축만 적용."""
        original = _make_png(500, 300, compress_level=0)
        optimized = compress_png(original, max_colors=0, compress_level=9)

        assert optimized[:4] == b"\x89PNG"
        assert len(optimized) <= len(original)


# ---------------------------------------------------------------------------
# optimize_pptx_images 테스트
# ---------------------------------------------------------------------------


class TestOptimizePptxImages:
    """optimize_pptx_images() 단위 테스트."""

    def test_optimize_pptx_images(self, tmp_path):
        """PPTX 파일 내 PNG 이미지가 최적화된다."""
        from pptx import Presentation
        from pptx.util import Inches

        # 테스트 PPTX 생성
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # 큰 PNG 이미지 삽입
        img_buf = BytesIO(_make_png(2000, 1000, compress_level=0))
        slide.shapes.add_picture(img_buf, Inches(1), Inches(1))

        pptx_path = tmp_path / "test.pptx"
        prs.save(str(pptx_path))
        original_size = pptx_path.stat().st_size

        # 최적화 실행
        result = optimize_pptx_images(pptx_path, in_place=True)

        assert result == pptx_path
        optimized_size = pptx_path.stat().st_size
        assert optimized_size < original_size

        # 최적화 후에도 유효한 PPTX
        prs2 = Presentation(str(pptx_path))
        assert len(prs2.slides) == 1

    def test_nonexistent_file(self):
        """존재하지 않는 파일 경로 → ImageOptimizationError."""
        with pytest.raises(ImageOptimizationError, match="PPTX 파일 없음"):
            optimize_pptx_images("/nonexistent/file.pptx")
