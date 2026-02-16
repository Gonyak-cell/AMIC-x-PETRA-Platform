"""E2E 파이프라인 테스트 — TITAN/COVENANT/FULL 듀얼 출력."""

from pathlib import Path

import pytest
from pptx import Presentation

from src.design_renderer.im_document import IMDocumentData
from src.design_renderer.pipeline import IMPipeline, PipelineResult
from src.design_renderer.security import SecurityOptions


class TestPipelinePptxOnly:
    """PPTX 전용 생성."""

    def test_titan_pptx(self, titan_data: IMDocumentData, tmp_output: Path):
        """TITAN 프리셋 PPTX 생성."""
        pipeline = IMPipeline()
        result = pipeline.generate_pptx(
            titan_data, output_path=tmp_output / "titan.pptx"
        )
        assert isinstance(result, PipelineResult)
        assert result.pptx_path is not None
        assert result.pptx_path.exists()
        assert result.total_pptx_slides > 0
        # TITAN은 9개 섹션
        assert len(result.section_results) == 9

    def test_covenant_pptx(self, covenant_data: IMDocumentData, tmp_output: Path):
        """COVENANT 프리셋 PPTX 생성."""
        pipeline = IMPipeline()
        result = pipeline.generate_pptx(
            covenant_data, output_path=tmp_output / "covenant.pptx"
        )
        assert result.pptx_path is not None
        assert result.pptx_path.exists()
        # COVENANT은 10개 섹션
        assert len(result.section_results) == 10

    def test_full_pptx(self, full_data: IMDocumentData, tmp_output: Path):
        """FULL 프리셋 PPTX 생성."""
        pipeline = IMPipeline()
        result = pipeline.generate_pptx(
            full_data, output_path=tmp_output / "full.pptx"
        )
        assert result.pptx_path is not None
        assert result.pptx_path.exists()
        # FULL은 19개 섹션
        assert len(result.section_results) == 19
        assert result.total_pptx_slides >= 19  # 최소 섹션당 1슬라이드


class TestPipelineResult:
    """PipelineResult 검증."""

    def test_section_results_populated(
        self, titan_data: IMDocumentData, tmp_output: Path
    ):
        """SectionResult 목록이 올바르게 채워짐."""
        pipeline = IMPipeline()
        result = pipeline.generate_pptx(
            titan_data, output_path=tmp_output / "test.pptx"
        )
        for sr in result.section_results:
            assert sr.section_id != ""
            # toc_divider는 current_section 없이 호출되므로 0일 수 있음
            if sr.section_id != "toc_divider":
                assert sr.pptx_slide_count >= 1
                assert sr.html_slide_count >= 1

    def test_elapsed_time(self, titan_data: IMDocumentData, tmp_output: Path):
        """실행 시간이 기록됨."""
        pipeline = IMPipeline()
        result = pipeline.generate_pptx(
            titan_data, output_path=tmp_output / "test.pptx"
        )
        assert result.elapsed_seconds > 0


class TestPipelineContinueOnError:
    """에러 처리."""

    def test_continue_on_error_default(
        self, titan_data: IMDocumentData, tmp_output: Path
    ):
        """continue_on_error=True → 실패 섹션 건너뛰고 계속."""
        pipeline = IMPipeline(continue_on_error=True)
        result = pipeline.generate_pptx(
            titan_data, output_path=tmp_output / "test.pptx"
        )
        # 정상 데이터이므로 에러 없어야 함
        assert len(result.errors) == 0


class TestPipelineSecurity:
    """보안 옵션 적용."""

    def test_pptx_watermark(self, titan_data: IMDocumentData, tmp_output: Path):
        """PPTX 워터마크 적용."""
        security = SecurityOptions(
            watermark_text="CONFIDENTIAL",
            pptx_read_only=True,
        )
        pipeline = IMPipeline(security=security)
        result = pipeline.generate_pptx(
            titan_data, output_path=tmp_output / "secure.pptx"
        )
        assert result.pptx_path is not None

        # PPTX 열어서 워터마크 텍스트 확인
        prs = Presentation(str(result.pptx_path))
        watermark_found = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            if "CONFIDENTIAL" in run.text:
                                watermark_found = True
        assert watermark_found

    def test_pptx_edit_restriction(
        self, titan_data: IMDocumentData, tmp_output: Path
    ):
        """PPTX 편집 제한."""
        from pptx.oxml.ns import qn

        security = SecurityOptions(
            pptx_read_only=True,
            pptx_edit_password="secret123",
        )
        pipeline = IMPipeline(security=security)
        result = pipeline.generate_pptx(
            titan_data, output_path=tmp_output / "restricted.pptx"
        )

        prs = Presentation(str(result.pptx_path))
        verifier = prs.element.find(qn("p:modifyVerifier"))
        assert verifier is not None


class TestPipelineDualOutput:
    """PPTX + PDF 듀얼 출력."""

    @pytest.mark.requires_browser
    def test_dual_output(self, titan_data: IMDocumentData, tmp_output: Path):
        """PPTX + PDF 동시 생성."""
        pipeline = IMPipeline()
        result = pipeline.generate(
            titan_data,
            pptx_path=tmp_output / "dual.pptx",
            pdf_path=tmp_output / "dual.pdf",
        )
        assert result.pptx_path is not None
        assert result.pdf_path is not None
        assert result.pptx_path.exists()
        assert result.pdf_path.exists()
        assert result.total_pptx_slides > 0
        assert result.total_pdf_pages > 0
