"""폴백 슬라이드 렌더러 테스트.

> 마지막 수정: 2026-02-11 10:30:00
"""

import pytest
from pptx import Presentation

from src.design_renderer.design_tokens import DEFAULT_TOKENS
from src.design_renderer.pptx_engine.slide_factory import SlideFactory
from src.design_renderer.pptx_engine.template_manager import TemplateManager
from src.design_renderer.section_renderers.fallback import (
    render_fallback_slide_html,
    render_fallback_slide_pptx,
)


@pytest.fixture
def prs() -> Presentation:
    manager = TemplateManager(tokens=DEFAULT_TOKENS)
    return manager.new_presentation()


@pytest.fixture
def factory(prs: Presentation) -> SlideFactory:
    manager = TemplateManager(tokens=DEFAULT_TOKENS)
    return SlideFactory(manager, prs=prs, tokens=DEFAULT_TOKENS)


class TestRenderFallbackSlidePptx:
    """PPTX 폴백 슬라이드 테스트."""

    def test_creates_slide(self, factory, prs):
        slide = render_fallback_slide_pptx(
            factory, "financial_analysis", "데이터 없음", prs=prs
        )
        assert slide is not None
        assert len(prs.slides) >= 1

    def test_contains_error_text(self, factory, prs):
        slide = render_fallback_slide_pptx(
            factory,
            "market_overview",
            "시장 데이터 로딩 실패",
            prs=prs,
            show_error_detail=True,
        )
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        texts.append(run.text)
        full_text = " ".join(texts)
        assert "데이터를 불러올 수 없습니다" in full_text
        assert "market_overview" in full_text

    def test_long_error_truncated(self, factory, prs):
        long_error = "x" * 500
        slide = render_fallback_slide_pptx(
            factory,
            "test_section",
            long_error,
            prs=prs,
            show_error_detail=True,
        )
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        texts.append(run.text)
        full_text = " ".join(texts)
        assert "..." in full_text
        assert len(full_text) < 600  # 원본 500자보다 짧게 잘림

    def test_with_custom_tokens(self, factory, prs):
        slide = render_fallback_slide_pptx(
            factory, "cover", "에러", prs=prs, tokens=DEFAULT_TOKENS
        )
        assert slide is not None


class TestRenderFallbackSlideHtml:
    """HTML 폴백 슬라이드 테스트."""

    def test_returns_html_string(self):
        html = render_fallback_slide_html("financial_analysis", "데이터 없음")
        assert isinstance(html, str)
        assert "<div" in html

    def test_contains_section_id(self):
        html = render_fallback_slide_html(
            "market_overview",
            "에러",
            show_error_detail=True,
        )
        assert "market_overview" in html

    def test_contains_error_message(self):
        html = render_fallback_slide_html("test", "시장 데이터 실패")
        assert "데이터를 불러올 수 없습니다" in html

    def test_html_escapes_error(self):
        html = render_fallback_slide_html(
            "test",
            "<script>alert('xss')</script>",
            show_error_detail=True,
        )
        assert "<script>" not in html
        assert "&lt;script&gt;" in html

    def test_long_error_truncated(self):
        html = render_fallback_slide_html(
            "test",
            "y" * 500,
            show_error_detail=True,
        )
        assert "..." in html

    def test_has_fallback_class(self):
        html = render_fallback_slide_html("test", "에러")
        assert "fallback-slide" in html
