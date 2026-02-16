"""PPTX 엔진 테스트 — 템플릿/팩토리/shape/스타일."""

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from src.design_renderer.design_tokens import DEFAULT_TOKENS
from src.design_renderer.pptx_engine.slide_factory import SlideFactory
from src.design_renderer.pptx_engine.style_applier import (
    add_watermark,
    apply_presentation_style,
    apply_run_style,
    set_edit_restriction,
)
from src.design_renderer.pptx_engine.template_manager import TemplateManager


@pytest.fixture
def manager() -> TemplateManager:
    """TemplateManager 인스턴스."""
    return TemplateManager(tokens=DEFAULT_TOKENS)


@pytest.fixture
def prs(manager: TemplateManager) -> Presentation:
    """새 Presentation."""
    return manager.new_presentation()


@pytest.fixture
def factory(manager: TemplateManager, prs: Presentation) -> SlideFactory:
    """SlideFactory 인스턴스."""
    return SlideFactory(manager, prs=prs, tokens=DEFAULT_TOKENS)


class TestTemplateManager:
    """TemplateManager 기본 동작."""

    def test_new_presentation(self, manager: TemplateManager):
        """새 Presentation 생성."""
        prs = manager.new_presentation()
        assert prs is not None
        assert hasattr(prs, "slides")

    def test_presentation_has_layouts(self, prs: Presentation):
        """Presentation에 슬라이드 레이아웃 존재."""
        assert len(prs.slide_layouts) > 0


class TestSlideFactory:
    """SlideFactory 슬라이드 생성."""

    def test_add_cover_slide(self, factory: SlideFactory, prs: Presentation):
        """표지 슬라이드 생성."""
        slide = factory.add_cover_slide(
            project_name="Project TEST",
            subtitle="Information Memorandum",
            date="2026-02-08",
        )
        assert slide is not None
        assert len(prs.slides) >= 1

    def test_add_disclaimer_slide(self, factory: SlideFactory, prs: Presentation):
        """면책조항 슬라이드 생성."""
        slide = factory.add_disclaimer_slide(
            text="테스트 면책조항 텍스트",
            title="Disclaimer",
        )
        assert slide is not None

    def test_add_content_slide(self, factory: SlideFactory, prs: Presentation):
        """콘텐츠 슬라이드 생성."""
        slide = factory.add_content_slide(title="테스트 슬라이드")
        assert slide is not None

    def test_add_contact_slide(self, factory: SlideFactory, prs: Presentation):
        """연락처 슬라이드 생성."""
        contacts = [
            {"name": "홍길동", "title": "MD", "email": "hong@test.com"},
        ]
        slide = factory.add_contact_slide(contacts=contacts)
        assert slide is not None

    def test_multiple_slides(self, factory: SlideFactory, prs: Presentation):
        """여러 슬라이드 순차 생성."""
        factory.add_cover_slide(project_name="Test")
        factory.add_content_slide(title="Page 1")
        factory.add_content_slide(title="Page 2")
        factory.add_contact_slide()
        assert len(prs.slides) == 4


class TestStyleApplier:
    """스타일 적용."""

    def test_apply_presentation_style(self, factory: SlideFactory, prs: Presentation):
        """전체 Presentation 스타일 적용 (에러 없음)."""
        factory.add_cover_slide(project_name="Test")
        factory.add_content_slide(title="Content")
        apply_presentation_style(prs, tokens=DEFAULT_TOKENS)
        # 에러 없이 완료되면 통과

    def test_set_edit_restriction_read_only(
        self, factory: SlideFactory, prs: Presentation
    ):
        """읽기 전용 편집 제한 → modifyVerifier XML 존재."""
        factory.add_cover_slide(project_name="Test")
        set_edit_restriction(prs, read_only=True)
        verifier = prs.element.find(qn("p:modifyVerifier"))
        assert verifier is not None

    def test_set_edit_restriction_with_password(
        self, factory: SlideFactory, prs: Presentation
    ):
        """비밀번호 편집 제한 → SHA-512 해시 존재."""
        factory.add_cover_slide(project_name="Test")
        set_edit_restriction(prs, read_only=True, password="test1234")
        verifier = prs.element.find(qn("p:modifyVerifier"))
        assert verifier is not None
        assert verifier.get("cryptAlgorithmSid") == "14"  # SHA-512
        assert verifier.get("hashData") != ""
        assert verifier.get("saltData") != ""


class TestWatermark:
    """워터마크 기능."""

    def test_add_watermark(self, factory: SlideFactory, prs: Presentation):
        """워터마크 추가 → 슬라이드에 텍스트 shape 존재."""
        factory.add_cover_slide(project_name="Test")
        factory.add_content_slide(title="Page 1")
        factory.add_content_slide(title="Page 2")

        add_watermark(prs, text="CONFIDENTIAL", skip_first_slide=True)

        # 두 번째, 세 번째 슬라이드에 워터마크 shape 존재
        for i in [1, 2]:
            slide = prs.slides[i]
            watermark_found = False
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            if "CONFIDENTIAL" in run.text:
                                watermark_found = True
            assert watermark_found, f"슬라이드 {i+1}에 워터마크 없음"

    def test_watermark_skip_first(self, factory: SlideFactory, prs: Presentation):
        """skip_first_slide=True → 첫 슬라이드에 워터마크 없음."""
        factory.add_cover_slide(project_name="Test")
        factory.add_content_slide(title="Page 1")

        add_watermark(prs, text="SECRET", skip_first_slide=True)

        # 첫 슬라이드에는 워터마크 없음
        cover_slide = prs.slides[0]
        for shape in cover_slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        assert "SECRET" not in run.text
