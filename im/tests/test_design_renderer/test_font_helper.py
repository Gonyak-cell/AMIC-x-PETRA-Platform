"""font_helper 테스트 — a:ea (동아시아) 폰트 설정 검증.

> 마지막 수정: 2026-02-11 10:00:00
"""

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from src.design_renderer.pptx_engine.font_helper import (
    ensure_ea_fonts_on_presentation,
    ensure_ea_fonts_on_slide,
    set_font_with_ea,
)

_NSMAP = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def _qn(tag: str) -> str:
    prefix, local = tag.split(":")
    return f"{{{_NSMAP[prefix]}}}{local}"


def _get_ea_typeface(run) -> str | None:
    """run의 a:ea typeface 속성을 반환. 없으면 None."""
    rpr = run._r.find(_qn("a:rPr"))
    if rpr is None:
        return None
    ea = rpr.find(_qn("a:ea"))
    if ea is None:
        return None
    return ea.get("typeface")


@pytest.fixture
def prs() -> Presentation:
    """빈 Presentation."""
    return Presentation()


@pytest.fixture
def slide_with_textbox(prs: Presentation):
    """텍스트박스가 포함된 슬라이드."""
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    txbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    tf = txbox.text_frame
    run = tf.paragraphs[0].add_run()
    run.text = "테스트 한글 텍스트"
    run.font.size = Pt(12)
    return slide, run


class TestSetFontWithEa:
    """set_font_with_ea 단위 테스트."""

    def test_sets_latin_font(self, slide_with_textbox):
        """Latin 폰트(a:latin)가 설정되는지 확인."""
        _, run = slide_with_textbox
        set_font_with_ea(run, "Pretendard")
        assert run.font.name == "Pretendard"

    def test_sets_ea_font_same_as_latin(self, slide_with_textbox):
        """ea_font 미지정 시 Latin과 동일한 폰트가 a:ea에 설정."""
        _, run = slide_with_textbox
        set_font_with_ea(run, "Pretendard")
        assert _get_ea_typeface(run) == "Pretendard"

    def test_sets_ea_font_different(self, slide_with_textbox):
        """ea_font를 별도 지정하면 a:ea에 해당 폰트 설정."""
        _, run = slide_with_textbox
        set_font_with_ea(run, "Inter", ea_font="Noto Sans KR")
        assert run.font.name == "Inter"
        assert _get_ea_typeface(run) == "Noto Sans KR"

    def test_overwrites_existing_ea(self, slide_with_textbox):
        """기존 a:ea가 있으면 덮어쓴다."""
        _, run = slide_with_textbox
        set_font_with_ea(run, "Inter", ea_font="Noto Sans KR")
        assert _get_ea_typeface(run) == "Noto Sans KR"

        # 다시 설정
        set_font_with_ea(run, "Inter", ea_font="Pretendard")
        assert _get_ea_typeface(run) == "Pretendard"

        # a:ea가 중복 생성되지 않는지 확인
        rpr = run._r.find(_qn("a:rPr"))
        ea_elements = rpr.findall(_qn("a:ea"))
        assert len(ea_elements) == 1

    def test_mono_font(self, slide_with_textbox):
        """IBM Plex Mono도 a:ea가 설정되는지 확인."""
        _, run = slide_with_textbox
        set_font_with_ea(run, "IBM Plex Mono")
        assert run.font.name == "IBM Plex Mono"
        assert _get_ea_typeface(run) == "IBM Plex Mono"

    def test_xml_structure(self, slide_with_textbox):
        """생성된 XML이 올바른 구조인지 확인."""
        _, run = slide_with_textbox
        set_font_with_ea(run, "Pretendard")

        r_elem = run._r
        rpr = r_elem.find(_qn("a:rPr"))
        assert rpr is not None

        ea = rpr.find(_qn("a:ea"))
        assert ea is not None
        assert ea.tag == _qn("a:ea")
        assert ea.get("typeface") == "Pretendard"


class TestEnsureEaFontsOnSlide:
    """ensure_ea_fonts_on_slide 테스트."""

    def test_adds_ea_to_runs_without_ea(self, prs: Presentation):
        """a:ea가 없는 run에 폰트를 추가."""
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)

        txbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = txbox.text_frame.paragraphs[0].add_run()
        run.text = "한글 텍스트"
        run.font.name = "Pretendard"  # a:latin만 설정됨

        count = ensure_ea_fonts_on_slide(slide, ea_font="Pretendard")
        assert count >= 1
        assert _get_ea_typeface(run) == "Pretendard"

    def test_skips_runs_with_existing_ea(self, prs: Presentation):
        """이미 a:ea가 있는 run은 건너뛴다."""
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)

        txbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = txbox.text_frame.paragraphs[0].add_run()
        run.text = "한글"
        set_font_with_ea(run, "Inter", ea_font="Noto Sans KR")

        count = ensure_ea_fonts_on_slide(slide, ea_font="Pretendard")
        assert count == 0
        # 기존 설정이 유지되어야 함
        assert _get_ea_typeface(run) == "Noto Sans KR"

    def test_multiple_runs(self, prs: Presentation):
        """여러 run이 있는 경우 모두 처리."""
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)

        txbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
        tf = txbox.text_frame

        runs = []
        for i in range(3):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            run = p.add_run()
            run.text = f"텍스트 {i}"
            run.font.name = "Pretendard"
            runs.append(run)

        count = ensure_ea_fonts_on_slide(slide, ea_font="Pretendard")
        assert count >= 3
        for run in runs:
            assert _get_ea_typeface(run) == "Pretendard"

    def test_empty_slide(self, prs: Presentation):
        """빈 슬라이드는 0 반환."""
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        count = ensure_ea_fonts_on_slide(slide, ea_font="Pretendard")
        assert count == 0


class TestEnsureEaFontsOnPresentation:
    """ensure_ea_fonts_on_presentation 테스트."""

    def test_processes_all_slides(self, prs: Presentation):
        """모든 슬라이드를 처리."""
        layout = prs.slide_layouts[0]

        for _ in range(3):
            slide = prs.slides.add_slide(layout)
            txbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
            run = txbox.text_frame.paragraphs[0].add_run()
            run.text = "텍스트"
            run.font.name = "Pretendard"

        total = ensure_ea_fonts_on_presentation(prs, ea_font="Pretendard")
        assert total >= 3

    def test_empty_presentation(self):
        """슬라이드가 없는 프레젠테이션은 0 반환."""
        prs = Presentation()
        total = ensure_ea_fonts_on_presentation(prs, ea_font="Pretendard")
        assert total == 0

    def test_mixed_slides(self, prs: Presentation):
        """일부 run에만 a:ea가 없는 혼합 슬라이드."""
        layout = prs.slide_layouts[0]

        # 슬라이드 1: a:ea 있음
        slide1 = prs.slides.add_slide(layout)
        txbox1 = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run1 = txbox1.text_frame.paragraphs[0].add_run()
        run1.text = "이미 설정됨"
        set_font_with_ea(run1, "Inter", ea_font="Noto Sans KR")

        # 슬라이드 2: a:ea 없음
        slide2 = prs.slides.add_slide(layout)
        txbox2 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run2 = txbox2.text_frame.paragraphs[0].add_run()
        run2.text = "설정 필요"
        run2.font.name = "Pretendard"

        total = ensure_ea_fonts_on_presentation(prs, ea_font="Pretendard")
        assert total >= 1  # 슬라이드 2의 run만 처리

        # 슬라이드 1의 기존 설정이 유지됨
        assert _get_ea_typeface(run1) == "Noto Sans KR"
        # 슬라이드 2는 새로 설정됨
        assert _get_ea_typeface(run2) == "Pretendard"
